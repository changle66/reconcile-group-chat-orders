from __future__ import annotations

import base64
import json
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import core
import finance_materials
import reconcile


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl6qWQAAAAASUVORK5CYII="
)


class FinanceMaterialsTests(unittest.TestCase):
    def _start_fixture(self, root: Path) -> Path:
        source = root / "source"
        finance_export = source / "finance"
        photos = finance_export / "photos"
        photos.mkdir(parents=True)
        for name in ("passport-a.png", "wechat.png", "passport-repeat.png", "line.png"):
            # PNG readers ignore bytes after IEND; the suffix keeps each
            # synthetic image valid while giving it a distinct content hash.
            (photos / name).write_bytes(PNG_1X1 + name.encode("utf-8"))
        (finance_export / "result.json").write_text(
            json.dumps(
                {
                    "id": "finance-group",
                    "name": "财务资料群",
                    "messages": [
                        {
                            "id": 1,
                            "type": "message",
                            "date_unixtime": "1780000000",
                            "from": "资料员",
                            "from_id": "user:staff",
                            "text": "护照",
                            "photo": "photos/passport-a.png",
                        },
                        {
                            "id": 2,
                            "type": "message",
                            "date_unixtime": "1780000060",
                            "from": "资料员",
                            "from_id": "user:staff",
                            "text": "微信资料",
                            "photo": "photos/wechat.png",
                        },
                        {
                            "id": 3,
                            "type": "message",
                            "date_unixtime": "1780000120",
                            "from": "资料员",
                            "from_id": "user:staff",
                            "text": "重复护照",
                            "photo": "photos/passport-repeat.png",
                        },
                        {
                            "id": 4,
                            "type": "message",
                            "date_unixtime": "1780000180",
                            "from": "资料员",
                            "from_id": "user:staff",
                            "text": "LINE 资料",
                            "photo": "photos/line.png",
                        },
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        other_export = source / "other"
        other_export.mkdir(parents=True)
        (other_export / "result.json").write_text(
            json.dumps(
                {
                    "id": "other-group",
                    "name": "普通换汇群",
                    "messages": [
                        {
                            "id": 1,
                            "type": "message",
                            "date_unixtime": "1780000000",
                            "from": "Alice",
                            "from_id": "user:alice",
                            "text": "不应进入财务资料模式",
                        }
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        work = root / "run"
        report = reconcile.start_run(
            reconcile.parse_args(
                [
                    "start",
                    str(source),
                    "--work",
                    str(work),
                    "--mode",
                    "finance",
                    "--no-ocr-candidates",
                ]
            )
        )
        self.assertEqual(report["groups"], 1)
        self.assertEqual(report["selected_groups"][0]["group_name"], "财务资料群")
        run = reconcile._load_run(work)
        self.assertEqual(run["group_mode"], finance_materials.GROUP_MODE)
        self.assertEqual(run["group_name_contains"], finance_materials.GROUP_NAME_MARKER)
        decision = reconcile._load_json(reconcile._decision_path(work, run["groups"][0]))
        self.assertEqual(decision["contract_version"], finance_materials.DECISION_CONTRACT)
        return work

    def _page(self, work: Path) -> dict:
        page = reconcile.review_command(
            Namespace(work=work, review_action="next", group="财务资料群", limit=500)
        )
        self.assertIn("open_people", page)
        self.assertNotIn("open_orders", page)
        return page

    def _confirmed_person(self) -> dict:
        return {
            "id": "P001",
            "name": "张三",
            "surname": "ZHANG",
            "given_names": "SAN",
            "nationality": "CHINESE",
            "birth_date": "1990-01-02",
            "documents": [
                {
                    "type": "passport",
                    "country_code": "CHN",
                    "number": "E01234567",
                    "media_labels": ["M0001", "M0003"],
                }
            ],
            "accounts": [
                {
                    "platform": "WeChat",
                    "account_id": "wxid_full_value",
                    "phone": "+86 0138 0000 0000",
                    "media_labels": ["M0002"],
                },
                {
                    "platform": "LINE",
                    "account_id": "line-full-id",
                    "phone": "",
                    "media_labels": ["M0004"],
                },
            ],
            "source_messages": [],
            "association_status": "confirmed",
            "note": "",
        }

    def _apply(
        self,
        root: Path,
        work: Path,
        page: dict,
        *,
        people: list[dict],
        media_decisions: dict[str, dict],
        batch_id: str,
        media_observations: dict[str, dict] | None = None,
        media_view_metrics: dict[str, int] | None = None,
    ) -> dict:
        batch_path = root / f"{batch_id}.json"
        payload = {
            "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
            "batch_id": batch_id,
            "base_fingerprint": page["semantic_fingerprint"],
            "page_commit": {
                "page_start": page["page_start"],
                "page_end": page["page_end"],
                "page_token": page["page_token"],
            },
            "open_people": [],
            "media_decisions": media_decisions,
            "people": people,
        }
        if media_observations is not None:
            payload["media_observations"] = media_observations
        if media_view_metrics is not None:
            payload["media_view_metrics"] = media_view_metrics
        core.atomic_json(batch_path, payload)
        return reconcile.review_command(
            Namespace(
                work=work,
                review_action="apply-batch",
                group="财务资料群",
                input=batch_path,
            )
        )

    def test_finance_mode_merges_repeated_passport_and_embeds_all_originals(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work = self._start_fixture(root)
            page = self._page(work)
            self.assertEqual(page["media_queue"]["mode"], "finance_materials")
            self.assertEqual(
                page["media_queue"]["recommended_parallel_limit"],
                reconcile.FINANCE_MEDIA_BATCH_LIMIT,
            )
            self.assertEqual(page["media_queue"]["queued_images"], 4)
            self.assertEqual(len(page["media_queue"]["batches"]), 1)
            self.assertEqual(
                [
                    item["label"]
                    for item in page["media_queue"]["batches"][0]["items"]
                ],
                ["M0001", "M0002", "M0003", "M0004"],
            )
            media_decisions = {
                "M0001": {"classification": "document", "viewed_original": True},
                "M0002": {"classification": "chat_profile", "viewed_original": True},
                "M0003": {"classification": "document", "viewed_original": True},
                "M0004": {"classification": "chat_profile", "viewed_original": True},
            }
            applied = self._apply(
                root,
                work,
                page,
                people=[self._confirmed_person()],
                media_decisions=media_decisions,
                batch_id="finance-confirmed-001",
                media_observations={
                    label: {
                        "contract_version": reconcile.MEDIA_OBSERVATION_CONTRACT,
                        "classification": media_decisions[label]["classification"],
                        "review_status": "clear",
                        "viewed_original": True,
                        "recheck_reasons": [],
                        **(
                            {
                                "facts": {
                                    "holder": {
                                        "name": "张三",
                                        "surname": "ZHANG",
                                        "given_names": "SAN",
                                        "nationality": "CHINESE",
                                        "birth_date": "1990-01-02",
                                    },
                                    "document": {
                                        "type": "passport",
                                        "country_code": "CHN",
                                        "number": "E01234567",
                                    },
                                }
                            }
                            if label in {"M0001", "M0003"}
                            else {}
                        ),
                    }
                    for label in media_decisions
                },
                media_view_metrics={
                    "view_batches": 1,
                    "opened_images": 4,
                    "failed_images": 0,
                    "single_image_rechecks": 0,
                    "elapsed_ms": 400,
                },
            )
            self.assertEqual(applied["people"], 1)
            self.assertEqual(applied["documents"], 1)
            self.assertEqual(applied["accounts"], 2)
            self.assertEqual(applied["pending_people"], 0)
            self.assertEqual(applied["evidence_hashes_computed"], 4)
            self.assertEqual(applied["evidence_hashes_reused"], 0)
            self.assertEqual(applied["observation_cache_entries"], 4)
            self.assertEqual(applied["media_recheck_required"], 0)
            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="财务资料群")
            )
            self.assertEqual(sealed["evidence_hashes_computed"], 4)
            self.assertEqual(sealed["evidence_hashes_reused"], 0)
            output = root / "财务资料.xlsx"
            result = reconcile.finish_run(
                Namespace(
                    work=work,
                    output=output,
                    template=Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
                )
            )
            self.assertEqual(result["people"], 1)
            self.assertEqual(result["documents"], 1)
            self.assertEqual(result["accounts"], 2)
            self.assertEqual(result["material_images"], 4)
            self.assertEqual(result["pending_people"], 0)

            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                self.assertEqual(len(workbook.sheetnames), 1)
                worksheet = workbook[workbook.sheetnames[0]]
                self.assertEqual(
                    [worksheet.cell(1, column).value for column in range(1, len(finance_materials.HEADERS) + 1)],
                    finance_materials.HEADERS,
                )
                self.assertEqual(worksheet.cell(2, 6).value, "E01234567")
                self.assertEqual(worksheet.cell(2, 8).value, "1990-01-02")
                self.assertEqual(worksheet.cell(2, 11).value, "wxid_full_value\nline-full-id")
                self.assertEqual(worksheet.cell(2, 12).value, "+86 0138 0000 0000\n")
                self.assertEqual(worksheet.cell(2, 6).number_format, "@")
                self.assertEqual(worksheet.cell(2, 12).number_format, "@")
                self.assertEqual(len(worksheet._images), 4)
            finally:
                workbook.close()

    def test_finance_mode_rejects_duplicate_passport_across_people_atomically(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work = self._start_fixture(root)
            page = self._page(work)
            first = self._confirmed_person()
            first["documents"][0]["media_labels"] = ["M0001"]
            first["accounts"] = []
            second = self._confirmed_person()
            second["id"] = "P002"
            second["name"] = "李四"
            second["surname"] = "LI"
            second["given_names"] = "SI"
            second["documents"][0]["number"] = "e 01234567"
            second["documents"][0]["media_labels"] = ["M0003"]
            second["accounts"] = []
            before_run = reconcile._load_run(work)
            decision_path = reconcile._decision_path(work, before_run["groups"][0])
            before = decision_path.read_bytes()
            with self.assertRaisesRegex(ValueError, "appears under multiple people"):
                self._apply(
                    root,
                    work,
                    page,
                    people=[first, second],
                    media_decisions={
                        "M0001": {"classification": "document", "viewed_original": True},
                        "M0002": {"classification": "reference"},
                        "M0003": {"classification": "document", "viewed_original": True},
                        "M0004": {"classification": "reference"},
                    },
                    batch_id="finance-duplicate-001",
                )
            self.assertEqual(decision_path.read_bytes(), before)

    def test_unmatched_chat_profile_is_preserved_as_pending_row(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work = self._start_fixture(root)
            page = self._page(work)
            pending = {
                "id": "PENDING001",
                "name": "",
                "surname": "",
                "given_names": "",
                "nationality": "",
                "birth_date": "",
                "documents": [],
                "accounts": [
                    {
                        "platform": "WhatsApp",
                        "account_id": "",
                        "phone": "+855 012 345 678",
                        "media_labels": ["M0002"],
                    }
                ],
                "source_messages": [],
                "association_status": "pending",
                "note": "资料页与哪一本证件对应不明确",
            }
            self._apply(
                root,
                work,
                page,
                people=[pending],
                media_decisions={
                    "M0001": {"classification": "reference"},
                    "M0002": {"classification": "chat_profile", "viewed_original": True},
                    "M0003": {"classification": "reference"},
                    "M0004": {"classification": "reference"},
                },
                batch_id="finance-pending-001",
            )
            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="财务资料群")
            )
            self.assertEqual(sealed["pending_people"], 1)
            output = root / "待确认财务资料.xlsx"
            result = reconcile.finish_run(
                Namespace(
                    work=work,
                    output=output,
                    template=Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
                )
            )
            self.assertEqual(result["pending_people"], 1)
            self.assertEqual(result["risk_report"]["flagged_groups"], 1)
            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                worksheet = workbook[workbook.sheetnames[0]]
                self.assertEqual(worksheet.cell(2, 12).value, "+855 012 345 678")
                self.assertEqual(worksheet.cell(2, 14).value, "待确认")
            finally:
                workbook.close()

    def test_non_passport_identity_documents_are_accepted_in_one_person_row(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work = self._start_fixture(root)
            page = self._page(work)
            person = {
                "id": "P002",
                "name": "示例人员",
                "surname": "",
                "given_names": "",
                "nationality": "THAI",
                "birth_date": "",
                "documents": [
                    {
                        "type": "身份证",
                        "country_code": "THA",
                        "number": "ID00001234",
                        "media_labels": ["M0001"],
                    },
                    {
                        "type": "签证",
                        "country_code": "THA",
                        "number": "VISA0005678",
                        "media_labels": ["M0002"],
                    },
                    {
                        "type": "居留证",
                        "country_code": "THA",
                        "number": "RES00009012",
                        "media_labels": ["M0003"],
                    },
                    {
                        "type": "驾照",
                        "country_code": "THA",
                        "number": "DL00003456",
                        "media_labels": ["M0004"],
                    },
                ],
                "accounts": [],
                "source_messages": [],
                "association_status": "confirmed",
                "note": "",
            }
            applied = self._apply(
                root,
                work,
                page,
                people=[person],
                media_decisions={
                    label: {"classification": "document", "viewed_original": True}
                    for label in ("M0001", "M0002", "M0003", "M0004")
                },
                batch_id="finance-other-documents-001",
            )
            self.assertEqual(applied["people"], 1)
            self.assertEqual(applied["documents"], 4)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="财务资料群")
            )
            output = root / "其他证件.xlsx"
            reconcile.finish_run(
                Namespace(
                    work=work,
                    output=output,
                    template=Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
                )
            )
            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                worksheet = workbook[workbook.sheetnames[0]]
                self.assertEqual(worksheet.cell(2, 4).value, "身份证\n签证\n居留证\n驾照")
                self.assertEqual(
                    worksheet.cell(2, 6).value,
                    "ID00001234\nVISA0005678\nRES00009012\nDL00003456",
                )
                self.assertEqual(worksheet.cell(2, 6).number_format, "@")
                self.assertEqual(len(worksheet._images), 4)
            finally:
                workbook.close()

    def test_ios_manifest_route_passes_finance_group_pattern_to_existing_reader(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            backup = root / "ios-backup"
            backup.mkdir()
            work = root / "run"

            def fake_extract(arguments: list[str]) -> int:
                self.assertIn("--group-pattern", arguments)
                pattern_index = arguments.index("--group-pattern") + 1
                self.assertEqual(arguments[pattern_index], "财务资料群")
                self.assertNotIn("--all-group-chats", arguments)
                output_index = arguments.index("-o") + 1
                output = Path(arguments[output_index])
                output.parent.mkdir(parents=True, exist_ok=True)
                output.write_text(
                    json.dumps(
                        {
                            "contract_version": core.NORMALIZED_CONTRACT,
                            "timezone": "Asia/Bangkok",
                            "source_fingerprint": "sha256:" + "0" * 64,
                            "groups": [
                                {
                                    "group_key": "line:finance-group",
                                    "group_name": "财务资料群",
                                    "platform": "line",
                                    "messages": [],
                                }
                            ],
                        },
                        ensure_ascii=False,
                    ),
                    encoding="utf-8",
                )
                return 0

            with patch("reconcile._line_has_matching_group", return_value=True), patch(
                "reconcile.extract_line_ios.main", side_effect=fake_extract
            ):
                documents = reconcile._line_documents(
                    work,
                    [backup],
                    contains=finance_materials.GROUP_NAME_MARKER,
                    group_mode=finance_materials.GROUP_MODE,
                    timezone_name="Asia/Bangkok",
                    roster_path=root / "roster.yaml",
                    self_name="LINE_SELF",
                )
            self.assertEqual(len(documents), 1)
            self.assertEqual(documents[0]["groups"][0]["group_name"], "财务资料群")


if __name__ == "__main__":
    unittest.main()
