from __future__ import annotations

import copy
import json
import subprocess
import sys
import tempfile
import unittest
from argparse import Namespace
from datetime import datetime
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

import build_workbook
import check_workbook
import core
import reconcile


class ReconcileWorkflowTests(unittest.TestCase):
    def test_redmi_whatsapp_export_uses_bracket_timestamps_and_parent_group_name(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            export = root / "WhatsApp Chat - 小额出🐱🐱🐱群1"
            export.mkdir()
            source = export / "_chat.txt"
            source.write_text(
                "[2026/9/4 15:53:27] Alice: 第一行\n"
                "续行\n"
                "\u200e[2026/9/4 16:03:39] QQ~x: 齐\n",
                encoding="utf-8",
            )

            telegram, whatsapp = reconcile.normalize_exports.discover_sources([root])
            self.assertEqual(telegram, [])
            self.assertEqual(whatsapp, [source.resolve()])
            self.assertEqual(
                reconcile.normalize_exports.whatsapp_group_name(source),
                "小额出🐱🐱🐱群1",
            )
            blocks = reconcile.normalize_exports.split_whatsapp_blocks(source)
            self.assertEqual(
                blocks,
                [
                    (datetime(2026, 9, 4, 15, 53, 27), "Alice: 第一行\n续行", 1),
                    (datetime(2026, 9, 4, 16, 3, 39), "QQ~x: 齐", 3),
                ],
            )

            attachment = export / "00000444-PHOTO-2026-09-04-12-43-33.jpg"
            attachment.write_bytes(b"fresh-image")
            sender, text = reconcile.normalize_exports.whatsapp_sender_and_text(
                "~ 刘超雄:<附件：00000444-PHOTO-2026-09-04-12-43-33.jpg>"
            )
            self.assertEqual(sender, "~ 刘超雄")
            media, warnings, fingerprint_paths = reconcile.normalize_exports.whatsapp_media(
                text, source
            )
            self.assertEqual(warnings, [])
            self.assertEqual(fingerprint_paths, [])
            self.assertEqual(len(media), 1)
            self.assertEqual(media[0]["availability"], "available")
            self.assertEqual(media[0]["path"], str(attachment.resolve()))

    def _start_fixture(
        self,
        root: Path,
        *,
        messages: list[dict] | None = None,
        extra_messages: list[dict] | None = None,
        controlled: bool = False,
        commit_page: bool = True,
        accounting_date: str | None = None,
        accounting_from: str | None = None,
        accounting_to: str | None = None,
        ocr_candidates: bool | None = False,
    ) -> tuple[Path, dict, dict]:
        export = root / "ChatExport"
        photos = export / "photos"
        photos.mkdir(parents=True)
        (photos / "customer.jpg").write_bytes(b"customer-evidence")
        (photos / "staff.jpg").write_bytes(b"staff-evidence")
        raw = {
            "id": "fixture-small",
            "name": "测试小额群",
            "messages": messages
            or [
                {
                    "id": 1,
                    "type": "message",
                    "date_unixtime": "1780000000",
                    "from": "Alice",
                    "from_id": "user:alice",
                    "text": "付款100元",
                    "photo": "photos/customer.jpg",
                },
                {
                    "id": 2,
                    "type": "message",
                    "date_unixtime": "1780000060",
                    "from": "QQ财务3",
                    "from_id": "user:staff",
                    "text": "回500铢",
                    "photo": "photos/staff.jpg",
                },
            ],
        }
        raw["messages"].extend(extra_messages or [])
        (export / "result.json").write_text(
            json.dumps(raw, ensure_ascii=False), encoding="utf-8"
        )
        work = root / "run"
        report = reconcile.start_run(
            Namespace(
                inputs=[export],
                work=work,
                contains="小额",
                date=accounting_date,
                accounting_from=accounting_from,
                accounting_to=accounting_to,
                timezone="Asia/Bangkok",
                roster=Path(__file__).resolve().parent.parent / "config" / "roster.yaml",
                line_backup=[],
                line_self_name="LINE_SELF",
                ocr_candidates=ocr_candidates,
            )
        )
        run = reconcile._load_run(work)
        self.assertEqual(report["runtime_contract"], reconcile.RUNTIME_CONTRACT)
        self.assertEqual(report["amount_policy"], reconcile.AMOUNT_POLICY)
        self.assertEqual(
            report["workbook_compatibility"],
            reconcile.WORKBOOK_COMPATIBILITY,
        )
        self.assertEqual(run["runtime_contract"], reconcile.RUNTIME_CONTRACT)
        self.assertEqual(run["amount_policy"], reconcile.AMOUNT_POLICY)
        self.assertEqual(
            run["workbook_compatibility"],
            reconcile.WORKBOOK_COMPATIBILITY,
        )
        if accounting_date is None:
            self.assertNotIn("accounting_date", report)
            self.assertNotIn("accounting_date", run)
        else:
            self.assertEqual(report["accounting_date"], accounting_date)
            self.assertEqual(run["accounting_date"], accounting_date)
        if accounting_from is None and accounting_to is None:
            self.assertNotIn("accounting_from", report)
            self.assertNotIn("accounting_from", run)
            self.assertNotIn("accounting_to", report)
            self.assertNotIn("accounting_to", run)
        else:
            self.assertEqual(report["accounting_from"], accounting_from)
            self.assertEqual(report["accounting_to"], accounting_to)
            self.assertEqual(run["accounting_from"], accounting_from)
            self.assertEqual(run["accounting_to"], accounting_to)
        normalized = reconcile._load_snapshot(work, run)
        self.assertEqual(report["groups"], 1)
        self.assertTrue(
            all(
                isinstance(media.get("blob_sha256"), str)
                and len(media["blob_sha256"]) == 64
                for message in normalized["groups"][0]["messages"]
                for media in message["media"]
                if media.get("availability") == "available"
            )
        )
        page = reconcile.review_command(
            Namespace(work=work, review_action="next", group=None, limit=500)
        )
        if commit_page:
            page_commit_path = root / "fixture-page-commit.json"
            core.atomic_json(
                page_commit_path,
                {
                    "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
                    "batch_id": "fixture-page-commit",
                    "base_fingerprint": page["semantic_fingerprint"],
                    "page_commit": {
                        "page_start": page["page_start"],
                        "page_end": page["page_end"],
                        "page_token": page["page_token"],
                    },
                    "open_orders": [],
                },
            )
            reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=page_commit_path,
                )
            )
        if not controlled:
            # Compatibility-only fixture. New business regression tests must
            # opt into controlled editing and submit semantics with apply-batch.
            run = reconcile._load_run(work)
            run_group = run["groups"][0]
            decision_path = reconcile._decision_path(work, run_group)
            decision = reconcile._load_json(decision_path)
            run_group.pop("edit_mode", None)
            decision.pop("edit_control", None)
            core.atomic_json(decision_path, decision)
            core.atomic_json(reconcile._run_path(work), run)
        return work, run, normalized

    @staticmethod
    def _synthetic_media_group(
        root: Path,
        *,
        count: int,
        identical: bool = False,
    ) -> dict:
        media_root = root / "synthetic-media"
        media_root.mkdir(parents=True, exist_ok=True)
        messages: list[dict] = []
        shared_path = media_root / "shared.jpg"
        if identical:
            shared_path.write_bytes(b"identical-image")
        for index in range(count):
            path = shared_path if identical else media_root / f"image-{index + 1}.jpg"
            if not identical:
                path.write_bytes(f"unique-image-{index + 1}".encode())
            messages.append(
                {
                    "message_id": f"message-{index + 1}",
                    "sender_name": "Alice",
                    "sender_id": "user:alice",
                    "role": "客户候选",
                    "media": [
                        {
                            "media_id": f"media-{index + 1}",
                            "kind": "image",
                            "mime_type": "image/jpeg",
                            "availability": "available",
                            "path": str(path.resolve()),
                            "byte_size": path.stat().st_size,
                            "blob_sha256": core.sha256_file(path),
                        }
                    ],
                }
            )
        return {"messages": messages}

    def _start_large_order_fixture(
        self,
        root: Path,
        specs: list[tuple[str, str, str, str, str, str, str]],
        *,
        group_name: str = "曼谷固定换汇群",
        include_empty_group: bool = False,
    ) -> tuple[Path, dict[str, dict], list[dict]]:
        source_root = root / "source"
        export = source_root / "large"
        photos = export / "photos"
        photos.mkdir(parents=True)
        timestamp = int(
            datetime.fromisoformat("2026-08-31T09:00:00+07:00").timestamp()
        )
        messages: list[dict] = []
        media_decisions: dict[str, dict] = {}
        orders: list[dict] = []
        for position, spec in enumerate(specs, start=1):
            text, fund_type, direction, source_amount, rate, operator, target_amount = spec
            source_currency, target_currency = direction.split("->", maxsplit=1)
            payment_file = f"payment-{position}.jpg"
            payout_file = f"payout-{position}.jpg"
            (photos / payment_file).write_bytes(f"payment-{position}".encode())
            (photos / payout_file).write_bytes(f"payout-{position}".encode())
            payment_message_id = position * 2 - 1
            payout_message_id = position * 2
            messages.extend(
                [
                    {
                        "id": payment_message_id,
                        "type": "message",
                        "date_unixtime": str(timestamp + payment_message_id * 60),
                        "from": "Alice",
                        "from_id": "user:alice",
                        "text": text,
                        "photo": f"photos/{payment_file}",
                    },
                    {
                        "id": payout_message_id,
                        "type": "message",
                        "date_unixtime": str(timestamp + payout_message_id * 60),
                        "from": "QQ财务3",
                        "from_id": "user:staff",
                        "text": f"已回款 {target_amount} {target_currency}",
                        "photo": f"photos/{payout_file}",
                    },
                ]
            )
            payment_label = f"M{payment_message_id:04d}"
            payout_label = f"M{payout_message_id:04d}"
            payment_source = f"S{payment_message_id:05d}"
            payout_source = f"S{payout_message_id:05d}"
            media_decisions[payment_label] = {
                "classification": "fund",
                "viewed_original": True,
                "entries": [
                    {
                        "amount": source_amount,
                        "currency": source_currency,
                        "payee": (
                            f"100-0-XXX{position:03d}"
                            if source_currency == "THB"
                            else f"customer-payee-{position}"
                        ),
                    }
                ],
            }
            media_decisions[payout_label] = {
                "classification": "fund",
                "viewed_original": True,
                "entries": [
                    {
                        "amount": target_amount,
                        "currency": target_currency,
                        "payee": (
                            f"200-0-XXX{position:03d}"
                            if target_currency == "THB"
                            else f"staff-payee-{position}"
                        ),
                    }
                ],
            }
            orders.append(
                {
                    "id": f"L{position:03d}",
                    "entry_ids": [f"{payment_label}.1", f"{payout_label}.1"],
                    "source_messages": [payment_source, payout_source],
                    "fund_type": fund_type,
                    "customer_nickname": "Alice",
                    "direction": direction,
                    "pricing": {
                        "source_messages": [payment_source],
                        "terms": {"rate": rate, "operator": operator},
                        "expected": {"kind": "explicit", "amount": target_amount},
                    },
                }
            )
        (export / "result.json").write_text(
            json.dumps(
                {"id": "large-orders", "name": group_name, "messages": messages},
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        if include_empty_group:
            empty_export = source_root / "empty"
            empty_export.mkdir(parents=True)
            (empty_export / "result.json").write_text(
                json.dumps(
                    {
                        "id": "large-empty",
                        "name": "当天无换汇群",
                        "messages": [
                            {
                                "id": 1,
                                "type": "message",
                                "date_unixtime": str(timestamp),
                                "from": "Alice",
                                "from_id": "user:alice",
                                "text": "早上好",
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
        work = root / "run"
        reconcile.start_run(
            reconcile.parse_args(
                [
                    "start",
                    str(source_root),
                    "--work",
                    str(work),
                    "--mode",
                    "large",
                    "--date",
                    "2026-08-31",
                    "--no-ocr-candidates",
                ]
            )
        )
        return work, media_decisions, orders

    def _apply_and_seal_large_group(
        self,
        root: Path,
        work: Path,
        *,
        group_name: str,
        media_decisions: dict[str, dict],
        orders: list[dict],
        batch_id: str,
    ) -> tuple[dict, dict]:
        page = reconcile.review_command(
            Namespace(work=work, review_action="next", group=group_name, limit=500)
        )
        batch_path = root / f"{batch_id}.json"
        core.atomic_json(
            batch_path,
            {
                "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
                "batch_id": batch_id,
                "base_fingerprint": page["semantic_fingerprint"],
                "page_commit": {
                    "page_start": page["page_start"],
                    "page_end": page["page_end"],
                    "page_token": page["page_token"],
                },
                "open_orders": [],
                "media_decisions": media_decisions,
                "orders": orders,
            },
        )
        applied = reconcile.review_command(
            Namespace(
                work=work,
                review_action="apply-batch",
                group=group_name,
                input=batch_path,
            )
        )
        sealed = reconcile.review_command(
            Namespace(work=work, review_action="seal", group=group_name)
        )
        return applied, sealed

    def test_start_large_mode_selects_every_group_except_small_and_finance(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source"
            timestamp = str(
                int(datetime.fromisoformat("2026-08-31T08:00:00+07:00").timestamp())
            )
            groups = [
                ("large-a", "曼谷固定换汇一群"),
                ("small", "QQ小额出🐱群"),
                ("finance", "财务资料群"),
                ("large-b", "长期合作换汇群"),
            ]
            for group_id, group_name in groups:
                export = source / group_id
                export.mkdir(parents=True)
                (export / "result.json").write_text(
                    json.dumps(
                        {
                            "id": group_id,
                            "name": group_name,
                            "messages": [
                                {
                                    "id": 1,
                                    "type": "message",
                                    "date_unixtime": timestamp,
                                    "from": "Alice",
                                    "from_id": "user:alice",
                                    "text": "今日换汇记录",
                                }
                            ],
                        },
                        ensure_ascii=False,
                    ),
                    encoding="utf-8",
                )

            work = root / "run"
            args = reconcile.parse_args(
                [
                    "start",
                    str(source),
                    "--work",
                    str(work),
                    "--mode",
                    "large",
                    "--date",
                    "2026-08-31",
                    "--from",
                    "2026-08-31 05:00",
                    "--to",
                    "2026-09-01 05:00",
                ]
            )
            report = reconcile.start_run(args)
            run = reconcile._load_run(work)
            normalized = reconcile._load_snapshot(work, run)

            expected_names = ["曼谷固定换汇一群", "长期合作换汇群"]
            self.assertEqual(
                sorted(item["group_name"] for item in report["selected_groups"]),
                sorted(expected_names),
            )
            self.assertEqual(
                sorted(group["group_name"] for group in normalized["groups"]),
                sorted(expected_names),
            )
            self.assertEqual(run["group_mode"], "large")
            self.assertEqual(run["accounting_date"], "2026-08-31")
            self.assertEqual(run["accounting_from"], "2026-08-31 05:00")
            self.assertEqual(run["accounting_to"], "2026-09-01 05:00")
            self.assertEqual(report["runtime_contract"], reconcile.RUNTIME_CONTRACT)
            self.assertEqual(report["amount_policy"], reconcile.AMOUNT_POLICY)
            self.assertEqual(
                report["workbook_compatibility"],
                reconcile.WORKBOOK_COMPATIBILITY,
            )

    def test_start_large_mode_requires_one_accounting_date(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source"
            source.mkdir()
            work = root / "run"
            args = reconcile.parse_args(
                [
                    "start",
                    str(source),
                    "--work",
                    str(work),
                    "--mode",
                    "large",
                ]
            )
            with self.assertRaisesRegex(ValueError, "large.*--date"):
                reconcile.start_run(args)
            self.assertFalse(work.exists())

    def test_large_group_review_uses_full_order_contract(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            specs = [
                (
                    "微信付款10000人民币，按4.72换47200泰铢",
                    "wechat",
                    "CNY->THB",
                    "10000",
                    "4.72",
                    "multiply",
                    "47200",
                )
            ]
            work, media_decisions, orders = self._start_large_order_fixture(
                root, specs
            )
            run = reconcile._load_run(work)
            decision = reconcile._load_json(
                reconcile._decision_path(work, run["groups"][0])
            )
            self.assertEqual(
                decision["contract_version"],
                reconcile.large_daily.DECISION_CONTRACT,
            )
            self.assertEqual(decision["orders"], [])
            self.assertEqual(decision["open_orders"], [])
            self.assertNotIn("exchanges", decision)

            page = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="曼谷固定换汇群",
                    limit=500,
                )
            )
            self.assertEqual(page["open_orders"], [])
            self.assertNotIn("open_exchanges", page)
            missing_type_orders = copy.deepcopy(orders)
            missing_type_orders[0].pop("fund_type")
            invalid_batch_path = root / "large-missing-fund-type.json"
            core.atomic_json(
                invalid_batch_path,
                {
                    "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
                    "batch_id": "large-missing-fund-type",
                    "base_fingerprint": page["semantic_fingerprint"],
                    "page_commit": {
                        "page_start": page["page_start"],
                        "page_end": page["page_end"],
                        "page_token": page["page_token"],
                    },
                    "open_orders": [],
                    "media_decisions": media_decisions,
                    "orders": missing_type_orders,
                },
            )
            with self.assertRaisesRegex(ValueError, "fund_type"):
                reconcile.review_command(
                    Namespace(
                        work=work,
                        review_action="apply-batch",
                        group="曼谷固定换汇群",
                        input=invalid_batch_path,
                    )
                )
            applied, sealed = self._apply_and_seal_large_group(
                root,
                work,
                group_name="曼谷固定换汇群",
                media_decisions=media_decisions,
                orders=orders,
                batch_id="large-orders-001",
            )
            self.assertEqual(applied["orders"], 1)
            self.assertEqual(applied["open_orders"], 0)
            self.assertEqual(applied["fund_entries"], 2)
            self.assertTrue(sealed["sealed"])
            self.assertEqual(sealed["orders"], 1)

    def test_large_group_multiple_fund_images_compile_to_one_full_order(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            specs = [
                (
                    "微信付款10000人民币，按4.72换47200泰铢",
                    "wechat",
                    "CNY->THB",
                    "10000",
                    "4.72",
                    "multiply",
                    "47200",
                )
            ]
            work, media_decisions, orders = self._start_large_order_fixture(
                root,
                specs,
                group_name="固定合作换汇群",
            )
            self._apply_and_seal_large_group(
                root,
                work,
                group_name="固定合作换汇群",
                media_decisions=media_decisions,
                orders=orders,
                batch_id="large-detail-001",
            )
            output = root / "single-full-order.xlsx"
            result = reconcile.finish_run(
                Namespace(
                    work=work,
                    output=output,
                    template=Path(__file__).resolve().parent.parent
                    / "assets"
                    / "模版.xlsx",
                )
            )
            self.assertEqual(result["orders"], 1)
            self.assertEqual(result["summary_rows"], 1)
            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                worksheet = workbook.worksheets[0]
                headers = [
                    worksheet.cell(1, column).value
                    for column in range(1, len(core.HEADERS) + 1)
                ]
                self.assertEqual(headers, core.HEADERS)
                row_types = [
                    row[0]
                    for row in worksheet.iter_rows(min_row=2, values_only=True)
                ]
                self.assertEqual(row_types[:3], ["订单汇总", "客户付款", "内部回款"])
                payee_column = core.HEADERS.index("收款方")
                detail_rows = list(
                    worksheet.iter_rows(min_row=2, max_row=4, values_only=True)
                )
                self.assertEqual(detail_rows[1][payee_column], "customer-payee-1")
                self.assertEqual(detail_rows[2][payee_column], "200-0-XXX001")
            finally:
                workbook.close()

    def test_legacy_large_exchange_contract_still_finishes(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            specs = [
                (
                    "微信付款10000人民币，按4.72换47200泰铢",
                    "wechat",
                    "CNY->THB",
                    "10000",
                    "4.72",
                    "multiply",
                    "47200",
                )
            ]
            work, media_decisions, _ = self._start_large_order_fixture(root, specs)
            run = reconcile._load_run(work)
            run_group = run["groups"][0]
            decision_path = reconcile._decision_path(work, run_group)
            decision = reconcile._load_json(decision_path)
            decision["contract_version"] = (
                reconcile.large_daily.LEGACY_DECISION_CONTRACT
            )
            for field in (
                "orders",
                "open_orders",
                "balance_links",
                "settlement_allocations",
            ):
                decision.pop(field, None)
            decision["exchanges"] = []
            decision["open_exchanges"] = []
            decision["edit_control"] = reconcile._new_edit_control(
                reconcile._semantic_fingerprint(decision)
            )
            core.atomic_json(decision_path, decision)

            page = reconcile.review_command(
                Namespace(work=work, review_action="next", group=None, limit=500)
            )
            batch_path = root / "legacy-large.json"
            core.atomic_json(
                batch_path,
                {
                    "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
                    "batch_id": "legacy-large-001",
                    "base_fingerprint": page["semantic_fingerprint"],
                    "page_commit": {
                        "page_start": page["page_start"],
                        "page_end": page["page_end"],
                        "page_token": page["page_token"],
                    },
                    "open_exchanges": [],
                    "media_decisions": media_decisions,
                    "exchanges": [
                        {
                            "id": "L001",
                            "source_messages": ["S00001", "S00002"],
                            "entry_ids": ["M0001.1", "M0002.1"],
                            "fund_type": "wechat",
                            "direction": "CNY->THB",
                            "source_amount": "10000",
                            "rate": "4.72",
                            "operator": "multiply",
                            "target_amount": "47200",
                        }
                    ],
                },
            )
            reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="曼谷固定换汇群",
                    input=batch_path,
                )
            )
            reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="seal",
                    group="曼谷固定换汇群",
                )
            )
            output = root / "legacy-large.xlsx"
            result = reconcile.finish_run(
                Namespace(
                    work=work,
                    output=output,
                    template=Path(__file__).resolve().parent.parent
                    / "assets"
                    / "模版.xlsx",
                )
            )
            self.assertEqual(result["exchanges"], 1)
            workbook = load_workbook(output, read_only=True, data_only=False)
            try:
                headers = [
                    workbook.worksheets[0].cell(1, column).value
                    for column in range(
                        1, len(reconcile.large_daily.LEGACY_HEADERS) + 1
                    )
                ]
                self.assertEqual(headers, reconcile.large_daily.LEGACY_HEADERS)
            finally:
                workbook.close()

    def test_large_summary_includes_pending_actual_flow_with_status(self) -> None:
        group_key = "line:pending-rent"
        summaries = reconcile.large_daily.compile_order_summaries(
            [
                {
                    "fund_type": "alipay",
                    "direction": "THB->CNY",
                    "payment_total": None,
                    "actual_payout_total": "91469",
                    "actual_rate": "4.97",
                    "rate_operator": "divide",
                    "reconciliation": {
                        "status": "pending",
                        "reason": "evidence_incomplete",
                    },
                }
            ],
            group_key=group_key,
        )

        self.assertEqual(len(summaries), 1)
        self.assertEqual(summaries[0]["fund_type_label"], "支付宝")
        self.assertEqual(summaries[0]["direction"], "THB->CNY")
        self.assertEqual(summaries[0]["rate_display"], "÷4.97")
        self.assertIsNone(summaries[0]["source_total"])
        self.assertEqual(summaries[0]["target_total"], "91469")
        self.assertEqual(summaries[0]["status_label"], "待确认")

    def test_large_group_finish_writes_daily_records_and_rate_grouped_totals(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            exchange_specs = [
                ("微信10000人民币按4.72换47200泰铢", "wechat", "CNY->THB", "10000", "4.72", "multiply", "47200"),
                ("微信5000人民币按4.72换23600泰铢", "wechat", "CNY->THB", "5000", "4.72", "multiply", "23600"),
                ("微信1000人民币按4.70换4700泰铢", "wechat", "CNY->THB", "1000", "4.70", "multiply", "4700"),
                ("支付宝2000人民币按4.70换9400泰铢", "alipay", "CNY->THB", "2000", "4.70", "multiply", "9400"),
                ("银行卡47000泰铢按4.70反向换10000人民币", "bank_card", "THB->CNY", "47000", "4.70", "divide", "10000"),
                ("32600泰铢按32.60换1000USDT", "usdt", "THB->USDT", "32600", "32.60", "divide", "1000"),
                ("银行卡付款100人民币，群聊未说明最终汇率", "bank_card", "CNY->THB", "100", "1", "multiply", "500"),
            ]
            work, media_decisions, orders = self._start_large_order_fixture(
                root,
                exchange_specs,
                include_empty_group=True,
            )
            orders[-1]["pricing"] = {
                "source_messages": ["S00013"],
                "expected": {"kind": "unknown", "reason": "not_stated"},
            }
            self._apply_and_seal_large_group(
                root,
                work,
                group_name="曼谷固定换汇群",
                media_decisions=media_decisions,
                orders=orders,
                batch_id="large-summary-001",
            )

            empty_page = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="当天无换汇群",
                    limit=500,
                )
            )
            empty_batch_path = root / "large-empty.json"
            core.atomic_json(
                empty_batch_path,
                {
                    "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
                    "batch_id": "large-empty-001",
                    "base_fingerprint": empty_page["semantic_fingerprint"],
                    "page_commit": {
                        "page_start": empty_page["page_start"],
                        "page_end": empty_page["page_end"],
                        "page_token": empty_page["page_token"],
                    },
                    "open_orders": [],
                },
            )
            reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="当天无换汇群",
                    input=empty_batch_path,
                )
            )
            reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="seal",
                    group="当天无换汇群",
                )
            )

            output = root / "大额群日换汇.xlsx"
            result = reconcile.finish_run(
                Namespace(
                    work=work,
                    output=output,
                    template=Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
                )
            )
            self.assertEqual(result["orders"], 7)
            self.assertEqual(result["pending_orders"], 1)
            self.assertEqual(result["summary_rows"], 6)
            self.assertEqual(result["groups"], 1)
            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                self.assertEqual(len(workbook.worksheets), 1)
                self.assertNotIn("期间说明", workbook.sheetnames)
                worksheet = workbook.worksheets[0]
                headers = [
                    worksheet.cell(1, column).value
                    for column in range(1, len(core.HEADERS) + 1)
                ]
                self.assertEqual(headers, core.HEADERS)
                summary_rows = [
                    tuple(row[: len(reconcile.large_daily.SUMMARY_HEADERS)])
                    for row in worksheet.iter_rows(values_only=True)
                    if row[0] in {"微信", "支付宝", "银行卡", "USDT"}
                ]
                self.assertEqual(
                    summary_rows,
                    [
                        ("微信", "CNY->THB", "乘", "×4.72", 2, 15000, 70800, "已确认"),
                        ("微信", "CNY->THB", "乘", "×4.7", 1, 1000, 4700, "已确认"),
                        ("支付宝", "CNY->THB", "乘", "×4.7", 1, 2000, 9400, "已确认"),
                        ("银行卡", "THB->CNY", "除", "÷4.7", 1, 47000, 10000, "已确认"),
                        ("银行卡", "CNY->THB", "乘", "待确认", 1, 100, 500, "待确认"),
                        ("USDT", "THB->USDT", "除", "÷32.6", 1, 32600, 1000, "已确认"),
                    ],
                )
                summary_header_row = next(
                    row
                    for row in range(1, worksheet.max_row + 1)
                    if worksheet.cell(row, 1).value == "资金类型"
                )
                summary_title_row = summary_header_row - 1
                self.assertEqual(
                    worksheet.cell(summary_title_row, 1).value,
                    "资金汇总",
                )
                self.assertIn(
                    f"A{summary_title_row}:H{summary_title_row}",
                    {str(item) for item in worksheet.merged_cells.ranges},
                )
                self.assertTrue(
                    str(worksheet.cell(summary_title_row, 1).fill.fgColor.rgb)
                    .upper()
                    .endswith(build_workbook.LARGE_SUMMARY_TITLE_FILL_RGB)
                )
                for column in range(1, len(reconcile.large_daily.SUMMARY_HEADERS) + 1):
                    header_cell = worksheet.cell(summary_header_row, column)
                    self.assertTrue(header_cell.font.bold)
                    self.assertTrue(
                        str(header_cell.font.color.rgb)
                        .upper()
                        .endswith(build_workbook.LARGE_SUMMARY_WHITE_FONT_RGB)
                    )
                    self.assertTrue(
                        str(header_cell.fill.fgColor.rgb)
                        .upper()
                        .endswith(build_workbook.LARGE_SUMMARY_HEADER_FILL_RGB)
                    )
                for offset in range(1, len(summary_rows) + 1):
                    expected_fill = (
                        build_workbook.LARGE_SUMMARY_BODY_FILL_RGB
                        if offset % 2 == 1
                        else build_workbook.LARGE_SUMMARY_BODY_ALT_FILL_RGB
                    )
                    for column in range(1, len(reconcile.large_daily.SUMMARY_HEADERS) + 1):
                        self.assertTrue(
                            str(
                                worksheet.cell(
                                    summary_header_row + offset,
                                    column,
                                ).fill.fgColor.rgb
                            )
                            .upper()
                            .endswith(expected_fill)
                        )
                    for column in (5, 6, 7):
                        self.assertTrue(
                            worksheet.cell(summary_header_row + offset, column).font.bold
                        )
                for column_letter, minimum_width in {
                    "A": 18,
                    "B": 23,
                    "C": 16,
                    "D": 13,
                    "E": 11,
                    "F": 18,
                    "G": 18,
                    "H": 14,
                }.items():
                    self.assertGreaterEqual(
                        worksheet.column_dimensions[column_letter].width,
                        minimum_width,
                    )
                self.assertFalse(worksheet.sheet_view.showGridLines)
            finally:
                workbook.close()

    def test_start_date_filters_messages_by_bangkok_calendar_day(self) -> None:
        def unix_time(value: str) -> str:
            return str(int(datetime.fromisoformat(value).timestamp()))

        messages = [
            {
                "id": 1,
                "type": "message",
                "date_unixtime": unix_time("2026-08-30T16:59:59+00:00"),
                "from": "Alice",
                "from_id": "user:alice",
                "text": "before",
            },
            {
                "id": 2,
                "type": "message",
                "date_unixtime": unix_time("2026-08-30T17:00:00+00:00"),
                "from": "Alice",
                "from_id": "user:alice",
                "text": "start",
            },
            {
                "id": 3,
                "type": "message",
                "date_unixtime": unix_time("2026-08-31T16:59:59+00:00"),
                "from": "Alice",
                "from_id": "user:alice",
                "text": "end",
            },
            {
                "id": 4,
                "type": "message",
                "date_unixtime": unix_time("2026-08-31T17:00:00+00:00"),
                "from": "Alice",
                "from_id": "user:alice",
                "text": "after",
            },
        ]
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            parsed = reconcile.parse_args(
                [
                    "start",
                    str(root / "source"),
                    "--work",
                    str(root / "parsed-run"),
                    "--date",
                    "2026-08-31",
                ]
            )
            self.assertEqual(parsed.date, "2026-08-31")
            _, run, normalized = self._start_fixture(
                root,
                messages=messages,
                controlled=True,
                commit_page=False,
                accounting_date="2026-08-31",
            )
            self.assertEqual(run["accounting_date"], "2026-08-31")
            self.assertEqual(
                [message["text"] for message in normalized["groups"][0]["messages"]],
                ["start", "end"],
            )
            self.assertEqual(normalized["statistics"]["messages"], 2)

    def test_start_rejects_invalid_date_before_creating_work_directory(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source"
            source.mkdir()
            work = root / "run"
            with self.assertRaisesRegex(ValueError, "YYYY-MM-DD"):
                reconcile.start_run(
                    Namespace(
                        inputs=[source],
                        work=work,
                        contains="小额",
                        date="2026-8-31",
                        timezone="Asia/Bangkok",
                        roster=Path(__file__).resolve().parent.parent
                        / "config"
                        / "roster.yaml",
                        line_backup=[],
                        line_android_backup=[],
                        line_self_name="LINE_SELF",
                    )
                )
            self.assertFalse(work.exists())

    def test_start_time_range_filters_start_inclusive_end_exclusive(self) -> None:
        def unix_time(value: str) -> str:
            return str(int(datetime.fromisoformat(value).timestamp()))

        messages = [
            {
                "id": 1,
                "type": "message",
                "date_unixtime": unix_time("2026-08-30T21:59:59+00:00"),
                "from": "Alice",
                "from_id": "user:alice",
                "text": "before",
            },
            {
                "id": 2,
                "type": "message",
                "date_unixtime": unix_time("2026-08-30T22:00:00+00:00"),
                "from": "Alice",
                "from_id": "user:alice",
                "text": "start",
            },
            {
                "id": 3,
                "type": "message",
                "date_unixtime": unix_time("2026-08-31T21:59:59+00:00"),
                "from": "Alice",
                "from_id": "user:alice",
                "text": "end",
            },
            {
                "id": 4,
                "type": "message",
                "date_unixtime": unix_time("2026-08-31T22:00:00+00:00"),
                "from": "Alice",
                "from_id": "user:alice",
                "text": "after",
            },
        ]
        accounting_from = "2026-08-31 05:00"
        accounting_to = "2026-09-01 05:00"
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            parsed = reconcile.parse_args(
                [
                    "start",
                    str(root / "source"),
                    "--work",
                    str(root / "parsed-run"),
                    "--from",
                    accounting_from,
                    "--to",
                    accounting_to,
                ]
            )
            self.assertEqual(parsed.accounting_from, accounting_from)
            self.assertEqual(parsed.accounting_to, accounting_to)
            _, run, normalized = self._start_fixture(
                root,
                messages=messages,
                controlled=True,
                commit_page=False,
                accounting_from=accounting_from,
                accounting_to=accounting_to,
            )
            self.assertEqual(run["accounting_from"], accounting_from)
            self.assertEqual(run["accounting_to"], accounting_to)
            self.assertEqual(
                [message["text"] for message in normalized["groups"][0]["messages"]],
                ["start", "end"],
            )
            self.assertEqual(normalized["statistics"]["messages"], 2)

    def test_start_rejects_invalid_time_window_before_creating_work_directory(self) -> None:
        cases = [
            (
                "missing end",
                None,
                "2026-08-31 05:00",
                None,
                "provided together",
            ),
            (
                "combined with date",
                "2026-08-31",
                "2026-08-31 05:00",
                "2026-09-01 05:00",
                "cannot be combined",
            ),
            (
                "reversed",
                None,
                "2026-09-01 05:00",
                "2026-08-31 05:00",
                "earlier than",
            ),
            (
                "invalid format",
                None,
                "2026-08-31T05:00",
                "2026-09-01 05:00",
                "YYYY-MM-DD HH:MM",
            ),
        ]
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source"
            source.mkdir()
            for index, (name, date, accounting_from, accounting_to, error) in enumerate(cases):
                work = root / f"run-{index}"
                with self.subTest(name=name), self.assertRaisesRegex(ValueError, error):
                    reconcile.start_run(
                        Namespace(
                            inputs=[source],
                            work=work,
                            contains="小额",
                            date=date,
                            accounting_from=accounting_from,
                            accounting_to=accounting_to,
                            timezone="Asia/Bangkok",
                            roster=Path(__file__).resolve().parent.parent
                            / "config"
                            / "roster.yaml",
                            line_backup=[],
                            line_android_backup=[],
                            line_self_name="LINE_SELF",
                        )
                    )
                self.assertFalse(work.exists())

    def _write_review_batch(
        self,
        root: Path,
        *,
        batch_id: str,
        base_fingerprint: str,
        media_decisions: dict | None = None,
        orders: list[dict] | None = None,
        **replacements: object,
    ) -> Path:
        payload: dict[str, object] = {
            "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
            "batch_id": batch_id,
            "base_fingerprint": base_fingerprint,
        }
        if media_decisions is not None:
            payload["media_decisions"] = media_decisions
        if orders is not None:
            payload["orders"] = orders
        payload.update(replacements)
        path = root / f"{batch_id}.json"
        core.atomic_json(path, payload)
        return path

    @staticmethod
    def _complete_batch_values(*, payment_side: str = "payment") -> tuple[dict, list[dict]]:
        return (
            {
                "M0001": {
                    "classification": "fund",
                    "viewed_original": True,
                    "entries": [
                        {
                            "amount": "100",
                            "currency": "CNY",
                            "payee": "张三",
                            "payee_state": "visible",
                            "side": payment_side,
                            "result": "completed",
                            "amount_state": "clear",
                        }
                    ],
                },
                "M0002": {
                    "classification": "fund",
                    "viewed_original": True,
                    "entries": [
                        {
                            "amount": "500",
                            "currency": "THB",
                            "payee": "206-4-xxx781",
                            "payee_state": "visible",
                            "side": "payout",
                            "result": "completed",
                            "amount_state": "clear",
                        }
                    ],
                },
            },
            [
                {
                    "id": "O001",
                    "entry_ids": ["M0001.1", "M0002.1"],
                    "source_messages": ["S00001", "S00002"],
                    "customer_nickname": "Alice",
                    "direction": "CNY->THB",
                    "pricing": {
                        "source_messages": ["S00001", "S00002"],
                        "terms": {"rate": "5", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "500"},
                    },
                }
            ],
        )

    def _write_complete_decision(
        self,
        work: Path,
        *,
        payee: str = "张三",
        through_batch: bool = False,
    ) -> Path:
        run = reconcile._load_run(work)
        run_group = run["groups"][0]
        path = reconcile._decision_path(work, run_group)
        decision = reconcile._load_json(path)
        decision["media_decisions"] = {
            "M0001": {
                "classification": "fund",
                "viewed_original": True,
                "entries": [
                    {
                        "amount": "100",
                        "currency": "CNY",
                        "payee": payee,
                        "payee_state": "visible",
                        "side": "payment",
                        "result": "completed",
                        "amount_state": "clear",
                    }
                ],
            },
            "M0002": {
                "classification": "fund",
                "viewed_original": True,
                "entries": [
                    {
                        "amount": "500",
                        "currency": "THB",
                        "payee": "206-4-xxx781",
                        "payee_state": "visible",
                        "side": "payout",
                        "result": "completed",
                        "amount_state": "clear",
                    }
                ],
            },
        }
        decision["orders"] = [
            {
                "id": "O001",
                "entry_ids": ["M0001.1", "M0002.1"],
                "source_messages": ["S00001", "S00002"],
                "customer_nickname": "Alice",
                "direction": "CNY->THB",
                "pricing": {
                    "source_messages": ["S00001", "S00002"],
                    "terms": {"rate": "5", "operator": "multiply"},
                    "expected": {"kind": "explicit", "amount": "500"},
                },
            }
        ]
        if through_batch:
            batch_path = work.parent / "complete-decision-batch.json"
            core.atomic_json(
                batch_path,
                {
                    "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
                    "batch_id": "complete-decision",
                    "base_fingerprint": reconcile._semantic_fingerprint(
                        reconcile._load_json(path)
                    ),
                    "media_decisions": decision["media_decisions"],
                    "orders": decision["orders"],
                },
            )
            reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=batch_path,
                )
            )
        else:
            core.atomic_json(path, decision)
        return path

    def _start_order_continuity_fixture(
        self,
        root: Path,
    ) -> tuple[Path, dict[str, dict]]:
        messages = [
            {
                "id": 1,
                "type": "message",
                "date_unixtime": "1780000000",
                "from": "QQ财务3",
                "from_id": "user:staff",
                "text": "按32.6，1000/32.6=31",
            },
            {
                "id": 2,
                "type": "message",
                "date_unixtime": "1780000060",
                "from": "Alice",
                "from_id": "user:alice",
                "text": "31U付款",
                "photo": "photos/customer.jpg",
            },
            {
                "id": 3,
                "type": "message",
                "date_unixtime": "1780000120",
                "from": "Alice",
                "from_id": "user:alice",
                "text": "我再多转16U，你们帮我付1500泰铢可以吗",
            },
            {
                "id": 4,
                "type": "message",
                "date_unixtime": "1780000180",
                "from": "QQ财务3",
                "from_id": "user:staff",
                "text": "可以",
            },
            {
                "id": 5,
                "type": "message",
                "date_unixtime": "1780000240",
                "from": "Alice",
                "from_id": "user:alice",
                "text": "16U付款",
                "photo": "photos/customer.jpg",
            },
            {
                "id": 6,
                "type": "message",
                "date_unixtime": "1780000300",
                "from": "QQ财务3",
                "from_id": "user:staff",
                "text": "1500泰铢已回",
                "photo": "photos/staff.jpg",
            },
        ]
        work, _, _ = self._start_fixture(
            root,
            messages=messages,
            controlled=True,
        )
        media_decisions = {
            "M0001": {
                "classification": "fund",
                "viewed_original": True,
                "entries": [
                    {
                        "amount": "31",
                        "currency": "USDT",
                        "payee": "TContinuationWallet",
                    }
                ],
            },
            "M0002": {
                "classification": "fund",
                "viewed_original": True,
                "entries": [
                    {
                        "amount": "16",
                        "currency": "USDT",
                        "payee": "TContinuationWallet",
                    }
                ],
            },
            "M0003": {
                "classification": "fund",
                "viewed_original": True,
                "entries": [
                    {
                        "amount": "1500",
                        "currency": "THB",
                        "payee": "79xxxx3249",
                    }
                ],
            },
        }
        return work, media_decisions

    def _split_continuation_orders(self) -> list[dict]:
        return [
            {
                "id": "O011",
                "entry_ids": ["M0001.1"],
                "source_messages": ["S00001", "S00002"],
                "customer_nickname": "Alice",
                "direction": "USDT->THB",
                "pricing": {
                    "source_messages": ["S00001"],
                    "terms": {"rate": "32.6", "operator": "multiply"},
                    "expected": {"kind": "explicit", "amount": "1000"},
                },
            },
            {
                "id": "O012",
                "entry_ids": ["M0002.1", "M0003.1"],
                "source_messages": ["S00003", "S00004", "S00005", "S00006"],
                "customer_nickname": "Alice",
                "direction": "USDT->THB",
                "pricing": {
                    "source_messages": ["S00003", "S00004"],
                    "expected": {"kind": "unknown", "reason": "not_stated"},
                },
            },
        ]

    def _apply_controlled_semantics(
        self,
        root: Path,
        work: Path,
        *,
        batch_id: str,
        media_decisions: dict[str, dict],
        orders: list[dict],
        **extra_fields: object,
    ) -> dict:
        run = reconcile._load_run(work)
        decision_path = reconcile._decision_path(work, run["groups"][0])
        batch = {
            "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
            "batch_id": batch_id,
            "base_fingerprint": reconcile._semantic_fingerprint(
                reconcile._load_json(decision_path)
            ),
            "media_decisions": media_decisions,
            "orders": orders,
            **extra_fields,
        }
        batch_path = root / f"{batch_id}.json"
        core.atomic_json(batch_path, batch)
        return reconcile.review_command(
            Namespace(
                work=work,
                review_action="apply-batch",
                group="测试小额群",
                input=batch_path,
            )
        )

    def test_order_continuity_risk_flags_split_without_mutating_or_blocking(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, media_decisions = self._start_order_continuity_fixture(root)
            report = self._apply_controlled_semantics(
                root,
                work,
                batch_id="continuity-split",
                media_decisions=media_decisions,
                orders=self._split_continuation_orders(),
            )

            self.assertEqual(report["order_continuity_review_candidate_count"], 1)
            self.assertEqual(
                report["order_continuity_review_candidates"],
                [
                    {
                        "earlier_order_id": "O011",
                        "later_order_id": "O012",
                        "customer_nickname": "Alice",
                        "direction": "USDT->THB",
                        "source_start": "S00001",
                        "source_end": "S00006",
                        "reason_codes": [
                            "earlier_payment_without_payout",
                            "earlier_confirmed_pricing",
                            "later_payment_and_payout",
                            "later_unknown_pricing",
                        ],
                    }
                ],
            )

            run = reconcile._load_run(work)
            decision_path = reconcile._decision_path(work, run["groups"][0])
            before_audit = decision_path.read_bytes()
            audit = reconcile.review_command(
                Namespace(work=work, review_action="audit", group="测试小额群")
            )
            self.assertEqual(before_audit, decision_path.read_bytes())
            self.assertEqual(
                audit["groups"][0]["order_continuity_review_candidate_count"],
                1,
            )

            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            self.assertTrue(sealed["sealed"])
            self.assertEqual(sealed["order_continuity_review_candidate_count"], 1)

    def test_order_continuity_risk_ignores_correctly_merged_order(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, media_decisions = self._start_order_continuity_fixture(root)
            report = self._apply_controlled_semantics(
                root,
                work,
                batch_id="continuity-merged",
                media_decisions=media_decisions,
                orders=[
                    {
                        "id": "O011",
                        "entry_ids": ["M0001.1", "M0002.1", "M0003.1"],
                        "source_messages": [
                            "S00001",
                            "S00002",
                            "S00003",
                            "S00004",
                            "S00005",
                            "S00006",
                        ],
                        "customer_nickname": "Alice",
                        "direction": "USDT->THB",
                        "pricing": {
                            "source_messages": ["S00001", "S00003", "S00004"],
                            "terms": {"rate": "32.6", "operator": "multiply"},
                            "expected": {"kind": "explicit", "amount": "1500"},
                        },
                    }
                ],
            )

            self.assertEqual(report["order_continuity_review_candidate_count"], 0)
            self.assertEqual(report["order_continuity_review_candidates"], [])

    def test_order_continuity_risk_ignores_two_complete_independent_orders(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            messages = [
                {
                    "id": 1,
                    "type": "message",
                    "date_unixtime": "1780000000",
                    "from": "Alice",
                    "from_id": "user:alice",
                    "text": "第一单31U，按32.5回1000泰铢",
                    "photo": "photos/customer.jpg",
                },
                {
                    "id": 2,
                    "type": "message",
                    "date_unixtime": "1780000060",
                    "from": "QQ财务3",
                    "from_id": "user:staff",
                    "text": "第一单已回1000泰铢",
                    "photo": "photos/staff.jpg",
                },
                {
                    "id": 3,
                    "type": "message",
                    "date_unixtime": "1780000120",
                    "from": "Alice",
                    "from_id": "user:alice",
                    "text": "新开第二单46U，按32.6回1500泰铢",
                    "photo": "photos/customer.jpg",
                },
                {
                    "id": 4,
                    "type": "message",
                    "date_unixtime": "1780000180",
                    "from": "QQ财务3",
                    "from_id": "user:staff",
                    "text": "第二单已回1500泰铢",
                    "photo": "photos/staff.jpg",
                },
            ]
            work, _, _ = self._start_fixture(
                root,
                messages=messages,
                controlled=True,
            )
            media_decisions = {
                "M0001": {
                    "classification": "fund",
                    "viewed_original": True,
                    "entries": [
                        {"amount": "31", "currency": "USDT", "payee": "TWallet1"}
                    ],
                },
                "M0002": {
                    "classification": "fund",
                    "viewed_original": True,
                    "entries": [
                        {"amount": "1000", "currency": "THB", "payee": "11xxxx1111"}
                    ],
                },
                "M0003": {
                    "classification": "fund",
                    "viewed_original": True,
                    "entries": [
                        {"amount": "46", "currency": "USDT", "payee": "TWallet2"}
                    ],
                },
                "M0004": {
                    "classification": "fund",
                    "viewed_original": True,
                    "entries": [
                        {"amount": "1500", "currency": "THB", "payee": "22xxxx2222"}
                    ],
                },
            }
            orders = [
                {
                    "id": "O001",
                    "entry_ids": ["M0001.1", "M0002.1"],
                    "source_messages": ["S00001", "S00002"],
                    "customer_nickname": "Alice",
                    "direction": "USDT->THB",
                    "pricing": {
                        "source_messages": ["S00001"],
                        "terms": {"rate": "32.5", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "1000"},
                    },
                },
                {
                    "id": "O002",
                    "entry_ids": ["M0003.1", "M0004.1"],
                    "source_messages": ["S00003", "S00004"],
                    "customer_nickname": "Alice",
                    "direction": "USDT->THB",
                    "pricing": {
                        "source_messages": ["S00003"],
                        "terms": {"rate": "32.6", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "1500"},
                    },
                },
            ]
            report = self._apply_controlled_semantics(
                root,
                work,
                batch_id="continuity-independent",
                media_decisions=media_decisions,
                orders=orders,
            )

            self.assertEqual(report["order_continuity_review_candidate_count"], 0)
            self.assertEqual(report["order_continuity_review_candidates"], [])

    def test_order_continuity_risk_skips_balance_linked_orders(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, media_decisions = self._start_order_continuity_fixture(root)
            report = self._apply_controlled_semantics(
                root,
                work,
                batch_id="continuity-balance-link",
                media_decisions=media_decisions,
                orders=self._split_continuation_orders(),
                balance_links=[
                    {
                        "source_order_id": "O011",
                        "target_order_id": "O012",
                        "kind": "shortfall_carryover",
                        "amount": "1",
                        "currency": "THB",
                        "source_messages": ["S00003"],
                        "already_in_expected": False,
                    }
                ],
            )

            self.assertEqual(report["order_continuity_review_candidate_count"], 0)
            self.assertEqual(report["order_continuity_review_candidates"], [])

    def test_review_next_is_read_only_until_page_commit(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, run, _ = self._start_fixture(
                root,
                controlled=True,
                commit_page=False,
            )
            decision_path = reconcile._decision_path(work, run["groups"][0])
            before = reconcile._load_json(decision_path)

            first = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="测试小额群",
                    limit=1,
                )
            )
            repeated = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="测试小额群",
                    limit=1,
                )
            )
            after = reconcile._load_json(decision_path)

            self.assertEqual(first["page_start"], 0)
            self.assertEqual(first["page_end"], 1)
            self.assertEqual(first["page_token"], repeated["page_token"])
            self.assertEqual(first["media_queue"], repeated["media_queue"])
            self.assertEqual(
                first["media_queue"]["contract_version"],
                reconcile.MEDIA_QUEUE_CONTRACT,
            )
            self.assertEqual(first["media_queue"]["mode"], "orders")
            self.assertEqual(
                first["media_queue"]["recommended_parallel_limit"],
                reconcile.ORDER_MEDIA_BATCH_LIMIT,
            )
            self.assertEqual(first["media_queue"]["queued_images"], 1)
            self.assertEqual(
                first["media_queue"]["batches"][0]["items"][0]["label"],
                "M0001",
            )
            self.assertEqual(
                first["media_queue"]["batches"][0]["items"][0]["message_label"],
                "S00001",
            )
            self.assertFalse(first["done"])
            self.assertEqual(after["reviewed_through"], 0)
            self.assertFalse(after["read_complete"])
            self.assertEqual(
                reconcile._semantic_fingerprint(before),
                reconcile._semantic_fingerprint(after),
            )

    def test_ocr_platform_defaults_and_manual_overrides_are_persisted(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source"
            source.mkdir()
            (source / "result.json").write_text(
                json.dumps(
                    {
                        "id": "ocr-platforms",
                        "name": "测试小额群",
                        "messages": [
                            {
                                "id": 1,
                                "type": "message",
                                "date_unixtime": "1780000000",
                                "from": "Alice",
                                "from_id": "user:alice",
                                "text": "仅测试平台默认值",
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            cases = [
                ("windows-auto", "Windows", [], False, "auto"),
                ("windows-enabled", "Windows", ["--ocr-candidates"], True, "enabled"),
                ("macos-auto", "Darwin", [], False, "auto"),
                ("macos-enabled", "Darwin", ["--ocr-candidates"], True, "enabled"),
                ("macos-disabled", "Darwin", ["--no-ocr-candidates"], False, "disabled"),
                ("linux-auto", "Linux", [], False, "auto"),
            ]
            for name, system_name, flags, expected_enabled, expected_requested in cases:
                with self.subTest(name=name):
                    work = root / name
                    args = reconcile.parse_args(
                        [
                            "start",
                            str(source),
                            "--work",
                            str(work),
                            *flags,
                        ]
                    )
                    with mock.patch.object(
                        reconcile.host_platform,
                        "system",
                        return_value=system_name,
                    ):
                        report = reconcile.start_run(args)
                    run = reconcile._load_run(work)
                    self.assertEqual(
                        report["ocr_candidates"],
                        run["ocr_candidates"],
                    )
                    self.assertEqual(
                        run["ocr_candidates"]["enabled"], expected_enabled
                    )
                    self.assertEqual(
                        run["ocr_candidates"]["requested"], expected_requested
                    )
                    self.assertFalse(run["ocr_candidates"]["default_enabled"])
                    self.assertFalse(reconcile._ocr_cache_path(work).exists())

    def test_disabled_ocr_does_not_invoke_worker_or_create_cache(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            with mock.patch.object(
                reconcile,
                "_invoke_ocr_worker",
                side_effect=AssertionError("disabled OCR must not invoke its worker"),
            ):
                work, _, _ = self._start_fixture(
                    root,
                    controlled=True,
                    commit_page=False,
                    ocr_candidates=False,
                )
                page = reconcile.review_command(
                    Namespace(
                        work=work,
                        review_action="next",
                        group="测试小额群",
                        limit=500,
                    )
                )
            self.assertEqual(page["media_queue"]["ocr_candidates"]["status"], "disabled")
            self.assertFalse(reconcile._ocr_cache_path(work).exists())
            self.assertTrue(
                all(
                    "ocr_candidate" not in item
                    for batch in page["media_queue"]["batches"]
                    for item in batch["items"]
                )
            )

    def test_enabled_ocr_candidates_are_cached_once_per_identical_hash(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            messages = [
                {
                    "id": index + 1,
                    "type": "message",
                    "date_unixtime": str(1780000000 + index * 60),
                    "from": "Alice",
                    "from_id": "user:alice",
                    "text": "重复发送同一张资金图",
                    "photo": "photos/customer.jpg",
                }
                for index in range(3)
            ]

            def fake_worker(requests: list[dict[str, str]]) -> dict:
                return {
                    "status": "ready",
                    "reason": None,
                    "backend": "fake-ocr",
                    "backend_version": "1.0",
                    "results": {
                        request["content_sha256"]: {
                            "contract_version": reconcile.OCR_CANDIDATE_CONTRACT,
                            "content_sha256": request["content_sha256"],
                            "status": "ok",
                            "backend": "fake-ocr",
                            "backend_version": "1.0",
                            "authoritative": False,
                            "text": "100,000 THB\nxxx-1234",
                            "average_confidence": 0.96,
                            "line_count": 2,
                            "elapsed_ms": 25,
                            "truncated": False,
                        }
                        for request in requests
                    },
                }

            with mock.patch.object(
                reconcile,
                "_invoke_ocr_worker",
                side_effect=fake_worker,
            ) as worker:
                work, run, _ = self._start_fixture(
                    root,
                    messages=messages,
                    controlled=True,
                    commit_page=False,
                    ocr_candidates=True,
                )
                first = reconcile.review_command(
                    Namespace(
                        work=work,
                        review_action="next",
                        group="测试小额群",
                        limit=500,
                    )
                )
                repeated = reconcile.review_command(
                    Namespace(
                        work=work,
                        review_action="next",
                        group="测试小额群",
                        limit=500,
                    )
                )

            self.assertEqual(worker.call_count, 1)
            self.assertEqual(len(worker.call_args.args[0]), 1)
            self.assertEqual(first["media_queue"], repeated["media_queue"])
            queue = first["media_queue"]
            self.assertEqual(queue["queued_images"], 1)
            self.assertEqual(queue["duplicate_alias_labels"], 2)
            self.assertEqual(queue["ocr_candidates"]["status"], "ready")
            self.assertEqual(queue["ocr_candidates"]["candidate_count"], 1)
            representative = queue["batches"][0]["items"][0]
            self.assertEqual(representative["same_content_labels"], ["M0002", "M0003"])
            self.assertEqual(
                representative["ocr_candidate"]["text"],
                "100,000 THB\nxxx-1234",
            )
            self.assertFalse(representative["ocr_candidate"]["authoritative"])
            self.assertTrue(reconcile._ocr_cache_path(work).is_file())
            decision = reconcile._load_json(
                reconcile._decision_path(work, run["groups"][0])
            )
            self.assertEqual(decision["media_decisions"], {})
            self.assertEqual(decision["reviewed_through"], 0)

    def test_adaptive_media_batch_policy_scales_queue_up_and_down(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            group = self._synthetic_media_group(
                Path(temporary),
                count=12,
            )

            def decision_with_metrics(**updates: int) -> dict:
                metrics = reconcile._empty_media_view_metrics()
                metrics.update(updates)
                return {
                    "contract_version": reconcile.DECISION_CONTRACT,
                    "media_decisions": {},
                    "media_observation_cache": reconcile._empty_media_observation_cache(),
                    "media_view_metrics": metrics,
                }

            default_queue = reconcile._review_media_queue(
                group,
                decision_with_metrics(),
                0,
                12,
            )
            fast_queue = reconcile._review_media_queue(
                group,
                decision_with_metrics(
                    view_batches=2,
                    opened_images=18,
                    elapsed_ms=20_000,
                ),
                0,
                12,
            )
            pressured_queue = reconcile._review_media_queue(
                group,
                decision_with_metrics(
                    view_batches=2,
                    opened_images=18,
                    failed_images=3,
                    elapsed_ms=20_000,
                ),
                0,
                12,
            )
            maximum_policy = reconcile._adaptive_media_batch_policy(
                decision_with_metrics(
                    view_batches=4,
                    opened_images=36,
                    elapsed_ms=32_000,
                ),
                finance_mode=False,
            )
            minimum_policy = reconcile._adaptive_media_batch_policy(
                decision_with_metrics(
                    view_batches=2,
                    opened_images=18,
                    failed_images=6,
                    elapsed_ms=20_000,
                ),
                finance_mode=False,
            )

            self.assertEqual(default_queue["recommended_parallel_limit"], 9)
            self.assertEqual(
                [len(batch["items"]) for batch in default_queue["batches"]],
                [9, 3],
            )
            self.assertEqual(fast_queue["recommended_parallel_limit"], 11)
            self.assertEqual(
                [len(batch["items"]) for batch in fast_queue["batches"]],
                [11, 1],
            )
            self.assertEqual(pressured_queue["recommended_parallel_limit"], 6)
            self.assertEqual(
                [len(batch["items"]) for batch in pressured_queue["batches"]],
                [6, 6],
            )
            self.assertEqual(maximum_policy["recommended_parallel_limit"], 12)
            self.assertEqual(minimum_policy["recommended_parallel_limit"], 4)
            finance_decision = decision_with_metrics(
                view_batches=2,
                opened_images=8,
                elapsed_ms=20_000,
            )
            finance_decision["contract_version"] = reconcile.finance_materials.DECISION_CONTRACT
            self.assertEqual(
                reconcile._adaptive_media_batch_policy(
                    finance_decision,
                    finance_mode=True,
                )["recommended_parallel_limit"],
                6,
            )

    def test_media_observation_cache_reuses_identical_images_without_merging_labels(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            messages = [
                {
                    "id": index + 1,
                    "type": "message",
                    "date_unixtime": str(1780000000 + index * 60),
                    "from": "Alice",
                    "from_id": "user:alice",
                    "text": "重复发送同一张付款图",
                    "photo": "photos/customer.jpg",
                }
                for index in range(3)
            ]
            work, run, _ = self._start_fixture(
                root,
                messages=messages,
                controlled=True,
                commit_page=False,
            )
            first_page = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="测试小额群",
                    limit=2,
                )
            )
            queue = first_page["media_queue"]
            self.assertEqual(queue["queued_images"], 1)
            self.assertEqual(queue["duplicate_alias_labels"], 1)
            self.assertEqual(
                queue["batches"][0]["items"][0]["same_content_labels"],
                ["M0002"],
            )
            digest = queue["batches"][0]["items"][0]["content_sha256"]
            fund_decision = {
                "classification": "fund",
                "viewed_original": True,
                "entries": [
                    {
                        "amount": "100",
                        "currency": "CNY",
                        "payee": "张三",
                    }
                ],
            }
            incomplete_observation_batch = self._write_review_batch(
                root,
                batch_id="observation-missing-label",
                base_fingerprint=first_page["semantic_fingerprint"],
                media_decisions={
                    "M0001": copy.deepcopy(fund_decision),
                    "M0002": copy.deepcopy(fund_decision),
                },
                media_observations={
                    "M0001": {
                        "contract_version": reconcile.MEDIA_OBSERVATION_CONTRACT,
                        "classification": "fund",
                        "review_status": "clear",
                        "viewed_original": True,
                        "recheck_reasons": [],
                    }
                },
            )
            with self.assertRaisesRegex(ValueError, "exactly one result"):
                reconcile.review_command(
                    Namespace(
                        work=work,
                        review_action="apply-batch",
                        group="测试小额群",
                        input=incomplete_observation_batch,
                    )
                )
            first_batch = self._write_review_batch(
                root,
                batch_id="observation-first",
                base_fingerprint=first_page["semantic_fingerprint"],
                media_decisions={
                    "M0001": copy.deepcopy(fund_decision),
                    "M0002": copy.deepcopy(fund_decision),
                },
                media_observations={
                    "M0001": {
                        "contract_version": reconcile.MEDIA_OBSERVATION_CONTRACT,
                        "classification": "fund",
                        "review_status": "clear",
                        "viewed_original": True,
                        "recheck_reasons": [],
                    },
                    "M0002": {"reuse_from": "M0001"},
                },
                media_view_metrics={
                    "view_batches": 1,
                    "opened_images": 1,
                    "failed_images": 0,
                    "single_image_rechecks": 0,
                    "elapsed_ms": 120,
                },
                page_commit={
                    "page_start": first_page["page_start"],
                    "page_end": first_page["page_end"],
                    "page_token": first_page["page_token"],
                },
                open_orders=[],
            )
            first_result = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=first_batch,
                )
            )
            self.assertEqual(first_result["media_observations_recorded"], 2)
            self.assertEqual(first_result["observation_cache_reuses"], 1)
            self.assertEqual(first_result["observation_cache_entries"], 1)

            second_page = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="测试小额群",
                    limit=1,
                )
            )
            second_queue = second_page["media_queue"]
            self.assertEqual(second_queue["queued_images"], 0)
            self.assertEqual(second_queue["cache_hit_labels"], 1)
            self.assertEqual(second_queue["cache_hits"][0]["labels"][0]["label"], "M0003")
            self.assertEqual(
                second_queue["cache_hits"][0]["observation"]["facts"]["entries"][0]["amount"],
                "100",
            )
            second_batch = self._write_review_batch(
                root,
                batch_id="observation-cache-hit",
                base_fingerprint=second_page["semantic_fingerprint"],
                media_decisions={"M0003": copy.deepcopy(fund_decision)},
                media_observations={"M0003": {"reuse_sha256": digest}},
                page_commit={
                    "page_start": second_page["page_start"],
                    "page_end": second_page["page_end"],
                    "page_token": second_page["page_token"],
                },
                open_orders=[],
            )
            second_result = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=second_batch,
                )
            )
            self.assertEqual(second_result["observation_cache_reuses"], 1)
            decision = reconcile._load_json(
                reconcile._decision_path(work, run["groups"][0])
            )
            self.assertEqual(sorted(decision["media_decisions"]), ["M0001", "M0002", "M0003"])
            self.assertEqual(
                decision["media_observation_cache"]["entries"][digest]["source_labels"],
                ["M0001", "M0002", "M0003"],
            )
            self.assertEqual(decision["orders"], [])

    def test_media_observation_recheck_queue_blocks_seal_until_resolved(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, run, _ = self._start_fixture(
                root,
                controlled=True,
                commit_page=False,
            )
            page = reconcile.review_command(
                Namespace(work=work, review_action="next", group="测试小额群", limit=500)
            )
            reference_decisions = {
                "M0001": {"classification": "reference", "note": "测试资料"},
                "M0002": {"classification": "reference", "note": "测试资料"},
            }
            batch_path = self._write_review_batch(
                root,
                batch_id="observation-needs-recheck",
                base_fingerprint=page["semantic_fingerprint"],
                media_decisions=reference_decisions,
                media_observations={
                    "M0001": {
                        "contract_version": reconcile.MEDIA_OBSERVATION_CONTRACT,
                        "classification": "reference",
                        "review_status": "recheck_required",
                        "viewed_original": True,
                        "recheck_reasons": ["small_text"],
                    },
                    "M0002": {
                        "contract_version": reconcile.MEDIA_OBSERVATION_CONTRACT,
                        "classification": "reference",
                        "review_status": "clear",
                        "viewed_original": True,
                        "recheck_reasons": [],
                    },
                },
                media_view_metrics={
                    "view_batches": 1,
                    "opened_images": 2,
                    "failed_images": 0,
                    "single_image_rechecks": 0,
                    "elapsed_ms": 200,
                },
                page_commit={
                    "page_start": page["page_start"],
                    "page_end": page["page_end"],
                    "page_token": page["page_token"],
                },
                open_orders=[],
            )
            applied = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=batch_path,
                )
            )
            self.assertEqual(applied["media_recheck_required"], 1)
            self.assertEqual(
                applied["media_recheck_queue"][0]["representative_label"],
                "M0001",
            )
            done_page = reconcile.review_command(
                Namespace(work=work, review_action="next", group="测试小额群", limit=500)
            )
            self.assertEqual(
                done_page["media_queue"]["pending_recheck_batches"][0]["items"][0]["label"],
                "M0001",
            )
            with self.assertRaisesRegex(ValueError, "media observation recheck"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

            decision_path = reconcile._decision_path(work, run["groups"][0])
            decision = reconcile._load_json(decision_path)
            resolved_path = self._write_review_batch(
                root,
                batch_id="observation-rechecked",
                base_fingerprint=reconcile._semantic_fingerprint(decision),
                media_decisions={"M0001": reference_decisions["M0001"]},
                media_observations={
                    "M0001": {
                        "contract_version": reconcile.MEDIA_OBSERVATION_CONTRACT,
                        "classification": "reference",
                        "review_status": "clear",
                        "viewed_original": True,
                        "recheck_reasons": [],
                    }
                },
                media_view_metrics={
                    "view_batches": 1,
                    "opened_images": 1,
                    "failed_images": 0,
                    "single_image_rechecks": 1,
                    "elapsed_ms": 80,
                },
            )
            resolved = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=resolved_path,
                )
            )
            self.assertEqual(resolved["media_recheck_required"], 0)
            self.assertEqual(
                resolved["media_view_performance"]["single_image_recheck_rate"],
                1 / 3,
            )
            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            self.assertTrue(sealed["sealed"])
            (root / "ChatExport" / "photos" / "customer.jpg").write_bytes(
                b"changed-after-observation"
            )
            with self.assertRaisesRegex(ValueError, "changed after the run snapshot"):
                reconcile.review_command(
                    Namespace(work=work, review_action="check", group="测试小额群")
                )

    def test_review_next_default_returns_largest_complete_page_within_budget(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            messages = [
                {
                    "id": index + 1,
                    "type": "message",
                    "date_unixtime": str(1780000000 + index * 60),
                    "from": "Alice",
                    "from_id": "user:alice",
                    "text": f"第{index + 1}条：" + "客户付款凭证与换汇上下文必须完整保留。" * 14,
                }
                for index in range(200)
            ]
            work, _, _ = self._start_fixture(
                root,
                messages=messages,
                controlled=True,
                commit_page=False,
            )

            page = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="测试小额群",
                    limit=None,
                )
            )
            rendered = json.dumps(
                page, ensure_ascii=False, separators=(",", ":")
            ) + "\n"
            self.assertGreater(page["page_end"], 0)
            self.assertLess(page["page_end"], len(messages))
            self.assertLessEqual(
                len(rendered), reconcile.DEFAULT_REVIEW_PAGE_OUTPUT_CHAR_BUDGET
            )

            one_more = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="测试小额群",
                    limit=page["page_end"] + 1,
                )
            )
            one_more_rendered = json.dumps(
                one_more, ensure_ascii=False, separators=(",", ":")
            ) + "\n"
            self.assertGreater(
                len(one_more_rendered), reconcile.DEFAULT_REVIEW_PAGE_OUTPUT_CHAR_BUDGET
            )

            legacy_page = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="测试小额群",
                    limit=160,
                )
            )
            self.assertEqual(legacy_page["page_end"], 160)

            completed = subprocess.run(
                [
                    sys.executable,
                    str(Path(reconcile.__file__).resolve()),
                    "review",
                    str(work),
                    "next",
                    "--group",
                    "测试小额群",
                ],
                check=True,
                capture_output=True,
                text=True,
                encoding="utf-8",
            )
            cli_page = json.loads(completed.stdout)
            self.assertEqual(cli_page["page_token"], page["page_token"])
            self.assertEqual(completed.stdout.count("\n"), 1)
            self.assertLessEqual(
                len(completed.stdout),
                reconcile.DEFAULT_REVIEW_PAGE_OUTPUT_CHAR_BUDGET,
            )

            batch_path = self._write_review_batch(
                root,
                batch_id="adaptive-page",
                base_fingerprint=page["semantic_fingerprint"],
                page_commit={
                    "page_start": page["page_start"],
                    "page_end": page["page_end"],
                    "page_token": page["page_token"],
                },
                open_orders=[],
            )
            applied = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=batch_path,
                )
            )
            self.assertEqual(applied["reviewed_through"], page["page_end"])

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            oversized_message = {
                "id": 1,
                "type": "message",
                "date_unixtime": "1780000000",
                "from": "Alice",
                "from_id": "user:alice",
                "text": "无法拆分的超长原始消息" * 2000,
            }
            work, _, _ = self._start_fixture(
                root,
                messages=[oversized_message],
                controlled=True,
                commit_page=False,
            )
            with self.assertRaisesRegex(
                ValueError, "cannot return even one complete message"
            ):
                reconcile.review_command(
                    Namespace(
                        work=work,
                        review_action="next",
                        group="测试小额群",
                        limit=None,
                    )
                )

    def test_page_commit_advances_and_carries_open_order_context(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, run, _ = self._start_fixture(
                root,
                controlled=True,
                commit_page=False,
            )
            decision_path = reconcile._decision_path(work, run["groups"][0])
            first_page = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="测试小额群",
                    limit=1,
                )
            )
            open_order = {
                "id": "P001",
                "start_message": "S00001",
                "source_messages": ["S00001"],
                "media_labels": ["M0001"],
                "customer_nickname": "Alice",
                "direction": "CNY->THB",
                "summary": "客户已付款100元，等待内部回款",
                "unresolved": ["等待内部回款", "等待确认最终汇率"],
            }
            first_commit = self._write_review_batch(
                root,
                batch_id="page-001",
                base_fingerprint=first_page["semantic_fingerprint"],
                page_commit={
                    "page_start": first_page["page_start"],
                    "page_end": first_page["page_end"],
                    "page_token": first_page["page_token"],
                },
                open_orders=[open_order],
            )
            applied = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=first_commit,
                )
            )
            self.assertEqual(applied["reviewed_through"], 1)
            self.assertFalse(applied["read_complete"])

            second_page = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="测试小额群",
                    limit=1,
                )
            )
            self.assertEqual(second_page["page_start"], 1)
            self.assertEqual(second_page["open_orders"][0]["id"], "P001")
            self.assertEqual(
                [message["label"] for message in second_page["carry_messages"]],
                ["S00001"],
            )

            second_commit = self._write_review_batch(
                root,
                batch_id="page-002",
                base_fingerprint=second_page["semantic_fingerprint"],
                page_commit={
                    "page_start": second_page["page_start"],
                    "page_end": second_page["page_end"],
                    "page_token": second_page["page_token"],
                },
                open_orders=[open_order],
            )
            reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=second_commit,
                )
            )
            decision = reconcile._load_json(decision_path)
            self.assertEqual(decision["reviewed_through"], 2)
            self.assertTrue(decision["read_complete"])
            with self.assertRaisesRegex(ValueError, "open_orders must be resolved"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_page_commit_requires_explicit_open_order_state_and_valid_token(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, run, _ = self._start_fixture(
                root,
                controlled=True,
                commit_page=False,
            )
            decision_path = reconcile._decision_path(work, run["groups"][0])
            page = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="测试小额群",
                    limit=1,
                )
            )
            missing_state = self._write_review_batch(
                root,
                batch_id="page-missing-state",
                base_fingerprint=page["semantic_fingerprint"],
                page_commit={
                    "page_start": page["page_start"],
                    "page_end": page["page_end"],
                    "page_token": page["page_token"],
                },
            )
            with self.assertRaisesRegex(ValueError, "must include the complete open_orders"):
                reconcile.review_command(
                    Namespace(
                        work=work,
                        review_action="apply-batch",
                        group="测试小额群",
                        input=missing_state,
                    )
                )

            wrong_token = self._write_review_batch(
                root,
                batch_id="page-wrong-token",
                base_fingerprint=page["semantic_fingerprint"],
                page_commit={
                    "page_start": page["page_start"],
                    "page_end": page["page_end"],
                    "page_token": "sha256:" + "0" * 64,
                },
                open_orders=[],
            )
            with self.assertRaisesRegex(ValueError, "token does not match"):
                reconcile.review_command(
                    Namespace(
                        work=work,
                        review_action="apply-batch",
                        group="测试小额群",
                        input=wrong_token,
                    )
                )
            unchanged = reconcile._load_json(decision_path)
            self.assertEqual(unchanged["reviewed_through"], 0)
            self.assertFalse(unchanged["read_complete"])

    def test_upgrade_checkpoints_preserves_legacy_decisions_but_resets_read_progress(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, run, _ = self._start_fixture(root, controlled=True)
            run = reconcile._load_run(work)
            run_group = run["groups"][0]
            decision_path = reconcile._decision_path(work, run_group)
            decision = reconcile._load_json(decision_path)
            media_decisions, orders = self._complete_batch_values()
            decision["contract_version"] = reconcile.LEGACY_DECISION_CONTRACT
            decision.pop("open_orders")
            decision["media_decisions"] = media_decisions
            decision["orders"] = orders
            legacy_fingerprint = reconcile._legacy_semantic_fingerprint_3_1(decision)
            decision["edit_control"] = {
                "mode": reconcile.LEGACY_EDIT_CONTROL_MODE,
                "approved_semantic_fingerprint": legacy_fingerprint,
                "batch_count": 0,
                "last_batch": None,
            }
            run_group["edit_mode"] = reconcile.LEGACY_EDIT_CONTROL_MODE
            core.atomic_json(decision_path, decision)
            core.atomic_json(reconcile._run_path(work), run)

            upgraded = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="upgrade-checkpoints",
                    group="测试小额群",
                )
            )

            self.assertEqual(upgraded["previous_reviewed_through"], 2)
            self.assertEqual(upgraded["reviewed_through"], 0)
            self.assertEqual(upgraded["preserved_media_decisions"], 2)
            self.assertEqual(upgraded["preserved_orders"], 1)
            decision = reconcile._load_json(decision_path)
            self.assertEqual(decision["contract_version"], reconcile.DECISION_CONTRACT)
            self.assertEqual(decision["open_orders"], [])
            self.assertFalse(decision["read_complete"])
            next_page = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="next",
                    group="测试小额群",
                    limit=1,
                )
            )
            self.assertEqual(next_page["page_start"], 0)

    def test_apply_batch_atomically_updates_a_controlled_decision(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, run, _ = self._start_fixture(root, controlled=True)
            decision_path = reconcile._decision_path(work, run["groups"][0])
            initial = reconcile._load_json(decision_path)
            base_fingerprint = reconcile._semantic_fingerprint(initial)
            media_decisions, orders = self._complete_batch_values()
            batch_path = self._write_review_batch(
                root,
                batch_id="fixture-001",
                base_fingerprint=base_fingerprint,
                media_decisions=media_decisions,
                orders=orders,
            )

            applied = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=batch_path,
                )
            )

            self.assertEqual(applied["batch_id"], "fixture-001")
            self.assertTrue(applied["controlled_editing"])
            self.assertEqual(applied["orders"], 1)
            self.assertEqual(applied["fund_entries"], 2)
            self.assertEqual(applied["evidence_hashes_computed"], 2)
            self.assertEqual(applied["evidence_hashes_reused"], 0)
            decision = reconcile._load_json(decision_path)
            self.assertEqual(
                decision["edit_control"]["approved_semantic_fingerprint"],
                reconcile._semantic_fingerprint(decision),
            )
            self.assertEqual(
                decision["edit_control"]["last_batch"]["id"],
                "fixture-001",
            )
            self.assertTrue(
                decision["media_decisions"]["M0001"]["evidence_sha256"]
            )
            replayed = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=batch_path,
                )
            )
            self.assertTrue(replayed["idempotent_replay"])
            self.assertEqual(replayed["evidence_hashes_computed"], 0)
            self.assertEqual(replayed["evidence_hashes_reused"], 2)

            decision = reconcile._load_json(decision_path)
            order_only_path = self._write_review_batch(
                root,
                batch_id="fixture-order-only",
                base_fingerprint=reconcile._semantic_fingerprint(decision),
                orders=copy.deepcopy(decision["orders"]),
            )
            order_only = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=order_only_path,
                )
            )
            self.assertEqual(order_only["evidence_hashes_computed"], 0)
            self.assertEqual(order_only["evidence_hashes_reused"], 2)
            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            self.assertTrue(sealed["sealed"])
            self.assertEqual(sealed["evidence_hashes_computed"], 2)
            self.assertEqual(sealed["evidence_hashes_reused"], 0)

    def test_controlled_decision_rejects_direct_semantic_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, run, _ = self._start_fixture(root, controlled=True)
            decision_path = reconcile._decision_path(work, run["groups"][0])
            decision = reconcile._load_json(decision_path)
            media_decisions, orders = self._complete_batch_values()
            batch_path = self._write_review_batch(
                root,
                batch_id="fixture-001",
                base_fingerprint=reconcile._semantic_fingerprint(decision),
                media_decisions=media_decisions,
                orders=orders,
            )
            reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="测试小额群",
                    input=batch_path,
                )
            )
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0001"]["entries"][0]["amount"] = "999"
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(
                ValueError,
                "edited outside review apply-batch",
            ):
                reconcile.review_command(
                    Namespace(work=work, review_action="check", group="测试小额群")
                )

    def test_invalid_apply_batch_does_not_partially_write_the_decision(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, run, _ = self._start_fixture(root, controlled=True)
            decision_path = reconcile._decision_path(work, run["groups"][0])
            initial = reconcile._load_json(decision_path)
            initial_fingerprint = reconcile._semantic_fingerprint(initial)
            media_decisions, orders = self._complete_batch_values(
                payment_side="payout"
            )
            batch_path = self._write_review_batch(
                root,
                batch_id="fixture-invalid",
                base_fingerprint=initial_fingerprint,
                media_decisions=media_decisions,
                orders=orders,
            )

            with self.assertRaisesRegex(
                ValueError,
                "customer sender.*ordinary fund entry.*payment",
            ):
                reconcile.review_command(
                    Namespace(
                        work=work,
                        review_action="apply-batch",
                        group="测试小额群",
                        input=batch_path,
                    )
                )

            unchanged = reconcile._load_json(decision_path)
            self.assertEqual(
                reconcile._semantic_fingerprint(unchanged),
                initial_fingerprint,
            )
            self.assertEqual(unchanged["media_decisions"], {})
            self.assertEqual(unchanged["orders"], [])

    def test_start_seal_and_finish_publish_only_the_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, _, _ = self._start_fixture(root, controlled=True)
            decision_path = self._write_complete_decision(work, through_batch=True)
            self.assertEqual(
                reconcile._load_json(decision_path)["contract_version"],
                "group-chat-decision/3.2",
            )
            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            self.assertTrue(sealed["sealed"])
            audit = reconcile.review_command(
                Namespace(work=work, review_action="audit", group=None)
            )
            self.assertEqual(audit["summary"]["groups"], 1)
            self.assertEqual(audit["summary"]["flagged_groups"], 0)
            output = root / "ledger.xlsx"
            result = reconcile.finish_run(
                Namespace(
                    work=work,
                    output=output,
                    template=Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
                )
            )
            self.assertTrue(output.is_file())
            self.assertEqual(result["pending_orders"], 0)
            self.assertEqual(result["risk_report"]["groups"], 1)
            self.assertEqual(result["risk_report"]["flagged_groups"], 0)
            self.assertFalse((work / "events").exists())
            self.assertFalse((work / "orders.json").exists())
            self.assertFalse((work / "simple_plan.json").exists())

    def test_generic_payee_is_rejected_at_group_seal(self) -> None:
        for payee in (
            "截图所示人民币收款方",
            "群内收款方",
            "泰铢收款账户",
            "USDT收款钱包",
            "银行卡收款方",
            "客户退款钱包",
        ):
            with self.subTest(payee=payee), tempfile.TemporaryDirectory() as temporary:
                work, _, _ = self._start_fixture(Path(temporary))
                self._write_complete_decision(work, payee=payee)
                with self.assertRaisesRegex(ValueError, "generic placeholder"):
                    reconcile.review_command(
                        Namespace(work=work, review_action="seal", group="测试小额群")
                    )

    def test_current_contract_infers_and_validates_payee_evidence_state(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            entry = decision["media_decisions"]["M0001"]["entries"][0]
            entry.pop("payee_state")
            core.atomic_json(decision_path, decision)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            decision = reconcile._load_json(decision_path)
            self.assertEqual(
                decision["media_decisions"]["M0001"]["entries"][0]["payee_state"],
                "visible",
            )

            entry = decision["media_decisions"]["M0001"]["entries"][0]
            entry["payee_state"] = "not_shown"
            entry["payee"] = "张三"
            core.atomic_json(decision_path, decision)
            with self.assertRaisesRegex(ValueError, "must be 未显示"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_transfer_payee_must_be_json_text(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0001"]["entries"][0]["payee"] = 6222021234567890
            core.atomic_json(decision_path, decision)
            with self.assertRaisesRegex(ValueError, "must be text"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_thai_transfer_records_only_the_visible_bank_account(self) -> None:
        for account in ("X-1342", "XXX-XXX-0116", "206-4-XXX781"):
            with self.subTest(account=account), tempfile.TemporaryDirectory() as temporary:
                work, _, _ = self._start_fixture(Path(temporary))
                decision_path = self._write_complete_decision(work)
                decision = reconcile._load_json(decision_path)
                decision["media_decisions"]["M0002"]["entries"][0]["payee"] = account
                core.atomic_json(decision_path, decision)
                sealed = reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )
                self.assertTrue(sealed["sealed"])

        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0002"]["entries"][0]["payee"] = "สมชาย ใจดี"
            core.atomic_json(decision_path, decision)
            with self.assertRaisesRegex(ValueError, "exact visible Thai bank account"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_legacy_2_3_decision_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["contract_version"] = "group-chat-decision/2.3"
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(ValueError, "unsupported decision contract"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_mass_unknown_payees_require_explicit_original_image_review(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            entries = decision["media_decisions"]["M0001"]["entries"]
            entries[0]["payee"] = "未显示"
            entries[0]["payee_state"] = "not_shown"
            for _ in range(9):
                entries.append(
                    {
                        "amount": "1",
                        "currency": "CNY",
                        "payee": "无法辨认",
                        "payee_state": "unreadable",
                        "side": "payment",
                        "result": "completed",
                        "amount_state": "clear",
                    }
                )
            decision["orders"][0]["entry_ids"] = [
                *(f"M0001.{index}" for index in range(1, 11)),
                "M0002.1",
            ]
            core.atomic_json(decision_path, decision)

            checked = reconcile.review_command(
                Namespace(work=work, review_action="check", group="测试小额群")
            )
            self.assertEqual(checked["transfer_payees"], 11)
            self.assertEqual(checked["unknown_payees"], 10)
            self.assertEqual(checked["unknown_payee_review_required"], 1)
            self.assertEqual(checked["unknown_payee_warning"], 1)
            self.assertEqual(checked["unreviewed_unknown_payees"], 9)
            with self.assertRaisesRegex(ValueError, "unreadable payee review"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

            decision = reconcile._load_json(decision_path)
            decision["unknown_payee_reviewed_entry_ids"] = [
                f"M0001.{index}" for index in range(2, 11)
            ]
            core.atomic_json(decision_path, decision)
            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            self.assertTrue(sealed["sealed"])

    def test_one_image_can_compile_multiple_fund_entries(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, normalized = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0001"]["entries"].append(
                {
                    "amount": "25",
                    "currency": "USD",
                    "payee": "张三",
                    "payee_state": "visible",
                    "side": "payment",
                    "result": "completed",
                    "amount_state": "clear",
                }
            )
            decision["orders"][0]["entry_ids"].insert(1, "M0001.2")
            core.atomic_json(decision_path, decision)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            sealed = reconcile._load_json(decision_path)
            events, _ = reconcile._compile_decisions_v3(
                normalized,
                {str(normalized["groups"][0]["group_key"]): sealed},
            )
            fund_events = [event for event in events if event["type"] in core.FUND_EVENT_TYPES]
            self.assertEqual(len(fund_events), 3)
            first_media_ids = [
                event["event_id"]
                for event in fund_events
                if event["media_id"].endswith("#media:0")
                and event["message_id"].endswith(":1")
            ]
            self.assertEqual(len(first_media_ids), 2)

    def test_seal_derives_normal_fund_defaults_from_sender_and_visible_fields(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            entry = decision["media_decisions"]["M0001"]["entries"][0]
            for field in (
                "side",
                "result",
                "amount_state",
                "payee_state",
                "amount_basis",
            ):
                entry.pop(field, None)
            core.atomic_json(decision_path, decision)

            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            entry = reconcile._load_json(decision_path)["media_decisions"]["M0001"]["entries"][0]
            self.assertEqual(entry["side"], "payment")
            self.assertEqual(entry["result"], "completed")
            self.assertEqual(entry["amount_state"], "clear")
            self.assertEqual(entry["payee_state"], "visible")
            self.assertEqual(entry["amount_basis"], "receiver_received")

        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0001"]["entries"][0]["amount_basis"] = (
                "payer_discounted_after_coupon"
            )
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(ValueError, "amount_basis is unsupported"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_seal_derives_failed_result_from_explicit_failure_text(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            entry = decision["media_decisions"]["M0001"]["entries"][0]
            entry.pop("result")
            entry["status_text"] = "交易已取消"
            core.atomic_json(decision_path, decision)

            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            entry = reconcile._load_json(decision_path)["media_decisions"]["M0001"]["entries"][0]
            self.assertEqual(entry["result"], "failed")

    def test_seal_rejects_internal_sender_recorded_as_customer_payment(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0002"]["entries"][0]["side"] = "payment"
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(
                ValueError,
                "internal sender.*ordinary fund entry.*payout",
            ):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_seal_rejects_customer_sender_recorded_as_internal_payout(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0001"]["entries"][0]["side"] = "payout"
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(
                ValueError,
                "customer sender.*ordinary fund entry.*payment",
            ):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_seal_rejects_guessed_side_for_unknown_sender_role(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(
                Path(temporary),
                messages=[
                    {
                        "id": 1,
                        "type": "message",
                        "date_unixtime": "1780000000",
                        "from": "统计机器人",
                        "from_id": "bot:stats",
                        "text": "付款100元",
                        "photo": "photos/customer.jpg",
                    },
                    {
                        "id": 2,
                        "type": "message",
                        "date_unixtime": "1780000060",
                        "from": "QQ财务3",
                        "from_id": "user:staff",
                        "text": "回500铢",
                        "photo": "photos/staff.jpg",
                    },
                ],
            )
            self._write_complete_decision(work)

            with self.assertRaisesRegex(
                ValueError,
                "unknown sender role.*side=unknown",
            ):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_currency_does_not_override_customer_and_internal_roles(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            internal_entry = decision["media_decisions"]["M0002"]["entries"][0]
            internal_entry.update(
                {
                    "amount": "15",
                    "currency": "USDT",
                    "payee": "TInternalPayoutWallet",
                    "side": "payout",
                }
            )
            order = decision["orders"][0]
            order["direction"] = "CNY->USDT"
            order["pricing"] = {
                "source_messages": ["S00001", "S00002"],
                "terms": {"rate": "0.15", "operator": "multiply"},
                "expected": {"kind": "explicit", "amount": "15"},
            }
            core.atomic_json(decision_path, decision)

            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            self.assertTrue(sealed["sealed"])

    def test_seal_allows_internal_sender_relay_with_explicit_chat_basis(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(
                Path(temporary),
                extra_messages=[
                    {
                        "id": 3,
                        "type": "message",
                        "date_unixtime": "1780000120",
                        "from": "QQ财务3",
                        "from_id": "user:staff",
                        "text": "这张是代客户转发的付款凭证",
                    }
                ],
            )
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            entry = decision["media_decisions"]["M0002"]["entries"][0]
            entry["side"] = "payment"
            entry["side_exception"] = {
                "kind": "relayed_customer_payment",
                "source_messages": ["S00003"],
                "detail": "内部人员代客户转发付款凭证",
            }
            decision["orders"][0]["source_messages"].append("S00003")
            core.atomic_json(decision_path, decision)

            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            self.assertTrue(sealed["sealed"])
            normalized = reconcile._load_snapshot(work, reconcile._load_run(work))
            events, plan = reconcile._compile_decisions_v3(
                normalized,
                {
                    str(normalized["groups"][0]["group_key"]): reconcile._load_json(
                        decision_path
                    )
                },
            )
            relay_event = next(
                event
                for event in events
                if event.get("flow_side") == "payment"
                and str(event.get("message_id") or "").endswith(":2")
            )
            self.assertEqual(
                relay_event["side_exception"]["kind"],
                "relayed_customer_payment",
            )
            reconcile.simple_ledger.compile_simple_ledger(
                normalized,
                events,
                plan["events_fingerprint"],
                plan,
            )

    def test_seal_auto_adds_side_exception_basis_to_order_sources(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(
                Path(temporary),
                extra_messages=[
                    {
                        "id": 3,
                        "type": "message",
                        "date_unixtime": "1780000120",
                        "from": "QQ财务3",
                        "from_id": "user:staff",
                        "text": "这张是代客户转发的付款凭证",
                    }
                ],
            )
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            entry = decision["media_decisions"]["M0002"]["entries"][0]
            entry["side"] = "payment"
            entry["side_exception"] = {
                "kind": "relayed_customer_payment",
                "source_messages": ["S00003"],
                "detail": "内部人员代客户转发付款凭证",
            }
            core.atomic_json(decision_path, decision)

            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            order = reconcile._load_json(decision_path)["orders"][0]
            self.assertIn("S00003", order["source_messages"])

    def test_usdt_106_customer_payment_and_cny_695_internal_payout_reach_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, _, _ = self._start_fixture(
                root,
                messages=[
                    {
                        "id": 1,
                        "type": "message",
                        "date_unixtime": "1780000000",
                        "from": "不忘初心-",
                        "from_id": "user:customer",
                        "text": "U换人民币，700元大概多少U？",
                    },
                    {
                        "id": 2,
                        "type": "message",
                        "date_unixtime": "1780000030",
                        "from": "QQ财务3",
                        "from_id": "user:staff",
                        "text": "按6.56",
                    },
                    {
                        "id": 3,
                        "type": "message",
                        "date_unixtime": "1780000060",
                        "from": "不忘初心-",
                        "from_id": "user:customer",
                        "text": "700/6.56=106",
                        "photo": "photos/customer.jpg",
                    },
                    {
                        "id": 4,
                        "type": "message",
                        "date_unixtime": "1780000090",
                        "from": "QQ财务3",
                        "from_id": "user:staff",
                        "text": "106×6.56=695",
                        "photo": "photos/staff.jpg",
                    },
                ],
            )
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            customer_entry = decision["media_decisions"]["M0001"]["entries"][0]
            customer_entry.update(
                {
                    "amount": "106",
                    "currency": "USDT",
                    "payee": "TQ7xCustomerDepositAddress",
                    "side": "payment",
                }
            )
            internal_entry = decision["media_decisions"]["M0002"]["entries"][0]
            internal_entry.update(
                {
                    "amount": "695",
                    "currency": "CNY",
                    "payee": "李四",
                    "side": "payment",
                }
            )
            order = decision["orders"][0]
            order.update(
                {
                    "source_messages": ["S00001", "S00002", "S00003", "S00004"],
                    "customer_nickname": "不忘初心-",
                    "direction": "USDT->CNY",
                    "pricing": {
                        "source_messages": ["S00001", "S00002", "S00003", "S00004"],
                        "terms": {"rate": "6.56", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "695"},
                    },
                }
            )
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(
                ValueError,
                "internal sender.*ordinary fund entry.*payout",
            ):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0002"]["entries"][0]["side"] = "payout"
            core.atomic_json(decision_path, decision)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )

            output = root / "usdt-cny-ledger.xlsx"
            result = reconcile.finish_run(
                Namespace(
                    work=work,
                    output=output,
                    template=Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
                )
            )
            self.assertEqual(result["pending_orders"], 0)

            workbook = load_workbook(output, read_only=True, data_only=False)
            try:
                worksheet = workbook.worksheets[0]
                rows = list(worksheet.iter_rows(min_row=2, values_only=True))
                summary = next(row for row in rows if row[0] == "订单汇总")
                self.assertEqual(summary[core.HEADERS.index("客户昵称")], "不忘初心-")
                self.assertEqual(summary[core.HEADERS.index("换汇方向")], "USDT->CNY")
                self.assertEqual(summary[core.HEADERS.index("付款合计")], 106)
                self.assertEqual(summary[core.HEADERS.index("汇率")], 6.56)
                self.assertEqual(summary[core.HEADERS.index("应回金额")], 695)
                self.assertEqual(summary[core.HEADERS.index("内部实际回款合计")], 695)
                self.assertIsNone(summary[core.HEADERS.index("核对结果")])
                customer_flow = next(row for row in rows if row[0] == "客户付款")
                internal_flow = next(row for row in rows if row[0] == "内部回款")
                self.assertEqual(
                    (
                        customer_flow[core.HEADERS.index("收款方实际到账金额")],
                        customer_flow[core.HEADERS.index("流水币种")],
                    ),
                    (106, "USDT"),
                )
                self.assertEqual(
                    (
                        internal_flow[core.HEADERS.index("收款方实际到账金额")],
                        internal_flow[core.HEADERS.index("流水币种")],
                    ),
                    (695, "CNY"),
                )
            finally:
                workbook.close()

    def test_seal_rejects_pending_result_for_normal_processing_status(self) -> None:
        for status_text in (
            "确认中",
            "处理中",
            "等待确认",
            "待区块确认",
            "已发送待区块确认",
            "Pending",
            "Processing",
            "Confirming",
            "Unconfirmed",
        ):
            with self.subTest(status_text=status_text), tempfile.TemporaryDirectory() as temporary:
                work, _, _ = self._start_fixture(Path(temporary))
                decision_path = self._write_complete_decision(work)
                decision = reconcile._load_json(decision_path)
                entry = decision["media_decisions"]["M0001"]["entries"][0]
                entry["status_text"] = status_text
                entry["result"] = "pending"
                core.atomic_json(decision_path, decision)

                with self.assertRaisesRegex(
                    ValueError,
                    "normal processing status.*result=completed",
                ):
                    reconcile.review_command(
                        Namespace(work=work, review_action="seal", group="测试小额群")
                    )

    def test_seal_rejects_failed_result_without_explicit_failure_text(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            entry = decision["media_decisions"]["M0001"]["entries"][0]
            entry["status_text"] = "已提交"
            entry["result"] = "failed"
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(
                ValueError,
                "result=failed requires explicit failure",
            ):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_seal_rejects_completed_result_for_explicit_failure_status(self) -> None:
        for status_text in ("交易失败", "已取消", "银行拒绝", "凭证作废", "交易无效"):
            with self.subTest(status_text=status_text), tempfile.TemporaryDirectory() as temporary:
                work, _, _ = self._start_fixture(Path(temporary))
                decision_path = self._write_complete_decision(work)
                decision = reconcile._load_json(decision_path)
                entry = decision["media_decisions"]["M0001"]["entries"][0]
                entry["status_text"] = status_text
                entry["result"] = "completed"
                core.atomic_json(decision_path, decision)

                with self.assertRaisesRegex(
                    ValueError,
                    "explicit failure status.*result=failed",
                ):
                    reconcile.review_command(
                        Namespace(work=work, review_action="seal", group="测试小额群")
                    )

    def test_seal_auto_adds_fund_entry_sources_to_order_source_messages(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["source_messages"] = ["S00002"]
            decision["orders"][0]["pricing"]["source_messages"] = ["S00002"]
            core.atomic_json(decision_path, decision)

            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            order = reconcile._load_json(decision_path)["orders"][0]
            self.assertEqual(order["source_messages"], ["S00002", "S00001"])

    def test_check_reports_mass_one_image_one_order_degradation_without_permanent_block(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            messages = [
                {
                    "id": position,
                    "type": "message",
                    "date_unixtime": str(1780000000 + position * 60),
                    "from": "Alice",
                    "from_id": "user:alice",
                    "text": "",
                    "photo": "photos/customer.jpg",
                }
                for position in range(1, 21)
            ]
            work, _, _ = self._start_fixture(Path(temporary), messages=messages)
            run = reconcile._load_run(work)
            decision_path = reconcile._decision_path(work, run["groups"][0])
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"] = {}
            decision["orders"] = []
            for position in range(1, 21):
                media_label = f"M{position:04d}"
                message_label = f"S{position:05d}"
                entry_id = f"{media_label}.1"
                decision["media_decisions"][media_label] = {
                    "classification": "fund",
                    "viewed_original": True,
                    "entries": [
                        {
                            "amount": str(position),
                            "currency": "CNY",
                            "payee": "陈甲",
                            "payee_state": "visible",
                            "side": "payment",
                            "result": "completed",
                            "amount_state": "clear",
                        }
                    ],
                }
                decision["orders"].append(
                    {
                        "id": f"O{position:03d}",
                        "entry_ids": [entry_id],
                        "source_messages": [message_label],
                        "customer_nickname": "Alice",
                        "direction": "",
                        "pricing": {
                            "source_messages": [message_label],
                            "expected": {"kind": "unknown", "reason": "not_stated"},
                        },
                        "note": "孤立凭证，缺少方向和计价上下文。",
                    }
                )
            core.atomic_json(decision_path, decision)

            report = reconcile.review_command(
                Namespace(work=work, review_action="check", group="测试小额群")
            )
            self.assertEqual(report["single_entry_orders"], 20)
            self.assertEqual(report["blank_direction_orders"], 20)
            self.assertEqual(report["unknown_pricing_orders"], 20)
            self.assertEqual(report["own_fund_source_only_orders"], 20)
            self.assertEqual(report["repeated_order_notes"], 19)
            self.assertEqual(report["mass_degenerate_orders"], 20)
            self.assertEqual(report["mass_degenerate_order_ratio"], 1.0)
            self.assertEqual(report["suspected_bulk_order_creation"], 1)

            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            self.assertTrue(sealed["sealed"])

    def test_seal_reports_multi_signal_degradation_even_with_filled_directions(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            messages = [
                {
                    "id": position,
                    "type": "message",
                    "date_unixtime": str(1780000000 + position * 60),
                    "from": "Alice",
                    "from_id": "user:alice",
                    "text": f"第{position}单付款凭证",
                    "photo": "photos/customer.jpg",
                }
                for position in range(1, 21)
            ]
            work, run, _ = self._start_fixture(
                Path(temporary),
                messages=messages,
            )
            decision_path = reconcile._decision_path(work, run["groups"][0])
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"] = {}
            decision["orders"] = []
            for position in range(1, 21):
                media_label = f"M{position:04d}"
                message_label = f"S{position:05d}"
                entry_id = f"{media_label}.1"
                decision["media_decisions"][media_label] = {
                    "classification": "fund",
                    "viewed_original": True,
                    "entries": [
                        {
                            "amount": str(position),
                            "currency": "CNY",
                            "payee": "陈甲",
                            "payee_state": "visible",
                            "side": "payment",
                            "result": "completed",
                            "amount_state": "clear",
                        }
                    ],
                }
                decision["orders"].append(
                    {
                        "id": f"O{position:03d}",
                        "entry_ids": [entry_id],
                        "source_messages": [message_label],
                        "customer_nickname": "Alice",
                        "direction": "CNY->THB",
                        "pricing": {
                            "source_messages": [message_label],
                            "expected": {
                                "kind": "unknown",
                                "reason": "not_stated",
                            },
                        },
                    }
                )
            core.atomic_json(decision_path, decision)

            report = reconcile.review_command(
                Namespace(work=work, review_action="check", group="测试小额群")
            )
            self.assertEqual(report["mass_degenerate_orders"], 0)
            self.assertEqual(report["single_unknown_pricing_orders"], 20)
            self.assertEqual(report["multi_signal_degenerate_orders"], 20)
            self.assertEqual(report["suspected_bulk_order_creation"], 1)

            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            self.assertTrue(sealed["sealed"])

    def test_many_single_entry_orders_pass_when_pricing_and_context_are_specific(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            messages = [
                {
                    "id": position,
                    "type": "message",
                    "date_unixtime": str(1780000000 + position * 60),
                    "from": "Alice",
                    "from_id": "user:alice",
                    "text": f"第{position}单独立付款，明确应回{position}泰铢",
                    "photo": "photos/customer.jpg",
                }
                for position in range(1, 13)
            ]
            work, run, _ = self._start_fixture(
                Path(temporary),
                messages=messages,
            )
            decision_path = reconcile._decision_path(work, run["groups"][0])
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"] = {}
            decision["orders"] = []
            for position in range(1, 13):
                media_label = f"M{position:04d}"
                message_label = f"S{position:05d}"
                entry_id = f"{media_label}.1"
                decision["media_decisions"][media_label] = {
                    "classification": "fund",
                    "viewed_original": True,
                    "entries": [
                        {
                            "amount": str(position),
                            "currency": "CNY",
                            "payee": "陈甲",
                            "payee_state": "visible",
                            "side": "payment",
                            "result": "completed",
                            "amount_state": "clear",
                        }
                    ],
                }
                decision["orders"].append(
                    {
                        "id": f"O{position:03d}",
                        "entry_ids": [entry_id],
                        "source_messages": [message_label],
                        "customer_nickname": "Alice",
                        "direction": "CNY->THB",
                        "pricing": {
                            "source_messages": [message_label],
                            "terms": {"rate": "1", "operator": "multiply"},
                            "expected": {
                                "kind": "explicit",
                                "amount": str(position),
                            },
                        },
                    }
                )
            core.atomic_json(decision_path, decision)

            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            self.assertTrue(sealed["sealed"])
            self.assertEqual(sealed["multi_signal_degenerate_orders"], 0)

    def test_seal_rejects_completed_order_without_meaningful_chat_context(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(
                Path(temporary),
                messages=[
                    {
                        "id": 1,
                        "type": "message",
                        "date_unixtime": "1780000000",
                        "from": "Alice",
                        "from_id": "user:alice",
                        "text": "",
                        "photo": "photos/customer.jpg",
                    },
                    {
                        "id": 2,
                        "type": "message",
                        "date_unixtime": "1780000060",
                        "from": "QQ财务3",
                        "from_id": "user:staff",
                        "text": "",
                        "photo": "photos/staff.jpg",
                    },
                ],
            )
            self._write_complete_decision(work)

            with self.assertRaisesRegex(
                ValueError,
                "completed ordinary order.*meaningful chat context",
            ):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_small_isolated_fund_order_requires_specific_note_then_can_seal(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(
                Path(temporary),
                messages=[
                    {
                        "id": 1,
                        "type": "message",
                        "date_unixtime": "1780000000",
                        "from": "Alice",
                        "from_id": "user:alice",
                        "text": "",
                        "photo": "photos/customer.jpg",
                    }
                ],
            )
            run = reconcile._load_run(work)
            decision_path = reconcile._decision_path(work, run["groups"][0])
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"] = {
                "M0001": {
                    "classification": "fund",
                    "viewed_original": True,
                    "entries": [
                        {
                            "amount": "88",
                            "currency": "CNY",
                            "payee": "陈甲",
                            "payee_state": "visible",
                            "side": "payment",
                            "result": "completed",
                            "amount_state": "clear",
                        }
                    ],
                }
            }
            decision["orders"] = [
                {
                    "id": "O001",
                    "entry_ids": ["M0001.1"],
                    "source_messages": ["S00001"],
                    "customer_nickname": "Alice",
                    "direction": "",
                    "pricing": {
                        "source_messages": ["S00001"],
                        "expected": {"kind": "unknown", "reason": "not_stated"},
                    },
                }
            ]
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(
                ValueError,
                "isolated fund evidence.*specific missing-context note",
            ):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["note"] = "聊天只剩客户付款凭证，未说明换汇方向、汇率和内部回款。"
            core.atomic_json(decision_path, decision)
            sealed = reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            self.assertTrue(sealed["sealed"])

    def test_seal_rejects_fund_entries_not_explicitly_assigned_to_an_order(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0001"]["entries"].append(
                {
                    "amount": "25",
                    "currency": "USD",
                    "payee": "张三",
                    "payee_state": "visible",
                    "side": "payment",
                    "result": "completed",
                    "amount_state": "clear",
                }
            )
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(ValueError, "must be explicitly assigned to an order"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_calculated_terms_require_an_operator(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            pricing = decision["orders"][0]["pricing"]
            pricing["terms"].pop("operator")
            pricing["expected"] = {"kind": "calculated_from_terms"}
            core.atomic_json(decision_path, decision)
            with self.assertRaisesRegex(ValueError, "terms.operator is required"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_calculated_from_terms_uses_the_group_formula(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, normalized = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["pricing"]["expected"] = {
                "kind": "calculated_from_terms"
            }
            core.atomic_json(decision_path, decision)

            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            sealed = reconcile._load_json(decision_path)
            events, plan = reconcile._compile_decisions_v3(
                normalized,
                {str(normalized["groups"][0]["group_key"]): sealed},
            )
            orders, _ = reconcile.simple_ledger.compile_simple_ledger(
                normalized,
                events,
                plan["events_fingerprint"],
                plan,
            )
            self.assertEqual(
                orders["groups"][0]["orders"][0]["expected_payout"],
                "500",
            )

    def test_explicit_expected_amount_wins_over_formula_diagnostic(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, normalized = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["pricing"]["terms"]["rate"] = "4.9"
            core.atomic_json(decision_path, decision)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            sealed = reconcile._load_json(decision_path)
            events, plan = reconcile._compile_decisions_v3(
                normalized,
                {str(normalized["groups"][0]["group_key"]): sealed},
            )
            orders, statistics = reconcile.simple_ledger.compile_simple_ledger(
                normalized,
                events,
                plan["events_fingerprint"],
                plan,
            )
            order = orders["groups"][0]["orders"][0]
            self.assertEqual(order["expected_payout"], "500")
            self.assertEqual(order["review_result"], "")
            self.assertEqual(order["reconciliation"]["status"], "matched")
            self.assertEqual(order["order_status"], "completed")
            self.assertEqual(statistics["pending_orders"], 0)
            self.assertIn("核对仍以明确应回为准", order["pricing_diagnostics"][0])
            self.assertIsNone(build_workbook.order_row_note(order))

    def test_actual_minus_authoritative_expected_is_written_to_result(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, normalized = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0002"]["entries"][0]["amount"] = "499.70"
            core.atomic_json(decision_path, decision)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            sealed = reconcile._load_json(decision_path)
            events, plan = reconcile._compile_decisions_v3(
                normalized,
                {str(normalized["groups"][0]["group_key"]): sealed},
            )
            orders, _ = reconcile.simple_ledger.compile_simple_ledger(
                normalized,
                events,
                plan["events_fingerprint"],
                plan,
            )
            order = orders["groups"][0]["orders"][0]
            self.assertEqual(order["reconciliation"]["status"], "short")
            self.assertEqual(order["reconciliation"]["difference"], "-0.3")
            self.assertEqual(order["review_result"], "少转 0.3 THB")
            summary = build_workbook.order_rows(order)[0]
            self.assertEqual(
                summary[core.HEADERS.index("核对结果")],
                "少转 0.3 THB",
            )

    def test_check_reports_and_seal_rejects_missing_pricing(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["orders"][0].pop("pricing")
            core.atomic_json(decision_path, decision)

            checked = reconcile.review_command(
                Namespace(work=work, review_action="check", group="测试小额群")
            )
            self.assertEqual(checked["pricing_scopes"], 1)
            self.assertEqual(checked["missing_pricing_scopes"], 1)
            self.assertFalse(checked["sealed"])

            with self.assertRaisesRegex(ValueError, "pricing is required before seal"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_explicit_pricing_requires_and_exports_the_confirmed_rate(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, normalized = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["pricing"].pop("terms")
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(ValueError, "confirmed exchange rate"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["pricing"]["terms"] = {
                "rate": "5",
                "operator": "multiply",
            }
            core.atomic_json(decision_path, decision)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            sealed = reconcile._load_json(decision_path)
            events, plan = reconcile._compile_decisions_v3(
                normalized,
                {str(normalized["groups"][0]["group_key"]): sealed},
            )
            orders, _ = reconcile.simple_ledger.compile_simple_ledger(
                normalized,
                events,
                plan["events_fingerprint"],
                plan,
            )
            summary = build_workbook.order_rows(orders["groups"][0]["orders"][0])[0]
            self.assertEqual(summary[core.HEADERS.index("汇率")], 5)

    def test_unknown_pricing_seals_without_guessing_values(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, normalized = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["pricing"] = {
                "source_messages": ["S00001", "S00002"],
                "expected": {"kind": "unknown", "reason": "not_stated"},
            }
            core.atomic_json(decision_path, decision)

            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            sealed = reconcile._load_json(decision_path)
            events, plan = reconcile._compile_decisions_v3(
                normalized,
                {str(normalized["groups"][0]["group_key"]): sealed},
            )
            compiled = plan["groups"][0]["orders"][0]
            self.assertEqual(compiled["pricing"]["expected"]["kind"], "unknown")
            orders, _ = reconcile.simple_ledger.compile_simple_ledger(
                normalized,
                events,
                plan["events_fingerprint"],
                plan,
            )
            order = orders["groups"][0]["orders"][0]
            self.assertIsNone(order["expected_payout"])
            self.assertEqual(order["review_result"], "待确认")
            self.assertIn("群聊未说明", order["reconciliation"]["detail"])
            self.assertIn("群聊未说明", build_workbook.order_row_note(order))
            summary = build_workbook.order_rows(order)[0]
            self.assertEqual(summary[core.HEADERS.index("汇率")], "待确认")

    def test_explicit_pricing_rejects_unknown_reason_field(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["pricing"]["expected"]["reason"] = "not_stated"
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(ValueError, "unsupported fields: reason"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_customer_id_is_not_part_of_the_v3_contract(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["customer_id"] = "user:alice"
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(ValueError, "unsupported fields: customer_id"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_multileg_order_requires_pricing_on_every_leg(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            order = decision["orders"][0]
            order["direction"] = ""
            order.pop("pricing")
            order["legs"] = [
                {
                    "leg_id": "thb",
                    "direction": "CNY->THB",
                    "allocation_amount": "60",
                    "pricing": {
                        "source_messages": ["S00001", "S00002"],
                        "terms": {"rate": "5", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "300"},
                    },
                    "payout_entry_ids": ["M0002.1"],
                    "recovery_entry_ids": [],
                },
                {
                    "leg_id": "usd",
                    "direction": "CNY->USD",
                    "allocation_amount": "40",
                    "payout_entry_ids": [],
                    "recovery_entry_ids": [],
                },
            ]
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(ValueError, r"legs\[1\].pricing is required"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["legs"][1]["pricing"] = {
                "source_messages": ["S00001", "S00002"],
                "expected": {"kind": "unknown", "reason": "not_stated"},
            }
            core.atomic_json(decision_path, decision)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            sealed = reconcile._load_json(decision_path)
            normalized = reconcile._load_snapshot(work, reconcile._load_run(work))
            _, plan = reconcile._compile_decisions_v3(
                normalized,
                {str(normalized["groups"][0]["group_key"]): sealed},
            )
            for leg in plan["groups"][0]["orders"][0]["legs"]:
                self.assertIn("pricing", leg)

    def test_multileg_pricing_compiles_each_authority_independently(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, normalized = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0002"]["entries"][0]["amount"] = "300"
            decision["media_decisions"]["M0002"]["entries"].append(
                {
                    "amount": "80",
                    "currency": "USD",
                    "payee": "Account-7788",
                    "payee_state": "visible",
                    "side": "payout",
                    "result": "completed",
                    "amount_state": "clear",
                }
            )
            order = decision["orders"][0]
            order["entry_ids"] = ["M0001.1", "M0002.1", "M0002.2"]
            order["direction"] = ""
            order.pop("pricing")
            order["legs"] = [
                {
                    "leg_id": "thb",
                    "direction": "CNY->THB",
                    "allocation_amount": "60",
                    "pricing": {
                        "source_messages": ["S00001", "S00002"],
                        "terms": {"rate": "5", "operator": "multiply"},
                        "expected": {"kind": "calculated_from_terms"},
                    },
                    "payout_entry_ids": ["M0002.1"],
                    "recovery_entry_ids": [],
                },
                {
                    "leg_id": "usd",
                    "direction": "CNY->USD",
                    "allocation_amount": "40",
                    "pricing": {
                        "source_messages": ["S00001", "S00002"],
                        "terms": {"rate": "2", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "80"},
                    },
                    "payout_entry_ids": ["M0002.2"],
                    "recovery_entry_ids": [],
                },
            ]
            core.atomic_json(decision_path, decision)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            sealed = reconcile._load_json(decision_path)
            events, plan = reconcile._compile_decisions_v3(
                normalized,
                {str(normalized["groups"][0]["group_key"]): sealed},
            )
            orders, statistics = reconcile.simple_ledger.compile_simple_ledger(
                normalized,
                events,
                plan["events_fingerprint"],
                plan,
            )
            order = orders["groups"][0]["orders"][0]
            self.assertEqual(order["review_result"], "")
            self.assertEqual(order["order_status"], "completed")
            self.assertEqual(statistics["pending_orders"], 0)
            self.assertEqual(
                [leg["expected_payout"] for leg in order["legs"]],
                ["300", "80"],
            )
            self.assertTrue(
                all(leg["reconciliation"]["status"] == "matched" for leg in order["legs"])
            )

    def test_same_direction_legs_require_distinct_display_labels(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, normalized = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0002"]["entries"][0]["amount"] = "300"
            decision["media_decisions"]["M0002"]["entries"].append(
                {
                    "amount": "200",
                    "currency": "THB",
                    "payee": "现金",
                    "payee_state": "cash",
                    "kind": "cash",
                    "side": "payout",
                    "result": "completed",
                    "amount_state": "clear",
                }
            )
            order = decision["orders"][0]
            order["entry_ids"] = ["M0001.1", "M0002.1", "M0002.2"]
            order["direction"] = ""
            order.pop("pricing")
            order["legs"] = [
                {
                    "leg_id": "thb_transfer",
                    "direction": "CNY->THB",
                    "allocation_amount": "60",
                    "pricing": {
                        "source_messages": ["S00001", "S00002"],
                        "terms": {"rate": "5", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "300"},
                    },
                    "payout_entry_ids": ["M0002.1"],
                    "recovery_entry_ids": [],
                },
                {
                    "leg_id": "thb_cash",
                    "direction": "CNY->THB",
                    "allocation_amount": "40",
                    "pricing": {
                        "source_messages": ["S00001", "S00002"],
                        "terms": {"rate": "5", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "200"},
                    },
                    "payout_entry_ids": ["M0002.2"],
                    "recovery_entry_ids": [],
                },
            ]
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(ValueError, "require display_label on every leg"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["legs"][0]["display_label"] = "CNY->THB（转账）"
            decision["orders"][0]["legs"][1]["display_label"] = "CNY->THB（现金）"
            core.atomic_json(decision_path, decision)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            sealed = reconcile._load_json(decision_path)
            events, plan = reconcile._compile_decisions_v3(
                normalized,
                {str(normalized["groups"][0]["group_key"]): sealed},
            )
            orders, _ = reconcile.simple_ledger.compile_simple_ledger(
                normalized,
                events,
                plan["events_fingerprint"],
                plan,
            )
            compiled = orders["groups"][0]["orders"][0]
            self.assertEqual(compiled["direction"], "CNY->THB（转账）\nCNY->THB（现金）")
            self.assertEqual(compiled["actual_rate_display"], "CNY->THB（转账）：5\nCNY->THB（现金）：5")
            self.assertEqual(
                [leg["display_label"] for leg in compiled["legs"]],
                ["CNY->THB（转账）", "CNY->THB（现金）"],
            )
            summary = build_workbook.order_rows(compiled)[0]
            self.assertEqual(summary[core.HEADERS.index("换汇方向")], compiled["direction"])
            self.assertEqual(summary[core.HEADERS.index("汇率")], compiled["actual_rate_display"])

    def test_previous_decision_contract_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["contract_version"] = "group-chat-decision/2.1"
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(ValueError, "unsupported decision contract"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_order_note_survives_decision_compile(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, normalized = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["note"] = "特殊安排写入订单汇总备注。"
            core.atomic_json(decision_path, decision)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            sealed = reconcile._load_json(decision_path)
            _, plan = reconcile._compile_decisions_v3(
                normalized,
                {str(normalized["groups"][0]["group_key"]): sealed},
            )
            self.assertEqual(
                plan["groups"][0]["orders"][0]["note"],
                "特殊安排写入订单汇总备注。",
            )

    def test_one_physical_payout_can_be_allocated_across_currency_orders(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, _, normalized = self._start_fixture(root)
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["media_decisions"]["M0001"]["entries"][0]["amount"] = "2000"
            decision["media_decisions"]["M0001"]["entries"].append(
                {
                    "amount": "630",
                    "currency": "USDT",
                    "payee": "TPayeeWalletAddress123456789",
                    "payee_state": "visible",
                    "side": "payment",
                    "result": "completed",
                    "amount_state": "clear",
                }
            )
            decision["media_decisions"]["M0002"]["entries"][0]["amount"] = "30000"
            decision["orders"] = [
                {
                    "id": "O001",
                    "entry_ids": ["M0001.1"],
                    "source_messages": ["S00001", "S00002"],
                    "customer_nickname": "Alice",
                    "direction": "CNY->THB",
                    "pricing": {
                        "source_messages": ["S00001", "S00002"],
                        "terms": {"rate": "4.89", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "9780"},
                    },
                },
                {
                    "id": "O002",
                    "entry_ids": ["M0001.2"],
                    "source_messages": ["S00001", "S00002"],
                    "customer_nickname": "Alice",
                    "direction": "USDT->THB",
                    "pricing": {
                        "source_messages": ["S00001", "S00002"],
                        "terms": {"rate": "32.1", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "20220"},
                    },
                },
            ]
            decision["settlement_allocations"] = [
                {
                    "entry_id": "M0002.1",
                    "allocations": [
                        {"order_id": "O001", "amount": "9780"},
                        {"order_id": "O002", "amount": "20220"},
                    ],
                    "source_messages": ["S00002"],
                }
            ]
            core.atomic_json(decision_path, decision)

            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="测试小额群")
            )
            sealed = reconcile._load_json(decision_path)
            events, plan = reconcile._compile_decisions_v3(
                normalized,
                {str(normalized["groups"][0]["group_key"]): sealed},
            )
            orders, statistics = reconcile.simple_ledger.compile_simple_ledger(
                normalized,
                events,
                plan["events_fingerprint"],
                plan,
            )

            first, second = orders["groups"][0]["orders"]
            self.assertEqual(first["actual_payout_total"], "9780")
            self.assertEqual(second["actual_payout_total"], "20220")
            self.assertEqual(first["actual_rate"], "4.89")
            self.assertEqual(second["actual_rate"], "32.1")
            self.assertEqual(first["review_result"], "")
            self.assertEqual(second["review_result"], "")
            self.assertEqual(statistics["unassigned_fund_images"], 0)
            for order, allocation in ((first, "9780"), (second, "20220")):
                flow = next(
                    item for item in order["flows"] if item.get("settlement_allocation")
                )
                self.assertEqual(flow["amount"], allocation)
                self.assertEqual(flow["source_amount"], "30000")
                self.assertEqual(build_workbook.flow_row_label(flow), "内部回款分摊")
                self.assertIn(f"本单计入 {allocation} THB", build_workbook.flow_row_note(flow))

            orders_path = root / "allocated-orders.json"
            workbook_path = root / "allocated-ledger.xlsx"
            core.atomic_json(orders_path, orders)
            workbook = build_workbook.build_workbook(
                build_workbook.validate_orders(orders),
                Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
            )
            try:
                workbook.save(workbook_path)
            finally:
                workbook.close()
            self.assertEqual(check_workbook.check(workbook_path, orders_path), [])

    def test_settlement_allocation_must_close_to_physical_payout(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, _, _ = self._start_fixture(Path(temporary))
            decision_path = self._write_complete_decision(work)
            decision = reconcile._load_json(decision_path)
            decision["orders"][0]["entry_ids"] = ["M0001.1"]
            decision["orders"].append(
                {
                    "id": "O002",
                    "entry_ids": ["M0001.1"],
                    "source_messages": ["S00001", "S00002"],
                    "customer_nickname": "Alice",
                    "direction": "CNY->THB",
                    "pricing": {
                        "source_messages": ["S00001", "S00002"],
                        "terms": {"rate": "1.99", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "199"},
                    },
                }
            )
            # Give the second order its own payment entry so only the allocation
            # closure, rather than duplicate assignment, is under test.
            decision["media_decisions"]["M0001"]["entries"].append(
                {
                    "amount": "100",
                    "currency": "CNY",
                    "payee": "张三",
                    "payee_state": "visible",
                    "side": "payment",
                    "result": "completed",
                    "amount_state": "clear",
                }
            )
            decision["orders"][1]["entry_ids"] = ["M0001.2"]
            decision["settlement_allocations"] = [
                {
                    "entry_id": "M0002.1",
                    "allocations": [
                        {"order_id": "O001", "amount": "300"},
                        {"order_id": "O002", "amount": "199"},
                    ],
                }
            ]
            core.atomic_json(decision_path, decision)

            with self.assertRaisesRegex(ValueError, "must equal source amount 500"):
                reconcile.review_command(
                    Namespace(work=work, review_action="seal", group="测试小额群")
                )

    def test_changed_snapshot_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work, run, _ = self._start_fixture(Path(temporary))
            snapshot = reconcile._snapshot_path(work)
            snapshot.write_text(snapshot.read_text(encoding="utf-8") + " ", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "snapshot changed"):
                reconcile._load_snapshot(work, run)


if __name__ == "__main__":
    unittest.main()
