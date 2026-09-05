#!/usr/bin/env python3
"""Validate and publish signed currency movements from internal store groups."""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU

import core


GROUP_MODE = "store-ledger"
GROUP_NAME_MARKER = "门店开票群"
DECISION_CONTRACT = "group-chat-store-ledger-decision/1.0"
OUTPUT_CONTRACT = "group-chat-store-ledger/1.0"
AMOUNT_POLICY = "store-signed-currency-flow/1.0"
ACCOUNTING_MODE = "store_signed_currency_flow"

MEDIA_CLASSIFICATIONS = frozenset({"voucher", "expense", "balance", "reference"})
RECORD_TYPES = frozenset(
    {"exchange", "cash_movement", "expense", "internal_transfer"}
)
POSTING_STATUSES = frozenset({"posted", "pending", "void"})
MOVEMENT_BASES = frozenset(
    {"ticket", "direct_reply_correction", "expense_semantics", "chat_explicit"}
)
MEDIA_ROLES = frozenset({"draft", "final", "duplicate", "supporting"})
SNAPSHOT_KINDS = frozenset({"checkpoint", "closing"})
TRANSFER_DIRECTIONS = frozenset({"send", "receive"})

ID_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9._-]{0,127}")
SOURCE_LABEL_RE = re.compile(r"S(\d{5})")
MEDIA_LABEL_RE = re.compile(r"M(\d{4})")
CURRENCY_RE = re.compile(r"[A-Z][A-Z0-9]{1,9}")
FORMULA_RE = re.compile(
    r"(?P<left>\d[\d,.]*)\s*(?P<operator>[*×xX/÷])\s*"
    r"(?P<right>\d[\d,.]*)\s*=\s*(?P<result>\d[\d,.]*)",
    re.IGNORECASE,
)


def _normalize_formula_text(value: object) -> str:
    """Normalize math glyphs commonly emitted by LINE/iOS emoji keyboards."""
    return (
        core.clean_text(value)
        .replace("\ufe0f", "")
        .replace("✖", "×")
        .replace("🟰", "=")
        .replace("＝", "=")
        .replace("➗", "÷")
    )


ILLEGAL_SHEET = re.compile(r"[\\/*?:\[\]]")

DETAIL_HEADERS = [
    "入账时间",
    "票面日期",
    "票号/记录号",
    "状态",
    "流水类型",
    "币种",
    "收支",
    "符号",
    "票面/说明金额",
    "实际入账金额（正收负支）",
    "汇率",
    "折算泰铢（参考）",
    "对方门店",
    "调拨编号",
    "发送人",
    "证据消息",
    "判定依据/异常",
    "调拨匹配",
    "代表图",
]
SUMMARY_HEADERS = [
    "日期",
    "币种",
    "期初余额",
    "收入合计",
    "支出合计",
    "净变动",
    "账面期末",
    "群聊期末盘点",
    "差额",
    "核对状态",
    "盘点消息",
]

STATUS_LABELS = {"posted": "已入账", "pending": "待完成", "void": "作废"}
RECORD_TYPE_LABELS = {
    "exchange": "货币置换",
    "cash_movement": "现金收支",
    "expense": "费用",
    "internal_transfer": "门店调拨",
}
BASIS_LABELS = {
    "ticket": "票面",
    "direct_reply_correction": "聊天更正",
    "expense_semantics": "费用语义",
    "chat_explicit": "聊天明确收支",
}
TRANSFER_MATCH_LABELS = {
    "matched": "已匹配",
    "unmatched": "未匹配",
    "mismatch": "金额或方向不一致",
    "not_posted": "未入账",
}

HEADER_FILL_RGB = "1F4E78"
HEADER_FONT_RGB = "FFFFFF"
SECTION_FILL_RGB = "17365D"
GRID_RGB = "A7B9C8"
BODY_ALT_FILL_RGB = "F4F8FB"
PENDING_FILL_RGB = "FFF2CC"
VOID_FILL_RGB = "E7E6E6"
ALERT_FILL_RGB = "F4CCCC"
MATCH_FILL_RGB = "E2F0D9"


def group_name_selected(value: object) -> bool:
    return core.normalize_name(GROUP_NAME_MARKER) in core.normalize_name(value)


def _clean_list(value: object, *, field: str, allow_empty: bool = True) -> list[str]:
    core.require(isinstance(value, list), f"{field} must be a list")
    cleaned = [core.clean_text(item) for item in value]
    core.require(all(cleaned), f"{field} cannot contain blank values")
    core.require(len(cleaned) == len(set(cleaned)), f"{field} repeats a value")
    if not allow_empty:
        core.require(bool(cleaned), f"{field} must not be empty")
    return cleaned


def _decimal_text(
    value: object,
    *,
    field: str,
    positive: bool = False,
    nonnegative: bool = False,
    optional: bool = False,
) -> str:
    if value in (None, ""):
        core.require(optional, f"{field} is required")
        return ""
    core.require(not isinstance(value, bool), f"{field} must be numeric")
    cleaned = str(value).strip().replace(",", "")
    try:
        number = Decimal(cleaned)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field} must be numeric") from exc
    core.require(number.is_finite(), f"{field} must be finite")
    if positive:
        core.require(number > 0, f"{field} must be greater than zero")
    if nonnegative:
        core.require(number >= 0, f"{field} must not be negative")
    normalized = format(number, "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return "0" if normalized in {"", "-0"} else normalized


def _currency(value: object, *, field: str) -> str:
    cleaned = core.clean_text(value).upper()
    core.require(CURRENCY_RE.fullmatch(cleaned) is not None, f"{field} is invalid")
    return cleaned


def _date(value: object, *, field: str, optional: bool = False) -> str:
    cleaned = core.clean_text(value)
    if not cleaned:
        core.require(optional, f"{field} is required")
        return ""
    try:
        parsed = datetime.strptime(cleaned, "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError(f"{field} must use YYYY-MM-DD") from exc
    core.require(parsed.isoformat() == cleaned, f"{field} must use YYYY-MM-DD")
    return cleaned


def _timestamp(value: object, *, field: str, optional: bool = False) -> str:
    cleaned = core.clean_text(value)
    if not cleaned:
        core.require(optional, f"{field} is required")
        return ""
    try:
        parsed = datetime.fromisoformat(cleaned.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError(f"{field} must be an ISO timestamp with timezone") from exc
    core.require(parsed.tzinfo is not None, f"{field} must include a timezone")
    return parsed.isoformat()


def _message_labels(group: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    return {
        f"S{position:05d}": message
        for position, message in enumerate(group.get("messages", []), start=1)
        if isinstance(message, Mapping)
    }


def media_inventory(
    group: Mapping[str, Any],
) -> dict[str, tuple[str, Mapping[str, Any], Mapping[str, Any]]]:
    inventory: dict[str, tuple[str, Mapping[str, Any], Mapping[str, Any]]] = {}
    position = 0
    for message_position, message in enumerate(group.get("messages", []), start=1):
        if not isinstance(message, Mapping):
            continue
        for media in message.get("media", []):
            if not isinstance(media, Mapping) or not core.media_is_evidence(media):
                continue
            position += 1
            inventory[f"M{position:04d}"] = (
                f"S{message_position:05d}",
                message,
                media,
            )
    return inventory


def _validate_source_messages(
    value: object,
    *,
    field: str,
    message_by_label: Mapping[str, Mapping[str, Any]],
    reviewed_through: int,
    allow_empty: bool = False,
) -> list[str]:
    labels = _clean_list(value, field=field, allow_empty=allow_empty)
    core.require(
        set(labels) <= set(message_by_label),
        f"{field} contains an unknown message label",
    )
    for label in labels:
        match = SOURCE_LABEL_RE.fullmatch(label)
        core.require(match is not None, f"{field} contains an invalid message label")
        core.require(
            int(match.group(1)) <= reviewed_through,
            f"{field} cites a message that has not been reviewed",
        )
    return labels


def _normalize_visible_movement(value: object, *, field: str) -> dict[str, str]:
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    allowed = {"currency", "sign", "amount", "rate", "thb_equivalent"}
    unknown = sorted(set(value) - allowed)
    core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
    sign = core.clean_text(value.get("sign"))
    core.require(sign in {"+", "-"}, f"{field}.sign must be + or -")
    result = {
        "currency": _currency(value.get("currency"), field=f"{field}.currency"),
        "sign": sign,
        "amount": _decimal_text(value.get("amount"), field=f"{field}.amount", positive=True),
    }
    for name in ("rate", "thb_equivalent"):
        normalized = _decimal_text(
            value.get(name), field=f"{field}.{name}", positive=True, optional=True
        )
        if normalized:
            result[name] = normalized
    return result


def _normalize_visible_balance(value: object, *, field: str) -> dict[str, str]:
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    unknown = sorted(set(value) - {"currency", "amount"})
    core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
    return {
        "currency": _currency(value.get("currency"), field=f"{field}.currency"),
        "amount": _decimal_text(
            value.get("amount"), field=f"{field}.amount", nonnegative=True
        ),
    }


def normalize_visible_facts(
    value: object,
    *,
    classification: str,
    field: str,
) -> dict[str, Any]:
    if value is None:
        return {}
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    allowed_by_class = {
        "voucher": {"voucher_number", "voucher_date", "description", "movements"},
        "expense": {"voucher_number", "voucher_date", "description", "movements"},
        "balance": {"description", "balances"},
        "reference": set(),
    }
    allowed = allowed_by_class[classification]
    unknown = sorted(set(value) - allowed)
    core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
    result: dict[str, Any] = {}
    if "voucher_number" in value:
        voucher_number = core.clean_text(value.get("voucher_number"))
        if voucher_number:
            result["voucher_number"] = voucher_number
    if "voucher_date" in value:
        voucher_date = _date(
            value.get("voucher_date"), field=f"{field}.voucher_date", optional=True
        )
        if voucher_date:
            result["voucher_date"] = voucher_date
    description = core.clean_text(value.get("description"))
    if description:
        result["description"] = description
    if "movements" in value:
        raw_movements = value.get("movements")
        core.require(isinstance(raw_movements, list), f"{field}.movements must be a list")
        result["movements"] = [
            _normalize_visible_movement(item, field=f"{field}.movements[{position}]")
            for position, item in enumerate(raw_movements)
        ]
    if "balances" in value:
        raw_balances = value.get("balances")
        core.require(isinstance(raw_balances, list), f"{field}.balances must be a list")
        result["balances"] = [
            _normalize_visible_balance(item, field=f"{field}.balances[{position}]")
            for position, item in enumerate(raw_balances)
        ]
        currencies = [item["currency"] for item in result["balances"]]
        core.require(
            len(currencies) == len(set(currencies)),
            f"{field}.balances repeats a currency",
        )
    return result


def _validate_media_decisions(
    group: Mapping[str, Any],
    decision: dict[str, Any],
    *,
    require_complete: bool,
    capture_hashes: bool,
    rehash_labels: set[str] | None,
) -> tuple[dict[str, Any], dict[str, set[str]]]:
    inventory = media_inventory(group)
    available = {
        label
        for label, (_, _, media) in inventory.items()
        if media.get("availability") == "available"
    }
    missing = set(inventory) - available
    media_decisions = decision.get("media_decisions")
    core.require(isinstance(media_decisions, dict), "media_decisions must be an object")
    unknown_labels = sorted(set(media_decisions) - available)
    core.require(
        not unknown_labels,
        f"store media decisions contain missing or unknown labels: {unknown_labels[:20]}",
    )
    if require_complete:
        unclassified = sorted(available - set(media_decisions))
        core.require(not unclassified, f"unclassified store media: {unclassified[:20]}")

    by_class = {classification: set() for classification in MEDIA_CLASSIFICATIONS}
    computed = 0
    reused = 0
    for label, raw in media_decisions.items():
        field = f"{group.get('group_key')}.media_decisions.{label}"
        core.require(isinstance(raw, dict), f"{field} must be an object")
        classification = core.clean_text(raw.get("classification")).casefold()
        core.require(
            classification in MEDIA_CLASSIFICATIONS,
            f"{field}.classification is unsupported",
        )
        raw["classification"] = classification
        by_class[classification].add(label)
        if classification == "reference":
            unknown = sorted(set(raw) - {"classification", "note"})
            core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
            continue
        unknown = sorted(
            set(raw)
            - {"classification", "viewed_original", "facts", "evidence_sha256", "note"}
        )
        core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
        core.require(
            raw.get("viewed_original") is True,
            f"{field}: open the original image before recording store facts",
        )
        raw["facts"] = normalize_visible_facts(
            raw.get("facts"), classification=classification, field=f"{field}.facts"
        )
        path = Path(str(inventory[label][2].get("path") or ""))
        core.require(path.is_file(), f"{field}: original media is unavailable: {path}")
        recorded_hash = core.clean_text(raw.get("evidence_sha256"))
        verify_hash = (
            rehash_labels is None
            or label in rehash_labels
            or (capture_hashes and not recorded_hash)
        )
        resolved_hash, was_computed = core.validate_cached_file_hash(
            path,
            recorded_hash,
            verify=verify_hash,
            changed_message=f"{field}: original store image changed after review",
        )
        snapshot_hash = core.clean_text(inventory[label][2].get("blob_sha256"))
        if was_computed and snapshot_hash:
            core.require(
                resolved_hash == snapshot_hash,
                f"{field}: original store image changed after the run snapshot was created",
            )
        computed += int(was_computed)
        reused += int(not was_computed and bool(recorded_hash))
        if not recorded_hash and capture_hashes:
            raw["evidence_sha256"] = resolved_hash
        elif require_complete:
            core.require(bool(recorded_hash), f"{field}: evidence hash has not been captured")
    return (
        {
            "available_media": len(available),
            "missing_media": len(missing),
            "classified_media": len(media_decisions),
            "voucher_media": len(by_class["voucher"]),
            "expense_media": len(by_class["expense"]),
            "balance_media": len(by_class["balance"]),
            "reference_media": len(by_class["reference"]),
            "evidence_hashes_computed": computed,
            "evidence_hashes_reused": reused,
        },
        by_class,
    )


def _validate_open_records(
    group: Mapping[str, Any],
    decision: dict[str, Any],
    *,
    reviewed_through: int,
    require_complete: bool,
) -> int:
    message_by_label = _message_labels(group)
    inventory = media_inventory(group)
    value = decision.get("open_records")
    core.require(isinstance(value, list), "open_records must be a list")
    ids: set[str] = set()
    for position, item in enumerate(value):
        field = f"{group.get('group_key')}.open_records[{position}]"
        core.require(isinstance(item, dict), f"{field} must be an object")
        unknown = sorted(
            set(item) - {"id", "source_messages", "media_labels", "summary", "unresolved"}
        )
        core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
        record_id = core.clean_text(item.get("id"))
        core.require(ID_RE.fullmatch(record_id) is not None, f"{field}.id is invalid")
        core.require(record_id not in ids, f"{field}.id is duplicated")
        ids.add(record_id)
        item["id"] = record_id
        item["source_messages"] = _validate_source_messages(
            item.get("source_messages"),
            field=f"{field}.source_messages",
            message_by_label=message_by_label,
            reviewed_through=reviewed_through,
        )
        media_labels = _clean_list(
            item.get("media_labels", []), field=f"{field}.media_labels"
        )
        core.require(
            set(media_labels) <= set(inventory),
            f"{field}.media_labels contains an unknown media label",
        )
        item["media_labels"] = media_labels
        summary = core.clean_text(item.get("summary"))
        core.require(bool(summary), f"{field}.summary is required")
        item["summary"] = summary
        unresolved = _clean_list(
            item.get("unresolved", []), field=f"{field}.unresolved"
        )
        item["unresolved"] = unresolved
    if decision.get("read_complete") or require_complete:
        core.require(not value, "open_records must be empty after the group is fully read")
    return len(value)


def _direct_reply_formula_exists(
    group: Mapping[str, Any],
    source_labels: Iterable[str],
    record_media_labels: set[str],
    corrected_amount: str,
) -> bool:
    message_by_label = _message_labels(group)
    label_by_id = {
        str(message.get("message_id") or ""): label
        for label, message in message_by_label.items()
    }
    inventory = media_inventory(group)
    media_by_source: dict[str, set[str]] = defaultdict(set)
    for label, (source_label, _, _) in inventory.items():
        media_by_source[source_label].add(label)
    for label in source_labels:
        message = message_by_label.get(label)
        if message is None:
            continue
        match = FORMULA_RE.search(_normalize_formula_text(message.get("text")))
        if match is None:
            continue
        formula_values = {
            _decimal_text(match.group(name), field="formula value")
            for name in ("left", "right", "result")
        }
        if corrected_amount not in formula_values:
            continue
        target_label = label_by_id.get(str(message.get("reply_to_message_id") or ""))
        if target_label and media_by_source[target_label] & record_media_labels:
            return True
    return False


def _validate_movement(
    value: object,
    *,
    field: str,
    group: Mapping[str, Any],
    message_by_label: Mapping[str, Mapping[str, Any]],
    reviewed_through: int,
    record_type: str,
    record_sources: set[str],
    record_media_labels: set[str],
    visible_movements: list[Mapping[str, Any]],
) -> dict[str, Any]:
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    allowed = {
        "currency",
        "sign",
        "amount",
        "rate",
        "thb_equivalent",
        "basis",
        "source_messages",
        "note",
    }
    unknown = sorted(set(value) - allowed)
    core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
    sign = core.clean_text(value.get("sign"))
    core.require(sign in {"+", "-"}, f"{field}.sign must be + or -")
    basis = core.clean_text(value.get("basis")).casefold()
    core.require(basis in MOVEMENT_BASES, f"{field}.basis is unsupported")
    sources = _validate_source_messages(
        value.get("source_messages"),
        field=f"{field}.source_messages",
        message_by_label=message_by_label,
        reviewed_through=reviewed_through,
    )
    core.require(
        set(sources) <= record_sources,
        f"{field}.source_messages must belong to the record source_messages",
    )
    currency = _currency(value.get("currency"), field=f"{field}.currency")
    amount = _decimal_text(value.get("amount"), field=f"{field}.amount", positive=True)
    if basis == "ticket":
        core.require(bool(record_media_labels), f"{field}: ticket basis requires ticket media")
        core.require(
            any(
                item.get("currency") == currency
                and item.get("sign") == sign
                and item.get("amount") == amount
                for item in visible_movements
            ),
            f"{field}: ticket movement must match a movement visible in this record's media",
        )
    if basis == "expense_semantics":
        core.require(record_type == "expense", f"{field}: expense_semantics requires expense")
    note = core.clean_text(value.get("note"))
    if basis == "direct_reply_correction":
        core.require(
            _direct_reply_formula_exists(
                group,
                sources,
                record_media_labels,
                amount,
            ),
            f"{field}: direct_reply_correction requires a direct image reply with a complete formula",
        )
        core.require(bool(note), f"{field}.note must explain the chat correction")
    result = {
        "currency": currency,
        "sign": sign,
        "amount": amount,
        "basis": basis,
        "source_messages": sources,
    }
    for name in ("rate", "thb_equivalent"):
        normalized = _decimal_text(
            value.get(name), field=f"{field}.{name}", positive=True, optional=True
        )
        if normalized:
            result[name] = normalized
    if note:
        result["note"] = note
    return result


def _validate_record(
    value: object,
    *,
    field: str,
    group: Mapping[str, Any],
    reviewed_through: int,
    by_class: Mapping[str, set[str]],
    media_decisions: Mapping[str, Any],
) -> dict[str, Any]:
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    allowed = {
        "id",
        "record_type",
        "posting_status",
        "posting_at",
        "voucher_number",
        "voucher_date",
        "source_messages",
        "media_roles",
        "representative_media_label",
        "movements",
        "transfer",
        "status_reason",
        "note",
    }
    unknown = sorted(set(value) - allowed)
    core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
    record_id = core.clean_text(value.get("id"))
    core.require(ID_RE.fullmatch(record_id) is not None, f"{field}.id is invalid")
    record_type = core.clean_text(value.get("record_type")).casefold()
    core.require(record_type in RECORD_TYPES, f"{field}.record_type is unsupported")
    posting_status = core.clean_text(value.get("posting_status")).casefold()
    core.require(posting_status in POSTING_STATUSES, f"{field}.posting_status is unsupported")
    posting_at = _timestamp(
        value.get("posting_at"), field=f"{field}.posting_at", optional=posting_status != "posted"
    )
    if posting_status != "posted":
        core.require(not posting_at, f"{field}.posting_at must be blank until actually posted")
    voucher_number = core.clean_text(value.get("voucher_number"))
    voucher_date = _date(
        value.get("voucher_date"), field=f"{field}.voucher_date", optional=True
    )
    message_by_label = _message_labels(group)
    sources = _validate_source_messages(
        value.get("source_messages"),
        field=f"{field}.source_messages",
        message_by_label=message_by_label,
        reviewed_through=reviewed_through,
    )

    raw_media_roles = value.get("media_roles", [])
    core.require(isinstance(raw_media_roles, list), f"{field}.media_roles must be a list")
    media_roles: list[dict[str, str]] = []
    media_labels: set[str] = set()
    for position, raw_role in enumerate(raw_media_roles):
        role_field = f"{field}.media_roles[{position}]"
        core.require(isinstance(raw_role, Mapping), f"{role_field} must be an object")
        unknown_role = sorted(set(raw_role) - {"label", "role"})
        core.require(not unknown_role, f"{role_field} has unsupported fields")
        label = core.clean_text(raw_role.get("label"))
        core.require(MEDIA_LABEL_RE.fullmatch(label) is not None, f"{role_field}.label is invalid")
        core.require(label not in media_labels, f"{field}.media_roles repeats {label}")
        role = core.clean_text(raw_role.get("role")).casefold()
        core.require(role in MEDIA_ROLES, f"{role_field}.role is unsupported")
        allowed_classes = {"expense"} if record_type == "expense" else {"voucher"}
        core.require(
            any(label in by_class[item] for item in allowed_classes),
            f"{role_field}.label has the wrong media classification",
        )
        media_labels.add(label)
        media_roles.append({"label": label, "role": role})

    representative = core.clean_text(value.get("representative_media_label"))
    if media_labels:
        core.require(
            representative in media_labels,
            f"{field}.representative_media_label must select this record's media",
        )
        role_by_label = {item["label"]: item["role"] for item in media_roles}
        core.require(
            role_by_label[representative] != "duplicate",
            f"{field}.representative_media_label cannot select duplicate media",
        )
    else:
        core.require(not representative, f"{field}.representative_media_label requires media")

    visible_movements = [
        movement
        for label in media_labels
        for movement in media_decisions[label].get("facts", {}).get("movements", [])
        if isinstance(movement, Mapping)
    ]

    raw_movements = value.get("movements")
    core.require(isinstance(raw_movements, list) and raw_movements, f"{field}.movements is required")
    movements = [
        _validate_movement(
            item,
            field=f"{field}.movements[{position}]",
            group=group,
            message_by_label=message_by_label,
            reviewed_through=reviewed_through,
            record_type=record_type,
            record_sources=set(sources),
            record_media_labels=media_labels,
            visible_movements=visible_movements,
        )
        for position, item in enumerate(raw_movements)
    ]

    transfer_value = value.get("transfer")
    transfer: dict[str, str] | None = None
    if record_type == "internal_transfer":
        core.require(isinstance(transfer_value, Mapping), f"{field}.transfer is required")
        allowed_transfer = {
            "transfer_id",
            "counterparty_group_name",
            "counterparty_group_key",
            "direction",
        }
        unknown_transfer = sorted(set(transfer_value) - allowed_transfer)
        core.require(
            not unknown_transfer,
            f"{field}.transfer has unsupported fields: {', '.join(unknown_transfer)}",
        )
        transfer_id = core.clean_text(transfer_value.get("transfer_id"))
        core.require(ID_RE.fullmatch(transfer_id) is not None, f"{field}.transfer.transfer_id is invalid")
        counterparty_name = core.clean_text(transfer_value.get("counterparty_group_name"))
        core.require(bool(counterparty_name), f"{field}.transfer.counterparty_group_name is required")
        direction = core.clean_text(transfer_value.get("direction")).casefold()
        core.require(direction in TRANSFER_DIRECTIONS, f"{field}.transfer.direction is unsupported")
        expected_sign = "-" if direction == "send" else "+"
        core.require(
            all(item["sign"] == expected_sign for item in movements),
            f"{field}: transfer direction disagrees with movement signs",
        )
        transfer = {
            "transfer_id": transfer_id,
            "counterparty_group_name": counterparty_name,
            "direction": direction,
        }
        counterparty_key = core.clean_text(transfer_value.get("counterparty_group_key"))
        if counterparty_key:
            transfer["counterparty_group_key"] = counterparty_key
    else:
        core.require(transfer_value in (None, {}), f"{field}.transfer is only for internal_transfer")

    status_reason = core.clean_text(value.get("status_reason"))
    if posting_status != "posted":
        core.require(bool(status_reason), f"{field}.status_reason is required when not posted")
    note = core.clean_text(value.get("note"))
    result: dict[str, Any] = {
        "id": record_id,
        "record_type": record_type,
        "posting_status": posting_status,
        "posting_at": posting_at,
        "voucher_number": voucher_number,
        "voucher_date": voucher_date,
        "source_messages": sources,
        "media_roles": media_roles,
        "representative_media_label": representative,
        "movements": movements,
        "status_reason": status_reason,
        "note": note,
    }
    if transfer is not None:
        result["transfer"] = transfer
    return result


def _validate_balance_snapshot(
    value: object,
    *,
    field: str,
    group: Mapping[str, Any],
    reviewed_through: int,
    balance_media: set[str],
) -> dict[str, Any]:
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    allowed = {"id", "date", "kind", "source_messages", "media_labels", "balances", "note"}
    unknown = sorted(set(value) - allowed)
    core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
    snapshot_id = core.clean_text(value.get("id"))
    core.require(ID_RE.fullmatch(snapshot_id) is not None, f"{field}.id is invalid")
    snapshot_date = _date(value.get("date"), field=f"{field}.date")
    kind = core.clean_text(value.get("kind")).casefold()
    core.require(kind in SNAPSHOT_KINDS, f"{field}.kind is unsupported")
    message_by_label = _message_labels(group)
    sources = _validate_source_messages(
        value.get("source_messages"),
        field=f"{field}.source_messages",
        message_by_label=message_by_label,
        reviewed_through=reviewed_through,
    )
    zone = ZoneInfo("Asia/Bangkok")
    core.require(
        any(
            datetime.fromisoformat(str(message_by_label[label]["timestamp"]))
            .astimezone(zone)
            .date()
            .isoformat()
            == snapshot_date
            for label in sources
        ),
        f"{field}.date must match at least one source message date",
    )
    media_labels = _clean_list(
        value.get("media_labels", []), field=f"{field}.media_labels"
    )
    core.require(
        set(media_labels) <= balance_media,
        f"{field}.media_labels must be classified as balance",
    )
    raw_balances = value.get("balances")
    core.require(isinstance(raw_balances, list) and raw_balances, f"{field}.balances is required")
    balances = [
        _normalize_visible_balance(item, field=f"{field}.balances[{position}]")
        for position, item in enumerate(raw_balances)
    ]
    currencies = [item["currency"] for item in balances]
    core.require(len(currencies) == len(set(currencies)), f"{field}.balances repeats a currency")
    return {
        "id": snapshot_id,
        "date": snapshot_date,
        "kind": kind,
        "source_messages": sources,
        "media_labels": media_labels,
        "balances": balances,
        "note": core.clean_text(value.get("note")),
    }


def validate_decision(
    normalized: Mapping[str, Any],
    group: Mapping[str, Any],
    decision: dict[str, Any],
    expected: Mapping[str, Any],
    *,
    require_complete: bool,
    capture_hashes: bool,
    rehash_labels: set[str] | None = None,
) -> dict[str, Any]:
    core.require(
        decision.get("contract_version") == DECISION_CONTRACT,
        f"unsupported store-ledger decision contract; expected {DECISION_CONTRACT}",
    )
    unknown_top = sorted(set(decision) - set(expected))
    core.require(not unknown_top, "store decision has unsupported fields: " + ", ".join(unknown_top))
    for field in (
        "normalized_source_fingerprint",
        "group_fingerprint",
        "group_key",
        "group_name",
        "platform",
        "message_count",
        "evidence_media_count",
        "group_mode",
    ):
        core.require(
            decision.get(field) == expected.get(field),
            f"{group.get('group_key')}: generated {field} was edited",
        )
    messages = list(group.get("messages", []))
    reviewed_through = decision.get("reviewed_through")
    core.require(
        isinstance(reviewed_through, int) and 0 <= reviewed_through <= len(messages),
        "invalid reviewed_through",
    )
    core.require(isinstance(decision.get("read_complete"), bool), "read_complete must be boolean")
    core.require(
        decision.get("read_complete") is (reviewed_through == len(messages)),
        "read_complete must exactly match reviewed_through",
    )
    if require_complete:
        core.require(decision.get("read_complete") is True, "store group chronology is incomplete")

    media_stats, by_class = _validate_media_decisions(
        group,
        decision,
        require_complete=require_complete,
        capture_hashes=capture_hashes,
        rehash_labels=rehash_labels,
    )
    open_count = _validate_open_records(
        group,
        decision,
        reviewed_through=reviewed_through,
        require_complete=require_complete,
    )

    records_value = decision.get("records")
    core.require(isinstance(records_value, list), "records must be a list")
    records: list[dict[str, Any]] = []
    ids: set[str] = set()
    voucher_numbers: set[str] = set()
    transfer_ids: set[str] = set()
    assigned_record_media: set[str] = set()
    pending_records = 0
    void_records = 0
    for position, raw_record in enumerate(records_value):
        record = _validate_record(
            raw_record,
            field=f"{group.get('group_key')}.records[{position}]",
            group=group,
            reviewed_through=reviewed_through,
            by_class=by_class,
            media_decisions=decision["media_decisions"],
        )
        core.require(record["id"] not in ids, f"record id is duplicated: {record['id']}")
        ids.add(record["id"])
        if record["voucher_number"]:
            canonical_number = re.sub(r"\s+", "", record["voucher_number"]).casefold()
            core.require(
                canonical_number not in voucher_numbers,
                f"voucher number must be one logical record: {record['voucher_number']}",
            )
            voucher_numbers.add(canonical_number)
        for media_role in record["media_roles"]:
            label = media_role["label"]
            core.require(label not in assigned_record_media, f"store media assigned twice: {label}")
            assigned_record_media.add(label)
        if record.get("transfer"):
            transfer_id = record["transfer"]["transfer_id"]
            core.require(transfer_id not in transfer_ids, f"transfer id repeats in one group: {transfer_id}")
            transfer_ids.add(transfer_id)
        pending_records += int(record["posting_status"] == "pending")
        void_records += int(record["posting_status"] == "void")
        records.append(record)
    accounting_window = normalized.get("accounting_window")
    if isinstance(accounting_window, Mapping):
        window_start = datetime.fromisoformat(str(accounting_window.get("start")))
        window_end = datetime.fromisoformat(str(accounting_window.get("end")))
        message_by_label = _message_labels(group)
        for record in records:
            core.require(
                any(
                    not message_by_label[label].get("accounting_context_only")
                    for label in record["source_messages"]
                ),
                f"{record['id']}: context-only messages cannot create an output record",
            )
            if record["posting_status"] != "posted":
                continue
            posting_at = datetime.fromisoformat(record["posting_at"])
            core.require(
                window_start <= posting_at.astimezone(window_start.tzinfo) < window_end,
                f"{record['id']}: posted record is outside this task's accounting window",
            )
    decision["records"] = records

    snapshots_value = decision.get("balance_snapshots")
    core.require(isinstance(snapshots_value, list), "balance_snapshots must be a list")
    snapshots: list[dict[str, Any]] = []
    snapshot_ids: set[str] = set()
    closing_dates: set[str] = set()
    assigned_balance_media: set[str] = set()
    for position, raw_snapshot in enumerate(snapshots_value):
        snapshot = _validate_balance_snapshot(
            raw_snapshot,
            field=f"{group.get('group_key')}.balance_snapshots[{position}]",
            group=group,
            reviewed_through=reviewed_through,
            balance_media=by_class["balance"],
        )
        core.require(snapshot["id"] not in snapshot_ids, f"balance snapshot id repeats: {snapshot['id']}")
        snapshot_ids.add(snapshot["id"])
        if snapshot["kind"] == "closing":
            core.require(
                snapshot["date"] not in closing_dates,
                f"only one official closing is allowed per date: {snapshot['date']}",
            )
            closing_dates.add(snapshot["date"])
        for label in snapshot["media_labels"]:
            core.require(label not in assigned_balance_media, f"balance media assigned twice: {label}")
            assigned_balance_media.add(label)
        snapshots.append(snapshot)
    if isinstance(accounting_window, Mapping):
        message_by_label = _message_labels(group)
        for snapshot in snapshots:
            core.require(
                any(
                    not message_by_label[label].get("accounting_context_only")
                    for label in snapshot["source_messages"]
                ),
                f"{snapshot['id']}: context-only messages cannot create a balance snapshot",
            )
    decision["balance_snapshots"] = snapshots

    expected_record_media = by_class["voucher"] | by_class["expense"]
    unassigned_record_media = expected_record_media - assigned_record_media
    unassigned_balance_media = by_class["balance"] - assigned_balance_media
    if require_complete:
        core.require(
            not unassigned_record_media,
            f"voucher or expense media must be assigned to a logical record: {sorted(unassigned_record_media)[:20]}",
        )
        core.require(
            not unassigned_balance_media,
            f"balance media must be assigned to a snapshot: {sorted(unassigned_balance_media)[:20]}",
        )
    return {
        **media_stats,
        "records": len(records),
        "posted_records": sum(item["posting_status"] == "posted" for item in records),
        "pending_store_records": pending_records,
        "void_records": void_records,
        "open_records": open_count,
        "balance_snapshots": len(snapshots),
        "official_closings": len(closing_dates),
        "unassigned_store_media": len(unassigned_record_media) + len(unassigned_balance_media),
    }


def _merge_by_id(
    candidate: dict[str, Any],
    batch: Mapping[str, Any],
    *,
    collection: str,
    updates_field: str,
    removals_field: str,
) -> None:
    updates = batch.get(updates_field, [])
    core.require(isinstance(updates, list), f"review batch {updates_field} must be a list")
    update_ids: list[str] = []
    for position, item in enumerate(updates):
        core.require(isinstance(item, Mapping), f"review batch {updates_field}[{position}] must be an object")
        item_id = core.clean_text(item.get("id"))
        core.require(bool(item_id), f"review batch {updates_field}[{position}].id is required")
        update_ids.append(item_id)
    core.require(len(update_ids) == len(set(update_ids)), f"review batch {updates_field} repeats an id")
    removals = _clean_list(
        batch.get(removals_field, []), field=f"review batch {removals_field}"
    )
    core.require(
        not (set(update_ids) & set(removals)),
        f"review batch cannot update and remove the same {collection} id",
    )
    current = candidate.get(collection)
    core.require(isinstance(current, list), f"decision {collection} must be a list")
    existing_ids = {
        core.clean_text(item.get("id"))
        for item in current
        if isinstance(item, Mapping)
    }
    missing = sorted(set(removals) - existing_ids)
    core.require(not missing, f"review batch removes unknown {collection} ids: {missing[:20]}")
    if removals:
        current[:] = [
            item
            for item in current
            if not isinstance(item, Mapping) or core.clean_text(item.get("id")) not in set(removals)
        ]
    positions = {
        core.clean_text(item.get("id")): position
        for position, item in enumerate(current)
        if isinstance(item, Mapping)
    }
    for item in updates:
        item_id = core.clean_text(item.get("id"))
        replacement = deepcopy(dict(item))
        if item_id in positions:
            current[positions[item_id]] = replacement
        else:
            positions[item_id] = len(current)
            current.append(replacement)


def merge_review_batch(candidate: dict[str, Any], batch: Mapping[str, Any]) -> dict[str, Any]:
    _merge_by_id(
        candidate,
        batch,
        collection="records",
        updates_field="records",
        removals_field="remove_record_ids",
    )
    _merge_by_id(
        candidate,
        batch,
        collection="balance_snapshots",
        updates_field="balance_snapshots",
        removals_field="remove_balance_snapshot_ids",
    )
    if "open_records" in batch:
        core.require(isinstance(batch["open_records"], list), "review batch open_records must be a list")
        candidate["open_records"] = deepcopy(batch["open_records"])
    return candidate


def _source_message_rows(
    group: Mapping[str, Any], labels: Iterable[str]
) -> list[dict[str, str]]:
    message_by_label = _message_labels(group)
    rows: list[dict[str, str]] = []
    for label in labels:
        message = message_by_label[str(label)]
        rows.append(
            {
                "label": str(label),
                "timestamp": core.clean_text(message.get("timestamp")),
                "sender": core.clean_text(
                    message.get("sender_name") or message.get("sender_id")
                ),
                "text": core.clean_text(message.get("text")),
            }
        )
    return rows


def _message_summary(rows: Iterable[Mapping[str, Any]]) -> str:
    parts: list[str] = []
    for row in rows:
        text = core.clean_text(row.get("text"))
        if len(text) > 160:
            text = text[:157] + "..."
        prefix = " ".join(
            item
            for item in (
                core.clean_text(row.get("label")),
                core.clean_text(row.get("timestamp")),
                core.clean_text(row.get("sender")),
            )
            if item
        )
        parts.append(f"{prefix}: {text}" if text else prefix)
    return "\n".join(parts)


def _record_media(
    record: Mapping[str, Any],
    inventory: Mapping[str, tuple[str, Mapping[str, Any], Mapping[str, Any]]],
    media_decisions: Mapping[str, Any],
) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    for item in record.get("media_roles", []):
        label = str(item["label"])
        media = inventory[label][2]
        decision = media_decisions[label]
        result.append(
            {
                "label": label,
                "role": str(item["role"]),
                "path": str(media.get("path") or ""),
                "sha256": str(decision.get("evidence_sha256") or ""),
            }
        )
    return result


def _snapshot_media(
    snapshot: Mapping[str, Any],
    inventory: Mapping[str, tuple[str, Mapping[str, Any], Mapping[str, Any]]],
    media_decisions: Mapping[str, Any],
) -> list[dict[str, str]]:
    return [
        {
            "label": label,
            "path": str(inventory[label][2].get("path") or ""),
            "sha256": str(media_decisions[label].get("evidence_sha256") or ""),
        }
        for label in snapshot.get("media_labels", [])
    ]


def _signed_totals(record: Mapping[str, Any]) -> dict[str, Decimal]:
    totals: dict[str, Decimal] = defaultdict(Decimal)
    for movement in record.get("movements", []):
        amount = Decimal(str(movement["amount"]))
        totals[str(movement["currency"])] += amount if movement["sign"] == "+" else -amount
    return dict(totals)


def _counterparty_matches(
    left_group: Mapping[str, Any],
    left_record: Mapping[str, Any],
    right_group: Mapping[str, Any],
) -> bool:
    transfer = left_record.get("transfer")
    if not isinstance(transfer, Mapping):
        return False
    expected_name = core.normalize_name(transfer.get("counterparty_group_name"))
    if expected_name != core.normalize_name(right_group.get("group_name")):
        return False
    expected_key = core.clean_text(transfer.get("counterparty_group_key"))
    return not expected_key or expected_key == core.clean_text(right_group.get("group_key"))


def _annotate_transfer_matches(groups: list[dict[str, Any]]) -> dict[str, int]:
    by_transfer: dict[str, list[tuple[dict[str, Any], dict[str, Any]]]] = defaultdict(list)
    for group in groups:
        for record in group.get("records", []):
            transfer = record.get("transfer")
            if isinstance(transfer, Mapping):
                by_transfer[str(transfer["transfer_id"])].append((group, record))
    statistics = {"matched_transfers": 0, "unmatched_transfers": 0, "mismatched_transfers": 0}
    for items in by_transfer.values():
        if any(record.get("posting_status") != "posted" for _, record in items):
            for _, record in items:
                record["transfer"]["match_status"] = "not_posted"
            continue
        status = "unmatched"
        if len(items) == 2:
            (left_group, left), (right_group, right) = items
            directions = {left["transfer"]["direction"], right["transfer"]["direction"]}
            counterpart_ok = _counterparty_matches(left_group, left, right_group) and _counterparty_matches(
                right_group, right, left_group
            )
            totals: dict[str, Decimal] = defaultdict(Decimal)
            for record in (left, right):
                for currency, amount in _signed_totals(record).items():
                    totals[currency] += amount
            if (
                left_group.get("group_key") != right_group.get("group_key")
                and directions == TRANSFER_DIRECTIONS
                and counterpart_ok
                and totals
                and all(amount == 0 for amount in totals.values())
            ):
                status = "matched"
            else:
                status = "mismatch"
        elif len(items) > 2:
            status = "mismatch"
        for _, record in items:
            record["transfer"]["match_status"] = status
        if status == "matched":
            statistics["matched_transfers"] += 1
        elif status == "unmatched":
            statistics["unmatched_transfers"] += 1
        else:
            statistics["mismatched_transfers"] += 1
    return statistics


def _local_date(timestamp: str) -> str:
    return (
        datetime.fromisoformat(timestamp)
        .astimezone(ZoneInfo("Asia/Bangkok"))
        .date()
        .isoformat()
    )


def _daily_reconciliation(group: Mapping[str, Any]) -> list[dict[str, Any]]:
    flow_by_date: dict[str, dict[str, dict[str, Decimal]]] = defaultdict(
        lambda: defaultdict(lambda: {"income": Decimal(0), "expense": Decimal(0)})
    )
    for record in group.get("records", []):
        if record.get("posting_status") != "posted":
            continue
        date = _local_date(str(record["posting_at"]))
        for movement in record.get("movements", []):
            currency = str(movement["currency"])
            amount = Decimal(str(movement["amount"]))
            bucket = "income" if movement["sign"] == "+" else "expense"
            flow_by_date[date][currency][bucket] += amount

    closings: dict[str, Mapping[str, Any]] = {}
    snapshots_by_date: dict[str, list[Mapping[str, Any]]] = defaultdict(list)
    for snapshot in group.get("balance_snapshots", []):
        snapshots_by_date[str(snapshot["date"])].append(snapshot)
        if snapshot.get("kind") == "closing":
            closings[str(snapshot["date"])] = snapshot
    dates = sorted(set(flow_by_date) | set(snapshots_by_date))
    rows: list[dict[str, Any]] = []
    prior_date: str | None = None
    prior_closing: dict[str, Decimal] = {}
    for date in dates:
        consecutive = False
        if prior_date is not None:
            previous = datetime.strptime(prior_date, "%Y-%m-%d").date()
            current = datetime.strptime(date, "%Y-%m-%d").date()
            consecutive = (current - previous).days == 1
        opening = prior_closing if consecutive else {}
        closing = closings.get(date)
        closing_values = {
            str(item["currency"]): Decimal(str(item["amount"]))
            for item in closing.get("balances", [])
        } if closing else {}
        currencies = sorted(set(flow_by_date.get(date, {})) | set(opening) | set(closing_values))
        checkpoint_currencies = {
            str(item["currency"])
            for snapshot in snapshots_by_date.get(date, [])
            for item in snapshot.get("balances", [])
        }
        currencies = sorted(set(currencies) | checkpoint_currencies)
        closing_message = core.clean_text(closing.get("source_message_summary")) if closing else ""
        for currency in currencies:
            opening_value = opening.get(currency)
            income = flow_by_date.get(date, {}).get(currency, {}).get("income", Decimal(0))
            expense = flow_by_date.get(date, {}).get(currency, {}).get("expense", Decimal(0))
            net = income - expense
            calculated = opening_value + net if opening_value is not None else None
            observed = closing_values.get(currency)
            difference = observed - calculated if observed is not None and calculated is not None else None
            if observed is None:
                status = "缺少期末盘点"
            elif opening_value is None:
                status = "期初待确认"
            elif difference == 0:
                status = "核对一致"
            else:
                status = "存在差额"
            rows.append(
                {
                    "date": date,
                    "currency": currency,
                    "opening_balance": _decimal_or_none(opening_value),
                    "income_total": _decimal_text(income, field="income_total", nonnegative=True),
                    "expense_total": _decimal_text(expense, field="expense_total", nonnegative=True),
                    "net_change": _decimal_text(net, field="net_change"),
                    "calculated_closing": _decimal_or_none(calculated),
                    "observed_closing": _decimal_or_none(observed),
                    "difference": _decimal_or_none(difference),
                    "status": status,
                    "closing_message": closing_message,
                }
            )
        prior_date = date
        prior_closing = closing_values if closing else {}
    return rows


def _decimal_or_none(value: Decimal | None) -> str | None:
    if value is None:
        return None
    return _decimal_text(value, field="derived decimal")


def compile_ledger(
    normalized: Mapping[str, Any],
    decisions: Mapping[str, Mapping[str, Any]],
) -> tuple[dict[str, Any], dict[str, int]]:
    groups: list[dict[str, Any]] = []
    for group in normalized.get("groups", []):
        if not isinstance(group, Mapping):
            continue
        group_key = core.clean_text(group.get("group_key"))
        decision = decisions.get(group_key)
        if decision is None:
            continue
        inventory = media_inventory(group)
        media_decisions = decision.get("media_decisions", {})
        records: list[dict[str, Any]] = []
        for raw_record in decision.get("records", []):
            record = deepcopy(dict(raw_record))
            source_rows = _source_message_rows(group, record["source_messages"])
            record["source_message_rows"] = source_rows
            record["source_message_summary"] = _message_summary(source_rows)
            record["senders"] = list(
                dict.fromkeys(row["sender"] for row in source_rows if row["sender"])
            )
            media = _record_media(record, inventory, media_decisions)
            record["media"] = media
            representative = record.pop("representative_media_label", "")
            record["representative_media"] = next(
                (item for item in media if item["label"] == representative), None
            )
            record.pop("media_roles", None)
            for movement in record.get("movements", []):
                movement["actual_amount"] = (
                    ("-" if movement["sign"] == "-" else "") + movement["amount"]
                    if record["posting_status"] == "posted"
                    else None
                )
            records.append(record)
        snapshots: list[dict[str, Any]] = []
        for raw_snapshot in decision.get("balance_snapshots", []):
            snapshot = deepcopy(dict(raw_snapshot))
            source_rows = _source_message_rows(group, snapshot["source_messages"])
            snapshot["source_message_rows"] = source_rows
            snapshot["source_message_summary"] = _message_summary(source_rows)
            snapshot["media"] = _snapshot_media(snapshot, inventory, media_decisions)
            snapshot.pop("media_labels", None)
            snapshots.append(snapshot)
        groups.append(
            {
                "group_key": group_key,
                "group_name": group.get("group_name"),
                "platform": group.get("platform"),
                "records": records,
                "balance_snapshots": snapshots,
            }
        )
    transfer_stats = _annotate_transfer_matches(groups)
    for group in groups:
        group["daily_reconciliation"] = _daily_reconciliation(group)
    ledger = {
        "contract_version": OUTPUT_CONTRACT,
        "accounting_mode": ACCOUNTING_MODE,
        "timezone": normalized.get("timezone"),
        "normalized_source_fingerprint": normalized.get("source_fingerprint"),
        "groups": groups,
    }
    statistics = {
        "groups": len(groups),
        "records": sum(len(group["records"]) for group in groups),
        "posted_records": sum(
            record["posting_status"] == "posted"
            for group in groups
            for record in group["records"]
        ),
        "pending_records": sum(
            record["posting_status"] == "pending"
            for group in groups
            for record in group["records"]
        ),
        "void_records": sum(
            record["posting_status"] == "void"
            for group in groups
            for record in group["records"]
        ),
        "daily_rows": sum(len(group["daily_reconciliation"]) for group in groups),
        "balance_mismatches": sum(
            row["status"] == "存在差额"
            for group in groups
            for row in group["daily_reconciliation"]
        ),
        **transfer_stats,
    }
    return ledger, statistics


def _validate_media_output(media: object, *, field: str) -> None:
    core.require(isinstance(media, Mapping), f"{field} must be an object")
    path = Path(str(media.get("path") or ""))
    core.require(path.is_file(), f"{field} source image is unavailable: {path}")
    expected_hash = core.clean_text(media.get("sha256"))
    core.require(bool(expected_hash), f"{field} source image hash is missing")
    core.require(core.sha256_file(path) == expected_hash, f"{field} source image changed after review")


def validate_ledger(value: object) -> list[dict[str, Any]]:
    core.require(isinstance(value, dict), "store ledger top level must be an object")
    core.require(value.get("contract_version") == OUTPUT_CONTRACT, "unsupported store ledger")
    core.require(value.get("accounting_mode") == ACCOUNTING_MODE, "store ledger mode mismatch")
    core.require(value.get("timezone") == "Asia/Bangkok", "store ledger timezone must be Asia/Bangkok")
    groups = value.get("groups")
    core.require(isinstance(groups, list) and groups, "store ledger groups must be nonempty")
    keys: set[str] = set()
    for group_position, group in enumerate(groups):
        field = f"store groups[{group_position}]"
        core.require(isinstance(group, dict), f"{field} must be an object")
        key = core.clean_text(group.get("group_key"))
        core.require(bool(key) and key not in keys, f"{field}.group_key must be unique")
        keys.add(key)
        core.require(isinstance(group.get("records"), list), f"{field}.records must be a list")
        core.require(
            isinstance(group.get("balance_snapshots"), list),
            f"{field}.balance_snapshots must be a list",
        )
        core.require(
            isinstance(group.get("daily_reconciliation"), list),
            f"{field}.daily_reconciliation must be a list",
        )
        for record_position, record in enumerate(group["records"]):
            record_field = f"{field}.records[{record_position}]"
            core.require(isinstance(record, dict), f"{record_field} must be an object")
            core.require(record.get("posting_status") in POSTING_STATUSES, f"{record_field} status invalid")
            core.require(isinstance(record.get("movements"), list) and record["movements"], f"{record_field} movements required")
            for movement in record["movements"]:
                expected_actual = (
                    ("-" if movement["sign"] == "-" else "") + movement["amount"]
                    if record["posting_status"] == "posted"
                    else None
                )
                core.require(
                    movement.get("actual_amount") == expected_actual,
                    f"{record_field} actual amount disagrees with posting status",
                )
            media_items = record.get("media")
            core.require(isinstance(media_items, list), f"{record_field}.media must be a list")
            for media_position, media in enumerate(media_items):
                _validate_media_output(media, field=f"{record_field}.media[{media_position}]")
            representative = record.get("representative_media")
            if representative is not None:
                core.require(representative in media_items, f"{record_field} representative media is invalid")
            transfer = record.get("transfer")
            if isinstance(transfer, Mapping):
                core.require(
                    transfer.get("match_status") in TRANSFER_MATCH_LABELS,
                    f"{record_field} transfer match status is invalid",
                )
        for snapshot_position, snapshot in enumerate(group["balance_snapshots"]):
            snapshot_field = f"{field}.balance_snapshots[{snapshot_position}]"
            core.require(isinstance(snapshot, dict), f"{snapshot_field} must be an object")
            media_items = snapshot.get("media")
            core.require(isinstance(media_items, list), f"{snapshot_field}.media must be a list")
            for media_position, media in enumerate(media_items):
                _validate_media_output(media, field=f"{snapshot_field}.media[{media_position}]")
    return groups


def _sheet_name(value: object, used: set[str]) -> str:
    cleaned = ILLEGAL_SHEET.sub("_", core.clean_text(value)) or GROUP_NAME_MARKER
    cleaned = cleaned[:31]
    candidate = cleaned
    suffix = 1
    while candidate.casefold() in used:
        suffix += 1
        marker = f"-{suffix}"
        candidate = cleaned[: 31 - len(marker)] + marker
    used.add(candidate.casefold())
    return candidate


def _excel_number(value: object) -> int | float | None:
    if value in (None, ""):
        return None
    number = Decimal(str(value))
    return int(number) if number == number.to_integral() else float(number)


def _record_anomaly(record: Mapping[str, Any], movement: Mapping[str, Any]) -> str:
    values = [
        BASIS_LABELS.get(str(movement.get("basis")), str(movement.get("basis") or "")),
        core.clean_text(movement.get("note")),
        core.clean_text(record.get("status_reason")),
        core.clean_text(record.get("note")),
    ]
    transfer = record.get("transfer")
    if isinstance(transfer, Mapping):
        match_status = str(transfer.get("match_status") or "")
        if match_status in {"unmatched", "mismatch"}:
            values.append(TRANSFER_MATCH_LABELS[match_status])
    return "；".join(dict.fromkeys(value for value in values if value))


def detail_rows(
    group: Mapping[str, Any],
) -> list[tuple[list[Any], Mapping[str, Any] | None]]:
    records = sorted(
        group.get("records", []),
        key=lambda record: (
            core.clean_text(record.get("posting_at")) or "9999",
            core.clean_text(record.get("source_message_rows", [{}])[0].get("timestamp"))
            if record.get("source_message_rows")
            else "",
            core.clean_text(record.get("id")),
        ),
    )
    rows: list[tuple[list[Any], Mapping[str, Any] | None]] = []
    for record in records:
        transfer = record.get("transfer") if isinstance(record.get("transfer"), Mapping) else {}
        for position, movement in enumerate(record.get("movements", [])):
            rows.append(
                (
                    [
                        core.clean_text(record.get("posting_at")),
                        core.clean_text(record.get("voucher_date")),
                        core.clean_text(record.get("voucher_number")) or core.clean_text(record.get("id")),
                        STATUS_LABELS.get(str(record.get("posting_status")), ""),
                        RECORD_TYPE_LABELS.get(str(record.get("record_type")), ""),
                        movement.get("currency"),
                        "收入" if movement.get("sign") == "+" else "支出",
                        movement.get("sign"),
                        _excel_number(movement.get("amount")),
                        _excel_number(movement.get("actual_amount")),
                        _excel_number(movement.get("rate")),
                        _excel_number(movement.get("thb_equivalent")),
                        core.clean_text(transfer.get("counterparty_group_name")),
                        core.clean_text(transfer.get("transfer_id")),
                        "\n".join(record.get("senders", [])),
                        core.clean_text(record.get("source_message_summary")),
                        _record_anomaly(record, movement),
                        TRANSFER_MATCH_LABELS.get(str(transfer.get("match_status") or ""), ""),
                        "",
                    ],
                    record.get("representative_media") if position == 0 else None,
                )
            )
    return rows


def summary_rows(group: Mapping[str, Any]) -> list[list[Any]]:
    return [
        [
            row.get("date"),
            row.get("currency"),
            _excel_number(row.get("opening_balance")),
            _excel_number(row.get("income_total")),
            _excel_number(row.get("expense_total")),
            _excel_number(row.get("net_change")),
            _excel_number(row.get("calculated_closing")),
            _excel_number(row.get("observed_closing")),
            _excel_number(row.get("difference")),
            row.get("status"),
            row.get("closing_message"),
        ]
        for row in group.get("daily_reconciliation", [])
    ]


def _add_image(
    worksheet: Any,
    *,
    row: int,
    column: int,
    media: Mapping[str, Any],
) -> None:
    path = Path(str(media.get("path") or ""))
    image = ExcelImage(str(path))
    box_width = 180
    box_height = 112
    scale = min(box_width / image.width, box_height / image.height)
    width = max(1, int(image.width * scale))
    height = max(1, int(image.height * scale))
    image.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=column - 1,
            colOff=pixels_to_EMU(6 + (box_width - width) // 2),
            row=row - 1,
            rowOff=pixels_to_EMU(5 + (box_height - height) // 2),
        ),
        ext=XDRPositiveSize2D(cx=pixels_to_EMU(width), cy=pixels_to_EMU(height)),
    )
    worksheet.add_image(image)


def build_workbook(ledger: object) -> Workbook:
    groups = validate_ledger(ledger)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    thin = Side(style="thin", color=GRID_RGB)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = [
        24, 13, 18, 12, 14, 10, 10, 8, 18, 22, 13, 18, 24, 18, 20, 54, 36, 18, 34
    ]
    image_column = len(DETAIL_HEADERS)
    for group in groups:
        worksheet = workbook.create_sheet(_sheet_name(group.get("group_name"), used))
        worksheet.sheet_state = "visible"
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        for column, header in enumerate(DETAIL_HEADERS, start=1):
            cell = worksheet.cell(1, column, header)
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL_RGB)
            cell.font = Font(color=HEADER_FONT_RGB, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            worksheet.column_dimensions[get_column_letter(column)].width = widths[column - 1]
        worksheet.row_dimensions[1].height = 34

        rows = detail_rows(group)
        for row_index, (values, media) in enumerate(rows, start=2):
            status = values[3]
            for column, value in enumerate(values, start=1):
                cell = worksheet.cell(row_index, column, value)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
                if row_index % 2 == 1:
                    cell.fill = PatternFill("solid", fgColor=BODY_ALT_FILL_RGB)
            if status == "待完成":
                for cell in worksheet[row_index]:
                    cell.fill = PatternFill("solid", fgColor=PENDING_FILL_RGB)
            elif status == "作废":
                for cell in worksheet[row_index]:
                    cell.fill = PatternFill("solid", fgColor=VOID_FILL_RGB)
            worksheet.row_dimensions[row_index].height = 94 if media else 32
            if media:
                _add_image(worksheet, row=row_index, column=image_column, media=media)
        detail_end = max(1, len(rows) + 1)
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(DETAIL_HEADERS))}{detail_end}"

        title_row = detail_end + 2
        header_row = title_row + 1
        worksheet.merge_cells(
            start_row=title_row,
            start_column=1,
            end_row=title_row,
            end_column=len(DETAIL_HEADERS),
        )
        title = worksheet.cell(title_row, 1, "每日余额核对")
        title.fill = PatternFill("solid", fgColor=SECTION_FILL_RGB)
        title.font = Font(color=HEADER_FONT_RGB, bold=True, size=12)
        title.alignment = Alignment(horizontal="left", vertical="center")
        for column in range(1, len(DETAIL_HEADERS) + 1):
            worksheet.cell(title_row, column).fill = PatternFill("solid", fgColor=SECTION_FILL_RGB)
            worksheet.cell(title_row, column).border = border
        for column, header in enumerate(SUMMARY_HEADERS, start=1):
            cell = worksheet.cell(header_row, column, header)
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL_RGB)
            cell.font = Font(color=HEADER_FONT_RGB, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for row_index, values in enumerate(summary_rows(group), start=header_row + 1):
            for column, value in enumerate(values, start=1):
                cell = worksheet.cell(row_index, column, value)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
            status = str(values[9] or "")
            if status == "核对一致":
                worksheet.cell(row_index, 10).fill = PatternFill("solid", fgColor=MATCH_FILL_RGB)
            elif status in {"存在差额", "缺少期末盘点", "期初待确认"}:
                for column in (9, 10):
                    worksheet.cell(row_index, column).fill = PatternFill("solid", fgColor=ALERT_FILL_RGB)
        worksheet.column_dimensions["K"].width = max(worksheet.column_dimensions["K"].width or 0, 54)
    core.require(bool(workbook.sheetnames), "store workbook must contain a group sheet")
    return workbook


def _image_anchor(image: Any) -> tuple[int, int] | None:
    anchor = getattr(image, "anchor", None)
    marker = getattr(anchor, "_from", None)
    if marker is None:
        return None
    return int(marker.row) + 1, int(marker.col) + 1


def _comparable(value: object) -> tuple[str, str]:
    if value in (None, ""):
        return ("blank", "")
    if isinstance(value, bool):
        return ("bool", str(value))
    if isinstance(value, (int, float, Decimal)):
        return ("number", _decimal_text(value, field="workbook value"))
    return ("text", str(value))


def check_workbook(workbook_path: Path, ledger: object) -> list[str]:
    errors: list[str] = []
    try:
        groups = validate_ledger(ledger)
    except (OSError, TypeError, ValueError) as exc:
        return [str(exc)]
    used: set[str] = set()
    expected_names = [_sheet_name(group.get("group_name"), used) for group in groups]
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        if getattr(workbook, "_external_links", []):
            errors.append("store workbook contains forbidden external links")
        if getattr(workbook, "vba_archive", None) is not None:
            errors.append("store workbook contains forbidden macros")
        if workbook.sheetnames != expected_names:
            errors.append(
                f"store sheet names mismatch: expected {expected_names!r}, got {workbook.sheetnames!r}"
            )
        for group, sheet_name in zip(groups, expected_names):
            worksheet = workbook[sheet_name]
            headers = [worksheet.cell(1, column).value for column in range(1, len(DETAIL_HEADERS) + 1)]
            if headers != DETAIL_HEADERS:
                errors.append(f"{sheet_name}: store detail header mismatch")
            if str(worksheet.freeze_panes or "") != "A2":
                errors.append(f"{sheet_name}: first row must be frozen at A2")
            expected_details = detail_rows(group)
            expected_images: Counter[tuple[int, int]] = Counter()
            for row_index, (values, media) in enumerate(expected_details, start=2):
                for column, expected in enumerate(values, start=1):
                    actual = worksheet.cell(row_index, column).value
                    if _comparable(actual) != _comparable(expected):
                        errors.append(
                            f"{sheet_name}!{worksheet.cell(row_index, column).coordinate}: expected {expected!r}, got {actual!r}"
                        )
                if media:
                    expected_images[(row_index, len(DETAIL_HEADERS))] += 1
            actual_images: Counter[tuple[int, int]] = Counter()
            for image in worksheet._images:
                anchor = _image_anchor(image)
                if anchor is None:
                    errors.append(f"{sheet_name}: image has no cell anchor")
                else:
                    actual_images[anchor] += 1
            if actual_images != expected_images:
                errors.append(
                    f"{sheet_name}: image anchors mismatch: expected {dict(expected_images)}, got {dict(actual_images)}"
                )
            detail_end = max(1, len(expected_details) + 1)
            title_row = detail_end + 2
            header_row = title_row + 1
            if worksheet.cell(title_row, 1).value != "每日余额核对":
                errors.append(f"{sheet_name}: daily reconciliation title is missing")
            summary_headers = [
                worksheet.cell(header_row, column).value
                for column in range(1, len(SUMMARY_HEADERS) + 1)
            ]
            if summary_headers != SUMMARY_HEADERS:
                errors.append(f"{sheet_name}: daily reconciliation header mismatch")
            for offset, values in enumerate(summary_rows(group), start=1):
                row_index = header_row + offset
                for column, expected in enumerate(values, start=1):
                    actual = worksheet.cell(row_index, column).value
                    if _comparable(actual) != _comparable(expected):
                        errors.append(
                            f"{sheet_name}!{worksheet.cell(row_index, column).coordinate}: expected {expected!r}, got {actual!r}"
                        )
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        errors.append(f"{sheet_name}!{cell.coordinate}: formulas are forbidden")
    finally:
        workbook.close()
    return errors
