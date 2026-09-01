#!/usr/bin/env python3
"""Build one worksheet per chat group from a simple-mode ledger."""

from __future__ import annotations

import argparse
import json
import re
import sys
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import core
import simple_ledger


ILLEGAL_SHEET = re.compile(r"[\\/*?:\[\]]")
PLATFORM_SHEET_PREFIXES = {
    "telegram": "TG",
    "whatsapp": "WA",
    "line": "LINE",
}
SUMMARY_FILL_RGB = "D9EAF7"
REVIEW_SHORT_FONT_RGB = "008000"
REVIEW_OVER_FONT_RGB = "C00000"


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("orders", type=Path)
    parser.add_argument("-o", "--output", required=True, type=Path)
    parser.add_argument(
        "--template",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
    )
    parser.add_argument("--force", action="store_true")
    return parser.parse_args(argv)


def _validate_reconciliation(value: object, *, field: str) -> bool:
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    status = core.clean_text(value.get("status")).casefold()
    core.require(
        status in {"matched", "short", "over", "pending", "composite"},
        f"{field}.status is unsupported",
    )
    if status == "pending":
        core.require(bool(core.clean_text(value.get("reason"))), f"{field}.reason is required")
        core.require(bool(core.clean_text(value.get("detail"))), f"{field}.detail is required")
        return True
    if status == "composite":
        items = value.get("items")
        core.require(isinstance(items, list) and items, f"{field}.items is required")
        pending = False
        for index, item in enumerate(items):
            item_field = f"{field}.items[{index}]"
            core.require(isinstance(item, Mapping), f"{item_field} must be an object")
            core.require(bool(core.clean_text(item.get("label"))), f"{item_field}.label is required")
            pending = _validate_reconciliation(
                item.get("reconciliation"),
                field=f"{item_field}.reconciliation",
            ) or pending
        return pending

    difference = core.parse_decimal(value.get("difference"), field=f"{field}.difference")
    core.require(difference is not None, f"{field}.difference is required")
    currency = core.normalize_currency(value.get("currency"), field=f"{field}.currency")
    assert currency is not None
    if status == "matched":
        core.require(
            abs(difference) <= core.currency_tolerance(currency),
            f"{field}.difference does not match status matched",
        )
    elif status == "short":
        core.require(difference < 0, f"{field}.difference must be negative for short")
    else:
        core.require(difference > 0, f"{field}.difference must be positive for over")
    return False


def validate_orders(data: object) -> list[dict[str, Any]]:
    core.require(isinstance(data, dict), "orders top level must be an object")
    core.require(
        data.get("contract_version") == core.ORDERS_CONTRACT,
        "unsupported orders contract_version",
    )
    core.require(data.get("rule_version") == core.RULE_VERSION, "orders rule_version mismatch")
    core.require(data.get("accounting_mode") == "simple", "orders must use simple mode")
    core.require(data.get("timezone") == "Asia/Bangkok", "orders timezone must be Asia/Bangkok")
    groups = data.get("groups")
    core.require(isinstance(groups, list) and groups, "orders groups must be a nonempty list")
    keys = [group.get("group_key") for group in groups if isinstance(group, dict)]
    core.require(
        len(keys) == len(groups) and all(keys) and len(keys) == len(set(keys)),
        "orders group keys are invalid",
    )
    for group_index, group in enumerate(groups):
        orders = group.get("orders")
        core.require(isinstance(orders, list), f"groups[{group_index}].orders must be a list")
        for order_index, order in enumerate(orders):
            core.require(
                isinstance(order, Mapping),
                f"groups[{group_index}].orders[{order_index}] must be an object",
            )
            order_field = f"groups[{group_index}].orders[{order_index}]"
            pending = _validate_reconciliation(
                order.get("reconciliation"),
                field=f"{order_field}.reconciliation",
            )
            core.require(
                order.get("review_result")
                == simple_ledger.format_reconciliation(order.get("reconciliation")),
                f"{order_field}.review_result does not match reconciliation",
            )
            if pending:
                core.require(
                    bool(order_row_note(order)),
                    f"{order_field}: pending reconciliation must produce a workbook note",
                )
            flows = order.get("flows")
            core.require(
                isinstance(flows, list),
                f"groups[{group_index}].orders[{order_index}].flows must be a list",
            )
            for flow_index, flow in enumerate(flows):
                core.require(
                    isinstance(flow, Mapping),
                    f"groups[{group_index}].orders[{order_index}].flows[{flow_index}] "
                    "must be an object",
                )
                payee_field = (
                    f"groups[{group_index}].orders[{order_index}].flows[{flow_index}].payee"
                )
                if (
                    data.get("payee_policy") == "thai_bank_account_only"
                    and flow.get("cash") is not True
                    and flow.get("currency") == "THB"
                ):
                    core.validate_thai_bank_account_payee(
                        flow.get("payee"),
                        field=payee_field,
                        payee_state=flow.get("payee_state"),
                    )
                else:
                    core.validate_payee(
                        flow.get("payee"),
                        field=payee_field,
                        cash=flow.get("cash") is True,
                        payee_state=flow.get("payee_state"),
                    )
    return groups


def sheet_name(value: object, used: set[str]) -> str:
    base = ILLEGAL_SHEET.sub("_", core.clean_text(value)).strip(" '") or "未命名群"
    base = base[:31]
    candidate = base
    index = 2
    while candidate.casefold() in used:
        suffix = f"_{index}"
        candidate = base[: 31 - len(suffix)] + suffix
        index += 1
    used.add(candidate.casefold())
    return candidate


def platform_sheet_prefix(value: object) -> str:
    platform = core.clean_text(value).casefold()
    if platform in PLATFORM_SHEET_PREFIXES:
        return PLATFORM_SHEET_PREFIXES[platform]
    fallback = re.sub(r"[^A-Z0-9]", "", core.clean_text(value).upper())[:8]
    return fallback or "CHAT"


def group_sheet_name(group: Mapping[str, Any], used: set[str]) -> str:
    prefix = platform_sheet_prefix(group.get("platform"))
    group_label = group.get("group_name") or group.get("group_key")
    return sheet_name(f"{prefix}-{core.clean_text(group_label)}", used)


def excel_number(value: object) -> int | float | str | None:
    if value in (None, ""):
        return None
    number = core.parse_decimal(value, field="workbook amount")
    assert number is not None
    if len(number.as_tuple().digits) > 15:
        return core.decimal_text(number)
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def excel_datetime(value: object) -> datetime | None:
    if value in (None, ""):
        return None
    parsed = datetime.fromisoformat(str(value))
    return parsed.replace(tzinfo=None)


def rate_cell(value: Mapping[str, Any]) -> int | float | str | None:
    display = value.get("actual_rate_display")
    if display not in (None, ""):
        display_text = str(display)
        if "\n" in display_text or "：" in display_text or display_text.startswith("÷"):
            return display_text
    rate = excel_number(value.get("actual_rate"))
    if rate is not None:
        return rate
    if value.get("pricing_basis") == "unknown":
        return "待确认"
    return None


def flow_row_label(flow: Mapping[str, Any]) -> str | None:
    if flow.get("settlement_allocation"):
        return "回款追回分摊" if flow.get("side") == "recovery" else "内部回款分摊"
    return core.clean_text(flow.get("flow_type")) or None


def flow_row_note(flow: Mapping[str, Any]) -> str | None:
    if flow.get("settlement_allocation"):
        source_amount = core.clean_text(flow.get("source_amount"))
        amount = core.clean_text(flow.get("amount"))
        currency = core.clean_text(flow.get("currency"))
        source_label = "原始合并追回" if flow.get("side") == "recovery" else "原始合并回款"
        return f"{source_label} {source_amount} {currency}；本单计入 {amount} {currency}。"
    if flow.get("side") == "unassigned":
        return "金额和币种按截图展示；尚未归入订单，未计入任何订单合计。"
    if flow.get("duplicate_of") or flow.get("status") == "failed":
        return None
    if not flow.get("included"):
        return "截图金额或币种尚未清楚确认，金额未计入。"
    return None


def order_row_note(order: Mapping[str, Any]) -> str | None:
    notes: list[str] = []
    manual = core.clean_text(order.get("note"))
    if manual:
        return manual

    def collect(value: object, label: str | None = None) -> None:
        if not isinstance(value, Mapping):
            return
        status = core.clean_text(value.get("status")).casefold()
        if status == "pending":
            detail = core.clean_text(value.get("detail"))
            if detail:
                notes.append(f"{label}：{detail}" if label else detail)
            return
        if status != "composite":
            return
        for item in value.get("items", []):
            if isinstance(item, Mapping):
                collect(
                    item.get("reconciliation"),
                    core.clean_text(item.get("label")) or None,
                )

    collect(order.get("reconciliation"))
    return "\n".join(dict.fromkeys(notes)) or None


def pricing_detail_rows(order: Mapping[str, Any]) -> list[list[Any]]:
    fee_labels = {
        "delivery_fee": "配送费",
        "service_fee": "手续费",
        "network_fee": "网络费用",
    }
    legs = order.get("legs", [])
    sources: list[tuple[str | None, list[Mapping[str, Any]], Mapping[str, Any] | None]] = []
    if legs:
        for leg in legs:
            sources.append(
                (
                    core.clean_text(leg.get("display_label"))
                    or core.clean_text(leg.get("direction"))
                    or None,
                    [item for item in leg.get("fees", []) if isinstance(item, Mapping)],
                    leg.get("rounding") if isinstance(leg.get("rounding"), Mapping) else None,
                )
            )
    else:
        sources.append(
            (
                core.clean_text(order.get("direction")) or None,
                [
                    item
                    for item in order.get("fee_adjustments", [])
                    if isinstance(item, Mapping)
                ],
                order.get("rounding")
                if isinstance(order.get("rounding"), Mapping)
                else None,
            )
        )

    rows: list[list[Any]] = []
    for direction, fees, rounding in sources:
        for fee in fees:
            kind = core.clean_text(fee.get("kind")).casefold()
            label = fee_labels.get(kind, kind or "费用")
            rows.append(
                [
                    label,
                    order.get("order_id"),
                    None,
                    direction,
                    None,
                    None,
                    excel_number(fee.get("amount")),
                    fee.get("currency"),
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                ]
            )
        if rounding:
            rows.append(
                [
                    "舍入",
                    order.get("order_id"),
                    None,
                    direction,
                    None,
                    None,
                    excel_number(rounding.get("unit")),
                    rounding.get("currency"),
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                ]
            )
    return rows


def order_rows(order: Mapping[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = [
        [
            "订单汇总",
            order.get("order_id"),
            order.get("customer_nickname"),
            order.get("direction"),
            excel_number(order.get("payment_total")),
            rate_cell(order),
            None,
            None,
            excel_number(order.get("expected_payout")),
            excel_number(order.get("actual_payout_total")),
            (
                simple_ledger.format_reconciliation(order.get("reconciliation"))
                if isinstance(order.get("reconciliation"), Mapping)
                else order.get("review_result")
            )
            or None,
            order_row_note(order),
            None,
            None,
        ]
    ]
    rows.extend(pricing_detail_rows(order))
    for flow in order.get("flows", []):
        if not flow.get("included") and not flow.get("display_in_workbook"):
            continue
        rows.append(
            [
                flow_row_label(flow),
                order.get("order_id"),
                None,
                flow.get("leg_display_label") or flow.get("leg_direction"),
                None,
                None,
                excel_number(flow.get("amount")),
                flow.get("currency"),
                None,
                None,
                None,
                flow_row_note(flow),
                flow.get("payee"),
                excel_datetime(flow.get("message_time")),
            ]
        )
    return rows


def snapshot_cell_style(source: Any) -> dict[str, Any]:
    return {
        "font": copy(source.font),
        "fill": copy(source.fill),
        "border": copy(source.border),
        "alignment": copy(source.alignment),
        "number_format": source.number_format,
        "protection": copy(source.protection),
    }


def apply_cell_style(destination: Any, style: Mapping[str, Any]) -> None:
    destination.font = copy(style["font"])
    destination.fill = copy(style["fill"])
    destination.border = copy(style["border"])
    destination.alignment = copy(style["alignment"])
    destination.number_format = str(style["number_format"])
    destination.protection = copy(style["protection"])


def template_styles(
    template_path: Path,
) -> tuple[list[Any], list[Any], dict[int, float | None], float | None]:
    workbook = load_workbook(template_path, read_only=False, data_only=False)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        header = [snapshot_cell_style(cell) for cell in sheet[1][: len(core.HEADERS)]]
        body = [snapshot_cell_style(cell) for cell in sheet[2][: len(core.HEADERS)]]
        widths = {
            index: sheet.column_dimensions[get_column_letter(index)].width
            for index in range(1, len(core.HEADERS) + 1)
        }
        return header, body, widths, sheet.row_dimensions[1].height
    finally:
        workbook.close()


def build_workbook(
    groups: Iterable[Mapping[str, Any]],
    template_path: Path,
) -> Workbook:
    group_list = list(groups)
    header_styles, body_styles, widths, header_height = template_styles(template_path)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for group in group_list:
        worksheet = workbook.create_sheet(group_sheet_name(group, used))
        worksheet.sheet_state = "visible"
        worksheet.freeze_panes = "A2"
        if header_height:
            worksheet.row_dimensions[1].height = header_height
        wrapped_headers = {"换汇方向", "汇率", "核对结果", "备注", "收款方"}
        wrapped_columns = {
            core.HEADERS.index(header) + 1 for header in wrapped_headers
        }
        payee_column = core.HEADERS.index("收款方") + 1
        review_column = core.HEADERS.index("核对结果") + 1
        review_column_letter = get_column_letter(review_column)
        for column, (value, style) in enumerate(zip(core.HEADERS, header_styles), start=1):
            cell = worksheet.cell(row=1, column=column, value=value)
            apply_cell_style(cell, style)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            if cell.font.name is None:
                cell.font = Font(
                    name="Arial",
                    size=10,
                    bold=cell.font.bold,
                    color=cell.font.color,
                )
            width = widths.get(column)
            if width is not None:
                worksheet.column_dimensions[get_column_letter(column)].width = width
        output_row = 2
        for order in group.get("orders", []):
            for values in order_rows(order):
                core.require(
                    len(values) == len(core.HEADERS),
                    f"workbook row has {len(values)} cells; expected {len(core.HEADERS)}",
                )
                for column, (value, style) in enumerate(zip(values, body_styles), start=1):
                    cell = worksheet.cell(row=output_row, column=column, value=value)
                    apply_cell_style(cell, style)
                    if column == payee_column:
                        cell.number_format = "@"
                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                        wrap_text=column in wrapped_columns,
                    )
                    if column == review_column and isinstance(value, str):
                        review_font = copy(cell.font)
                        if value.startswith("少转"):
                            review_font.color = REVIEW_SHORT_FONT_RGB
                        elif value.startswith("多转"):
                            review_font.color = REVIEW_OVER_FONT_RGB
                        cell.font = review_font
                    if cell.font.name is None:
                        fallback_font = copy(cell.font)
                        fallback_font.name = "Arial"
                        fallback_font.sz = 10
                        cell.font = fallback_font
                if values[0] == "订单汇总":
                    for cell in worksheet[output_row]:
                        cell.fill = PatternFill("solid", fgColor=SUMMARY_FILL_RGB)
                        bold_font = copy(cell.font)
                        bold_font.bold = True
                        cell.font = bold_font
                else:
                    for cell in worksheet[output_row]:
                        cell.fill = PatternFill()
                worksheet.cell(output_row, len(core.HEADERS)).number_format = (
                    "yyyy-mm-dd hh:mm:ss"
                )
                output_row += 1
        review_range = f"{review_column_letter}2:{review_column_letter}{max(2, output_row - 1)}"
        worksheet.conditional_formatting.add(
            review_range,
            FormulaRule(
                formula=[f'LEFT(${review_column_letter}2,2)="少转"'],
                font=Font(color=REVIEW_SHORT_FONT_RGB),
                stopIfTrue=True,
            ),
        )
        worksheet.conditional_formatting.add(
            review_range,
            FormulaRule(
                formula=[f'LEFT(${review_column_letter}2,2)="多转"'],
                font=Font(color=REVIEW_OVER_FONT_RGB),
                stopIfTrue=True,
            ),
        )
        last_column = get_column_letter(len(core.HEADERS))
        worksheet.auto_filter.ref = f"A1:{last_column}{max(1, output_row - 1)}"
    core.require(bool(workbook.sheetnames), "workbook must contain at least one group sheet")
    return workbook


def main(argv: list[str] | None = None) -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")
    sys.stderr.reconfigure(encoding="utf-8", errors="backslashreplace")
    args = parse_args(argv)
    output = args.output.resolve()
    if output.exists() and not args.force:
        print(f"Output already exists (use --force): {output}", file=sys.stderr)
        return 2
    try:
        orders = json.loads(args.orders.resolve().read_text(encoding="utf-8-sig"))
        groups = validate_orders(orders)
        workbook = build_workbook(groups, args.template.resolve())
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary = output.parent / f".{output.name}.tmp.xlsx"
        workbook.save(temporary)
        workbook.close()
        temporary.replace(output)
    except (OSError, ValueError, KeyError, TypeError, json.JSONDecodeError) as exc:
        print(f"Workbook build failed: {exc}", file=sys.stderr)
        return 2
    print(json.dumps({"output": str(output), "sheets": len(groups)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
