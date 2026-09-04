#!/usr/bin/env python3
"""Validate and publish identity and chat-account material from finance groups."""

from __future__ import annotations

import math
import re
from collections import Counter
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU

import core


GROUP_MODE = "finance"
GROUP_NAME_MARKER = "财务资料群"
DECISION_CONTRACT = "group-chat-finance-materials-decision/1.0"
OUTPUT_CONTRACT = "group-chat-finance-materials-ledger/1.0"

MEDIA_CLASSIFICATIONS = frozenset({"document", "chat_profile", "reference"})
DOCUMENT_TYPES = frozenset(
    {"passport", "identity_card", "visa", "residence_permit", "driver_license", "other"}
)
DOCUMENT_TYPE_ALIASES = {
    "passport": "passport",
    "护照": "passport",
    "identity_card": "identity_card",
    "identity card": "identity_card",
    "id_card": "identity_card",
    "id card": "identity_card",
    "身份证": "identity_card",
    "visa": "visa",
    "签证": "visa",
    "residence_permit": "residence_permit",
    "residence permit": "residence_permit",
    "居留证": "residence_permit",
    "driver_license": "driver_license",
    "driver license": "driver_license",
    "driving_license": "driver_license",
    "驾照": "driver_license",
    "驾驶证": "driver_license",
    "other": "other",
    "其他": "other",
}
DOCUMENT_TYPE_LABELS = {
    "passport": "护照",
    "identity_card": "身份证",
    "visa": "签证",
    "residence_permit": "居留证",
    "driver_license": "驾照",
    "other": "其他证件",
}
PLATFORM_ALIASES = {
    "wechat": "WeChat",
    "微信": "WeChat",
    "line": "LINE",
    "whatsapp": "WhatsApp",
    "whatsapp business": "WhatsApp",
    "telegram": "Telegram",
    "tg": "Telegram",
    "飞机": "Telegram",
    "qq": "QQ",
    "other": "Other",
    "其他": "Other",
}

HEADERS = [
    "姓名",
    "姓拼音",
    "名拼音",
    "证件类型",
    "国家代码",
    "护照号/证件号码",
    "国籍",
    "出生日期",
    "证件原图",
    "平台",
    "账号 ID",
    "手机号",
    "资料页原图",
    "关联状态",
    "备注",
]

TEXT_COLUMNS = frozenset({5, 6, 8, 11, 12})
DOCUMENT_IMAGE_COLUMN = HEADERS.index("证件原图") + 1
PROFILE_IMAGE_COLUMN = HEADERS.index("资料页原图") + 1
IMAGE_COLUMNS = frozenset({DOCUMENT_IMAGE_COLUMN, PROFILE_IMAGE_COLUMN})

HEADER_FILL_RGB = "1F4E78"
HEADER_FONT_RGB = "FFFFFF"
GRID_RGB = "A7B9C8"
PENDING_FILL_RGB = "FFF2CC"
BODY_ALT_FILL_RGB = "F4F8FB"

ILLEGAL_SHEET = re.compile(r"[\\/*?:\[\]]")
PERSON_ID_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9._-]{0,127}")
SOURCE_LABEL_RE = re.compile(r"S(\d{5})")
COUNTRY_CODE_RE = re.compile(r"[A-Z]{3}")


def group_name_selected(value: object) -> bool:
    return GROUP_NAME_MARKER.casefold() in str(value or "").casefold()


def normalize_document_type(value: object, *, field: str) -> str:
    cleaned = core.clean_text(value).casefold()
    normalized = DOCUMENT_TYPE_ALIASES.get(cleaned)
    core.require(normalized in DOCUMENT_TYPES, f"{field} is unsupported")
    return str(normalized)


def normalize_platform(value: object, *, field: str) -> str:
    cleaned = core.clean_text(value)
    normalized = PLATFORM_ALIASES.get(cleaned.casefold())
    core.require(normalized is not None, f"{field} is unsupported")
    return str(normalized)


def _message_labels(group: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    return {
        f"S{position:05d}": message
        for position, message in enumerate(group.get("messages", []), start=1)
        if isinstance(message, Mapping)
    }


def media_inventory(
    group: Mapping[str, Any],
) -> dict[str, tuple[str, Mapping[str, Any], Mapping[str, Any]]]:
    inventory: dict[str, tuple[str, Mapping[str, Any], Mapping[str, Any]]] = {}
    media_position = 0
    for message_position, message in enumerate(group.get("messages", []), start=1):
        if not isinstance(message, Mapping):
            continue
        for media in message.get("media", []):
            if not isinstance(media, Mapping) or not core.media_is_evidence(media):
                continue
            media_position += 1
            inventory[f"M{media_position:04d}"] = (
                f"S{message_position:05d}",
                message,
                media,
            )
    return inventory


def _clean_list(value: object, *, field: str) -> list[str]:
    core.require(isinstance(value, list), f"{field} must be a list")
    cleaned = [core.clean_text(item) for item in value]
    core.require(all(cleaned), f"{field} cannot contain blank values")
    core.require(len(cleaned) == len(set(cleaned)), f"{field} repeats a value")
    return cleaned


def _validate_source_messages(
    value: object,
    *,
    field: str,
    message_labels: set[str],
    reviewed_through: int,
) -> list[str]:
    labels = _clean_list(value, field=field)
    core.require(set(labels) <= message_labels, f"{field} contains an unknown message label")
    for label in labels:
        match = SOURCE_LABEL_RE.fullmatch(label)
        core.require(match is not None, f"{field} contains an invalid message label")
        core.require(
            int(match.group(1)) <= reviewed_through,
            f"{field} cites a message that has not been reviewed",
        )
    return labels


def _normalize_birth_date(value: object, *, field: str) -> str:
    cleaned = core.clean_text(value)
    if not cleaned:
        return ""
    try:
        datetime.strptime(cleaned, "%Y-%m-%d")
    except ValueError as exc:
        raise ValueError(f"{field} must use YYYY-MM-DD") from exc
    return cleaned


def _account_keys(platform: str, account_id: str, phone: str) -> list[tuple[str, str, str]]:
    keys: list[tuple[str, str, str]] = []
    if account_id:
        keys.append((platform.casefold(), "id", account_id.casefold()))
    if phone:
        canonical_phone = re.sub(r"[\s()\-–—]", "", phone).casefold()
        keys.append((platform.casefold(), "phone", canonical_phone))
    return keys


def _document_key(document_type: str, country_code: str, number: str) -> tuple[str, str, str]:
    canonical_number = re.sub(r"[\s\-–—]+", "", number).casefold()
    return document_type, country_code.casefold(), canonical_number


def validate_decision(
    normalized: Mapping[str, Any],
    group: Mapping[str, Any],
    decision: dict[str, Any],
    expected: Mapping[str, Any],
    *,
    require_complete: bool,
    capture_hashes: bool,
    rehash_labels: set[str] | None = None,
) -> dict[str, Any]:
    del normalized
    core.require(
        decision.get("contract_version") == DECISION_CONTRACT,
        f"unsupported finance-material decision contract; expected {DECISION_CONTRACT}",
    )
    unknown_top = sorted(set(decision) - set(expected))
    core.require(
        not unknown_top,
        "finance-material decision has unsupported fields: " + ", ".join(unknown_top),
    )
    for field in (
        "normalized_source_fingerprint",
        "group_fingerprint",
        "group_key",
        "group_name",
        "platform",
        "message_count",
        "evidence_media_count",
        "group_mode",
    ):
        core.require(
            decision.get(field) == expected.get(field),
            f"{group.get('group_key')}: generated {field} was edited",
        )

    messages = list(group.get("messages", []))
    reviewed_through = decision.get("reviewed_through")
    core.require(
        isinstance(reviewed_through, int) and 0 <= reviewed_through <= len(messages),
        "invalid reviewed_through",
    )
    core.require(isinstance(decision.get("read_complete"), bool), "read_complete must be boolean")
    core.require(
        decision.get("read_complete") is (reviewed_through == len(messages)),
        "read_complete must exactly match reviewed_through",
    )
    if require_complete:
        core.require(
            decision.get("read_complete") is True,
            "finance-material group chronology has not been read completely",
        )

    message_by_label = _message_labels(group)
    message_labels = set(message_by_label)
    inventory = media_inventory(group)
    available = {
        label for label, (_, _, media) in inventory.items() if media.get("availability") == "available"
    }
    missing = set(inventory) - available
    decisions = decision.get("media_decisions")
    core.require(isinstance(decisions, dict), "media_decisions must be an object keyed by M labels")
    unknown_media = sorted(set(decisions) - available)
    core.require(
        not unknown_media,
        f"finance-material media decisions contain missing or unknown labels: {unknown_media[:20]}",
    )
    if require_complete:
        unclassified = sorted(available - set(decisions))
        core.require(
            not unclassified,
            f"unclassified available finance-material media: {unclassified[:20]}",
        )

    labels_by_class: dict[str, set[str]] = {
        classification: set() for classification in MEDIA_CLASSIFICATIONS
    }
    evidence_hashes_computed = 0
    evidence_hashes_reused = 0
    for label, raw in decisions.items():
        field = f"{group.get('group_key')}.media_decisions.{label}"
        core.require(isinstance(raw, dict), f"{field} must be an object")
        classification = core.clean_text(raw.get("classification")).casefold()
        core.require(
            classification in MEDIA_CLASSIFICATIONS,
            f"{field}.classification must be document, chat_profile, or reference",
        )
        raw["classification"] = classification
        labels_by_class[classification].add(label)
        if classification == "reference":
            unknown = sorted(set(raw) - {"classification", "note"})
            core.require(not unknown, f"{field}: reference has unsupported fields: {', '.join(unknown)}")
            continue
        unknown = sorted(
            set(raw) - {"classification", "viewed_original", "evidence_sha256", "note"}
        )
        core.require(
            not unknown,
            f"{field}: material decision has unsupported fields: {', '.join(unknown)}",
        )
        core.require(
            raw.get("viewed_original") is True,
            f"{field}: open the original image before recording document or account facts",
        )
        media_path = Path(str(inventory[label][2].get("path") or ""))
        core.require(media_path.is_file(), f"{field}: original media file is unavailable: {media_path}")
        recorded_hash = core.clean_text(raw.get("evidence_sha256"))
        verify_hash = (
            rehash_labels is None
            or label in rehash_labels
            or (capture_hashes and not recorded_hash)
        )
        resolved_hash, computed = core.validate_cached_file_hash(
            media_path,
            recorded_hash,
            verify=verify_hash,
            changed_message=f"{field}: original material image changed after review",
        )
        snapshot_hash = core.clean_text(inventory[label][2].get("blob_sha256"))
        if computed and snapshot_hash:
            core.require(
                resolved_hash == snapshot_hash,
                f"{field}: original material image changed after the run snapshot was created",
            )
        if computed:
            evidence_hashes_computed += 1
        elif recorded_hash:
            evidence_hashes_reused += 1
        if not recorded_hash and capture_hashes:
            raw["evidence_sha256"] = resolved_hash
        elif require_complete:
            core.require(
                bool(recorded_hash),
                f"{field}: evidence hash has not been captured; run review seal",
            )

    open_people_value = decision.get("open_people")
    core.require(isinstance(open_people_value, list), "open_people must be a list")
    open_ids: set[str] = set()
    for position, item in enumerate(open_people_value):
        field = f"{group.get('group_key')}.open_people[{position}]"
        core.require(isinstance(item, dict), f"{field} must be an object")
        unknown = sorted(set(item) - {"id", "source_messages", "summary", "unresolved"})
        core.require(not unknown, f"{field}: unsupported fields: {', '.join(unknown)}")
        person_id = core.clean_text(item.get("id"))
        core.require(PERSON_ID_RE.fullmatch(person_id) is not None, f"{field}.id is invalid")
        core.require(person_id not in open_ids, f"{field}.id is duplicated")
        open_ids.add(person_id)
        item["id"] = person_id
        item["source_messages"] = _validate_source_messages(
            item.get("source_messages"),
            field=f"{field}.source_messages",
            message_labels=message_labels,
            reviewed_through=reviewed_through,
        )
        summary = core.clean_text(item.get("summary"))
        core.require(bool(summary), f"{field}.summary is required")
        item["summary"] = summary
        unresolved_value = item.get("unresolved", [])
        core.require(isinstance(unresolved_value, list), f"{field}.unresolved must be a list")
        item["unresolved"] = [core.clean_text(value) for value in unresolved_value]
        core.require(all(item["unresolved"]), f"{field}.unresolved cannot contain blanks")
    if decision.get("read_complete") or require_complete:
        core.require(not open_people_value, "open_people must be empty after the group is fully read")

    people = decision.get("people")
    core.require(isinstance(people, list), "people must be a list")
    person_ids: set[str] = set()
    assigned_material: set[str] = set()
    document_owner: dict[tuple[str, str, str], str] = {}
    account_owner: dict[tuple[str, str, str], str] = {}
    document_count = 0
    account_count = 0
    pending_people = 0

    for position, person in enumerate(people):
        field = f"{group.get('group_key')}.people[{position}]"
        core.require(isinstance(person, dict), f"{field} must be an object")
        allowed = {
            "id",
            "name",
            "surname",
            "given_names",
            "nationality",
            "birth_date",
            "documents",
            "accounts",
            "source_messages",
            "association_status",
            "note",
        }
        unknown = sorted(set(person) - allowed)
        core.require(not unknown, f"{field}: unsupported fields: {', '.join(unknown)}")
        person_id = core.clean_text(person.get("id"))
        core.require(PERSON_ID_RE.fullmatch(person_id) is not None, f"{field}.id is invalid")
        core.require(person_id not in person_ids, f"{field}.id is duplicated")
        person_ids.add(person_id)
        person["id"] = person_id
        for text_field in ("name", "surname", "given_names", "nationality"):
            person[text_field] = core.clean_text(person.get(text_field))
        person["birth_date"] = _normalize_birth_date(
            person.get("birth_date"), field=f"{field}.birth_date"
        )
        association_status = core.clean_text(person.get("association_status")).casefold()
        core.require(
            association_status in {"confirmed", "pending"},
            f"{field}.association_status must be confirmed or pending",
        )
        person["association_status"] = association_status
        note = core.clean_text(person.get("note"))
        person["note"] = note
        if association_status == "pending":
            pending_people += 1
            core.require(bool(note), f"{field}.note is required when association_status is pending")

        documents = person.get("documents")
        accounts = person.get("accounts")
        core.require(isinstance(documents, list), f"{field}.documents must be a list")
        core.require(isinstance(accounts, list), f"{field}.accounts must be a list")
        core.require(documents or accounts, f"{field} must contain at least one document or account")
        if accounts and not documents:
            core.require(
                association_status == "pending",
                f"{field}: an account without a document must remain pending",
            )

        derived_sources: list[str] = []
        local_document_keys: set[tuple[str, str, str]] = set()
        for document_position, document in enumerate(documents):
            document_field = f"{field}.documents[{document_position}]"
            core.require(isinstance(document, dict), f"{document_field} must be an object")
            unknown_document = sorted(
                set(document) - {"type", "country_code", "number", "media_labels"}
            )
            core.require(
                not unknown_document,
                f"{document_field}: unsupported fields: {', '.join(unknown_document)}",
            )
            document_type = normalize_document_type(
                document.get("type"), field=f"{document_field}.type"
            )
            country_code = core.clean_text(document.get("country_code")).upper()
            if country_code:
                core.require(
                    COUNTRY_CODE_RE.fullmatch(country_code) is not None,
                    f"{document_field}.country_code must be a three-letter code",
                )
            number = core.clean_text(document.get("number"))
            media_labels = _clean_list(
                document.get("media_labels"), field=f"{document_field}.media_labels"
            )
            core.require(
                bool(media_labels),
                f"{document_field}.media_labels must contain the original document image",
            )
            for label in media_labels:
                core.require(
                    label in labels_by_class["document"],
                    f"{document_field}.media_labels contains a label not classified as document: {label}",
                )
                core.require(
                    label not in assigned_material,
                    f"material image assigned to multiple people or records: {label}",
                )
                assigned_material.add(label)
                derived_sources.append(inventory[label][0])
            document["type"] = document_type
            document["country_code"] = country_code
            document["number"] = number
            document["media_labels"] = media_labels
            if association_status == "confirmed":
                core.require(bool(number), f"{document_field}.number is required for a confirmed record")
                if document_type == "passport":
                    core.require(
                        bool(country_code),
                        f"{document_field}.country_code is required for a confirmed passport",
                    )
                    core.require(
                        bool(person["name"] or (person["surname"] and person["given_names"])),
                        f"{field}: a confirmed passport requires the visible holder name",
                    )
                    core.require(
                        bool(person["surname"] and person["given_names"]),
                        f"{field}: a confirmed passport requires surname and given_names",
                    )
                    core.require(bool(person["nationality"]), f"{field}.nationality is required")
                    core.require(bool(person["birth_date"]), f"{field}.birth_date is required")
            if number:
                key = _document_key(document_type, country_code, number)
                core.require(key not in local_document_keys, f"{field} repeats the same document")
                local_document_keys.add(key)
                owner = document_owner.get(key)
                core.require(
                    owner in (None, person_id),
                    f"document {country_code}/{number} appears under multiple people",
                )
                document_owner[key] = person_id
            document_count += 1

        local_account_keys: set[tuple[str, str, str]] = set()
        for account_position, account in enumerate(accounts):
            account_field = f"{field}.accounts[{account_position}]"
            core.require(isinstance(account, dict), f"{account_field} must be an object")
            unknown_account = sorted(
                set(account) - {"platform", "account_id", "phone", "media_labels"}
            )
            core.require(
                not unknown_account,
                f"{account_field}: unsupported fields: {', '.join(unknown_account)}",
            )
            platform = normalize_platform(account.get("platform"), field=f"{account_field}.platform")
            account_id = core.clean_text(account.get("account_id"))
            phone = core.clean_text(account.get("phone"))
            media_value = account.get("media_labels", [])
            media_labels = _clean_list(media_value, field=f"{account_field}.media_labels")
            for label in media_labels:
                core.require(
                    label in labels_by_class["chat_profile"],
                    f"{account_field}.media_labels contains a label not classified as chat_profile: {label}",
                )
                core.require(
                    label not in assigned_material,
                    f"material image assigned to multiple people or records: {label}",
                )
                assigned_material.add(label)
                derived_sources.append(inventory[label][0])
            if association_status == "confirmed":
                core.require(
                    bool(account_id or phone),
                    f"{account_field} requires account_id or phone for a confirmed record",
                )
                core.require(
                    bool(media_labels),
                    f"{account_field}.media_labels is required for a confirmed record",
                )
            account["platform"] = platform
            account["account_id"] = account_id
            account["phone"] = phone
            account["media_labels"] = media_labels
            keys = _account_keys(platform, account_id, phone)
            for key in keys:
                core.require(key not in local_account_keys, f"{field} repeats the same chat account")
                local_account_keys.add(key)
                owner = account_owner.get(key)
                core.require(
                    owner in (None, person_id),
                    f"chat account {platform}/{key[2]} appears under multiple people",
                )
                account_owner[key] = person_id
            account_count += 1

        source_messages_value = person.get("source_messages", [])
        core.require(
            isinstance(source_messages_value, list),
            f"{field}.source_messages must be a list",
        )
        source_messages = [core.clean_text(value) for value in source_messages_value]
        core.require(all(source_messages), f"{field}.source_messages cannot contain blanks")
        for label in derived_sources:
            if label not in source_messages:
                source_messages.append(label)
        core.require(source_messages, f"{field}.source_messages is required")
        core.require(
            len(source_messages) == len(set(source_messages)),
            f"{field}.source_messages repeats a label",
        )
        person["source_messages"] = _validate_source_messages(
            source_messages,
            field=f"{field}.source_messages",
            message_labels=message_labels,
            reviewed_through=reviewed_through,
        )

    material_labels = labels_by_class["document"] | labels_by_class["chat_profile"]
    unassigned_material = sorted(material_labels - assigned_material)
    if require_complete:
        core.require(
            not unassigned_material,
            "document and chat_profile media must be assigned to a person: "
            + str(unassigned_material[:20]),
        )

    return {
        "available_media": len(available),
        "missing_media": len(missing),
        "classified_media": len(decisions),
        "reference_media": len(labels_by_class["reference"]),
        "document_media": len(labels_by_class["document"]),
        "profile_media": len(labels_by_class["chat_profile"]),
        "people": len(people),
        "open_people": len(open_people_value),
        "pending_people": pending_people,
        "documents": document_count,
        "accounts": account_count,
        "unassigned_material_media": len(unassigned_material),
        "evidence_hashes_computed": evidence_hashes_computed,
        "evidence_hashes_reused": evidence_hashes_reused,
    }


def merge_review_batch(candidate: dict[str, Any], batch: Mapping[str, Any]) -> dict[str, Any]:
    people_updates = batch.get("people", [])
    core.require(isinstance(people_updates, list), "review batch people must be a list")
    update_ids: list[str] = []
    for position, person in enumerate(people_updates):
        core.require(isinstance(person, Mapping), f"review batch people[{position}] must be an object")
        person_id = core.clean_text(person.get("id"))
        core.require(bool(person_id), f"review batch people[{position}].id is required")
        update_ids.append(person_id)
    core.require(len(update_ids) == len(set(update_ids)), "review batch people repeats a person id")

    remove_value = batch.get("remove_person_ids", [])
    core.require(isinstance(remove_value, list), "review batch remove_person_ids must be a list")
    remove_ids = [core.clean_text(value) for value in remove_value]
    core.require(all(remove_ids), "review batch remove_person_ids cannot contain blanks")
    core.require(len(remove_ids) == len(set(remove_ids)), "review batch remove_person_ids repeats an id")
    core.require(
        not (set(update_ids) & set(remove_ids)),
        "review batch cannot update and remove the same person id",
    )

    people = candidate.get("people")
    core.require(isinstance(people, list), "decision people must be a list")
    existing_ids = {
        core.clean_text(person.get("id"))
        for person in people
        if isinstance(person, Mapping)
    }
    missing_ids = sorted(set(remove_ids) - existing_ids)
    core.require(not missing_ids, f"review batch removes unknown person ids: {missing_ids[:20]}")
    if remove_ids:
        people[:] = [
            person
            for person in people
            if not isinstance(person, Mapping)
            or core.clean_text(person.get("id")) not in set(remove_ids)
        ]
    positions = {
        core.clean_text(person.get("id")): position
        for position, person in enumerate(people)
        if isinstance(person, Mapping)
    }
    for person in people_updates:
        person_id = core.clean_text(person.get("id"))
        replacement = deepcopy(dict(person))
        if person_id in positions:
            people[positions[person_id]] = replacement
        else:
            positions[person_id] = len(people)
            people.append(replacement)
    if "open_people" in batch:
        core.require(isinstance(batch["open_people"], list), "review batch open_people must be a list")
        candidate["open_people"] = deepcopy(batch["open_people"])
    return candidate


def compile_ledger(
    normalized: Mapping[str, Any],
    decisions: Mapping[str, Mapping[str, Any]],
) -> tuple[dict[str, Any], dict[str, int]]:
    groups_by_key = {
        str(group.get("group_key")): group
        for group in normalized.get("groups", [])
        if isinstance(group, Mapping)
    }
    output_groups: list[dict[str, Any]] = []
    document_count = 0
    account_count = 0
    pending_people = 0
    for group in normalized.get("groups", []):
        if not isinstance(group, Mapping):
            continue
        group_key = str(group.get("group_key") or "")
        if group_key not in decisions:
            continue
        decision = decisions[group_key]
        inventory = media_inventory(groups_by_key[group_key])
        media_decisions = decision.get("media_decisions", {})
        people_output: list[dict[str, Any]] = []
        for raw_person in decision.get("people", []):
            person = deepcopy(dict(raw_person))
            for document in person.get("documents", []):
                document["media"] = [
                    {
                        "label": label,
                        "path": str(inventory[label][2].get("path") or ""),
                        "sha256": str(media_decisions[label].get("evidence_sha256") or ""),
                    }
                    for label in document.pop("media_labels", [])
                ]
                document_count += 1
            for account in person.get("accounts", []):
                account["media"] = [
                    {
                        "label": label,
                        "path": str(inventory[label][2].get("path") or ""),
                        "sha256": str(media_decisions[label].get("evidence_sha256") or ""),
                    }
                    for label in account.pop("media_labels", [])
                ]
                account_count += 1
            pending_people += person.get("association_status") == "pending"
            people_output.append(person)
        output_groups.append(
            {
                "group_key": group_key,
                "group_name": group.get("group_name"),
                "platform": group.get("platform"),
                "people": people_output,
            }
        )
    ledger = {
        "contract_version": OUTPUT_CONTRACT,
        "accounting_mode": "finance_materials",
        "timezone": normalized.get("timezone"),
        "normalized_source_fingerprint": normalized.get("source_fingerprint"),
        "groups": output_groups,
    }
    statistics = {
        "groups": len(output_groups),
        "people": sum(len(group["people"]) for group in output_groups),
        "documents": document_count,
        "accounts": account_count,
        "pending_people": int(pending_people),
    }
    return ledger, statistics


def validate_ledger(value: object) -> list[dict[str, Any]]:
    core.require(isinstance(value, dict), "finance-material ledger top level must be an object")
    core.require(value.get("contract_version") == OUTPUT_CONTRACT, "unsupported finance-material ledger")
    core.require(value.get("accounting_mode") == "finance_materials", "finance-material mode mismatch")
    core.require(value.get("timezone") == "Asia/Bangkok", "finance-material timezone must be Asia/Bangkok")
    groups = value.get("groups")
    core.require(isinstance(groups, list) and groups, "finance-material groups must be nonempty")
    keys: list[str] = []
    for position, group in enumerate(groups):
        field = f"finance-material groups[{position}]"
        core.require(isinstance(group, dict), f"{field} must be an object")
        group_key = core.clean_text(group.get("group_key"))
        core.require(bool(group_key), f"{field}.group_key is required")
        keys.append(group_key)
        core.require(isinstance(group.get("people"), list), f"{field}.people must be a list")
        for person_position, person in enumerate(group["people"]):
            person_field = f"{field}.people[{person_position}]"
            core.require(isinstance(person, dict), f"{person_field} must be an object")
            for collection_name in ("documents", "accounts"):
                core.require(
                    isinstance(person.get(collection_name), list),
                    f"{person_field}.{collection_name} must be a list",
                )
            for item in [*person["documents"], *person["accounts"]]:
                core.require(isinstance(item, dict), f"{person_field} material entry must be an object")
                media_items = item.get("media")
                core.require(isinstance(media_items, list), f"{person_field} material media must be a list")
                for media in media_items:
                    core.require(isinstance(media, dict), f"{person_field} media must be an object")
                    path = Path(str(media.get("path") or ""))
                    core.require(path.is_file(), f"finance-material source image is unavailable: {path}")
                    expected_hash = core.clean_text(media.get("sha256"))
                    core.require(bool(expected_hash), f"finance-material source image hash is missing: {path}")
                    core.require(
                        core.sha256_file(path) == expected_hash,
                        f"finance-material source image changed after review: {path}",
                    )
    core.require(len(keys) == len(set(keys)), "finance-material group keys must be unique")
    return groups


def _sheet_name(value: object, used: set[str]) -> str:
    cleaned = ILLEGAL_SHEET.sub("_", core.clean_text(value)) or "财务资料群"
    cleaned = cleaned[:31]
    candidate = cleaned
    suffix = 1
    while candidate.casefold() in used:
        suffix += 1
        marker = f"-{suffix}"
        candidate = cleaned[: 31 - len(marker)] + marker
    used.add(candidate.casefold())
    return candidate


def _group_sheet_name(group: Mapping[str, Any], used: set[str]) -> str:
    prefix = {"telegram": "TG", "whatsapp": "WA", "line": "LINE"}.get(
        core.clean_text(group.get("platform")).casefold(), "CHAT"
    )
    return _sheet_name(f"{prefix}-{core.clean_text(group.get('group_name'))}", used)


def _aligned(values: list[object]) -> str:
    cleaned = [core.clean_text(value) for value in values]
    return "\n".join(cleaned) if any(cleaned) else ""


def person_row(person: Mapping[str, Any]) -> tuple[list[Any], list[Mapping[str, Any]], list[Mapping[str, Any]]]:
    documents = [item for item in person.get("documents", []) if isinstance(item, Mapping)]
    accounts = [item for item in person.get("accounts", []) if isinstance(item, Mapping)]
    document_media = [
        media
        for document in documents
        for media in document.get("media", [])
        if isinstance(media, Mapping)
    ]
    profile_media = [
        media
        for account in accounts
        for media in account.get("media", [])
        if isinstance(media, Mapping)
    ]
    status = "待确认" if person.get("association_status") == "pending" else "已确认"
    return (
        [
            core.clean_text(person.get("name")),
            core.clean_text(person.get("surname")),
            core.clean_text(person.get("given_names")),
            _aligned([DOCUMENT_TYPE_LABELS.get(str(item.get("type")), "其他证件") for item in documents]),
            _aligned([item.get("country_code") for item in documents]),
            _aligned([item.get("number") for item in documents]),
            core.clean_text(person.get("nationality")),
            core.clean_text(person.get("birth_date")),
            "",
            _aligned([item.get("platform") for item in accounts]),
            _aligned([item.get("account_id") for item in accounts]),
            _aligned([item.get("phone") for item in accounts]),
            "",
            status,
            core.clean_text(person.get("note")),
        ],
        document_media,
        profile_media,
    )


def _gallery_rows(count: int) -> int:
    return max(1, math.ceil(count / 2))


def _add_gallery(
    worksheet: Any,
    *,
    row: int,
    column: int,
    media_items: list[Mapping[str, Any]],
) -> None:
    box_width = 118
    box_height = 94
    gap = 7
    for position, media in enumerate(media_items):
        path = Path(str(media.get("path") or ""))
        # openpyxl only closes the metadata probe immediately when given a string path.
        image = ExcelImage(str(path))
        scale = min(box_width / image.width, box_height / image.height)
        width = max(1, int(image.width * scale))
        height = max(1, int(image.height * scale))
        gallery_column = position % 2
        gallery_row = position // 2
        x_offset = 5 + gallery_column * (box_width + gap) + (box_width - width) // 2
        y_offset = 5 + gallery_row * (box_height + gap) + (box_height - height) // 2
        image.anchor = OneCellAnchor(
            _from=AnchorMarker(
                col=column - 1,
                colOff=pixels_to_EMU(x_offset),
                row=row - 1,
                rowOff=pixels_to_EMU(y_offset),
            ),
            ext=XDRPositiveSize2D(
                cx=pixels_to_EMU(width),
                cy=pixels_to_EMU(height),
            ),
        )
        worksheet.add_image(image)


def build_workbook(ledger: object) -> Workbook:
    groups = validate_ledger(ledger)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    thin = Side(style="thin", color=GRID_RGB)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = [18, 16, 20, 14, 12, 24, 15, 15, 39, 16, 25, 24, 39, 13, 34]

    for group in groups:
        worksheet = workbook.create_sheet(_group_sheet_name(group, used))
        worksheet.sheet_state = "visible"
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        for column, header in enumerate(HEADERS, start=1):
            cell = worksheet.cell(1, column, header)
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL_RGB)
            cell.font = Font(color=HEADER_FONT_RGB, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            worksheet.column_dimensions[get_column_letter(column)].width = widths[column - 1]
        worksheet.row_dimensions[1].height = 28

        for row_index, person in enumerate(group.get("people", []), start=2):
            values, document_media, profile_media = person_row(person)
            for column, value in enumerate(values, start=1):
                cell = worksheet.cell(row_index, column, value)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
                if column in TEXT_COLUMNS:
                    cell.number_format = "@"
                if row_index % 2 == 1:
                    cell.fill = PatternFill("solid", fgColor=BODY_ALT_FILL_RGB)
            if person.get("association_status") == "pending":
                for cell in worksheet[row_index]:
                    cell.fill = PatternFill("solid", fgColor=PENDING_FILL_RGB)
            gallery_lines = max(_gallery_rows(len(document_media)), _gallery_rows(len(profile_media)))
            worksheet.row_dimensions[row_index].height = max(58, gallery_lines * 76)
            _add_gallery(
                worksheet,
                row=row_index,
                column=DOCUMENT_IMAGE_COLUMN,
                media_items=document_media,
            )
            _add_gallery(
                worksheet,
                row=row_index,
                column=PROFILE_IMAGE_COLUMN,
                media_items=profile_media,
            )
        last_column = get_column_letter(len(HEADERS))
        worksheet.auto_filter.ref = f"A1:{last_column}{max(1, worksheet.max_row)}"
    core.require(bool(workbook.sheetnames), "finance-material workbook must contain a group sheet")
    return workbook


def _comparable(value: object) -> str:
    return "" if value is None else str(value)


def _image_anchor(image: Any) -> tuple[int, int] | None:
    anchor = getattr(image, "anchor", None)
    marker = getattr(anchor, "_from", None)
    if marker is None:
        return None
    return int(marker.row) + 1, int(marker.col) + 1


def check_workbook(workbook_path: Path, ledger: object) -> list[str]:
    errors: list[str] = []
    try:
        groups = validate_ledger(ledger)
    except (OSError, TypeError, ValueError) as exc:
        return [str(exc)]
    expected_used: set[str] = set()
    expected_names = [_group_sheet_name(group, expected_used) for group in groups]
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        if getattr(workbook, "_external_links", []):
            errors.append("finance-material workbook contains forbidden external links")
        if getattr(workbook, "vba_archive", None) is not None:
            errors.append("finance-material workbook contains forbidden macros")
        if workbook.sheetnames != expected_names:
            errors.append(
                f"finance-material sheet names mismatch: expected {expected_names!r}, got {workbook.sheetnames!r}"
            )
        for group, sheet_name in zip(groups, expected_names):
            worksheet = workbook[sheet_name]
            headers = [worksheet.cell(1, column).value for column in range(1, len(HEADERS) + 1)]
            if headers != HEADERS:
                errors.append(f"{sheet_name}: finance-material header mismatch")
            if str(worksheet.freeze_panes or "") != "A2":
                errors.append(f"{sheet_name}: first row must be frozen at A2")
            expected_rows = [person_row(person) for person in group.get("people", [])]
            if max(0, worksheet.max_row - 1) != len(expected_rows):
                errors.append(
                    f"{sheet_name}: expected {len(expected_rows)} people rows, got {max(0, worksheet.max_row - 1)}"
                )
            expected_images: Counter[tuple[int, int]] = Counter()
            for row_index, (values, document_media, profile_media) in enumerate(expected_rows, start=2):
                for column, expected in enumerate(values, start=1):
                    cell = worksheet.cell(row_index, column)
                    if _comparable(cell.value) != _comparable(expected):
                        errors.append(
                            f"{sheet_name}!{cell.coordinate}: expected {expected!r}, got {cell.value!r}"
                        )
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        errors.append(f"{sheet_name}!{cell.coordinate}: formulas are forbidden")
                    if column in TEXT_COLUMNS and cell.number_format != "@":
                        errors.append(f"{sheet_name}!{cell.coordinate}: identifiers must use text format")
                expected_images[(row_index, DOCUMENT_IMAGE_COLUMN)] += len(document_media)
                expected_images[(row_index, PROFILE_IMAGE_COLUMN)] += len(profile_media)
            actual_images: Counter[tuple[int, int]] = Counter()
            for image in worksheet._images:
                anchor = _image_anchor(image)
                if anchor is None:
                    errors.append(f"{sheet_name}: image has no cell anchor")
                else:
                    actual_images[anchor] += 1
            if actual_images != expected_images:
                errors.append(
                    f"{sheet_name}: image anchors mismatch: expected {dict(expected_images)}, got {dict(actual_images)}"
                )
            if worksheet.max_column > len(HEADERS):
                for row in worksheet.iter_rows(min_col=len(HEADERS) + 1):
                    if any(cell.value not in (None, "") for cell in row):
                        errors.append(f"{sheet_name}: values exist beyond the final finance-material column")
                        break
    finally:
        workbook.close()
    return errors
