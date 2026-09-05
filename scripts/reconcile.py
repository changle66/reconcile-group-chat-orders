#!/usr/bin/env python3
"""Run fresh or cross-day group-chat reconciliation through review and finish."""

from __future__ import annotations

import argparse
import copy
import json
import platform as host_platform
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Iterable, Mapping
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook

import build_workbook
import check_workbook
import core
import extract_line_android_miui
import extract_line_ios
import finance_materials
import large_daily
import normalize_exports
import simple_ledger
import store_ledger


RUN_CONTRACT = "group-chat-reconcile-run/1.0"
RUNTIME_CONTRACT = "reconcile-current-runtime/1.0"
AMOUNT_POLICY = "receiver-actual-received/1.0"
WORKBOOK_COMPATIBILITY = "wps-xlsx-static-values/1.0"
DECISION_CONTRACT = "group-chat-decision/3.2"
REVIEW_PAGE_CONTRACT = "group-chat-review-page/2.1"
REVIEW_BATCH_CONTRACT = "group-chat-review-batch/1.1"
MEDIA_QUEUE_CONTRACT = "group-chat-media-queue/2.1"
MEDIA_OBSERVATION_CONTRACT = "group-chat-media-observation/1.0"
MEDIA_OBSERVATION_CACHE_CONTRACT = "group-chat-media-observation-cache/1.0"
MEDIA_BATCH_POLICY_CONTRACT = "group-chat-media-batch-policy/1.0"
OCR_CANDIDATE_CONFIG_CONTRACT = "group-chat-ocr-candidate-config/1.0"
OCR_CANDIDATE_CONTRACT = "group-chat-ocr-candidate/1.0"
OCR_CANDIDATE_CACHE_CONTRACT = "group-chat-ocr-candidate-cache/1.0"
OCR_WORKER_CONTRACT = "group-chat-ocr-worker/1.0"
EDIT_CONTROL_MODE = "review-apply-batch/1.1"
LEGACY_DECISION_CONTRACT = "group-chat-decision/3.1"
LEGACY_EDIT_CONTROL_MODE = "review-apply-batch/1.0"
NORMALIZER_VERSION = "reconcile-start/1.0"
DEFAULT_PAGE_SIZE = 200
MAX_PAGE_SIZE = 500
DEFAULT_REVIEW_PAGE_OUTPUT_CHAR_BUDGET = 14_000
ORDER_MEDIA_BATCH_LIMIT = 9
FINANCE_MEDIA_BATCH_LIMIT = 4
ORDER_MEDIA_BATCH_MIN = 4
ORDER_MEDIA_BATCH_MAX = 12
FINANCE_MEDIA_BATCH_MIN = 2
FINANCE_MEDIA_BATCH_MAX = 6
MEDIA_RECHECK_QUEUE_LIMIT = 20
OCR_CANDIDATE_TEXT_LIMIT = 800
OCR_CANDIDATE_PAGE_LIMIT = 20
OCR_WORKER_TIMEOUT_SECONDS = 60
MEDIA_REVIEW_STATUSES = frozenset(
    {"clear", "recheck_required", "rechecked_unreadable"}
)
MEDIA_RECHECK_REASONS = frozenset(
    {
        "small_text",
        "blurred",
        "cropped",
        "obscured",
        "label_mapping_uncertain",
        "conflicting_visible_fields",
        "thumbnail_only",
        "read_failure",
        "amount_unreadable",
        "payee_unreadable",
        "document_field_unreadable",
        "account_field_unreadable",
        "other",
    }
)
MEDIA_VIEW_METRIC_FIELDS = (
    "view_batches",
    "opened_images",
    "failed_images",
    "single_image_rechecks",
    "elapsed_ms",
    "observation_cache_reuses",
)
UNKNOWN_PAYEE_WARNING_MIN_TRANSFERS = 10
FLOW_SIDES = frozenset({"payment", "payment_refund", "payout", "recovery", "unknown"})
SIDE_EXCEPTION_SIDES = {
    "relayed_customer_payment": "payment",
    "relayed_internal_payout": "payout",
    "explicit_payment_refund": "payment_refund",
    "explicit_recovery": "recovery",
}
ENTRY_KINDS = frozenset({"transfer", "cash"})
ENTRY_RESULTS = frozenset({"completed", "failed", "pending", "not_shown", "unknown"})
AMOUNT_STATES = frozenset({"clear", "partial", "unreadable"})
AMOUNT_BASES = frozenset(
    {"receiver_received", "cash_face_value", "attempted_not_received"}
)
NORMAL_PROCESSING_STATUS_MARKERS = (
    "确认中",
    "处理中",
    "等待确认",
    "待区块确认",
    "pending",
    "processing",
    "confirming",
    "unconfirmed",
)
FAILURE_STATUS_MARKERS = (
    "失败",
    "取消",
    "拒绝",
    "作废",
    "无效",
    "风控导致未完成",
    "风控未完成",
    "风控未通过",
    "风控拦截",
    "failed",
    "declined",
    "rejected",
    "cancelled",
    "canceled",
    "invalid",
)
MISSING_CONTEXT_NOTE_MARKERS = (
    "客户",
    "方向",
    "付款",
    "回款",
    "汇率",
    "计价",
    "金额",
    "币种",
    "收款方",
    "上下文",
    "无法辨认",
    "未说明",
    "缺少",
)
PRICING_EXPECTED_KINDS = frozenset(
    {"explicit", "calculated_from_terms", "unknown"}
)
PRICING_UNKNOWN_REASONS = frozenset(
    {"not_stated", "conflicting_authority", "incomplete_formula"}
)
STATUS_CLASS = {
    "completed": "completed",
    "failed": "failed",
    "pending": "pending",
    "not_shown": "blank",
    "unknown": "unknown",
}
ORDER_LIFECYCLE_STATUSES = frozenset(
    {"completed", "pending_next_day", "cancelled"}
)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)

    start = commands.add_parser("start", help="create a new run from raw exports")
    start.add_argument("inputs", nargs="+", type=Path, help="raw export files or roots")
    start.add_argument("--work", required=True, type=Path, help="new, non-existing run directory")
    start.add_argument(
        "--mode",
        choices=("small", "large", "finance", "store-ledger"),
        default="small",
        help=(
            "small selects the existing name filter; large selects every group except 小额出 and 财务资料群 (门店开票群也按大额订单语义独立处理); "
            "finance selects 财务资料群 for identity and chat-account records; "
            "store-ledger selects 门店开票群 for signed internal currency movements"
        ),
    )
    start.add_argument("--contains", default="小额", help="required text in the parsed group name")
    start.add_argument("--date", help="optional accounting date in YYYY-MM-DD")
    start.add_argument(
        "--from",
        dest="accounting_from",
        metavar="YYYY-MM-DD HH:MM",
        help="optional inclusive accounting start time",
    )
    start.add_argument(
        "--to",
        dest="accounting_to",
        metavar="YYYY-MM-DD HH:MM",
        help="optional exclusive accounting end time",
    )
    start.add_argument("--timezone", default="Asia/Bangkok")
    start.add_argument(
        "--roster",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "config" / "roster.yaml",
    )
    start.add_argument("--line-backup", action="append", default=[], type=Path)
    start.add_argument("--line-android-backup", action="append", default=[], type=Path)
    start.add_argument("--line-self-name", default="LINE_SELF")
    start.add_argument(
        "--ocr-candidates",
        action=argparse.BooleanOptionalAction,
        default=None,
        help=(
            "enable or disable local OCR hints; default is disabled on all platforms"
        ),
    )

    review = commands.add_parser("review", help="read, validate, or seal one group review")
    review.add_argument("work", type=Path)
    review_actions = review.add_subparsers(dest="review_action", required=True)
    next_page = review_actions.add_parser("next", help="return the next compact chronological page")
    next_page.add_argument("--group", help="group key, run label, or unique exact group name")
    next_page.add_argument(
        "--limit",
        type=int,
        default=None,
        help="legacy exact message count; omit for an adaptive complete page",
    )
    status = review_actions.add_parser("status", help="show review progress")
    status.add_argument("--group", help="optional group selector")
    audit = review_actions.add_parser("audit", help="report review-risk diagnostics without writing files")
    audit.add_argument("--group", help="optional group selector")
    check = review_actions.add_parser("check", help="validate the current decision data")
    check.add_argument("--group", required=True)
    upgrade_checkpoints = review_actions.add_parser(
        "upgrade-checkpoints",
        help="preserve decisions but reset legacy read progress for safe page commits",
    )
    upgrade_checkpoints.add_argument("--group", required=True)
    seal = review_actions.add_parser("seal", help="validate and seal a completely reviewed group")
    seal.add_argument("--group", required=True)
    apply_batch = review_actions.add_parser(
        "apply-batch",
        help="atomically merge one schema-validated semantic review batch",
    )
    apply_batch.add_argument("--group", required=True)
    apply_batch.add_argument("--input", required=True, type=Path)

    finish = commands.add_parser("finish", help="compile, verify, and publish the final workbook")
    finish.add_argument("work", type=Path)
    finish.add_argument("-o", "--output", required=True, type=Path)
    finish.add_argument(
        "--template",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
    )
    finish.add_argument(
        "--replace-predecessor",
        action="store_true",
        help="atomically replace only the predecessor workbook registered by roll-forward",
    )

    roll_forward = commands.add_parser(
        "roll-forward",
        help="continue a finished run with the immediately following accounting period",
    )
    roll_forward.add_argument("previous_work", type=Path)
    roll_forward.add_argument("inputs", nargs="+", type=Path)
    roll_forward.add_argument("--previous-output", required=True, type=Path)
    roll_forward.add_argument("--work", required=True, type=Path)
    roll_forward.add_argument("--date", required=True)
    roll_forward.add_argument("--confirmed-amendments", type=Path)
    return parser.parse_args(argv)


def _run_path(work: Path) -> Path:
    return work / "run.json"


def _ocr_cache_path(work: Path) -> Path:
    return work / "cache" / "ocr_candidates.json"


def _host_platform_key(system_name: object | None = None) -> str:
    value = core.clean_text(
        host_platform.system() if system_name is None else system_name
    ).casefold()
    if value.startswith("win"):
        return "windows"
    if value in {"darwin", "mac", "macos", "osx"}:
        return "macos"
    if value == "linux":
        return "linux"
    return value or "other"


def _resolve_ocr_candidate_config(
    requested_value: object,
    *,
    system_name: object | None = None,
) -> dict[str, Any]:
    core.require(
        requested_value is None or isinstance(requested_value, bool),
        "OCR candidate setting must be automatic, enabled, or disabled",
    )
    platform_key = _host_platform_key(system_name)
    requested = (
        "auto"
        if requested_value is None
        else "enabled"
        if requested_value
        else "disabled"
    )
    default_enabled = False
    enabled = default_enabled if requested == "auto" else requested == "enabled"
    return {
        "contract_version": OCR_CANDIDATE_CONFIG_CONTRACT,
        "requested": requested,
        "host_platform": platform_key,
        "default_enabled": default_enabled,
        "enabled": enabled,
        "resolved_by": "platform_default" if requested == "auto" else "explicit",
    }


def _normalize_ocr_candidate_config(value: object, *, field: str) -> dict[str, Any]:
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    allowed = {
        "contract_version",
        "requested",
        "host_platform",
        "default_enabled",
        "enabled",
        "resolved_by",
    }
    unknown = sorted(set(value) - allowed)
    core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
    core.require(
        value.get("contract_version") == OCR_CANDIDATE_CONFIG_CONTRACT,
        f"{field} has an unsupported contract",
    )
    requested = core.clean_text(value.get("requested"))
    core.require(
        requested in {"auto", "enabled", "disabled"},
        f"{field}.requested must be auto, enabled, or disabled",
    )
    platform_key = core.clean_text(value.get("host_platform")).casefold()
    core.require(bool(platform_key), f"{field}.host_platform is required")
    default_enabled = value.get("default_enabled")
    enabled = value.get("enabled")
    core.require(
        isinstance(default_enabled, bool) and isinstance(enabled, bool),
        f"{field}.default_enabled and enabled must be booleans",
    )
    resolved_by = core.clean_text(value.get("resolved_by"))
    core.require(
        resolved_by in {"platform_default", "explicit", "legacy_safe_default"},
        f"{field}.resolved_by is invalid",
    )
    if resolved_by == "platform_default":
        core.require(requested == "auto", f"{field} platform default must use auto")
        core.require(enabled == default_enabled, f"{field} platform default is inconsistent")
    elif resolved_by == "explicit":
        core.require(requested != "auto", f"{field} explicit setting cannot use auto")
        core.require(
            enabled == (requested == "enabled"),
            f"{field} explicit setting is inconsistent",
        )
    return {
        "contract_version": OCR_CANDIDATE_CONFIG_CONTRACT,
        "requested": requested,
        "host_platform": platform_key,
        "default_enabled": default_enabled,
        "enabled": enabled,
        "resolved_by": resolved_by,
    }


def _ocr_candidate_config_for_run(run: Mapping[str, Any]) -> dict[str, Any]:
    value = run.get("ocr_candidates")
    if value is None:
        # Old work directories must never begin loading OCR merely because they
        # are resumed on a different host.
        return {
            "contract_version": OCR_CANDIDATE_CONFIG_CONTRACT,
            "requested": "disabled",
            "host_platform": "legacy",
            "default_enabled": False,
            "enabled": False,
            "resolved_by": "legacy_safe_default",
        }
    return _normalize_ocr_candidate_config(value, field="run.ocr_candidates")


def _snapshot_path(work: Path) -> Path:
    return work / "snapshot" / "normalized.json"


def _load_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8-sig"))
    core.require(isinstance(value, dict), f"{path}: top level must be an object")
    return value


def _load_run(work: Path) -> dict[str, Any]:
    path = _run_path(work)
    core.require(path.is_file(), f"run manifest does not exist: {path}")
    run = _load_json(path)
    core.require(run.get("contract_version") == RUN_CONTRACT, "unsupported run contract")
    core.require(isinstance(run.get("groups"), list), "run groups must be a list")
    expected_amount_policy = (
        store_ledger.AMOUNT_POLICY
        if core.clean_text(run.get("group_mode")).casefold() == store_ledger.GROUP_MODE
        else AMOUNT_POLICY
    )
    for field, expected in (
        ("runtime_contract", RUNTIME_CONTRACT),
        ("amount_policy", expected_amount_policy),
        ("workbook_compatibility", WORKBOOK_COMPATIBILITY),
    ):
        if field in run:
            core.require(run.get(field) == expected, f"run {field} is unsupported")
    if "ocr_candidates" in run:
        run["ocr_candidates"] = _normalize_ocr_candidate_config(
            run["ocr_candidates"], field="run.ocr_candidates"
        )
    return run


def _load_snapshot(work: Path, run: Mapping[str, Any]) -> dict[str, Any]:
    path = _snapshot_path(work)
    core.require(path.is_file(), f"normalized snapshot does not exist: {path}")
    expected = core.clean_text(run.get("snapshot_sha256"))
    actual = core.sha256_file(path)
    core.require(actual == expected, "normalized snapshot changed after start; create a new run")
    normalized = core.load_normalized(path)
    core.require(
        normalized.get("source_fingerprint") == run.get("normalized_source_fingerprint"),
        "run manifest and normalized snapshot disagree",
    )
    return normalized


def _source_document(
    inputs: Iterable[Path],
    *,
    zone: ZoneInfo,
    roster: Mapping[str, Any],
) -> dict[str, Any] | None:
    try:
        telegram_sources, whatsapp_sources = normalize_exports.discover_sources(inputs)
    except ValueError as exc:
        if str(exc) == "no Telegram result.json or WhatsApp .txt exports found":
            return None
        raise
    groups: list[dict[str, Any]] = []
    warnings: list[str] = []
    fingerprint_paths: list[Path] = []
    for source in telegram_sources:
        group, local_warnings, local_paths = normalize_exports.normalize_telegram(
            source,
            zone=zone,
            roster=dict(roster),
        )
        groups.append(group)
        warnings.extend(local_warnings)
        fingerprint_paths.extend(local_paths)
    for source in whatsapp_sources:
        group, local_warnings, local_paths = normalize_exports.normalize_whatsapp(
            source,
            zone=zone,
            roster=dict(roster),
        )
        groups.append(group)
        warnings.extend(local_warnings)
        fingerprint_paths.extend(local_paths)
    group_keys = [str(group["group_key"]) for group in groups]
    core.require(len(group_keys) == len(set(group_keys)), "duplicate group export; provide one snapshot per group")
    groups.sort(key=lambda item: (str(item.get("platform")), str(item.get("group_name")), str(item.get("group_key"))))
    source_fingerprint = core.fingerprint_files(
        fingerprint_paths,
        context={
            "normalizer": NORMALIZER_VERSION,
            "timezone": str(zone.key),
            "roster": roster,
            "groups": group_keys,
        },
    )
    messages = [message for group in groups for message in group.get("messages", [])]
    return {
        "contract_version": core.NORMALIZED_CONTRACT,
        "timezone": str(zone.key),
        "source_fingerprint": source_fingerprint,
        "groups": groups,
        "statistics": {
            "groups": len(groups),
            "messages": len(messages),
            "excluded_messages": sum(bool(message.get("excluded_from_accounting")) for message in messages),
            "media": sum(len(message.get("media", [])) for message in messages),
            "source_files": len(telegram_sources) + len(whatsapp_sources),
            "fingerprinted_files": len({path.resolve() for path in fingerprint_paths}),
        },
        "warnings": warnings,
        "skipped_files": [],
    }


def _generated_path(path: Path) -> bool:
    return any(
        marker.casefold() in part.casefold()
        for part in path.parts
        for marker in normalize_exports.GENERATED_PATH_MARKERS
    )


def _discover_line_backups(inputs: Iterable[Path], explicit: Iterable[Path]) -> list[Path]:
    candidates: set[Path] = set()
    for raw in explicit:
        candidates.add(extract_line_ios.find_device_dir(raw).resolve())
    for raw in inputs:
        path = raw.resolve()
        if path.is_file() and path.name.casefold() == "manifest.db":
            candidates.add(path.parent)
        elif path.is_dir():
            direct = path / "Manifest.db"
            if direct.is_file() and not _generated_path(direct):
                candidates.add(path)
            for manifest in path.rglob("Manifest.db"):
                if not _generated_path(manifest):
                    candidates.add(manifest.parent.resolve())
    return sorted(candidates, key=lambda item: str(item).casefold())


def _discover_line_android_backups(
    inputs: Iterable[Path], explicit: Iterable[Path]
) -> list[Path]:
    candidates: set[Path] = set()
    for raw in explicit:
        candidates.add(extract_line_android_miui.find_backup_file(raw).resolve())
    for raw in inputs:
        path = raw.resolve()
        if path.is_file():
            if extract_line_android_miui.is_line_android_backup(path) and not _generated_path(path):
                candidates.add(path)
            continue
        if not path.is_dir():
            continue
        for backup in path.rglob("*.bak"):
            if not _generated_path(backup) and extract_line_android_miui.is_line_android_backup(backup):
                candidates.add(backup.resolve())
    return sorted(candidates, key=lambda item: str(item).casefold())


def _large_group_name(value: object) -> bool:
    normalized = core.normalize_name(value)
    return all(
        core.normalize_name(excluded) not in normalized
        for excluded in ("小额出", "财务资料群")
    )


def _group_name_selected(value: object, *, group_mode: str, contains: str) -> bool:
    if group_mode == "large":
        return _large_group_name(value)
    if group_mode == finance_materials.GROUP_MODE:
        return finance_materials.group_name_selected(value)
    if group_mode == store_ledger.GROUP_MODE:
        return store_ledger.group_name_selected(value)
    return contains.casefold() in str(value or "").casefold()


def _line_has_matching_group(
    device_dir: Path,
    contains: str,
    *,
    group_mode: str,
) -> bool:
    manifest = extract_line_ios.sqlite_ro(device_dir / "Manifest.db")
    line_db = None
    group_db = None
    try:
        line_row = extract_line_ios.manifest_row_for_suffix(manifest, "/Messages/Line.sqlite")
        group_row = extract_line_ios.manifest_row_for_suffix(manifest, "/Messages/UnifiedGroup.sqlite")
        line_path = extract_line_ios.physical_backup_file(device_dir, str(line_row["fileID"]))
        group_path = extract_line_ios.physical_backup_file(device_dir, str(group_row["fileID"]))
        line_db = extract_line_ios.sqlite_ro(line_path)
        group_db = extract_line_ios.sqlite_ro(group_path)
        groups = extract_line_ios.available_groups(line_db, group_db)
        return any(
            _group_name_selected(
                group.get("group_name"),
                group_mode=group_mode,
                contains=contains,
            )
            for group in groups
        )
    finally:
        if line_db is not None:
            line_db.close()
        if group_db is not None:
            group_db.close()
        manifest.close()


def _line_documents(
    work: Path,
    backups: Iterable[Path],
    *,
    contains: str,
    group_mode: str,
    timezone_name: str,
    roster_path: Path,
    self_name: str,
) -> list[dict[str, Any]]:
    documents: list[dict[str, Any]] = []
    for index, backup in enumerate(backups, start=1):
        if not _line_has_matching_group(backup, contains, group_mode=group_mode):
            continue
        output = work / "snapshot" / "sources" / f"line_{index:02d}.json"
        media_dir = work / "snapshot" / "line_media" / f"backup_{index:02d}"
        extraction_args = [
                str(backup),
                "-o",
                str(output),
                "--timezone",
                timezone_name,
                "--self-name",
                self_name,
                "--roster",
                str(roster_path),
                "--media-dir",
                str(media_dir),
            ]
        if group_mode == "large":
            extraction_args.append("--all-group-chats")
        else:
            extraction_args.extend(["--group-pattern", re.escape(contains)])
        exit_code = extract_line_ios.main(extraction_args)
        core.require(exit_code == 0, f"LINE extraction failed for {backup}")
        documents.append(core.load_normalized(output))
    return documents


def _line_android_documents(
    work: Path,
    backups: Iterable[Path],
    *,
    contains: str,
    group_mode: str,
    timezone_name: str,
    roster_path: Path,
) -> list[dict[str, Any]]:
    documents: list[dict[str, Any]] = []
    for index, backup in enumerate(backups, start=1):
        groups = extract_line_android_miui.available_groups_for_backup(backup)
        if not any(
            _group_name_selected(
                group.get("group_name"),
                group_mode=group_mode,
                contains=contains,
            )
            for group in groups
        ):
            continue
        output = work / "snapshot" / "sources" / f"line_android_{index:02d}.json"
        media_dir = work / "snapshot" / "line_android_media" / f"backup_{index:02d}"
        extraction_args = [
                str(backup),
                "-o",
                str(output),
                "--timezone",
                timezone_name,
                "--roster",
                str(roster_path),
                "--media-dir",
                str(media_dir),
            ]
        if group_mode == "large":
            extraction_args.append("--all-group-chats")
        else:
            extraction_args.extend(["--group-pattern", re.escape(contains)])
        exit_code = extract_line_android_miui.main(extraction_args)
        core.require(exit_code == 0, f"LINE Android extraction failed for {backup}")
        documents.append(core.load_normalized(output))
    return documents


def _merge_documents(
    documents: list[dict[str, Any]],
    *,
    contains: str,
    group_mode: str,
    timezone_name: str,
) -> dict[str, Any]:
    core.require(bool(documents), "no supported raw chat exports were found")
    groups: list[dict[str, Any]] = []
    warnings: list[str] = []
    source_fingerprints: list[str] = []
    seen: set[str] = set()
    for document in documents:
        core.require(document.get("timezone") == timezone_name, "normalized source timezone mismatch")
        source_fingerprints.append(str(document.get("source_fingerprint") or ""))
        warnings.extend(str(item) for item in document.get("warnings", []))
        for group in document.get("groups", []):
            group_name = group.get("group_name")
            selected = _group_name_selected(
                group_name,
                group_mode=group_mode,
                contains=contains,
            )
            if not selected:
                continue
            key = str(group.get("group_key") or "")
            core.require(key and key not in seen, f"duplicate group snapshot: {key}")
            seen.add(key)
            groups.append(copy.deepcopy(group))
    if group_mode == "large":
        core.require(
            bool(groups),
            "no large groups remain after excluding 小额出 and 财务资料群",
        )
    else:
        core.require(bool(groups), f"no parsed group name contains {contains!r}")
    groups.sort(key=lambda item: (str(item.get("platform")), str(item.get("group_name")), str(item.get("group_key"))))
    messages = [message for group in groups for message in group.get("messages", [])]
    media = [item for message in messages for item in message.get("media", [])]
    source_fingerprint = core.fingerprint_json(
        {
            "normalizer": NORMALIZER_VERSION,
            "contains": contains,
            "group_mode": group_mode,
            "sources": sorted(source_fingerprints),
            "groups": [group["group_key"] for group in groups],
        }
    )
    return {
        "contract_version": core.NORMALIZED_CONTRACT,
        "timezone": timezone_name,
        "source_fingerprint": source_fingerprint,
        "groups": groups,
        "statistics": {
            "groups": len(groups),
            "messages": len(messages),
            "excluded_messages": sum(bool(message.get("excluded_from_accounting")) for message in messages),
            "media": len(media),
            "available_media": sum(item.get("availability") == "available" for item in media),
            "missing_media": sum(item.get("availability") == "missing" for item in media),
            "normalized_inputs": len(documents),
        },
        "warnings": warnings,
        "skipped_files": [],
    }


def _decision_filename(group_key: str) -> str:
    return f"decision_{group_key.replace(':', '-')}.json"


def _empty_media_observation_cache() -> dict[str, Any]:
    return {
        "contract_version": MEDIA_OBSERVATION_CACHE_CONTRACT,
        "observation_contract": MEDIA_OBSERVATION_CONTRACT,
        "entries": {},
    }


def _empty_media_view_metrics() -> dict[str, int]:
    return {field: 0 for field in MEDIA_VIEW_METRIC_FIELDS}


def _decision_template(
    normalized: Mapping[str, Any],
    group: Mapping[str, Any],
    *,
    group_mode: str = "small",
) -> dict[str, Any]:
    evidence = _media_inventory(group)
    contract_version = (
        finance_materials.DECISION_CONTRACT
        if group_mode == finance_materials.GROUP_MODE
        else store_ledger.DECISION_CONTRACT
        if group_mode == store_ledger.GROUP_MODE
        else large_daily.DECISION_CONTRACT
        if group_mode == "large"
        else DECISION_CONTRACT
    )
    decision = {
        "contract_version": contract_version,
        "normalized_source_fingerprint": normalized.get("source_fingerprint"),
        "group_fingerprint": core.fingerprint_json(group),
        "group_key": group.get("group_key"),
        "group_name": group.get("group_name"),
        "platform": group.get("platform"),
        "message_count": len(group.get("messages", [])),
        "evidence_media_count": len(evidence),
        "reviewed_through": 0,
        "read_complete": not bool(group.get("messages", [])),
        "sealed": False,
        "sealed_decision_fingerprint": None,
        "media_decisions": {},
        "media_observation_cache": _empty_media_observation_cache(),
        "media_view_metrics": _empty_media_view_metrics(),
    }
    if group_mode == finance_materials.GROUP_MODE:
        decision.update(
            {
                "group_mode": finance_materials.GROUP_MODE,
                "people": [],
                "open_people": [],
            }
        )
    elif group_mode == store_ledger.GROUP_MODE:
        decision.update(
            {
                "group_mode": store_ledger.GROUP_MODE,
                "records": [],
                "open_records": [],
                "balance_snapshots": [],
            }
        )
    elif group_mode == "large":
        decision.update(
            {
                "group_mode": "large",
                "orders": [],
                "open_orders": [],
                "balance_links": [],
                "settlement_allocations": [],
                "unknown_payee_reviewed_entry_ids": [],
            }
        )
    else:
        decision.update(
            {
                "orders": [],
                "open_orders": [],
                "balance_links": [],
                "settlement_allocations": [],
                "unknown_payee_reviewed_entry_ids": [],
            }
        )
    decision["edit_control"] = _new_edit_control(_semantic_fingerprint(decision))
    return decision


def _platform_prefix(platform: object) -> str:
    return {"Telegram": "TG", "WhatsApp": "WA", "LINE": "LINE"}.get(str(platform), "CHAT")


def _normalize_accounting_date(value: object) -> str | None:
    if value is None:
        return None
    cleaned = core.clean_text(value)
    error = "--date must be a valid date in YYYY-MM-DD format"
    core.require(bool(cleaned), error)
    try:
        parsed = datetime.strptime(cleaned, "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError(error) from exc
    core.require(parsed.isoformat() == cleaned, error)
    return cleaned


def _normalize_accounting_time(value: object, *, option: str) -> str | None:
    if value is None:
        return None
    cleaned = core.clean_text(value)
    error = f"{option} must be a valid time in YYYY-MM-DD HH:MM format"
    core.require(bool(cleaned), error)
    try:
        parsed = datetime.strptime(cleaned, "%Y-%m-%d %H:%M")
    except ValueError as exc:
        raise ValueError(error) from exc
    core.require(parsed.strftime("%Y-%m-%d %H:%M") == cleaned, error)
    return cleaned


def _filter_normalized_window(
    normalized: dict[str, Any],
    *,
    start_at: datetime,
    end_at: datetime,
    fingerprint_fields: Mapping[str, object],
    preserve_reply_context: bool = False,
) -> dict[str, Any]:
    core.require(start_at.tzinfo is not None and end_at.tzinfo is not None, "accounting window requires timezone")
    filtered = copy.deepcopy(normalized)
    for group in filtered.get("groups", []):
        messages = list(group.get("messages", []))
        in_window_ids = {
            str(message.get("message_id") or "")
            for message in messages
            if start_at
            <= datetime.fromisoformat(str(message.get("timestamp"))).astimezone(start_at.tzinfo)
            < end_at
        }
        retained_ids = set(in_window_ids)
        if preserve_reply_context:
            changed = True
            while changed:
                changed = False
                for message in messages:
                    message_id = str(message.get("message_id") or "")
                    reply_id = str(message.get("reply_to_message_id") or "")
                    if (
                        (message_id in retained_ids and reply_id and reply_id not in retained_ids)
                        or (reply_id in retained_ids and message_id not in retained_ids)
                    ):
                        retained_ids.update(item for item in (message_id, reply_id) if item)
                        changed = True
        retained: list[dict[str, Any]] = []
        for message in messages:
            message_id = str(message.get("message_id") or "")
            if message_id not in retained_ids:
                continue
            if message_id not in in_window_ids:
                message["accounting_context_only"] = True
                message["excluded_from_accounting"] = True
            retained.append(message)
        group["messages"] = retained
    messages = [message for group in filtered.get("groups", []) for message in group.get("messages", [])]
    media = [item for message in messages for item in message.get("media", [])]
    filtered["source_fingerprint"] = core.fingerprint_json(
        {
            "base_source_fingerprint": normalized.get("source_fingerprint"),
            **fingerprint_fields,
        }
    )
    statistics = dict(normalized.get("statistics", {}))
    statistics.update(
        {
            "groups": len(filtered.get("groups", [])),
            "messages": len(messages),
            "excluded_messages": sum(
                bool(message.get("excluded_from_accounting")) for message in messages
            ),
            "media": len(media),
            "available_media": sum(item.get("availability") == "available" for item in media),
            "missing_media": sum(item.get("availability") == "missing" for item in media),
        }
    )
    filtered["statistics"] = statistics
    filtered["accounting_window"] = {
        "start": start_at.isoformat(),
        "end": end_at.isoformat(),
        "reply_context_preserved": preserve_reply_context,
    }
    return filtered


def _filter_normalized_date(
    normalized: dict[str, Any],
    *,
    accounting_date: str,
    timezone_name: str,
    preserve_reply_context: bool = False,
) -> dict[str, Any]:
    zone = ZoneInfo(timezone_name)
    start_at = datetime.strptime(accounting_date, "%Y-%m-%d").replace(tzinfo=zone)
    return _filter_normalized_window(
        normalized,
        start_at=start_at,
        end_at=start_at + timedelta(days=1),
        fingerprint_fields={
            "accounting_date": accounting_date,
            "timezone": timezone_name,
        },
        preserve_reply_context=preserve_reply_context,
    )


def _filter_normalized_time_range(
    normalized: dict[str, Any],
    *,
    accounting_from: str,
    accounting_to: str,
    timezone_name: str,
    preserve_reply_context: bool = False,
) -> dict[str, Any]:
    zone = ZoneInfo(timezone_name)
    return _filter_normalized_window(
        normalized,
        start_at=datetime.strptime(accounting_from, "%Y-%m-%d %H:%M").replace(tzinfo=zone),
        end_at=datetime.strptime(accounting_to, "%Y-%m-%d %H:%M").replace(tzinfo=zone),
        fingerprint_fields={
            "accounting_from": accounting_from,
            "accounting_to": accounting_to,
            "timezone": timezone_name,
        },
        preserve_reply_context=preserve_reply_context,
    )


def _capture_normalized_media_hashes(normalized: dict[str, Any]) -> dict[str, int]:
    """Capture immutable content hashes once for selected evidence media."""

    by_path: dict[str, str] = {}
    computed = 0
    reused = 0
    for group in normalized.get("groups", []):
        for message in group.get("messages", []):
            for media in message.get("media", []):
                if (
                    not isinstance(media, dict)
                    or not core.media_is_evidence(media)
                    or media.get("availability") != "available"
                ):
                    continue
                kind = core.clean_text(media.get("kind")).casefold()
                mime_type = core.clean_text(media.get("mime_type")).casefold()
                if kind != "image" and not mime_type.startswith("image/"):
                    continue
                path = Path(str(media.get("path") or "")).resolve()
                core.require(path.is_file(), f"available media file is unavailable: {path}")
                expected_size = media.get("byte_size")
                actual_size = path.stat().st_size
                core.require(
                    expected_size == actual_size,
                    f"media size changed while creating the snapshot: {path}",
                )
                path_key = str(path).casefold()
                digest = by_path.get(path_key)
                if digest is None:
                    digest = core.sha256_file(path)
                    by_path[path_key] = digest
                    computed += 1
                else:
                    reused += 1
                media["blob_sha256"] = digest
    return {
        "media_content_hashes_computed": computed,
        "media_content_hashes_reused": reused,
    }


def _stable_timestamp(value: object) -> str:
    parsed = datetime.fromisoformat(str(value))
    return parsed.replace(microsecond=0).isoformat()


def _message_source_identity(
    group: Mapping[str, Any], message: Mapping[str, Any]
) -> str:
    platform = core.clean_text(group.get("platform")).casefold()
    message_id = core.clean_text(message.get("message_id"))
    if platform not in {"telegram", "line"} or not message_id:
        return ""
    return f"{core.clean_text(group.get('group_key'))}|{message_id}"


def _assign_stable_identities(normalized: dict[str, Any]) -> None:
    """Add cross-snapshot identities without changing the normalized adapter contract."""

    for group in normalized.get("groups", []):
        occurrences: dict[str, int] = {}
        slot_occurrences: dict[str, int] = {}
        for message in group.get("messages", []):
            sender = core.normalize_name(
                message.get("sender_name") or message.get("sender_id") or ""
            )
            material = {
                "group_key": core.clean_text(group.get("group_key")),
                "timestamp": _stable_timestamp(message.get("timestamp")),
                "sender": sender,
                "text": core.clean_text(message.get("text")),
            }
            base = core.fingerprint_json(material)
            occurrences[base] = occurrences.get(base, 0) + 1
            slot_base = core.fingerprint_json(
                {
                    "group_key": material["group_key"],
                    "timestamp": material["timestamp"],
                    "sender": material["sender"],
                }
            )
            slot_occurrences[slot_base] = slot_occurrences.get(slot_base, 0) + 1
            stable_message_id = core.clean_text(message.get("stable_message_id")) or (
                "stable-message:"
                + core.stable_token(base, occurrences[base], length=32)
            )
            message["stable_message_id"] = stable_message_id
            message["stable_message_slot_base"] = core.clean_text(
                message.get("stable_message_slot_base")
            ) or slot_base
            message["stable_message_slot_id"] = core.clean_text(
                message.get("stable_message_slot_id")
            ) or (
                "stable-slot:"
                + core.stable_token(
                    slot_base, slot_occurrences[slot_base], length=32
                )
            )
            source_identity = core.clean_text(
                message.get("source_message_identity")
            ) or _message_source_identity(group, message)
            if source_identity:
                message["source_message_identity"] = source_identity
            for position, media in enumerate(message.get("media", []), start=1):
                if not isinstance(media, dict):
                    continue
                media["stable_media_id"] = core.clean_text(
                    media.get("stable_media_id")
                ) or f"{stable_message_id}#media:{position:03d}"


def _message_conflict_fields(
    left: Mapping[str, Any], right: Mapping[str, Any]
) -> list[str]:
    conflicts: list[str] = []
    try:
        left_time = _stable_timestamp(left.get("timestamp"))
        right_time = _stable_timestamp(right.get("timestamp"))
    except (TypeError, ValueError):
        left_time = core.clean_text(left.get("timestamp"))
        right_time = core.clean_text(right.get("timestamp"))
    if left_time != right_time:
        conflicts.append("timestamp")
    for field in ("sender_name", "sender_id"):
        left_value = core.normalize_name(left.get(field))
        right_value = core.normalize_name(right.get(field))
        if left_value and right_value and left_value != right_value:
            conflicts.append(field)
    if core.clean_text(left.get("text")) != core.clean_text(right.get("text")):
        conflicts.append("text")
    return conflicts


def _reply_relation_identity(
    group: Mapping[str, Any], message: Mapping[str, Any]
) -> tuple[str, str]:
    """Return the resolved stable target and the adapter-level target id."""

    raw_target = core.clean_text(message.get("reply_to_message_id"))
    if not raw_target:
        return "", ""
    for candidate in group.get("messages", []):
        if not isinstance(candidate, Mapping):
            continue
        candidate_ids = {
            core.clean_text(candidate.get("message_id")),
            core.clean_text(candidate.get("original_message_id")),
        }
        if raw_target in candidate_ids:
            return core.clean_text(candidate.get("stable_message_id")), raw_target
    return "", raw_target


def _merge_reply_relation(
    target_group: Mapping[str, Any],
    target_message: dict[str, Any],
    incoming_group: Mapping[str, Any],
    incoming_message: Mapping[str, Any],
) -> None:
    left_stable, left_raw = _reply_relation_identity(target_group, target_message)
    right_stable, right_raw = _reply_relation_identity(
        incoming_group, incoming_message
    )
    if left_raw and right_raw:
        same_relation = (
            bool(left_stable)
            and bool(right_stable)
            and left_stable == right_stable
        ) or left_raw == right_raw
        core.require(
            same_relation,
            "same stable message identity has conflicting reply relationships",
        )
    elif not left_raw and right_raw:
        target_message["reply_to_message_id"] = right_raw


def _merge_message_media(
    target: dict[str, Any], incoming: Mapping[str, Any]
) -> None:
    target_media = target.get("media", [])
    incoming_media = incoming.get("media", [])
    core.require(
        isinstance(target_media, list) and isinstance(incoming_media, list),
        "message media must be lists",
    )
    for position, incoming_item in enumerate(incoming_media):
        if not isinstance(incoming_item, Mapping):
            continue
        if position >= len(target_media):
            target_media.append(copy.deepcopy(dict(incoming_item)))
            continue
        target_item = target_media[position]
        core.require(isinstance(target_item, dict), "message media item must be an object")
        left_hash = core.clean_text(target_item.get("blob_sha256"))
        right_hash = core.clean_text(incoming_item.get("blob_sha256"))
        core.require(
            not left_hash or not right_hash or left_hash == right_hash,
            "same stable media identity has conflicting content hashes",
        )
        if (
            target_item.get("availability") != "available"
            and incoming_item.get("availability") == "available"
        ):
            stable_media_id = target_item.get("stable_media_id")
            target_item.clear()
            target_item.update(copy.deepcopy(dict(incoming_item)))
            if stable_media_id:
                target_item["stable_media_id"] = stable_media_id
        elif not left_hash and right_hash:
            target_item["blob_sha256"] = right_hash
    target["media"] = target_media


def _canonicalize_group_messages(group: dict[str, Any]) -> None:
    raw_to_stable: dict[str, str] = {}
    for message in group.get("messages", []):
        raw_id = core.clean_text(message.get("message_id"))
        stable_id = core.clean_text(message.get("stable_message_id"))
        if raw_id and stable_id:
            raw_to_stable[raw_id] = stable_id
    for message in group.get("messages", []):
        stable_id = core.clean_text(message.get("stable_message_id"))
        core.require(bool(stable_id), "stable_message_id is missing")
        reply_id = core.clean_text(message.get("reply_to_message_id"))
        if not core.clean_text(message.get("original_message_id")):
            message["original_message_id"] = core.clean_text(message.get("message_id"))
        message["message_id"] = stable_id
        if reply_id:
            message["reply_to_message_id"] = raw_to_stable.get(reply_id, reply_id)
        for media in message.get("media", []):
            if not isinstance(media, dict):
                continue
            stable_media_id = core.clean_text(media.get("stable_media_id"))
            core.require(bool(stable_media_id), "stable_media_id is missing")
            if not core.clean_text(media.get("original_media_id")):
                media["original_media_id"] = core.clean_text(media.get("media_id"))
            media["media_id"] = stable_media_id
    for sequence, message in enumerate(group["messages"], start=1):
        message["source_sequence"] = sequence


def _merge_normalized_snapshots(
    predecessor: dict[str, Any], incoming: dict[str, Any]
) -> tuple[dict[str, Any], dict[str, int]]:
    core.require(
        predecessor.get("timezone") == incoming.get("timezone") == "Asia/Bangkok",
        "predecessor and incoming snapshots must use Asia/Bangkok",
    )
    left = copy.deepcopy(predecessor)
    right = copy.deepcopy(incoming)
    _assign_stable_identities(left)
    _assign_stable_identities(right)
    groups = {
        core.clean_text(group.get("group_key")): group
        for group in left.get("groups", [])
        if isinstance(group, dict)
    }
    deduplicated = 0
    richer_media = 0
    appended = 0
    for incoming_group in right.get("groups", []):
        core.require(isinstance(incoming_group, dict), "incoming group must be an object")
        key = core.clean_text(incoming_group.get("group_key"))
        if key not in groups:
            groups[key] = copy.deepcopy(incoming_group)
            appended += len(incoming_group.get("messages", []))
            continue
        target_group = groups[key]
        core.require(
            core.normalize_name(target_group.get("group_name"))
            == core.normalize_name(incoming_group.get("group_name")),
            f"group identity conflict for {key}: group name changed",
        )
        core.require(
            core.clean_text(target_group.get("platform")).casefold()
            == core.clean_text(incoming_group.get("platform")).casefold(),
            f"group identity conflict for {key}: platform changed",
        )
        stable_index = {
            core.clean_text(message.get("stable_message_id")): message
            for message in target_group.get("messages", [])
            if isinstance(message, dict)
        }
        source_index = {
            core.clean_text(message.get("source_message_identity")): message
            for message in target_group.get("messages", [])
            if isinstance(message, dict)
            and core.clean_text(message.get("source_message_identity"))
        }
        slot_index = {
            core.clean_text(message.get("stable_message_slot_id")): message
            for message in target_group.get("messages", [])
            if isinstance(message, dict)
            and core.clean_text(message.get("stable_message_slot_id"))
        }
        target_slot_counts: dict[str, int] = {}
        incoming_slot_counts: dict[str, int] = {}
        for message in target_group.get("messages", []):
            if isinstance(message, Mapping):
                slot_base = core.clean_text(
                    message.get("stable_message_slot_base")
                )
                if slot_base:
                    target_slot_counts[slot_base] = (
                        target_slot_counts.get(slot_base, 0) + 1
                    )
        for message in incoming_group.get("messages", []):
            if isinstance(message, Mapping):
                slot_base = core.clean_text(
                    message.get("stable_message_slot_base")
                )
                if slot_base:
                    incoming_slot_counts[slot_base] = (
                        incoming_slot_counts.get(slot_base, 0) + 1
                    )
        for incoming_message in incoming_group.get("messages", []):
            core.require(isinstance(incoming_message, Mapping), "incoming message must be an object")
            stable_id = core.clean_text(incoming_message.get("stable_message_id"))
            source_identity = core.clean_text(
                incoming_message.get("source_message_identity")
            )
            target_message = source_index.get(source_identity) if source_identity else None
            if target_message is None:
                target_message = stable_index.get(stable_id)
            if target_message is None:
                slot_base = core.clean_text(
                    incoming_message.get("stable_message_slot_base")
                )
                if (
                    slot_base
                    and target_slot_counts.get(slot_base)
                    == incoming_slot_counts.get(slot_base)
                ):
                    target_message = slot_index.get(
                        core.clean_text(
                            incoming_message.get("stable_message_slot_id")
                        )
                    )
            if target_message is None:
                copied = copy.deepcopy(dict(incoming_message))
                target_group.setdefault("messages", []).append(copied)
                stable_index[stable_id] = copied
                if source_identity:
                    source_index[source_identity] = copied
                slot_id = core.clean_text(copied.get("stable_message_slot_id"))
                if slot_id:
                    slot_index[slot_id] = copied
                appended += 1
                continue
            conflicts = _message_conflict_fields(target_message, incoming_message)
            core.require(
                not conflicts,
                f"message identity conflict in {key}: {', '.join(conflicts)}",
            )
            _merge_reply_relation(
                target_group,
                target_message,
                incoming_group,
                incoming_message,
            )
            before_available = sum(
                isinstance(item, Mapping) and item.get("availability") == "available"
                for item in target_message.get("media", [])
            )
            _merge_message_media(target_message, incoming_message)
            after_available = sum(
                isinstance(item, Mapping) and item.get("availability") == "available"
                for item in target_message.get("media", [])
            )
            richer_media += max(0, after_available - before_available)
            # A message previously counted as context is authoritative when its
            # predecessor copy belonged to the actual accounting period.
            if not incoming_message.get("accounting_context_only"):
                target_message.pop("accounting_context_only", None)
                if not incoming_message.get("excluded_from_accounting"):
                    target_message["excluded_from_accounting"] = False
            deduplicated += 1
        target_group["source_files"] = list(
            dict.fromkeys(
                [str(item) for item in target_group.get("source_files", [])]
                + [str(item) for item in incoming_group.get("source_files", [])]
            )
        )
        merged_rate_candidates: list[Any] = []
        seen_rate_candidates: set[str] = set()
        for candidate in list(target_group.get("rate_candidates", [])) + list(
            incoming_group.get("rate_candidates", [])
        ):
            fingerprint = core.fingerprint_json(candidate)
            if fingerprint in seen_rate_candidates:
                continue
            seen_rate_candidates.add(fingerprint)
            merged_rate_candidates.append(copy.deepcopy(candidate))
        target_group["rate_candidates"] = merged_rate_candidates

    result_groups = list(groups.values())
    for group in result_groups:
        # Reply ids are only meaningful inside their own chat. LINE and other
        # adapters may reuse the same raw id in unrelated groups.
        raw_to_stable: dict[str, str] = {}
        for message in group.get("messages", []):
            if not isinstance(message, Mapping):
                continue
            stable_id = core.clean_text(message.get("stable_message_id"))
            for field in ("message_id", "original_message_id"):
                raw_id = core.clean_text(message.get(field))
                if raw_id:
                    raw_to_stable[raw_id] = stable_id
        for message in group.get("messages", []):
            reply_id = core.clean_text(message.get("reply_to_message_id"))
            if reply_id in raw_to_stable:
                message["reply_to_message_id"] = raw_to_stable[reply_id]
        _canonicalize_group_messages(group)
    result_groups.sort(
        key=lambda item: (
            str(item.get("platform")),
            str(item.get("group_name")),
            str(item.get("group_key")),
        )
    )
    messages = [
        message for group in result_groups for message in group.get("messages", [])
    ]
    media = [item for message in messages for item in message.get("media", [])]
    result = {
        "contract_version": core.NORMALIZED_CONTRACT,
        "timezone": "Asia/Bangkok",
        "source_fingerprint": core.fingerprint_json(
            {
                "normalizer": "roll-forward/1.0",
                "predecessor": predecessor.get("source_fingerprint"),
                "incoming": incoming.get("source_fingerprint"),
                "stable_messages": [
                    message.get("stable_message_id") for message in messages
                ],
            }
        ),
        "groups": result_groups,
        "statistics": {
            "groups": len(result_groups),
            "messages": len(messages),
            "excluded_messages": sum(
                bool(message.get("excluded_from_accounting")) for message in messages
            ),
            "media": len(media),
            "available_media": sum(
                isinstance(item, Mapping) and item.get("availability") == "available"
                for item in media
            ),
            "missing_media": sum(
                isinstance(item, Mapping) and item.get("availability") == "missing"
                for item in media
            ),
            "rate_candidates": sum(
                len(group.get("rate_candidates", [])) for group in result_groups
            ),
            "source_files": sum(
                len(group.get("source_files", [])) for group in result_groups
            ),
        },
        "warnings": list(
            dict.fromkeys(
                [str(item) for item in predecessor.get("warnings", [])]
                + [str(item) for item in incoming.get("warnings", [])]
            )
        ),
        "skipped_files": list(
            dict.fromkeys(
                [str(item) for item in predecessor.get("skipped_files", [])]
                + [str(item) for item in incoming.get("skipped_files", [])]
            )
        ),
    }
    return result, {
        "messages_added": appended,
        "messages_deduplicated": deduplicated,
        "media_enriched": richer_media,
    }


def start_run(args: argparse.Namespace) -> dict[str, Any]:
    work = args.work.resolve()
    core.require(not work.exists(), f"work directory already exists; start requires a new path: {work}")
    ocr_candidate_config = _resolve_ocr_candidate_config(
        getattr(args, "ocr_candidates", None)
    )
    inputs = [path.resolve() for path in args.inputs]
    for path in inputs:
        core.require(path.exists(), f"input does not exist: {path}")
    group_mode = core.clean_text(getattr(args, "mode", "small")).casefold()
    core.require(
        group_mode
        in {"small", "large", finance_materials.GROUP_MODE, store_ledger.GROUP_MODE},
        "--mode must be small, large, finance, or store-ledger",
    )
    contains = (
        finance_materials.GROUP_NAME_MARKER
        if group_mode == finance_materials.GROUP_MODE
        else store_ledger.GROUP_NAME_MARKER
        if group_mode == store_ledger.GROUP_MODE
        else core.clean_text(getattr(args, "contains", "小额"))
    )
    core.require(bool(contains), "--contains cannot be empty")
    zone = ZoneInfo(args.timezone)
    core.require(args.timezone == "Asia/Bangkok", "accounting timezone must be Asia/Bangkok")
    accounting_date = _normalize_accounting_date(getattr(args, "date", None))
    accounting_from = _normalize_accounting_time(
        getattr(args, "accounting_from", None), option="--from"
    )
    accounting_to = _normalize_accounting_time(
        getattr(args, "accounting_to", None), option="--to"
    )
    core.require(
        (accounting_from is None) == (accounting_to is None),
        "--from and --to must be provided together",
    )
    if group_mode != "large":
        core.require(
            accounting_date is None or accounting_from is None,
            "--date cannot be combined with --from or --to outside large mode",
        )
    if group_mode == "large":
        core.require(
            accounting_date is not None,
            "--mode large requires one accounting date label via --date",
        )
    if accounting_from is not None and accounting_to is not None:
        core.require(
            datetime.strptime(accounting_from, "%Y-%m-%d %H:%M")
            < datetime.strptime(accounting_to, "%Y-%m-%d %H:%M"),
            "--from must be earlier than --to",
        )
    roster_path = args.roster.resolve()
    roster = core.load_roster(roster_path)
    work.mkdir(parents=True)
    (work / "snapshot" / "sources").mkdir(parents=True)
    (work / "decisions").mkdir()

    documents: list[dict[str, Any]] = []
    chat_document = _source_document(inputs, zone=zone, roster=roster)
    if chat_document is not None:
        documents.append(chat_document)
    line_backups = _discover_line_backups(inputs, args.line_backup)
    documents.extend(
        _line_documents(
            work,
            line_backups,
            contains=contains,
            group_mode=group_mode,
            timezone_name=args.timezone,
            roster_path=roster_path,
            self_name=args.line_self_name,
        )
    )
    line_android_backups = _discover_line_android_backups(
        inputs, getattr(args, "line_android_backup", [])
    )
    documents.extend(
        _line_android_documents(
            work,
            line_android_backups,
            contains=contains,
            group_mode=group_mode,
            timezone_name=args.timezone,
            roster_path=roster_path,
        )
    )
    normalized = _merge_documents(
        documents,
        contains=contains,
        group_mode=group_mode,
        timezone_name=args.timezone,
    )
    if accounting_from is not None and accounting_to is not None:
        normalized = _filter_normalized_time_range(
            normalized,
            accounting_from=accounting_from,
            accounting_to=accounting_to,
            timezone_name=args.timezone,
            preserve_reply_context=True,
        )
    elif accounting_date is not None:
        normalized = _filter_normalized_date(
            normalized,
            accounting_date=accounting_date,
            timezone_name=args.timezone,
            preserve_reply_context=True,
        )
    media_hash_statistics = _capture_normalized_media_hashes(normalized)
    _assign_stable_identities(normalized)
    snapshot_path = _snapshot_path(work)
    core.atomic_json(snapshot_path, normalized)
    core.load_normalized(snapshot_path)

    group_reports: list[dict[str, Any]] = []
    used_labels: set[str] = set()
    for group in normalized["groups"]:
        decision = _decision_template(normalized, group, group_mode=group_mode)
        filename = _decision_filename(str(group["group_key"]))
        decision_path = work / "decisions" / filename
        core.atomic_json(decision_path, decision)
        base_label = f"[{_platform_prefix(group.get('platform'))}] {group.get('group_name')}"
        run_label = base_label
        if run_label in used_labels:
            run_label = f"{base_label} ({core.stable_token(group.get('group_key'), length=6)})"
        used_labels.add(run_label)
        media_count = len(_media_inventory(group))
        group_reports.append(
            {
                "group_key": group["group_key"],
                "group_name": group.get("group_name"),
                "platform": group.get("platform"),
                "run_label": run_label,
                "messages": len(group.get("messages", [])),
                "evidence_media": media_count,
                "decision_file": f"decisions/{filename}",
                "edit_mode": EDIT_CONTROL_MODE,
                "group_mode": group_mode,
            }
        )
    amount_policy = (
        store_ledger.AMOUNT_POLICY
        if group_mode == store_ledger.GROUP_MODE
        else AMOUNT_POLICY
    )
    run = {
        "contract_version": RUN_CONTRACT,
        "runtime_contract": RUNTIME_CONTRACT,
        "amount_policy": amount_policy,
        "workbook_compatibility": WORKBOOK_COMPATIBILITY,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "accounting_timezone": args.timezone,
        "group_mode": group_mode,
        "status": "reviewing",
        "normalized_source_fingerprint": normalized["source_fingerprint"],
        "snapshot_sha256": core.sha256_file(snapshot_path),
        "input_roots": [str(path) for path in inputs],
        "roster_path": str(roster_path),
        "roster_sha256": core.sha256_file(roster_path),
        "line_self_name": args.line_self_name,
        "line_backups": [str(path) for path in line_backups],
        "line_android_backups": [str(path) for path in line_android_backups],
        "ocr_candidates": ocr_candidate_config,
        "groups": group_reports,
    }
    if group_mode == "large":
        run["group_name_excludes"] = [
            "小额出",
            "财务资料群",
        ]
    else:
        run["group_name_contains"] = contains
    if accounting_date is not None:
        run["accounting_date"] = accounting_date
    if accounting_from is not None and accounting_to is not None:
        run["accounting_from"] = accounting_from
        run["accounting_to"] = accounting_to
    core.atomic_json(_run_path(work), run)
    result = {
        "work": str(work),
        "runtime_contract": RUNTIME_CONTRACT,
        "amount_policy": amount_policy,
        "workbook_compatibility": WORKBOOK_COMPATIBILITY,
        "groups": len(group_reports),
        "messages": normalized["statistics"]["messages"],
        "evidence_media": sum(item["evidence_media"] for item in group_reports),
        "selected_groups": group_reports,
        "ocr_candidates": ocr_candidate_config,
        **media_hash_statistics,
    }
    if accounting_date is not None:
        result["accounting_date"] = accounting_date
    if accounting_from is not None and accounting_to is not None:
        result["accounting_from"] = accounting_from
        result["accounting_to"] = accounting_to
    return result


def _group_index(normalized: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    return {str(group["group_key"]): group for group in normalized.get("groups", [])}


def _message_labels(group: Mapping[str, Any]) -> tuple[dict[str, str], dict[str, Mapping[str, Any]]]:
    label_by_id: dict[str, str] = {}
    message_by_label: dict[str, Mapping[str, Any]] = {}
    for index, message in enumerate(group.get("messages", []), start=1):
        label = f"S{index:05d}"
        message_id = str(message.get("message_id") or "")
        label_by_id[message_id] = label
        message_by_label[label] = message
    return label_by_id, message_by_label


def _media_inventory(group: Mapping[str, Any]) -> dict[str, tuple[Mapping[str, Any], Mapping[str, Any]]]:
    inventory: dict[str, tuple[Mapping[str, Any], Mapping[str, Any]]] = {}
    index = 0
    for message in group.get("messages", []):
        for media in message.get("media", []):
            if isinstance(media, Mapping) and core.media_is_evidence(media):
                index += 1
                inventory[f"M{index:04d}"] = (message, media)
    return inventory


def _select_run_group(run: Mapping[str, Any], selector: str | None, *, prefer_incomplete: bool = False) -> Mapping[str, Any]:
    groups = [item for item in run.get("groups", []) if isinstance(item, Mapping)]
    if selector:
        exact = [
            item
            for item in groups
            if selector in {str(item.get("group_key")), str(item.get("run_label"))}
            or selector == str(item.get("group_name"))
        ]
        core.require(len(exact) == 1, f"group selector must match exactly one group: {selector!r}")
        return exact[0]
    core.require(prefer_incomplete, "--group is required")
    for item in groups:
        if not bool(item.get("sealed")):
            return item
    raise ValueError("all selected groups are already sealed")


def _decision_path(work: Path, run_group: Mapping[str, Any]) -> Path:
    relative = Path(str(run_group.get("decision_file") or ""))
    core.require(not relative.is_absolute() and ".." not in relative.parts, "invalid decision path in run manifest")
    return work / relative


def _semantic_fingerprint(decision: Mapping[str, Any]) -> str:
    def with_media_review_state(material: dict[str, Any]) -> dict[str, Any]:
        # Older controlled work directories did not contain these fields. Keep
        # their approved fingerprints valid until a new batch adds the state.
        for field in ("media_observation_cache", "media_view_metrics"):
            if field in decision:
                material[field] = decision.get(field)
        return material

    if decision.get("contract_version") == finance_materials.DECISION_CONTRACT:
        return core.fingerprint_json(
            with_media_review_state({
                "reviewed_through": decision.get("reviewed_through"),
                "read_complete": decision.get("read_complete"),
                "media_decisions": decision.get("media_decisions"),
                "people": decision.get("people"),
                "open_people": decision.get("open_people"),
            })
        )
    if decision.get("contract_version") == store_ledger.DECISION_CONTRACT:
        return core.fingerprint_json(
            with_media_review_state({
                "reviewed_through": decision.get("reviewed_through"),
                "read_complete": decision.get("read_complete"),
                "media_decisions": decision.get("media_decisions"),
                "records": decision.get("records"),
                "open_records": decision.get("open_records"),
                "balance_snapshots": decision.get("balance_snapshots"),
            })
        )
    if decision.get("contract_version") == large_daily.LEGACY_DECISION_CONTRACT:
        return core.fingerprint_json(
            with_media_review_state({
                "reviewed_through": decision.get("reviewed_through"),
                "read_complete": decision.get("read_complete"),
                "media_decisions": decision.get("media_decisions"),
                "exchanges": decision.get("exchanges"),
                "open_exchanges": decision.get("open_exchanges"),
                "unknown_payee_reviewed_entry_ids": decision.get(
                    "unknown_payee_reviewed_entry_ids"
                ),
            })
        )
    return core.fingerprint_json(
        with_media_review_state({
            "reviewed_through": decision.get("reviewed_through"),
            "read_complete": decision.get("read_complete"),
            "media_decisions": decision.get("media_decisions"),
            "orders": decision.get("orders"),
            "open_orders": decision.get("open_orders"),
            "balance_links": decision.get("balance_links"),
            "settlement_allocations": decision.get("settlement_allocations"),
            "unknown_payee_reviewed_entry_ids": decision.get(
                "unknown_payee_reviewed_entry_ids"
            ),
        })
    )


def _open_field_for_decision(decision: Mapping[str, Any]) -> str:
    contract_version = decision.get("contract_version")
    if contract_version == finance_materials.DECISION_CONTRACT:
        return "open_people"
    if contract_version == store_ledger.DECISION_CONTRACT:
        return "open_records"
    if contract_version == large_daily.LEGACY_DECISION_CONTRACT:
        return "open_exchanges"
    return "open_orders"


def _expected_contracts_for_run_group(run_group: Mapping[str, Any]) -> set[str]:
    group_mode = core.clean_text(run_group.get("group_mode")).casefold()
    if group_mode == finance_materials.GROUP_MODE:
        return {finance_materials.DECISION_CONTRACT}
    if group_mode == store_ledger.GROUP_MODE:
        return {store_ledger.DECISION_CONTRACT}
    if group_mode == "large":
        return {large_daily.DECISION_CONTRACT, large_daily.LEGACY_DECISION_CONTRACT}
    return {DECISION_CONTRACT}


def _legacy_semantic_fingerprint_3_1(decision: Mapping[str, Any]) -> str:
    return core.fingerprint_json(
        {
            "media_decisions": decision.get("media_decisions"),
            "orders": decision.get("orders"),
            "balance_links": decision.get("balance_links"),
            "settlement_allocations": decision.get("settlement_allocations"),
            "unknown_payee_reviewed_entry_ids": decision.get(
                "unknown_payee_reviewed_entry_ids"
            ),
        }
    )


def _new_edit_control(approved_fingerprint: str) -> dict[str, Any]:
    return {
        "mode": EDIT_CONTROL_MODE,
        "approved_semantic_fingerprint": approved_fingerprint,
        "batch_count": 0,
        "last_batch": None,
    }


def _controlled_editing(run_group: Mapping[str, Any]) -> bool:
    return core.clean_text(run_group.get("edit_mode")) == EDIT_CONTROL_MODE


def _require_controlled_semantics(
    run_group: Mapping[str, Any],
    decision: Mapping[str, Any],
) -> bool:
    if not _controlled_editing(run_group):
        return False
    control = decision.get("edit_control")
    core.require(
        isinstance(control, Mapping),
        "controlled decision is missing edit_control; use review apply-batch",
    )
    unknown = sorted(
        set(control)
        - {"mode", "approved_semantic_fingerprint", "batch_count", "last_batch"}
    )
    core.require(
        not unknown,
        f"edit_control has unsupported fields: {', '.join(unknown)}",
    )
    core.require(
        core.clean_text(control.get("mode")) == EDIT_CONTROL_MODE,
        "unsupported decision edit_control mode",
    )
    batch_count = control.get("batch_count")
    core.require(
        isinstance(batch_count, int) and batch_count >= 0,
        "edit_control.batch_count must be a non-negative integer",
    )
    last_batch = control.get("last_batch")
    if batch_count == 0:
        core.require(last_batch is None, "edit_control.last_batch must be empty before the first batch")
    else:
        core.require(isinstance(last_batch, Mapping), "edit_control.last_batch is required")
        unknown_last = sorted(
            set(last_batch)
            - {
                "id",
                "input_fingerprint",
                "base_fingerprint",
                "result_fingerprint",
                "applied_at",
            }
        )
        core.require(
            not unknown_last,
            f"edit_control.last_batch has unsupported fields: {', '.join(unknown_last)}",
        )
    current = _semantic_fingerprint(decision)
    approved = core.clean_text(control.get("approved_semantic_fingerprint"))
    core.require(
        approved == current,
        "decision semantic fields were edited outside review apply-batch; "
        "restore the last approved decision or submit a batch from its approved fingerprint",
    )
    if isinstance(last_batch, Mapping):
        core.require(
            core.clean_text(last_batch.get("result_fingerprint")) == current,
            "edit_control.last_batch does not match the approved decision",
        )
    return True


def _normalize_media_view_metrics(value: object, *, field: str) -> dict[str, int]:
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    unknown = sorted(set(value) - set(MEDIA_VIEW_METRIC_FIELDS))
    core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
    normalized: dict[str, int] = {}
    for name in MEDIA_VIEW_METRIC_FIELDS:
        raw = value.get(name, 0)
        core.require(
            isinstance(raw, int) and not isinstance(raw, bool) and raw >= 0,
            f"{field}.{name} must be a non-negative integer",
        )
        normalized[name] = raw
    core.require(
        normalized["failed_images"] <= normalized["opened_images"],
        f"{field}.failed_images cannot exceed opened_images",
    )
    core.require(
        normalized["single_image_rechecks"] <= normalized["opened_images"],
        f"{field}.single_image_rechecks cannot exceed opened_images",
    )
    core.require(
        normalized["opened_images"] == 0 or normalized["view_batches"] > 0,
        f"{field}.view_batches is required when images were opened",
    )
    return normalized


def _merge_media_view_metrics(
    decision: dict[str, Any],
    batch_metrics: object,
) -> None:
    if batch_metrics is None:
        return
    current = _normalize_media_view_metrics(
        decision.get("media_view_metrics", _empty_media_view_metrics()),
        field="decision.media_view_metrics",
    )
    update = _normalize_media_view_metrics(
        batch_metrics,
        field="review batch media_view_metrics",
    )
    decision["media_view_metrics"] = {
        field: current[field] + update[field]
        for field in MEDIA_VIEW_METRIC_FIELDS
    }


def _prune_media_observation_sources(
    decision: dict[str, Any],
    labels: set[str],
) -> None:
    if not labels or "media_observation_cache" not in decision:
        return
    cache = decision.get("media_observation_cache")
    core.require(isinstance(cache, dict), "media_observation_cache must be an object")
    entries = cache.get("entries")
    core.require(isinstance(entries, dict), "media_observation_cache.entries must be an object")
    for digest in list(entries):
        entry = entries[digest]
        core.require(isinstance(entry, dict), f"media_observation_cache.entries.{digest} must be an object")
        source_labels = entry.get("source_labels")
        core.require(
            isinstance(source_labels, list),
            f"media_observation_cache.entries.{digest}.source_labels must be a list",
        )
        remaining = [str(label) for label in source_labels if str(label) not in labels]
        if remaining:
            entry["source_labels"] = remaining
        else:
            del entries[digest]


def _load_review_batch(path: Path) -> dict[str, Any]:
    core.require(path.is_file(), f"review batch does not exist: {path}")
    batch = _load_json(path)
    allowed = {
        "contract_version",
        "batch_id",
        "base_fingerprint",
        "media_decisions",
        "media_observations",
        "media_view_metrics",
        "remove_media_labels",
        "orders",
        "remove_order_ids",
        "open_orders",
        "page_commit",
        "balance_links",
        "settlement_allocations",
        "unknown_payee_reviewed_entry_ids",
        "exchanges",
        "remove_exchange_ids",
        "open_exchanges",
        "people",
        "remove_person_ids",
        "open_people",
        "records",
        "remove_record_ids",
        "open_records",
        "balance_snapshots",
        "remove_balance_snapshot_ids",
    }
    unknown = sorted(set(batch) - allowed)
    core.require(
        not unknown,
        f"review batch has unsupported fields: {', '.join(unknown)}",
    )
    core.require(
        batch.get("contract_version") == REVIEW_BATCH_CONTRACT,
        f"unsupported review batch contract; expected {REVIEW_BATCH_CONTRACT}",
    )
    batch_id = core.clean_text(batch.get("batch_id"))
    core.require(
        bool(re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,127}", batch_id)),
        "review batch_id must use 1..128 ASCII letters, digits, dot, underscore, or hyphen",
    )
    batch["batch_id"] = batch_id
    base_fingerprint = core.clean_text(batch.get("base_fingerprint"))
    core.require(
        bool(re.fullmatch(r"sha256:[0-9a-f]{64}", base_fingerprint)),
        "review batch base_fingerprint must be a sha256 fingerprint from review status, next, or check",
    )
    batch["base_fingerprint"] = base_fingerprint
    page_commit = batch.get("page_commit")
    if page_commit is not None:
        core.require(isinstance(page_commit, Mapping), "review batch page_commit must be an object")
        unknown_page_fields = sorted(
            set(page_commit) - {"page_start", "page_end", "page_token"}
        )
        core.require(
            not unknown_page_fields,
            f"review batch page_commit has unsupported fields: {', '.join(unknown_page_fields)}",
        )
        page_start = page_commit.get("page_start")
        page_end = page_commit.get("page_end")
        core.require(
            isinstance(page_start, int) and isinstance(page_end, int),
            "review batch page_commit page_start and page_end must be integers",
        )
        page_token = core.clean_text(page_commit.get("page_token"))
        core.require(
            bool(re.fullmatch(r"sha256:[0-9a-f]{64}", page_token)),
            "review batch page_commit page_token must come from review next",
        )
        page_commit["page_token"] = page_token
        core.require(
            any(
                field in batch
                for field in ("open_orders", "open_exchanges", "open_people", "open_records")
            ),
            "review batch that commits a page must include the complete open_orders, "
            "open_exchanges, open_people, or open_records list",
        )
    if "open_orders" in batch:
        core.require(isinstance(batch["open_orders"], list), "review batch open_orders must be a list")
    if "open_exchanges" in batch:
        core.require(
            isinstance(batch["open_exchanges"], list),
            "review batch open_exchanges must be a list",
        )
    if "open_people" in batch:
        core.require(isinstance(batch["open_people"], list), "review batch open_people must be a list")
    if "open_records" in batch:
        core.require(isinstance(batch["open_records"], list), "review batch open_records must be a list")
    if "media_observations" in batch:
        core.require(
            isinstance(batch["media_observations"], Mapping),
            "review batch media_observations must be an object keyed by M labels",
        )
        media_updates = batch.get("media_decisions", {})
        core.require(
            isinstance(media_updates, Mapping),
            "review batch media_decisions must be an object when media_observations are supplied",
        )
        core.require(
            set(batch["media_observations"]) == set(media_updates),
            "review batch media_observations must contain exactly one result for every updated media label",
        )
    if "media_view_metrics" in batch:
        core.require(
            isinstance(batch["media_view_metrics"], Mapping),
            "review batch media_view_metrics must be an object",
        )
        core.require(
            "observation_cache_reuses" not in batch["media_view_metrics"],
            "review batch media_view_metrics.observation_cache_reuses is maintained by the script",
        )
        batch["media_view_metrics"] = _normalize_media_view_metrics(
            batch["media_view_metrics"],
            field="review batch media_view_metrics",
        )
    return batch


def _merge_review_batch(
    decision: Mapping[str, Any],
    batch: Mapping[str, Any],
) -> dict[str, Any]:
    candidate = copy.deepcopy(dict(decision))
    candidate["sealed"] = False
    candidate["sealed_decision_fingerprint"] = None

    media_updates = batch.get("media_decisions", {})
    core.require(isinstance(media_updates, Mapping), "review batch media_decisions must be an object")
    remove_media_value = batch.get("remove_media_labels", [])
    core.require(isinstance(remove_media_value, list), "review batch remove_media_labels must be a list")
    remove_media = [str(item) for item in remove_media_value]
    core.require(len(remove_media) == len(set(remove_media)), "review batch remove_media_labels repeats a label")
    core.require(
        not (set(media_updates) & set(remove_media)),
        "review batch cannot update and remove the same media label",
    )
    candidate_media = candidate.get("media_decisions")
    core.require(isinstance(candidate_media, dict), "decision media_decisions must be an object")
    missing_media = sorted(set(remove_media) - set(candidate_media))
    core.require(not missing_media, f"review batch removes unknown media labels: {missing_media[:20]}")
    for label in remove_media:
        del candidate_media[label]
    for label, value in media_updates.items():
        candidate_media[str(label)] = copy.deepcopy(value)

    _prune_media_observation_sources(
        candidate,
        {str(label) for label in media_updates} | set(remove_media),
    )
    _merge_media_view_metrics(candidate, batch.get("media_view_metrics"))

    if candidate.get("contract_version") == finance_materials.DECISION_CONTRACT:
        incompatible = (
            "orders",
            "remove_order_ids",
            "open_orders",
            "balance_links",
            "settlement_allocations",
            "unknown_payee_reviewed_entry_ids",
            "exchanges",
            "remove_exchange_ids",
            "open_exchanges",
            "records",
            "remove_record_ids",
            "open_records",
            "balance_snapshots",
            "remove_balance_snapshot_ids",
        )
        core.require(
            not any(field in batch for field in incompatible),
            "finance-material review batches must use people and open_people",
        )
        return finance_materials.merge_review_batch(candidate, batch)

    if candidate.get("contract_version") == store_ledger.DECISION_CONTRACT:
        incompatible = (
            "orders",
            "remove_order_ids",
            "open_orders",
            "balance_links",
            "settlement_allocations",
            "unknown_payee_reviewed_entry_ids",
            "exchanges",
            "remove_exchange_ids",
            "open_exchanges",
            "people",
            "remove_person_ids",
            "open_people",
        )
        core.require(
            not any(field in batch for field in incompatible),
            "store-ledger review batches must use records, open_records, and balance_snapshots",
        )
        return store_ledger.merge_review_batch(candidate, batch)

    if candidate.get("contract_version") == large_daily.LEGACY_DECISION_CONTRACT:
        core.require(
            not any(
                field in batch
                for field in (
                    "orders",
                    "remove_order_ids",
                    "open_orders",
                    "balance_links",
                    "settlement_allocations",
                    "people",
                    "remove_person_ids",
                    "open_people",
                    "records",
                    "remove_record_ids",
                    "open_records",
                    "balance_snapshots",
                    "remove_balance_snapshot_ids",
                )
            ),
            "large-group review batches must use exchanges and open_exchanges",
        )
        exchange_updates = batch.get("exchanges", [])
        core.require(
            isinstance(exchange_updates, list),
            "review batch exchanges must be a list",
        )
        update_ids: list[str] = []
        for position, exchange in enumerate(exchange_updates):
            core.require(
                isinstance(exchange, Mapping),
                f"review batch exchanges[{position}] must be an object",
            )
            exchange_id = core.clean_text(exchange.get("id"))
            core.require(
                bool(exchange_id),
                f"review batch exchanges[{position}].id is required",
            )
            update_ids.append(exchange_id)
        core.require(
            len(update_ids) == len(set(update_ids)),
            "review batch exchanges repeats an exchange id",
        )
        remove_value = batch.get("remove_exchange_ids", [])
        core.require(
            isinstance(remove_value, list),
            "review batch remove_exchange_ids must be a list",
        )
        remove_ids = [str(item) for item in remove_value]
        core.require(
            len(remove_ids) == len(set(remove_ids)),
            "review batch remove_exchange_ids repeats an exchange id",
        )
        core.require(
            not (set(update_ids) & set(remove_ids)),
            "review batch cannot update and remove the same exchange id",
        )
        candidate_exchanges = candidate.get("exchanges")
        core.require(isinstance(candidate_exchanges, list), "decision exchanges must be a list")
        existing_ids = {
            core.clean_text(item.get("id"))
            for item in candidate_exchanges
            if isinstance(item, Mapping)
        }
        missing_ids = sorted(set(remove_ids) - existing_ids)
        core.require(
            not missing_ids,
            f"review batch removes unknown exchange ids: {missing_ids[:20]}",
        )
        if remove_ids:
            candidate_exchanges[:] = [
                item
                for item in candidate_exchanges
                if not isinstance(item, Mapping)
                or core.clean_text(item.get("id")) not in set(remove_ids)
            ]
        positions = {
            core.clean_text(item.get("id")): position
            for position, item in enumerate(candidate_exchanges)
            if isinstance(item, Mapping)
        }
        for exchange in exchange_updates:
            exchange_id = core.clean_text(exchange.get("id"))
            replacement = copy.deepcopy(dict(exchange))
            if exchange_id in positions:
                candidate_exchanges[positions[exchange_id]] = replacement
            else:
                positions[exchange_id] = len(candidate_exchanges)
                candidate_exchanges.append(replacement)
        for field in ("open_exchanges", "unknown_payee_reviewed_entry_ids"):
            if field in batch:
                core.require(
                    isinstance(batch[field], list),
                    f"review batch {field} must be a list",
                )
                candidate[field] = copy.deepcopy(batch[field])
        return candidate

    core.require(
        not any(
            field in batch
            for field in (
                "exchanges",
                "remove_exchange_ids",
                "open_exchanges",
                "people",
                "remove_person_ids",
                "open_people",
                "records",
                "remove_record_ids",
                "open_records",
                "balance_snapshots",
                "remove_balance_snapshot_ids",
            )
        ),
        "order-based review batches must use orders and open_orders",
    )
    order_updates = batch.get("orders", [])
    core.require(isinstance(order_updates, list), "review batch orders must be a list")
    update_order_ids: list[str] = []
    for position, order in enumerate(order_updates):
        core.require(isinstance(order, Mapping), f"review batch orders[{position}] must be an object")
        order_id = core.clean_text(order.get("id"))
        core.require(bool(order_id), f"review batch orders[{position}].id is required")
        update_order_ids.append(order_id)
    core.require(
        len(update_order_ids) == len(set(update_order_ids)),
        "review batch orders repeats an order id",
    )
    remove_orders_value = batch.get("remove_order_ids", [])
    core.require(isinstance(remove_orders_value, list), "review batch remove_order_ids must be a list")
    remove_order_ids = [str(item) for item in remove_orders_value]
    core.require(
        len(remove_order_ids) == len(set(remove_order_ids)),
        "review batch remove_order_ids repeats an order id",
    )
    core.require(
        not (set(update_order_ids) & set(remove_order_ids)),
        "review batch cannot update and remove the same order id",
    )
    candidate_orders = candidate.get("orders")
    core.require(isinstance(candidate_orders, list), "decision orders must be a list")
    existing_order_ids = {
        core.clean_text(order.get("id"))
        for order in candidate_orders
        if isinstance(order, Mapping)
    }
    missing_orders = sorted(set(remove_order_ids) - existing_order_ids)
    core.require(not missing_orders, f"review batch removes unknown order ids: {missing_orders[:20]}")
    if remove_order_ids:
        candidate_orders[:] = [
            order
            for order in candidate_orders
            if not isinstance(order, Mapping)
            or core.clean_text(order.get("id")) not in set(remove_order_ids)
        ]
    order_positions = {
        core.clean_text(order.get("id")): position
        for position, order in enumerate(candidate_orders)
        if isinstance(order, Mapping)
    }
    for order in order_updates:
        order_id = core.clean_text(order.get("id"))
        replacement = copy.deepcopy(dict(order))
        if order_id in order_positions:
            existing_order = candidate_orders[order_positions[order_id]]
            if (
                "lifecycle" not in replacement
                and isinstance(existing_order, Mapping)
                and isinstance(existing_order.get("lifecycle"), Mapping)
            ):
                replacement["lifecycle"] = copy.deepcopy(existing_order["lifecycle"])
            candidate_orders[order_positions[order_id]] = replacement
        else:
            order_positions[order_id] = len(candidate_orders)
            candidate_orders.append(replacement)

    for field in (
        "open_orders",
        "balance_links",
        "settlement_allocations",
        "unknown_payee_reviewed_entry_ids",
    ):
        if field in batch:
            core.require(isinstance(batch[field], list), f"review batch {field} must be a list")
            candidate[field] = copy.deepcopy(batch[field])
    return candidate


def _media_content_hash_for_label(
    group: Mapping[str, Any],
    decision: Mapping[str, Any],
    label: str,
    *,
    inventory: Mapping[str, tuple[Mapping[str, Any], Mapping[str, Any]]] | None = None,
    verify_file: bool,
) -> tuple[str, bool]:
    if inventory is None:
        inventory = _media_inventory(group)
    core.require(label in inventory, f"unknown media label for content hash: {label}")
    _, media = inventory[label]
    core.require(media.get("availability") == "available", f"media is unavailable: {label}")
    snapshot_hash = core.clean_text(media.get("blob_sha256"))
    core.require(
        not snapshot_hash or re.fullmatch(r"[0-9a-f]{64}", snapshot_hash) is not None,
        f"{label}: snapshot media hash is invalid",
    )
    media_decisions = decision.get("media_decisions", {})
    recorded_hash = ""
    if isinstance(media_decisions, Mapping):
        media_decision = media_decisions.get(label)
        if isinstance(media_decision, Mapping):
            recorded_hash = core.clean_text(media_decision.get("evidence_sha256"))
    core.require(
        not recorded_hash or re.fullmatch(r"[0-9a-f]{64}", recorded_hash) is not None,
        f"{label}: recorded evidence hash is invalid",
    )
    if snapshot_hash and recorded_hash:
        core.require(
            snapshot_hash == recorded_hash,
            f"{label}: evidence hash disagrees with the fixed media snapshot",
        )
    trusted = recorded_hash or snapshot_hash
    if trusted and not verify_file:
        return trusted, False
    path = Path(str(media.get("path") or ""))
    core.require(path.is_file(), f"{label}: original media file is unavailable: {path}")
    actual = core.sha256_file(path)
    if snapshot_hash:
        core.require(
            actual == snapshot_hash,
            f"{label}: original media changed after the run snapshot was created",
        )
    if recorded_hash:
        core.require(
            actual == recorded_hash,
            f"{label}: original media changed after review",
        )
    return actual, True


def _order_observation_facts(media_decision: Mapping[str, Any]) -> dict[str, Any]:
    if core.clean_text(media_decision.get("classification")).casefold() != "fund":
        return {}
    result: list[dict[str, Any]] = []
    for raw_entry in media_decision.get("entries", []):
        core.require(isinstance(raw_entry, Mapping), "fund observation entries must be objects")
        entry: dict[str, Any] = {}
        for field in (
            "amount",
            "amount_text",
            "amount_basis",
            "currency",
            "payee",
            "payee_state",
            "kind",
            "status_text",
            "amount_state",
        ):
            if field in raw_entry and raw_entry.get(field) not in (None, ""):
                entry[field] = copy.deepcopy(raw_entry.get(field))
        result.append(entry)
    return {"entries": result}


def _finance_observation_target(
    decision: Mapping[str, Any],
    label: str,
    classification: str,
) -> tuple[dict[str, Any], int] | None:
    matches: list[tuple[dict[str, Any], int]] = []
    for person in decision.get("people", []):
        if not isinstance(person, Mapping):
            continue
        if classification == "document":
            for document in person.get("documents", []):
                if not isinstance(document, Mapping):
                    continue
                labels = document.get("media_labels", [])
                if isinstance(labels, list) and label in labels:
                    holder = {
                        field: core.clean_text(person.get(field))
                        for field in (
                            "name",
                            "surname",
                            "given_names",
                            "nationality",
                            "birth_date",
                        )
                        if core.clean_text(person.get(field))
                    }
                    document_facts = {
                        field: core.clean_text(document.get(field))
                        for field in ("type", "country_code", "number")
                        if core.clean_text(document.get(field))
                    }
                    matches.append(
                        ({"holder": holder, "document": document_facts}, len(labels))
                    )
        elif classification == "chat_profile":
            for account in person.get("accounts", []):
                if not isinstance(account, Mapping):
                    continue
                labels = account.get("media_labels", [])
                if isinstance(labels, list) and label in labels:
                    account_facts = {
                        field: core.clean_text(account.get(field))
                        for field in ("platform", "account_id", "phone")
                        if core.clean_text(account.get(field))
                    }
                    matches.append(({"account": account_facts}, len(labels)))
    core.require(
        len(matches) <= 1,
        f"{label}: finance material is assigned to multiple observation targets",
    )
    return matches[0] if matches else None


def _normalize_finance_observation_facts(
    value: object,
    *,
    classification: str,
    canonical: dict[str, Any],
    default_to_canonical: bool,
    field: str,
) -> dict[str, Any]:
    if value is None:
        if not default_to_canonical:
            return {}
        if classification == "document":
            return (
                {"document": copy.deepcopy(canonical["document"])}
                if canonical.get("document")
                else {}
            )
        return copy.deepcopy(canonical)
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    allowed_sections = {
        "document": {"holder", "document"},
        "chat_profile": {"account"},
        "reference": set(),
    }[classification]
    unknown_sections = sorted(set(value) - allowed_sections)
    core.require(
        not unknown_sections,
        f"{field} has unsupported sections: {', '.join(unknown_sections)}",
    )
    allowed_fields = {
        "holder": {"name", "surname", "given_names", "nationality", "birth_date"},
        "document": {"type", "country_code", "number"},
        "account": {"platform", "account_id", "phone"},
    }
    normalized: dict[str, Any] = {}
    for section, raw_section in value.items():
        core.require(isinstance(raw_section, Mapping), f"{field}.{section} must be an object")
        unknown_fields = sorted(set(raw_section) - allowed_fields[section])
        core.require(
            not unknown_fields,
            f"{field}.{section} has unsupported fields: {', '.join(unknown_fields)}",
        )
        section_values = {
            name: core.clean_text(raw)
            for name, raw in raw_section.items()
            if core.clean_text(raw)
        }
        if section == "document" and "type" in section_values:
            section_values["type"] = finance_materials.normalize_document_type(
                section_values["type"], field=f"{field}.document.type"
            )
        if section == "document" and "country_code" in section_values:
            section_values["country_code"] = section_values["country_code"].upper()
        if section == "account" and "platform" in section_values:
            section_values["platform"] = finance_materials.normalize_platform(
                section_values["platform"], field=f"{field}.account.platform"
            )
        canonical_section = canonical.get(section, {})
        if canonical_section:
            for name, item in section_values.items():
                core.require(
                    canonical_section.get(name) == item,
                    f"{field}.{section}.{name} disagrees with the submitted finance record",
                )
        if section_values:
            normalized[section] = section_values
    return normalized


def _normalize_full_media_observation(
    value: object,
    *,
    label: str,
    decision: Mapping[str, Any],
    field: str,
) -> dict[str, Any]:
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    unknown = sorted(
        set(value)
        - {
            "contract_version",
            "classification",
            "review_status",
            "viewed_original",
            "recheck_reasons",
            "facts",
            "note",
        }
    )
    core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
    core.require(
        value.get("contract_version") == MEDIA_OBSERVATION_CONTRACT,
        f"{field}.contract_version must be {MEDIA_OBSERVATION_CONTRACT}",
    )
    media_decisions = decision.get("media_decisions")
    core.require(isinstance(media_decisions, Mapping), "media_decisions must be an object")
    media_decision = media_decisions.get(label)
    core.require(isinstance(media_decision, Mapping), f"{field} cites an unclassified media label")
    classification = core.clean_text(value.get("classification")).casefold()
    decision_classification = core.clean_text(media_decision.get("classification")).casefold()
    core.require(
        classification == decision_classification,
        f"{field}.classification must match media_decisions.{label}",
    )
    finance_mode = decision.get("contract_version") == finance_materials.DECISION_CONTRACT
    store_mode = decision.get("contract_version") == store_ledger.DECISION_CONTRACT
    allowed_classifications = (
        finance_materials.MEDIA_CLASSIFICATIONS
        if finance_mode
        else store_ledger.MEDIA_CLASSIFICATIONS
        if store_mode
        else {"fund", "reference"}
    )
    core.require(
        classification in allowed_classifications,
        f"{field}.classification is unsupported for this group mode",
    )
    review_status = core.clean_text(value.get("review_status")).casefold()
    core.require(
        review_status in MEDIA_REVIEW_STATUSES,
        f"{field}.review_status must be clear, recheck_required, or rechecked_unreadable",
    )
    viewed_original = value.get("viewed_original")
    core.require(isinstance(viewed_original, bool), f"{field}.viewed_original must be boolean")
    if classification not in {"reference"}:
        core.require(
            viewed_original is True,
            f"{field}.viewed_original must be true for {classification}",
        )
    reasons_value = value.get("recheck_reasons", [])
    core.require(isinstance(reasons_value, list), f"{field}.recheck_reasons must be a list")
    reasons = [core.clean_text(reason).casefold() for reason in reasons_value]
    core.require(all(reasons), f"{field}.recheck_reasons cannot contain blanks")
    core.require(len(reasons) == len(set(reasons)), f"{field}.recheck_reasons repeats a reason")
    unsupported_reasons = sorted(set(reasons) - MEDIA_RECHECK_REASONS)
    core.require(
        not unsupported_reasons,
        f"{field}.recheck_reasons has unsupported values: {', '.join(unsupported_reasons)}",
    )
    if review_status == "clear":
        core.require(not reasons, f"{field}.recheck_reasons must be empty when review_status is clear")
    else:
        core.require(bool(reasons), f"{field}.recheck_reasons is required for {review_status}")

    if finance_mode:
        target = _finance_observation_target(decision, label, classification)
        canonical = target[0] if target is not None else {}
        facts = _normalize_finance_observation_facts(
            value.get("facts"),
            classification=classification,
            canonical=canonical,
            default_to_canonical=target is not None and target[1] == 1,
            field=f"{field}.facts",
        )
        if classification != "reference" and review_status == "clear":
            core.require(
                bool(facts),
                f"{field}.facts must include at least one visible document or account field",
            )
    elif store_mode:
        canonical = store_ledger.normalize_visible_facts(
            media_decision.get("facts"),
            classification=classification,
            field=f"{field}.facts",
        )
        supplied_facts = value.get("facts")
        if supplied_facts is not None:
            supplied = store_ledger.normalize_visible_facts(
                supplied_facts,
                classification=classification,
                field=f"{field}.facts",
            )
            core.require(
                supplied == canonical,
                f"{field}.facts must match media_decisions.{label}.facts",
            )
        facts = canonical
        if classification != "reference" and review_status == "clear":
            core.require(
                bool(facts),
                f"{field}.facts must include at least one fact visible in the store image",
            )
    else:
        canonical = _order_observation_facts(media_decision)
        supplied_facts = value.get("facts")
        if supplied_facts is not None:
            core.require(isinstance(supplied_facts, Mapping), f"{field}.facts must be an object")
            core.require(
                dict(supplied_facts) == canonical,
                f"{field}.facts must match the normalized visible fields in media_decisions.{label}",
            )
        facts = canonical
        derived_recheck_reasons: list[str] = []
        if classification == "fund":
            fact_entries = facts.get("entries", [])
            if any(
                entry.get("amount_state") in {"partial", "unreadable"}
                for entry in fact_entries
                if isinstance(entry, Mapping)
            ):
                derived_recheck_reasons.append("amount_unreadable")
            if any(
                entry.get("payee_state") == "unreadable"
                for entry in fact_entries
                if isinstance(entry, Mapping)
            ):
                derived_recheck_reasons.append("payee_unreadable")
        if derived_recheck_reasons:
            core.require(
                review_status != "clear",
                f"{field}.review_status cannot be clear while visible fund fields are unreadable",
            )
            for reason in derived_recheck_reasons:
                if reason not in reasons:
                    reasons.append(reason)

    observation = {
        "contract_version": MEDIA_OBSERVATION_CONTRACT,
        "classification": classification,
        "review_status": review_status,
        "viewed_original": viewed_original,
        "recheck_reasons": reasons,
        "facts": facts,
    }
    note = core.clean_text(value.get("note"))
    if note:
        observation["note"] = note
    return observation


def _apply_media_observations(
    decision: dict[str, Any],
    batch: Mapping[str, Any],
    group: Mapping[str, Any],
) -> dict[str, int]:
    updates = batch.get("media_observations")
    if updates is None:
        return {
            "media_observations_recorded": 0,
            "observation_cache_reuses": 0,
            "observation_hashes_computed": 0,
            "observation_hashes_reused": 0,
        }
    assert isinstance(updates, Mapping)
    inventory = _media_inventory(group)
    label_hashes: dict[str, str] = {}
    hashes_computed = 0
    hashes_reused = 0
    for label in updates:
        digest, computed = _media_content_hash_for_label(
            group,
            decision,
            str(label),
            inventory=inventory,
            verify_file=False,
        )
        label_hashes[str(label)] = digest
        if computed:
            hashes_computed += 1
        else:
            hashes_reused += 1

    cache = decision.setdefault("media_observation_cache", _empty_media_observation_cache())
    core.require(isinstance(cache, dict), "media_observation_cache must be an object")
    core.require(
        cache.get("contract_version") == MEDIA_OBSERVATION_CACHE_CONTRACT,
        "unsupported media_observation_cache contract",
    )
    core.require(
        cache.get("observation_contract") == MEDIA_OBSERVATION_CONTRACT,
        "unsupported media observation contract in cache",
    )
    entries = cache.get("entries")
    core.require(isinstance(entries, dict), "media_observation_cache.entries must be an object")

    resolved: dict[str, dict[str, Any]] = {}
    reuse_count = 0
    for raw_label, raw_observation in updates.items():
        label = str(raw_label)
        core.require(isinstance(raw_observation, Mapping), f"media_observations.{label} must be an object")
        if "reuse_from" in raw_observation:
            core.require(
                set(raw_observation) == {"reuse_from"},
                f"media_observations.{label}.reuse_from cannot be combined with other fields",
            )
            source_label = core.clean_text(raw_observation.get("reuse_from"))
            core.require(
                source_label in resolved,
                f"media_observations.{label}.reuse_from must cite an earlier full observation in this batch",
            )
            core.require(
                label_hashes[label] == label_hashes.get(source_label),
                f"media_observations.{label}.reuse_from is only allowed for byte-identical media",
            )
            resolved[label] = copy.deepcopy(resolved[source_label])
            resolved[label] = _normalize_full_media_observation(
                resolved[label],
                label=label,
                decision=decision,
                field=f"media_observations.{label}",
            )
            reuse_count += 1
            continue
        if "reuse_sha256" in raw_observation:
            core.require(
                set(raw_observation) == {"reuse_sha256"},
                f"media_observations.{label}.reuse_sha256 cannot be combined with other fields",
            )
            digest = core.clean_text(raw_observation.get("reuse_sha256"))
            core.require(
                re.fullmatch(r"[0-9a-f]{64}", digest) is not None,
                f"media_observations.{label}.reuse_sha256 must be a lowercase SHA-256",
            )
            core.require(
                digest == label_hashes[label],
                f"media_observations.{label}.reuse_sha256 does not match the current media",
            )
            cache_entry = entries.get(digest)
            core.require(
                isinstance(cache_entry, Mapping),
                f"media_observations.{label}.reuse_sha256 is not present in the observation cache",
            )
            cached_observation = cache_entry.get("observation")
            core.require(
                isinstance(cached_observation, Mapping),
                f"media_observations.{label}.reuse_sha256 has an invalid cached observation",
            )
            core.require(
                cached_observation.get("review_status") != "recheck_required",
                f"media_observations.{label} cannot reuse an observation that still requires recheck",
            )
            resolved[label] = _normalize_full_media_observation(
                cached_observation,
                label=label,
                decision=decision,
                field=f"media_observations.{label}",
            )
            reuse_count += 1
            continue
        resolved[label] = _normalize_full_media_observation(
            raw_observation,
            label=label,
            decision=decision,
            field=f"media_observations.{label}",
        )

    for label, observation in resolved.items():
        digest = label_hashes[label]
        existing = entries.get(digest)
        source_labels: list[str] = []
        if isinstance(existing, Mapping):
            existing_labels = existing.get("source_labels", [])
            core.require(
                isinstance(existing_labels, list),
                f"media_observation_cache.entries.{digest}.source_labels must be a list",
            )
            source_labels = [str(item) for item in existing_labels]
        source_labels.append(label)
        entries[digest] = {
            "observation": copy.deepcopy(observation),
            "source_labels": sorted(set(source_labels)),
        }
    metrics = decision.setdefault("media_view_metrics", _empty_media_view_metrics())
    normalized_metrics = _normalize_media_view_metrics(
        metrics,
        field="decision.media_view_metrics",
    )
    normalized_metrics["observation_cache_reuses"] += reuse_count
    decision["media_view_metrics"] = normalized_metrics
    return {
        "media_observations_recorded": len(resolved),
        "observation_cache_reuses": reuse_count,
        "observation_hashes_computed": hashes_computed,
        "observation_hashes_reused": hashes_reused,
    }


def _media_view_performance(decision: Mapping[str, Any]) -> dict[str, Any]:
    metrics = _normalize_media_view_metrics(
        decision.get("media_view_metrics", _empty_media_view_metrics()),
        field="decision.media_view_metrics",
    )
    opened = metrics["opened_images"]
    return {
        **metrics,
        "average_ms_per_opened_image": (
            metrics["elapsed_ms"] / opened if opened else 0.0
        ),
        "failed_image_rate": metrics["failed_images"] / opened if opened else 0.0,
        "single_image_recheck_rate": (
            metrics["single_image_rechecks"] / opened if opened else 0.0
        ),
    }


def _adaptive_media_batch_policy(
    decision: Mapping[str, Any],
    *,
    finance_mode: bool,
) -> dict[str, Any]:
    default_limit = (
        FINANCE_MEDIA_BATCH_LIMIT if finance_mode else ORDER_MEDIA_BATCH_LIMIT
    )
    minimum_limit = FINANCE_MEDIA_BATCH_MIN if finance_mode else ORDER_MEDIA_BATCH_MIN
    maximum_limit = FINANCE_MEDIA_BATCH_MAX if finance_mode else ORDER_MEDIA_BATCH_MAX
    performance = _media_view_performance(decision)
    opened = int(performance["opened_images"])
    view_batches = int(performance["view_batches"])
    failed_rate = float(performance["failed_image_rate"])
    recheck_rate = float(performance["single_image_recheck_rate"])
    elapsed_ms = int(performance["elapsed_ms"])
    average_batch_ms = elapsed_ms / view_batches if view_batches else 0.0

    recommended = default_limit
    reason = "insufficient_samples"
    if opened > 0 and (
        failed_rate >= 0.30
        or recheck_rate >= 0.40
        or (average_batch_ms >= 90_000 and elapsed_ms > 0)
    ):
        recommended = minimum_limit
        reason = "severe_failure_recheck_or_latency"
    elif opened > 0 and (
        failed_rate >= 0.15
        or recheck_rate >= 0.25
        or (average_batch_ms >= 60_000 and elapsed_ms > 0)
    ):
        recommended = max(
            minimum_limit,
            default_limit - (2 if finance_mode else 3),
        )
        reason = "high_failure_recheck_or_latency"
    elif opened >= default_limit and (
        failed_rate >= 0.05
        or recheck_rate >= 0.12
        or (average_batch_ms >= 35_000 and elapsed_ms > 0)
    ):
        recommended = max(
            minimum_limit,
            default_limit - (1 if finance_mode else 2),
        )
        reason = "moderate_failure_recheck_or_latency"
    elif (
        opened >= default_limit * 4
        and elapsed_ms > 0
        and failed_rate == 0
        and recheck_rate <= 0.02
        and average_batch_ms <= 10_000
    ):
        recommended = maximum_limit
        reason = "long_sustained_fast_clear_batches"
    elif (
        opened >= default_limit * 2
        and elapsed_ms > 0
        and failed_rate == 0
        and recheck_rate <= 0.03
        and average_batch_ms <= 15_000
    ):
        recommended = min(maximum_limit, default_limit + 2)
        reason = "sustained_fast_clear_batches"
    elif (
        opened >= default_limit
        and elapsed_ms > 0
        and failed_rate <= 0.02
        and recheck_rate <= 0.06
        and average_batch_ms <= 25_000
    ):
        recommended = min(maximum_limit, default_limit + 1)
        reason = "fast_clear_batches"
    elif opened >= default_limit:
        reason = "stable_default"

    return {
        "contract_version": MEDIA_BATCH_POLICY_CONTRACT,
        "strategy": "adaptive-v1",
        "minimum_parallel_limit": minimum_limit,
        "default_parallel_limit": default_limit,
        "maximum_parallel_limit": maximum_limit,
        "recommended_parallel_limit": recommended,
        "reason": reason,
        "sample_opened_images": opened,
        "sample_view_batches": view_batches,
        "average_batch_elapsed_ms": average_batch_ms,
        "failed_image_rate": failed_rate,
        "single_image_recheck_rate": recheck_rate,
    }


def _validate_media_observation_cache(
    group: Mapping[str, Any],
    decision: dict[str, Any],
    *,
    require_complete: bool,
    verify_files: bool,
) -> dict[str, Any]:
    cache = decision.get("media_observation_cache")
    if cache is None:
        return {
            "observation_cache_entries": 0,
            "observation_cache_labels": 0,
            "media_recheck_required": 0,
            "media_recheck_queue": [],
            "media_view_performance": _media_view_performance(decision),
        }
    core.require(isinstance(cache, dict), "media_observation_cache must be an object")
    unknown_cache = sorted(
        set(cache) - {"contract_version", "observation_contract", "entries"}
    )
    core.require(
        not unknown_cache,
        "media_observation_cache has unsupported fields: " + ", ".join(unknown_cache),
    )
    core.require(
        cache.get("contract_version") == MEDIA_OBSERVATION_CACHE_CONTRACT,
        "unsupported media_observation_cache contract",
    )
    core.require(
        cache.get("observation_contract") == MEDIA_OBSERVATION_CONTRACT,
        "unsupported media observation contract in cache",
    )
    entries = cache.get("entries")
    core.require(isinstance(entries, dict), "media_observation_cache.entries must be an object")
    inventory = _media_inventory(group)
    pending: list[dict[str, Any]] = []
    label_count = 0
    verified_paths: set[str] = set()
    for digest, raw_entry in entries.items():
        field = f"media_observation_cache.entries.{digest}"
        core.require(
            re.fullmatch(r"[0-9a-f]{64}", str(digest)) is not None,
            f"{field}: key must be a lowercase SHA-256",
        )
        core.require(isinstance(raw_entry, dict), f"{field} must be an object")
        unknown_entry = sorted(set(raw_entry) - {"observation", "source_labels"})
        core.require(not unknown_entry, f"{field} has unsupported fields: {', '.join(unknown_entry)}")
        source_value = raw_entry.get("source_labels")
        core.require(isinstance(source_value, list) and source_value, f"{field}.source_labels is required")
        source_labels = [str(label) for label in source_value]
        core.require(
            len(source_labels) == len(set(source_labels)),
            f"{field}.source_labels repeats a label",
        )
        core.require(
            source_labels == sorted(source_labels),
            f"{field}.source_labels must be sorted",
        )
        normalized_observation: dict[str, Any] | None = None
        for label in source_labels:
            core.require(label in inventory, f"{field}.source_labels contains unknown media: {label}")
            _, source_media = inventory[label]
            source_path_key = str(Path(str(source_media.get("path") or "")).resolve()).casefold()
            media_decision = decision.get("media_decisions", {}).get(label, {})
            evidence_was_verified = bool(
                isinstance(media_decision, Mapping)
                and core.clean_text(media_decision.get("evidence_sha256"))
            )
            should_verify = (
                verify_files
                and not evidence_was_verified
                and source_path_key not in verified_paths
            )
            current_hash, _ = _media_content_hash_for_label(
                group,
                decision,
                label,
                inventory=inventory,
                verify_file=should_verify,
            )
            core.require(current_hash == digest, f"{field} is indexed under the wrong media hash")
            if verify_files and (should_verify or evidence_was_verified):
                verified_paths.add(source_path_key)
            normalized = _normalize_full_media_observation(
                raw_entry.get("observation"),
                label=label,
                decision=decision,
                field=f"{field}.observation",
            )
            if normalized_observation is None:
                normalized_observation = normalized
            else:
                core.require(
                    normalized == normalized_observation,
                    f"{field}.observation is not compatible with all source labels",
                )
        assert normalized_observation is not None
        raw_entry["observation"] = normalized_observation
        raw_entry["source_labels"] = source_labels
        label_count += len(source_labels)
        if normalized_observation["review_status"] == "recheck_required":
            representative = source_labels[0]
            _, media = inventory[representative]
            pending.append(
                {
                    "content_sha256": digest,
                    "representative_label": representative,
                    "covered_labels": source_labels,
                    "path": media.get("path"),
                    "recheck_reasons": normalized_observation["recheck_reasons"],
                }
            )
    core.require(
        not require_complete or not pending,
        "media observation recheck is still required: "
        + ", ".join(item["representative_label"] for item in pending[:20]),
    )
    return {
        "observation_cache_entries": len(entries),
        "observation_cache_labels": label_count,
        "media_recheck_required": len(pending),
        "media_recheck_queue": pending[:MEDIA_RECHECK_QUEUE_LIMIT],
        "media_view_performance": _media_view_performance(decision),
    }


def _validate_entry(
    entry: dict[str, Any],
    *,
    field: str,
    sender_role: str,
    source_message_label: str,
    message_labels: set[str],
) -> None:
    allowed = {
        "amount",
        "amount_text",
        "amount_basis",
        "currency",
        "payee",
        "payee_state",
        "kind",
        "side",
        "side_exception",
        "result",
        "status_text",
        "amount_state",
        "note",
    }
    unknown = sorted(set(entry) - allowed)
    core.require(not unknown, f"{field}: unsupported fields: {', '.join(unknown)}")
    kind = core.clean_text(entry.get("kind")).casefold() or "transfer"
    core.require(kind in ENTRY_KINDS, f"{field}.kind must be transfer or cash")
    entry["kind"] = kind
    exception = entry.get("side_exception")
    side = core.clean_text(entry.get("side")).casefold()
    if not side:
        exception_kind_hint = (
            core.clean_text(exception.get("kind")).casefold()
            if isinstance(exception, Mapping)
            else ""
        )
        side = SIDE_EXCEPTION_SIDES.get(
            exception_kind_hint,
            {
                "内部人员": "payout",
                "客户候选": "payment",
            }.get(sender_role, "unknown"),
        )
    core.require(side in FLOW_SIDES, f"{field}.side is unsupported")
    exception_kind: str | None = None
    if exception not in (None, ""):
        core.require(isinstance(exception, dict), f"{field}.side_exception must be an object")
        unknown_exception_fields = sorted(
            set(exception) - {"kind", "source_messages", "detail"}
        )
        core.require(
            not unknown_exception_fields,
            f"{field}.side_exception: unsupported fields: {', '.join(unknown_exception_fields)}",
        )
        exception_kind = core.clean_text(exception.get("kind")).casefold()
        core.require(
            exception_kind in SIDE_EXCEPTION_SIDES,
            f"{field}.side_exception.kind is unsupported",
        )
        core.require(
            SIDE_EXCEPTION_SIDES[exception_kind] == side,
            f"{field}.side_exception.kind does not support side={side}",
        )
        raw_exception_sources = exception.get("source_messages")
        core.require(
            isinstance(raw_exception_sources, list) and raw_exception_sources,
            f"{field}.side_exception.source_messages is required",
        )
        exception_sources = [str(item) for item in raw_exception_sources]
        core.require(
            len(exception_sources) == len(set(exception_sources)),
            f"{field}.side_exception.source_messages repeats a label",
        )
        core.require(
            set(exception_sources) <= message_labels,
            f"{field}.side_exception.source_messages contains an unknown label",
        )
        core.require(
            any(label != source_message_label for label in exception_sources),
            f"{field}.side_exception must cite chat context beyond the fund image itself",
        )
        detail = core.clean_text(exception.get("detail"))
        core.require(bool(detail), f"{field}.side_exception.detail is required")
        exception["kind"] = exception_kind
        exception["source_messages"] = exception_sources
        exception["detail"] = detail

    if side in {"payment", "payout"}:
        expected_side = {
            "内部人员": "payout",
            "客户候选": "payment",
        }.get(sender_role)
        if expected_side is None:
            core.require(
                side == "unknown",
                f"{field}: unknown sender role ordinary fund entry must use side=unknown",
            )
        elif side != expected_side:
            required_exception = (
                "relayed_customer_payment"
                if sender_role == "内部人员"
                else "relayed_internal_payout"
            )
            role_name = "internal" if sender_role == "内部人员" else "customer"
            core.require(
                exception_kind == required_exception,
                f"{field}: {role_name} sender ordinary fund entry must use "
                f"side={expected_side} unless side_exception.kind={required_exception}",
            )
        else:
            core.require(
                exception_kind is None,
                f"{field}.side_exception is only allowed when ordinary side differs from sender role",
            )
    elif side == "payment_refund":
        core.require(
            exception_kind == "explicit_payment_refund",
            f"{field}: side=payment_refund requires side_exception.kind=explicit_payment_refund",
        )
    elif side == "recovery":
        core.require(
            exception_kind == "explicit_recovery",
            f"{field}: side=recovery requires side_exception.kind=explicit_recovery",
        )
    else:
        core.require(
            exception_kind is None,
            f"{field}.side_exception is not allowed when side=unknown",
        )
    entry["side"] = side
    status_text = core.clean_text(entry.get("status_text"))
    if status_text:
        entry["status_text"] = status_text
    normalized_status_text = status_text.casefold()
    failure_text = f"{status_text} {core.clean_text(entry.get('note'))}".casefold()
    result = core.clean_text(entry.get("result")).casefold()
    if not result:
        result = (
            "failed"
            if any(marker in failure_text for marker in FAILURE_STATUS_MARKERS)
            else "completed"
        )
    core.require(result in ENTRY_RESULTS, f"{field}.result is unsupported")
    if any(marker in normalized_status_text for marker in NORMAL_PROCESSING_STATUS_MARKERS):
        core.require(
            result == "completed",
            f"{field}: normal processing status must use result=completed",
        )
    if any(marker in normalized_status_text for marker in FAILURE_STATUS_MARKERS):
        core.require(
            result == "failed",
            f"{field}: explicit failure status must use result=failed",
        )
    if result == "failed":
        core.require(
            any(marker in failure_text for marker in FAILURE_STATUS_MARKERS),
            f"{field}: result=failed requires explicit failure, cancellation, rejection, invalidation, or risk-control text",
        )
    entry["result"] = result
    amount_state = core.clean_text(entry.get("amount_state")).casefold()
    if not amount_state:
        has_amount = entry.get("amount") not in (None, "")
        has_currency = entry.get("currency") not in (None, "")
        amount_state = "clear" if has_amount and has_currency else "partial" if has_amount or has_currency else "unreadable"
    core.require(amount_state in AMOUNT_STATES, f"{field}.amount_state is unsupported")
    entry["amount_state"] = amount_state
    if entry.get("amount") not in (None, ""):
        amount = core.parse_decimal(entry.get("amount"), field=f"{field}.amount")
        core.require(amount is not None and amount >= 0, f"{field}.amount cannot be negative")
        entry["amount"] = core.decimal_text(amount)
    if entry.get("currency") not in (None, ""):
        entry["currency"] = core.normalize_currency(entry.get("currency"), field=f"{field}.currency")
    if amount_state == "clear":
        core.require(entry.get("amount") not in (None, ""), f"{field}.amount is required when amount_state is clear")
        core.require(entry.get("currency") not in (None, ""), f"{field}.currency is required when amount_state is clear")
        expected_amount_basis = (
            "cash_face_value"
            if kind == "cash"
            else "receiver_received"
            if result == "completed"
            else "attempted_not_received"
        )
        amount_basis = core.clean_text(entry.get("amount_basis")).casefold()
        if not amount_basis:
            amount_basis = expected_amount_basis
        core.require(
            amount_basis in AMOUNT_BASES,
            f"{field}.amount_basis is unsupported",
        )
        core.require(
            amount_basis == expected_amount_basis,
            f"{field}.amount_basis must be {expected_amount_basis}",
        )
        entry["amount_basis"] = amount_basis
    elif entry.get("amount_basis") not in (None, ""):
        amount_basis = core.clean_text(entry.get("amount_basis")).casefold()
        core.require(
            amount_basis in AMOUNT_BASES,
            f"{field}.amount_basis is unsupported",
        )
        entry["amount_basis"] = amount_basis
    if entry.get("payee_state") in (None, ""):
        payee_text = core.clean_text(entry.get("payee"))
        entry["payee_state"] = (
            "cash"
            if kind == "cash"
            else "not_shown"
            if payee_text == "未显示"
            else "unreadable"
            if payee_text == "无法辨认"
            else "visible"
        )
    if kind == "transfer" and entry.get("currency") == "THB":
        entry["payee"] = core.validate_thai_bank_account_payee(
            entry.get("payee"),
            field=f"{field}.payee",
            payee_state=entry.get("payee_state"),
        )
    else:
        entry["payee"] = core.validate_payee(
            entry.get("payee"),
            field=f"{field}.payee",
            cash=kind == "cash",
            payee_state=entry.get("payee_state"),
            require_state=True,
        )
    if entry.get("payee_state") not in (None, ""):
        entry["payee_state"] = core.clean_text(entry.get("payee_state")).casefold()


def _validate_network_fees(value: object, *, field: str) -> None:
    if value in (None, "", []):
        return
    core.require(isinstance(value, list), f"{field} must be a list")
    for position, fee in enumerate(value):
        core.require(isinstance(fee, Mapping), f"{field}[{position}] must be an object")
        if core.clean_text(fee.get("kind")).casefold() == "network_fee":
            core.require(
                fee.get("customer_requested") is True,
                f"{field}[{position}].customer_requested must be true for a network fee",
            )


def _validate_pricing_scope(
    value: object,
    *,
    field: str,
    source_labels: set[str],
    payout_currency: str | None,
    require_complete: bool,
) -> bool:
    """Validate one v3 pricing authority. Return True only when it is missing."""

    if value in (None, ""):
        core.require(not require_complete, f"{field} is required before seal")
        return True
    core.require(isinstance(value, dict), f"{field} must be an object")
    pricing = value
    unknown = sorted(set(pricing) - {"source_messages", "terms", "expected"})
    core.require(not unknown, f"{field}: unsupported fields: {', '.join(unknown)}")

    raw_sources = pricing.get("source_messages")
    core.require(
        isinstance(raw_sources, list) and raw_sources,
        f"{field}.source_messages is required",
    )
    cited = [str(item) for item in raw_sources]
    core.require(
        len(cited) == len(set(cited)),
        f"{field}.source_messages repeats a label",
    )
    core.require(
        set(cited) <= source_labels,
        f"{field}.source_messages must cite labels already present in order.source_messages",
    )

    terms_value = pricing.get("terms")
    terms: dict[str, Any] | None = None
    if terms_value not in (None, ""):
        core.require(isinstance(terms_value, dict), f"{field}.terms must be an object")
        terms = terms_value
        unknown_terms = sorted(set(terms) - {"rate", "operator", "fees", "rounding"})
        core.require(
            not unknown_terms,
            f"{field}.terms: unsupported fields: {', '.join(unknown_terms)}",
        )
        rate = core.parse_decimal(terms.get("rate"), field=f"{field}.terms.rate")
        core.require(rate is not None and rate > 0, f"{field}.terms.rate must be positive")
        terms["rate"] = core.decimal_text(rate)
        operator = core.clean_text(terms.get("operator")).casefold()
        core.require(
            operator in simple_ledger.RATE_OPERATORS,
            f"{field}.terms.operator is required and must be multiply or divide",
        )
        terms["operator"] = operator
        _validate_network_fees(terms.get("fees"), field=f"{field}.terms.fees")
        rounding = terms.get("rounding")
        if rounding not in (None, ""):
            core.require(isinstance(rounding, dict), f"{field}.terms.rounding must be an object")
            unknown_rounding = sorted(set(rounding) - {"unit", "mode", "currency"})
            core.require(
                not unknown_rounding,
                f"{field}.terms.rounding: unsupported fields: {', '.join(unknown_rounding)}",
            )
            unit = core.parse_decimal(
                rounding.get("unit"),
                field=f"{field}.terms.rounding.unit",
            )
            core.require(unit is not None and unit > 0, f"{field}.terms.rounding.unit must be positive")
            rounding["unit"] = core.decimal_text(unit)
            mode = core.clean_text(rounding.get("mode")).casefold() or "half_up"
            core.require(
                mode in simple_ledger.ROUNDING_MODES,
                f"{field}.terms.rounding.mode is unsupported",
            )
            rounding["mode"] = mode
            if rounding.get("currency") not in (None, ""):
                rounding["currency"] = core.normalize_currency(
                    rounding.get("currency"),
                    field=f"{field}.terms.rounding.currency",
                )
                if payout_currency:
                    core.require(
                        rounding["currency"] == payout_currency,
                        f"{field}.terms.rounding.currency must use {payout_currency}",
                    )

    expected = pricing.get("expected")
    core.require(isinstance(expected, dict), f"{field}.expected is required")
    kind = core.clean_text(expected.get("kind")).casefold()
    core.require(kind in PRICING_EXPECTED_KINDS, f"{field}.expected.kind is unsupported")
    expected["kind"] = kind
    if kind == "explicit":
        unknown_expected = sorted(set(expected) - {"kind", "amount"})
        core.require(
            not unknown_expected,
            f"{field}.expected: unsupported fields: {', '.join(unknown_expected)}",
        )
        amount = core.parse_decimal(expected.get("amount"), field=f"{field}.expected.amount")
        core.require(amount is not None and amount >= 0, f"{field}.expected.amount cannot be negative")
        expected["amount"] = core.decimal_text(amount)
    elif kind == "calculated_from_terms":
        unknown_expected = sorted(set(expected) - {"kind"})
        core.require(
            not unknown_expected,
            f"{field}.expected: unsupported fields: {', '.join(unknown_expected)}",
        )
        core.require(terms is not None, f"{field}.terms is required for calculated_from_terms")
    else:
        unknown_expected = sorted(set(expected) - {"kind", "reason"})
        core.require(
            not unknown_expected,
            f"{field}.expected: unsupported fields: {', '.join(unknown_expected)}",
        )
        reason = core.clean_text(expected.get("reason")).casefold()
        core.require(
            reason in PRICING_UNKNOWN_REASONS,
            f"{field}.expected.reason is required and unsupported",
        )
        expected["reason"] = reason
    if require_complete and kind != "unknown":
        core.require(
            terms is not None,
            f"{field}.terms with the confirmed exchange rate is required before seal",
        )
    return False


def _validate_open_orders(
    group: Mapping[str, Any],
    decision: dict[str, Any],
    *,
    reviewed_through: int,
    require_complete: bool,
    require_fund_type: bool = False,
) -> int:
    open_orders = decision.get("open_orders")
    core.require(isinstance(open_orders, list), "open_orders must be a list")
    if require_complete:
        core.require(
            not open_orders,
            "open_orders must be resolved, cancelled, or converted to final orders before seal",
        )

    _, message_by_label = _message_labels(group)
    message_positions = {
        label: position
        for position, label in enumerate(message_by_label, start=1)
    }
    label_by_message_id, _ = _message_labels(group)
    inventory = _media_inventory(group)
    media_source_labels = {
        media_label: label_by_message_id.get(str(message.get("message_id") or ""), "")
        for media_label, (message, _) in inventory.items()
    }
    seen_ids: set[str] = set()
    allowed_fields = {
        "id",
        "start_message",
        "source_messages",
        "media_labels",
        "customer_nickname",
        "direction",
        "rate",
        "operator",
        "summary",
        "unresolved",
    }
    if require_fund_type:
        allowed_fields.add("fund_type")
    for position, item in enumerate(open_orders):
        field = f"open_orders[{position}]"
        core.require(isinstance(item, dict), f"{field} must be an object")
        unknown_fields = sorted(set(item) - allowed_fields)
        core.require(
            not unknown_fields,
            f"{field} has unsupported fields: {', '.join(unknown_fields)}",
        )
        open_id = core.clean_text(item.get("id"))
        core.require(bool(open_id), f"{field}.id is required")
        core.require(open_id not in seen_ids, f"open_orders repeats id {open_id}")
        seen_ids.add(open_id)
        item["id"] = open_id

        start_message = core.clean_text(item.get("start_message"))
        core.require(start_message in message_positions, f"{field}.start_message is unknown")
        source_value = item.get("source_messages")
        core.require(
            isinstance(source_value, list) and source_value,
            f"{field}.source_messages must be a non-empty list",
        )
        source_messages = [str(value) for value in source_value]
        core.require(
            len(source_messages) == len(set(source_messages)),
            f"{field}.source_messages repeats a label",
        )
        core.require(
            set(source_messages) <= set(message_positions),
            f"{field}.source_messages contains an unknown label",
        )
        core.require(
            start_message in source_messages,
            f"{field}.source_messages must include start_message",
        )
        core.require(
            all(message_positions[label] <= reviewed_through for label in source_messages),
            f"{field}.source_messages cannot cite an uncommitted page",
        )
        item["start_message"] = start_message
        item["source_messages"] = source_messages

        media_value = item.get("media_labels", [])
        core.require(isinstance(media_value, list), f"{field}.media_labels must be a list")
        media_labels = [str(value) for value in media_value]
        core.require(
            len(media_labels) == len(set(media_labels)),
            f"{field}.media_labels repeats a label",
        )
        core.require(
            set(media_labels) <= set(inventory),
            f"{field}.media_labels contains an unknown label",
        )
        core.require(
            all(media_source_labels[label] in source_messages for label in media_labels),
            f"{field}.source_messages must include every media source message",
        )
        item["media_labels"] = media_labels

        if "fund_type" in item:
            item["fund_type"] = large_daily.normalize_fund_type(
                item.get("fund_type"), field=f"{field}.fund_type"
            )

        if "customer_nickname" in item:
            item["customer_nickname"] = core.clean_text(item.get("customer_nickname"))
        direction = core.clean_text(item.get("direction"))
        if direction:
            item["direction"] = core.canonical_direction(direction, field=f"{field}.direction")
        elif "direction" in item:
            item["direction"] = ""

        rate_text = core.clean_text(item.get("rate"))
        operator = core.clean_text(item.get("operator")).casefold()
        if rate_text:
            rate = core.parse_decimal(rate_text, field=f"{field}.rate")
            core.require(rate is not None and rate > 0, f"{field}.rate must be positive")
            core.require(operator in {"multiply", "divide"}, f"{field}.operator is required with rate")
            item["rate"] = core.decimal_text(rate)
            item["operator"] = operator
        else:
            core.require(not operator, f"{field}.operator requires rate")
            item.pop("rate", None)
            item.pop("operator", None)

        summary = core.clean_text(item.get("summary"))
        core.require(bool(summary), f"{field}.summary is required")
        item["summary"] = summary
        unresolved_value = item.get("unresolved")
        core.require(
            isinstance(unresolved_value, list) and unresolved_value,
            f"{field}.unresolved must explain why the order is still open",
        )
        unresolved = [core.clean_text(value) for value in unresolved_value]
        core.require(all(unresolved), f"{field}.unresolved cannot contain blank items")
        item["unresolved"] = unresolved
    return len(open_orders)


def _order_continuity_review_candidates(
    facts: list[dict[str, Any]],
    *,
    advanced_relation_order_ids: set[str],
) -> list[dict[str, Any]]:
    """Return narrow, non-mutating candidates for a possible split continuation."""

    grouped: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for fact in facts:
        customer_key = core.clean_text(fact.get("customer_key"))
        direction = core.clean_text(fact.get("direction"))
        if not customer_key or not direction:
            continue
        grouped.setdefault((customer_key, direction), []).append(fact)

    candidates: list[dict[str, Any]] = []
    for related_orders in grouped.values():
        related_orders.sort(
            key=lambda item: (
                int(item["source_start_position"]),
                int(item["source_end_position"]),
                int(item["order_position"]),
            )
        )
        for earlier, later in zip(related_orders, related_orders[1:]):
            earlier_order_id = str(earlier["order_id"])
            later_order_id = str(later["order_id"])
            if (
                earlier_order_id in advanced_relation_order_ids
                or later_order_id in advanced_relation_order_ids
                or earlier["has_complex_flow"]
                or later["has_complex_flow"]
            ):
                continue
            if not (
                earlier["has_completed_payment"]
                and not earlier["has_completed_payout"]
                and earlier["has_confirmed_pricing"]
                and later["has_completed_payment"]
                and later["has_completed_payout"]
                and later["has_unknown_pricing"]
            ):
                continue
            candidates.append(
                {
                    "earlier_order_id": earlier_order_id,
                    "later_order_id": later_order_id,
                    "customer_nickname": earlier["customer_nickname"],
                    "direction": earlier["direction"],
                    "source_start": f"S{min(int(earlier['source_start_position']), int(later['source_start_position'])):05d}",
                    "source_end": f"S{max(int(earlier['source_end_position']), int(later['source_end_position'])):05d}",
                    "reason_codes": [
                        "earlier_payment_without_payout",
                        "earlier_confirmed_pricing",
                        "later_payment_and_payout",
                        "later_unknown_pricing",
                    ],
                }
            )
    return candidates


def _validate_open_exchanges(
    group: Mapping[str, Any],
    decision: dict[str, Any],
    *,
    reviewed_through: int,
    require_complete: bool,
) -> int:
    open_exchanges = decision.get("open_exchanges")
    core.require(isinstance(open_exchanges, list), "open_exchanges must be a list")
    if require_complete:
        core.require(
            not open_exchanges,
            "open_exchanges must be resolved or discarded before seal",
        )
    _, message_by_label = _message_labels(group)
    positions = {
        label: position for position, label in enumerate(message_by_label, start=1)
    }
    seen_ids: set[str] = set()
    allowed_fields = {
        "id",
        "start_message",
        "source_messages",
        "fund_type",
        "direction",
        "rate",
        "operator",
        "summary",
        "unresolved",
    }
    for position, item in enumerate(open_exchanges):
        field = f"open_exchanges[{position}]"
        core.require(isinstance(item, dict), f"{field} must be an object")
        unknown = sorted(set(item) - allowed_fields)
        core.require(not unknown, f"{field}: unsupported fields: {', '.join(unknown)}")
        exchange_id = core.clean_text(item.get("id"))
        core.require(bool(exchange_id), f"{field}.id is required")
        core.require(exchange_id not in seen_ids, f"open_exchanges repeats id {exchange_id}")
        seen_ids.add(exchange_id)
        item["id"] = exchange_id
        start_message = core.clean_text(item.get("start_message"))
        core.require(start_message in positions, f"{field}.start_message is unknown")
        source_value = item.get("source_messages")
        core.require(
            isinstance(source_value, list) and source_value,
            f"{field}.source_messages must be a non-empty list",
        )
        source_messages = [str(value) for value in source_value]
        core.require(
            len(source_messages) == len(set(source_messages)),
            f"{field}.source_messages repeats a label",
        )
        core.require(
            start_message in source_messages and set(source_messages) <= set(positions),
            f"{field}.source_messages must contain valid labels including start_message",
        )
        core.require(
            all(positions[label] <= reviewed_through for label in source_messages),
            f"{field}.source_messages cannot cite an uncommitted page",
        )
        item["start_message"] = start_message
        item["source_messages"] = source_messages
        if "fund_type" in item:
            fund_type = core.clean_text(item.get("fund_type")).casefold()
            core.require(fund_type in large_daily.FUND_TYPES, f"{field}.fund_type is unsupported")
            item["fund_type"] = fund_type
        if core.clean_text(item.get("direction")):
            item["direction"] = core.canonical_direction(
                item.get("direction"), field=f"{field}.direction"
            )
        rate_text = core.clean_text(item.get("rate"))
        operator = core.clean_text(item.get("operator")).casefold()
        if rate_text:
            rate = core.parse_decimal(rate_text, field=f"{field}.rate")
            core.require(rate is not None and rate > 0, f"{field}.rate must be positive")
            core.require(
                operator in large_daily.RATE_OPERATORS,
                f"{field}.operator is required with rate",
            )
            item["rate"] = core.decimal_text(rate)
            item["operator"] = operator
        else:
            core.require(not operator, f"{field}.operator requires rate")
            item.pop("rate", None)
            item.pop("operator", None)
        summary = core.clean_text(item.get("summary"))
        core.require(bool(summary), f"{field}.summary is required")
        item["summary"] = summary
        unresolved_value = item.get("unresolved")
        core.require(
            isinstance(unresolved_value, list) and unresolved_value,
            f"{field}.unresolved must explain why the exchange is still open",
        )
        unresolved = [core.clean_text(value) for value in unresolved_value]
        core.require(all(unresolved), f"{field}.unresolved cannot contain blank items")
        item["unresolved"] = unresolved
    return len(open_exchanges)


def _validate_large_media_decisions(
    group: Mapping[str, Any],
    decision: dict[str, Any],
    *,
    require_complete: bool,
    capture_hashes: bool,
    rehash_labels: set[str] | None,
) -> tuple[dict[str, Any], set[str], set[str]]:
    inventory = _media_inventory(group)
    label_by_message_id, message_by_label = _message_labels(group)
    message_labels = set(message_by_label)
    available = {
        label
        for label, (_, media) in inventory.items()
        if media.get("availability") == "available"
    }
    missing = set(inventory) - available
    media_decisions = decision.get("media_decisions")
    core.require(
        isinstance(media_decisions, dict),
        "media_decisions must be an object keyed by M labels",
    )
    unknown_labels = sorted(set(media_decisions) - available)
    core.require(
        not unknown_labels,
        f"media decisions contain missing or unknown labels: {unknown_labels[:10]}",
    )
    if require_complete:
        unclassified = sorted(available - set(media_decisions))
        core.require(not unclassified, f"unclassified available media: {unclassified[:20]}")

    entry_ids: set[str] = set()
    completed_entry_ids: set[str] = set()
    transfer_payees = 0
    unknown_payees: set[str] = set()
    unreadable_payees: set[str] = set()
    reference_count = 0
    fund_media_count = 0
    evidence_hashes_computed = 0
    evidence_hashes_reused = 0
    for label, raw in media_decisions.items():
        field = f"{group.get('group_key')}.media_decisions.{label}"
        core.require(isinstance(raw, dict), f"{field} must be an object")
        classification = core.clean_text(raw.get("classification")).casefold()
        core.require(
            classification in {"reference", "fund"},
            f"{field}.classification must be reference or fund",
        )
        if classification == "reference":
            unknown = sorted(set(raw) - {"classification", "note"})
            core.require(not unknown, f"{field}: reference has unsupported fields: {', '.join(unknown)}")
            reference_count += 1
            continue
        unknown = sorted(
            set(raw)
            - {"classification", "viewed_original", "evidence_sha256", "entries", "note"}
        )
        core.require(not unknown, f"{field}: fund decision has unsupported fields: {', '.join(unknown)}")
        core.require(
            raw.get("viewed_original") is True,
            f"{field}: open the original image before recording fund facts",
        )
        entries = raw.get("entries")
        core.require(
            isinstance(entries, list) and entries,
            f"{field}.entries must contain at least one fund entry",
        )
        message, media = inventory[label]
        media_path = Path(str(media.get("path") or ""))
        core.require(media_path.is_file(), f"{field}: original media file is unavailable: {media_path}")
        recorded_hash = core.clean_text(raw.get("evidence_sha256"))
        verify_hash = (
            rehash_labels is None
            or label in rehash_labels
            or (capture_hashes and not recorded_hash)
        )
        resolved_hash, computed = core.validate_cached_file_hash(
            media_path,
            recorded_hash,
            verify=verify_hash,
            changed_message=f"{field}: original fund evidence changed after review",
        )
        snapshot_hash = core.clean_text(media.get("blob_sha256"))
        if computed and snapshot_hash:
            core.require(
                resolved_hash == snapshot_hash,
                f"{field}: original fund evidence changed after the run snapshot was created",
            )
        if computed:
            evidence_hashes_computed += 1
        elif recorded_hash:
            evidence_hashes_reused += 1
        if not recorded_hash and capture_hashes:
            raw["evidence_sha256"] = resolved_hash
        elif require_complete:
            core.require(
                bool(recorded_hash),
                f"{field}: evidence hash has not been captured; run review seal",
            )
        source_message_label = label_by_message_id[str(message.get("message_id") or "")]
        for entry_position, entry in enumerate(entries, start=1):
            core.require(
                isinstance(entry, dict),
                f"{field}.entries[{entry_position - 1}] must be an object",
            )
            _validate_entry(
                entry,
                field=f"{field}.entries[{entry_position - 1}]",
                sender_role=core.clean_text(message.get("role")),
                source_message_label=source_message_label,
                message_labels=message_labels,
            )
            entry_id = f"{label}.{entry_position}"
            entry_ids.add(entry_id)
            if entry["result"] == "completed":
                completed_entry_ids.add(entry_id)
            if entry["kind"] == "transfer" and entry["result"] == "completed":
                transfer_payees += 1
                if entry["payee"] in core.UNKNOWN_PAYEES:
                    unknown_payees.add(entry_id)
                if entry.get("payee_state") == "unreadable":
                    unreadable_payees.add(entry_id)
        fund_media_count += 1

    reviewed_value = decision.get("unknown_payee_reviewed_entry_ids", [])
    core.require(
        isinstance(reviewed_value, list),
        "unknown_payee_reviewed_entry_ids must be a list",
    )
    reviewed = [str(item) for item in reviewed_value]
    core.require(
        len(reviewed) == len(set(reviewed)),
        "unknown_payee_reviewed_entry_ids repeats an entry",
    )
    core.require(
        set(reviewed) <= unknown_payees,
        "unknown_payee_reviewed_entry_ids may only cite unknown transfer payees",
    )
    unreviewed = unreadable_payees - set(reviewed)
    if require_complete:
        core.require(
            not unreviewed,
            f"unreadable payee review is required: {sorted(unreviewed)[:40]}",
        )
    warning = (
        transfer_payees >= UNKNOWN_PAYEE_WARNING_MIN_TRANSFERS
        and len(unknown_payees) * 2 > transfer_payees
    )
    return (
        {
            "available_media": len(available),
            "missing_media": len(missing),
            "classified_media": len(media_decisions),
            "reference_media": reference_count,
            "fund_media": fund_media_count,
            "fund_entries": len(entry_ids),
            "excluded_failed_or_incomplete_fund_entries": len(
                entry_ids - completed_entry_ids
            ),
            "transfer_payees": transfer_payees,
            "unknown_payees": len(unknown_payees),
            "unknown_payee_warning": int(warning),
            "unknown_payee_review_required": int(bool(unreadable_payees)),
            "unreviewed_unknown_payees": len(unreviewed),
            "evidence_hashes_computed": evidence_hashes_computed,
            "evidence_hashes_reused": evidence_hashes_reused,
        },
        entry_ids,
        completed_entry_ids,
    )


def _validate_large_decision(
    normalized: Mapping[str, Any],
    group: Mapping[str, Any],
    decision: dict[str, Any],
    *,
    require_complete: bool,
    capture_hashes: bool,
    rehash_labels: set[str] | None,
) -> dict[str, Any]:
    expected = _decision_template(normalized, group, group_mode="large")
    core.require(
        decision.get("contract_version") == large_daily.LEGACY_DECISION_CONTRACT,
        "unsupported legacy large decision contract; expected "
        f"{large_daily.LEGACY_DECISION_CONTRACT}",
    )
    for field in (
        "normalized_source_fingerprint",
        "group_fingerprint",
        "group_key",
        "group_name",
        "platform",
        "message_count",
        "evidence_media_count",
        "group_mode",
    ):
        core.require(
            decision.get(field) == expected.get(field),
            f"{group.get('group_key')}: generated {field} was edited",
        )
    messages = list(group.get("messages", []))
    reviewed_through = decision.get("reviewed_through")
    core.require(
        isinstance(reviewed_through, int) and 0 <= reviewed_through <= len(messages),
        "invalid reviewed_through",
    )
    core.require(isinstance(decision.get("read_complete"), bool), "read_complete must be boolean")
    core.require(
        decision.get("read_complete") is (reviewed_through == len(messages)),
        "read_complete must exactly match reviewed_through",
    )
    if require_complete:
        core.require(
            decision.get("read_complete") is True,
            "group chronology has not been read completely",
        )
    open_count = _validate_open_exchanges(
        group,
        decision,
        reviewed_through=reviewed_through,
        require_complete=require_complete,
    )
    media_statistics, entry_ids, completed_entry_ids = _validate_large_media_decisions(
        group,
        decision,
        require_complete=require_complete,
        capture_hashes=capture_hashes,
        rehash_labels=rehash_labels,
    )
    _, message_by_label = _message_labels(group)
    exchange_statistics = large_daily.validate_exchanges(
        decision.get("exchanges"),
        group_key=str(group.get("group_key") or ""),
        message_labels=set(message_by_label),
        fund_entry_ids=completed_entry_ids,
    )
    if require_complete:
        core.require(
            exchange_statistics["unassigned_exchange_entries"] == 0,
            "fund entries must be explicitly assigned to an exchange",
        )
    return {
        **media_statistics,
        **exchange_statistics,
        "open_exchanges": open_count,
        "unassigned_entries": exchange_statistics["unassigned_exchange_entries"],
    }


def _lifecycle_evidence_event(
    status: str,
    *,
    source_labels: list[str],
    message_by_label: Mapping[str, Mapping[str, Any]],
    reason: str,
) -> dict[str, Any]:
    cited = [label for label in source_labels if label in message_by_label]
    core.require(bool(cited), "order lifecycle requires at least one source message")
    latest = max(
        cited,
        key=lambda label: str(message_by_label[label].get("timestamp") or ""),
    )
    return {
        "status": status,
        "at": str(message_by_label[latest]["timestamp"]),
        "reason": reason,
        "source_messages": [latest],
    }


def _derived_order_lifecycle(
    *,
    source_labels: list[str],
    message_by_label: Mapping[str, Mapping[str, Any]],
    entry_records: Mapping[str, Mapping[str, Any]],
    entry_source_labels: Mapping[str, str],
    entry_ids: list[str],
) -> dict[str, Any]:
    completed_sides = {
        core.clean_text(entry_records[entry_id].get("side")).casefold()
        for entry_id in entry_ids
        if entry_records[entry_id].get("result") == "completed"
    }
    completed_labels = list(
        dict.fromkeys(
            entry_source_labels[entry_id]
            for entry_id in entry_ids
            if entry_records[entry_id].get("result") == "completed"
            and core.clean_text(entry_records[entry_id].get("side")).casefold()
            in {"payment", "payout"}
        )
    )
    cited_text = "\n".join(
        core.clean_text(message_by_label[label].get("text"))
        for label in source_labels
        if label in message_by_label
    ).casefold()
    has_cancel_marker = any(marker.casefold() in cited_text for marker in FAILURE_STATUS_MARKERS)
    all_failed_or_blank = all(
        core.clean_text(entry_records[entry_id].get("result")).casefold()
        in {"failed", "not_shown", "unknown"}
        for entry_id in entry_ids
    )
    signed_totals: dict[tuple[str, str], Any] = {}
    amount_unknown = False
    for entry_id in entry_ids:
        entry = entry_records[entry_id]
        if entry.get("result") != "completed":
            continue
        side = core.clean_text(entry.get("side")).casefold()
        currency = core.clean_text(entry.get("currency"))
        amount = core.parse_decimal(
            entry.get("amount"),
            field="order lifecycle amount",
            allow_none=True,
        )
        if not currency or amount is None:
            amount_unknown = True
            continue
        family = "payment" if side in {"payment", "payment_refund"} else "payout"
        sign = -1 if side in {"payment_refund", "recovery"} else 1
        key = (family, currency)
        signed_totals[key] = signed_totals.get(key, amount * 0) + amount * sign
    funds_fully_reversed = (
        not amount_unknown
        and bool(signed_totals)
        and all(amount == 0 for amount in signed_totals.values())
    )
    if has_cancel_marker and (all_failed_or_blank or funds_fully_reversed):
        status = "cancelled"
        reason = "聊天明确取消或失败，且没有待履行资金"
        evidence = source_labels
    elif {"payment", "payout"} <= completed_sides:
        status = "completed"
        reason = "付款与内部回款均已有完成依据"
        evidence = completed_labels or source_labels
    else:
        status = "pending_next_day"
        reason = "截至本期末尚未形成完整付款与内部回款闭环"
        evidence = source_labels
    event = _lifecycle_evidence_event(
        status,
        source_labels=evidence,
        message_by_label=message_by_label,
        reason=reason,
    )
    return {
        "status": status,
        "completion_at": event["at"] if status == "completed" else "",
        "reason": reason,
        "source_messages": event["source_messages"],
        "history": [event],
    }


def _normalize_order_lifecycle(
    value: object,
    *,
    field: str,
    source_labels: list[str],
    message_by_label: Mapping[str, Mapping[str, Any]],
    entry_records: Mapping[str, Mapping[str, Any]],
    entry_source_labels: Mapping[str, str],
    entry_ids: list[str],
    allow_derived_transition: bool = False,
) -> dict[str, Any]:
    derived = _derived_order_lifecycle(
        source_labels=source_labels,
        message_by_label=message_by_label,
        entry_records=entry_records,
        entry_source_labels=entry_source_labels,
        entry_ids=entry_ids,
    )
    if value in (None, ""):
        return derived
    core.require(isinstance(value, Mapping), f"{field} must be an object")
    unknown = sorted(
        set(value) - {"status", "completion_at", "reason", "source_messages", "history"}
    )
    core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
    status = core.clean_text(value.get("status")).casefold()
    core.require(status in ORDER_LIFECYCLE_STATUSES, f"{field}.status is unsupported")
    raw_sources = value.get("source_messages", [])
    core.require(isinstance(raw_sources, list), f"{field}.source_messages must be a list")
    lifecycle_sources = [str(item) for item in raw_sources]
    core.require(
        bool(lifecycle_sources)
        and len(lifecycle_sources) == len(set(lifecycle_sources))
        and set(lifecycle_sources) <= set(source_labels),
        f"{field}.source_messages must cite unique messages belonging to the order",
    )
    reason = core.clean_text(value.get("reason"))
    if status != "completed":
        core.require(bool(reason), f"{field}.reason is required when status is {status}")
    completion_at = core.clean_text(value.get("completion_at"))
    cited_timestamps = {
        str(message_by_label[label].get("timestamp") or "")
        for label in lifecycle_sources
    }
    if status == "completed":
        core.require(
            completion_at in cited_timestamps,
            f"{field}.completion_at must equal a cited source-message timestamp",
        )
    else:
        core.require(not completion_at, f"{field}.completion_at is only for completed orders")

    history_value = value.get("history", [])
    core.require(isinstance(history_value, list) and history_value, f"{field}.history is required")
    history: list[dict[str, Any]] = []
    prior_at = ""
    for position, raw_event in enumerate(history_value):
        event_field = f"{field}.history[{position}]"
        core.require(isinstance(raw_event, Mapping), f"{event_field} must be an object")
        unknown_event = sorted(
            set(raw_event) - {"status", "at", "reason", "source_messages"}
        )
        core.require(
            not unknown_event,
            f"{event_field} has unsupported fields: {', '.join(unknown_event)}",
        )
        event_status = core.clean_text(raw_event.get("status")).casefold()
        core.require(
            event_status in ORDER_LIFECYCLE_STATUSES,
            f"{event_field}.status is unsupported",
        )
        event_sources_value = raw_event.get("source_messages", [])
        core.require(
            isinstance(event_sources_value, list) and event_sources_value,
            f"{event_field}.source_messages is required",
        )
        event_sources = [str(item) for item in event_sources_value]
        core.require(
            len(event_sources) == len(set(event_sources))
            and set(event_sources) <= set(source_labels),
            f"{event_field}.source_messages must belong to the order",
        )
        event_at = core.clean_text(raw_event.get("at"))
        core.require(
            event_at
            in {
                str(message_by_label[label].get("timestamp") or "")
                for label in event_sources
            },
            f"{event_field}.at must equal a cited source-message timestamp",
        )
        core.require(event_at >= prior_at, f"{field}.history must be chronological")
        prior_at = event_at
        event_reason = core.clean_text(raw_event.get("reason"))
        if event_status != "completed":
            core.require(bool(event_reason), f"{event_field}.reason is required")
        history.append(
            {
                "status": event_status,
                "at": event_at,
                "reason": event_reason,
                "source_messages": event_sources,
            }
        )

    # A carried pending order normally omits lifecycle in the next review batch.
    # The merge retains its history and this transition closes it from new facts.
    if (
        allow_derived_transition
        and status == "pending_next_day"
        and derived["status"] in {"completed", "cancelled"}
    ):
        transition = derived["history"][-1]
        if transition["at"] > history[-1]["at"]:
            history.append(transition)
            status = derived["status"]
            completion_at = derived["completion_at"]
            reason = derived["reason"]
            lifecycle_sources = derived["source_messages"]

    core.require(history[-1]["status"] == status, f"{field}.history final status mismatch")
    if status == "completed":
        core.require(
            completion_at == history[-1]["at"],
            f"{field}.completion_at must equal the final completed history time",
        )
        core.require(
            derived["status"] == "completed",
            f"{field}: completed requires completed payment and payout evidence",
        )
    if status == "cancelled":
        core.require(
            derived["status"] == "cancelled",
            f"{field}: cancelled requires explicit cancellation/failure evidence and no live obligation",
        )
    return {
        "status": status,
        "completion_at": completion_at,
        "reason": reason,
        "source_messages": lifecycle_sources,
        "history": history,
    }


def _validate_decision(
    normalized: Mapping[str, Any],
    group: Mapping[str, Any],
    decision: dict[str, Any],
    *,
    require_complete: bool,
    capture_hashes: bool,
    rehash_labels: set[str] | None = None,
    auto_transition_order_ids: set[str] | None = None,
) -> dict[str, Any]:
    if decision.get("contract_version") == finance_materials.DECISION_CONTRACT:
        expected = _decision_template(
            normalized,
            group,
            group_mode=finance_materials.GROUP_MODE,
        )
        statistics = finance_materials.validate_decision(
            normalized,
            group,
            decision,
            expected,
            require_complete=require_complete,
            capture_hashes=capture_hashes,
            rehash_labels=rehash_labels,
        )
        return {
            **statistics,
            **_validate_media_observation_cache(
                group,
                decision,
                require_complete=require_complete,
                verify_files=rehash_labels is None,
            ),
        }
    if decision.get("contract_version") == store_ledger.DECISION_CONTRACT:
        expected = _decision_template(
            normalized,
            group,
            group_mode=store_ledger.GROUP_MODE,
        )
        statistics = store_ledger.validate_decision(
            normalized,
            group,
            decision,
            expected,
            require_complete=require_complete,
            capture_hashes=capture_hashes,
            rehash_labels=rehash_labels,
        )
        return {
            **statistics,
            **_validate_media_observation_cache(
                group,
                decision,
                require_complete=require_complete,
                verify_files=rehash_labels is None,
            ),
        }
    if decision.get("contract_version") == large_daily.LEGACY_DECISION_CONTRACT:
        statistics = _validate_large_decision(
            normalized,
            group,
            decision,
            require_complete=require_complete,
            capture_hashes=capture_hashes,
            rehash_labels=rehash_labels,
        )
        return {
            **statistics,
            **_validate_media_observation_cache(
                group,
                decision,
                require_complete=require_complete,
                verify_files=rehash_labels is None,
            ),
        }
    decision_contract = core.clean_text(decision.get("contract_version"))
    large_order_mode = decision_contract == large_daily.DECISION_CONTRACT
    expected_contract = large_daily.DECISION_CONTRACT if large_order_mode else DECISION_CONTRACT
    expected = _decision_template(
        normalized,
        group,
        group_mode="large" if large_order_mode else "small",
    )
    core.require(
        decision_contract == expected_contract,
        "unsupported decision contract; expected "
        f"{DECISION_CONTRACT} or {large_daily.DECISION_CONTRACT}",
    )
    generated_fields = [
        "normalized_source_fingerprint",
        "group_fingerprint",
        "group_key",
        "group_name",
        "platform",
        "message_count",
        "evidence_media_count",
    ]
    if large_order_mode:
        generated_fields.append("group_mode")
    for field in generated_fields:
        core.require(decision.get(field) == expected.get(field), f"{group.get('group_key')}: generated {field} was edited")
    messages = list(group.get("messages", []))
    reviewed_through = decision.get("reviewed_through")
    core.require(isinstance(reviewed_through, int) and 0 <= reviewed_through <= len(messages), "invalid reviewed_through")
    core.require(isinstance(decision.get("read_complete"), bool), "read_complete must be boolean")
    core.require(
        decision.get("read_complete") is (reviewed_through == len(messages)),
        "read_complete must exactly match reviewed_through",
    )
    if require_complete:
        core.require(decision.get("read_complete") is True and reviewed_through == len(messages), "group chronology has not been read completely")
    open_order_count = _validate_open_orders(
        group,
        decision,
        reviewed_through=reviewed_through,
        require_complete=require_complete,
        require_fund_type=large_order_mode,
    )

    inventory = _media_inventory(group)
    label_by_message_id, message_by_label = _message_labels(group)
    message_labels = set(message_by_label)
    available = {label for label, (_, media) in inventory.items() if media.get("availability") == "available"}
    missing = set(inventory) - available
    media_decisions = decision.get("media_decisions")
    core.require(isinstance(media_decisions, dict), "media_decisions must be an object keyed by M labels")
    unknown_labels = sorted(set(media_decisions) - available)
    core.require(not unknown_labels, f"media decisions contain missing or unknown labels: {unknown_labels[:10]}")
    if require_complete:
        unclassified = sorted(available - set(media_decisions))
        core.require(not unclassified, f"unclassified available media: {unclassified[:20]}")

    entry_ids: set[str] = set()
    entry_records: dict[str, dict[str, Any]] = {}
    entry_source_labels: dict[str, str] = {}
    transfer_payees = 0
    unknown_payee_entry_ids: set[str] = set()
    unreadable_payee_entry_ids: set[str] = set()
    reference_count = 0
    fund_media_count = 0
    evidence_hashes_computed = 0
    evidence_hashes_reused = 0
    for label, raw in media_decisions.items():
        field = f"{group.get('group_key')}.media_decisions.{label}"
        core.require(isinstance(raw, dict), f"{field} must be an object")
        classification = core.clean_text(raw.get("classification")).casefold()
        core.require(classification in {"reference", "fund"}, f"{field}.classification must be reference or fund")
        if classification == "reference":
            unknown = sorted(set(raw) - {"classification", "note"})
            core.require(not unknown, f"{field}: reference has unsupported fields: {', '.join(unknown)}")
            reference_count += 1
            continue
        unknown = sorted(set(raw) - {"classification", "viewed_original", "evidence_sha256", "entries", "note"})
        core.require(not unknown, f"{field}: fund decision has unsupported fields: {', '.join(unknown)}")
        core.require(raw.get("viewed_original") is True, f"{field}: open the original image before recording fund facts")
        entries = raw.get("entries")
        core.require(isinstance(entries, list) and entries, f"{field}.entries must contain at least one fund entry")
        message, media = inventory[label]
        media_path = Path(str(media.get("path") or ""))
        core.require(media_path.is_file(), f"{field}: original media file is unavailable: {media_path}")
        recorded_hash = core.clean_text(raw.get("evidence_sha256"))
        verify_hash = (
            rehash_labels is None
            or label in rehash_labels
            or (capture_hashes and not recorded_hash)
        )
        resolved_hash, computed = core.validate_cached_file_hash(
            media_path,
            recorded_hash,
            verify=verify_hash,
            changed_message=f"{field}: original fund evidence changed after review",
        )
        snapshot_hash = core.clean_text(media.get("blob_sha256"))
        if computed and snapshot_hash:
            core.require(
                resolved_hash == snapshot_hash,
                f"{field}: original fund evidence changed after the run snapshot was created",
            )
        if computed:
            evidence_hashes_computed += 1
        elif recorded_hash:
            evidence_hashes_reused += 1
        if not recorded_hash and capture_hashes:
            raw["evidence_sha256"] = resolved_hash
        elif require_complete:
            core.require(
                bool(recorded_hash),
                f"{field}: evidence hash has not been captured; run review seal",
            )
        for position, entry in enumerate(entries, start=1):
            core.require(isinstance(entry, dict), f"{field}.entries[{position - 1}] must be an object")
            source_message_label = label_by_message_id[
                str(message.get("message_id") or "")
            ]
            _validate_entry(
                entry,
                field=f"{field}.entries[{position - 1}]",
                sender_role=core.clean_text(message.get("role")),
                source_message_label=source_message_label,
                message_labels=message_labels,
            )
            entry_id = f"{label}.{position}"
            core.require(entry_id not in entry_ids, f"duplicate entry id: {entry_id}")
            entry_ids.add(entry_id)
            entry_records[entry_id] = entry
            entry_source_labels[entry_id] = source_message_label
            if entry["kind"] == "transfer":
                transfer_payees += 1
                if entry["payee"] in core.UNKNOWN_PAYEES:
                    unknown_payee_entry_ids.add(entry_id)
                if entry.get("payee_state") == "unreadable":
                    unreadable_payee_entry_ids.add(entry_id)
        fund_media_count += 1

    reviewed_unknown_value = decision.get("unknown_payee_reviewed_entry_ids", [])
    core.require(
        isinstance(reviewed_unknown_value, list),
        "unknown_payee_reviewed_entry_ids must be a list",
    )
    reviewed_unknown = [str(item) for item in reviewed_unknown_value]
    core.require(
        len(reviewed_unknown) == len(set(reviewed_unknown)),
        "unknown_payee_reviewed_entry_ids repeats an entry",
    )
    core.require(
        set(reviewed_unknown) <= unknown_payee_entry_ids,
        "unknown_payee_reviewed_entry_ids may only cite transfer entries marked 未显示 or 无法辨认",
    )
    unknown_payee_warning = (
        transfer_payees >= UNKNOWN_PAYEE_WARNING_MIN_TRANSFERS
        and len(unknown_payee_entry_ids) * 2 > transfer_payees
    )
    unknown_payee_review_required = bool(unreadable_payee_entry_ids)
    unreviewed_unknown_payees = unreadable_payee_entry_ids - set(reviewed_unknown)
    if require_complete and unreviewed_unknown_payees:
        core.require(
            not unreviewed_unknown_payees,
            "unreadable payee review is required: reopen only the cited unclear originals and list "
            "the still-unreadable entry IDs in unknown_payee_reviewed_entry_ids; remaining: "
            f"{sorted(unreviewed_unknown_payees)[:40]}",
        )

    del label_by_message_id
    orders = decision.get("orders")
    core.require(isinstance(orders, list), "orders must be a list")
    order_ids: set[str] = set()
    order_by_id: dict[str, dict[str, Any]] = {}
    assigned_entries: set[str] = set()
    pricing_scopes = 0
    missing_pricing_scopes: list[str] = []
    single_entry_orders = 0
    blank_direction_orders = 0
    unknown_pricing_orders = 0
    own_fund_source_only_orders = 0
    mass_degenerate_orders = 0
    single_unknown_pricing_orders = 0
    order_degradation_signals: list[tuple[bool, bool, bool, bool, str]] = []
    order_continuity_facts: list[dict[str, Any]] = []
    order_note_counts: dict[str, int] = {}
    order_fields = {
        "id",
        "entry_ids",
        "source_messages",
        "customer_nickname",
        "direction",
        "pricing",
        "note",
        "same_transactions",
        "legs",
        "lifecycle",
    }
    if large_order_mode:
        order_fields.add("fund_type")
    for position, order in enumerate(orders):
        field = f"{group.get('group_key')}.orders[{position}]"
        core.require(isinstance(order, dict), f"{field} must be an object")
        unknown = sorted(set(order) - order_fields)
        core.require(not unknown, f"{field}: unsupported fields: {', '.join(unknown)}")
        if large_order_mode:
            order["fund_type"] = large_daily.normalize_fund_type(
                order.get("fund_type"), field=f"{field}.fund_type"
            )
        for required_field in ("customer_nickname", "direction"):
            core.require(
                required_field in order,
                f"{field}.{required_field} must be explicitly supplied by the model",
            )
        order_note = core.clean_text(order.get("note"))
        if order_note:
            order["note"] = order_note
            order_note_counts[order_note] = order_note_counts.get(order_note, 0) + 1
        order_id = core.clean_text(order.get("id"))
        core.require(bool(order_id), f"{field}.id is required")
        core.require(order_id not in order_ids, f"{field}.id is duplicated")
        order_ids.add(order_id)
        order_by_id[order_id] = order
        raw_entries = order.get("entry_ids")
        core.require(isinstance(raw_entries, list) and raw_entries, f"{field}.entry_ids is required")
        refs = [str(item) for item in raw_entries]
        is_single_entry_order = len(refs) == 1
        if is_single_entry_order:
            single_entry_orders += 1
        core.require(len(refs) == len(set(refs)), f"{field}.entry_ids repeats an entry")
        for entry_id in refs:
            core.require(entry_id in entry_ids, f"{field}: unknown fund entry {entry_id}")
            core.require(entry_id not in assigned_entries, f"fund entry assigned to multiple orders: {entry_id}")
            assigned_entries.add(entry_id)
        raw_messages = order.get("source_messages", [])
        core.require(isinstance(raw_messages, list), f"{field}.source_messages must be a list")
        source_labels = [str(item) for item in raw_messages]
        core.require(len(source_labels) == len(set(source_labels)), f"{field}.source_messages repeats a label")
        derived_source_labels = [entry_source_labels[entry_id] for entry_id in refs]
        for entry_id in refs:
            side_exception = entry_records[entry_id].get("side_exception")
            if isinstance(side_exception, Mapping):
                derived_source_labels.extend(
                    str(item) for item in side_exception.get("source_messages", [])
                )
        pricing_candidates: list[object] = [order.get("pricing")]
        if isinstance(order.get("legs"), list):
            pricing_candidates.extend(
                leg.get("pricing")
                for leg in order["legs"]
                if isinstance(leg, Mapping)
            )
        for pricing_candidate in pricing_candidates:
            if isinstance(pricing_candidate, Mapping):
                derived_source_labels.extend(
                    str(item)
                    for item in pricing_candidate.get("source_messages", [])
                )
        for label in derived_source_labels:
            if label not in source_labels:
                source_labels.append(label)
        order["source_messages"] = source_labels
        core.require(set(source_labels) <= set(message_by_label), f"{field}.source_messages contains an unknown label")
        source_label_set = set(source_labels)
        own_fund_source_only = source_label_set == {
            entry_source_labels[entry_id] for entry_id in refs
        }
        if own_fund_source_only:
            own_fund_source_only_orders += 1
        for entry_id in refs:
            side_exception = entry_records[entry_id].get("side_exception")
            if isinstance(side_exception, Mapping):
                exception_sources = {
                    str(item) for item in side_exception.get("source_messages", [])
                }
                core.require(
                    exception_sources <= source_label_set,
                    f"{field}: side_exception.source_messages must belong to the order",
                )
        order_payout_currency: str | None = None
        blank_direction = not core.clean_text(order.get("direction"))
        if blank_direction:
            blank_direction_orders += 1
        else:
            order["direction"] = core.canonical_direction(order.get("direction"), field=f"{field}.direction")
            _, order_payout_currency = core.direction_currencies(
                order["direction"],
                field=f"{field}.direction",
            )
        same_transactions = order.get("same_transactions")
        if same_transactions not in (None, "", []):
            core.require(isinstance(same_transactions, list), f"{field}.same_transactions must be a list")
            duplicate_entries: set[str] = set()
            canonical_entries: set[str] = set()
            for pair_position, pair in enumerate(same_transactions):
                pair_field = f"{field}.same_transactions[{pair_position}]"
                core.require(isinstance(pair, Mapping), f"{pair_field} must be an object")
                unknown_pair_fields = sorted(set(pair) - {"entry_id", "same_as"})
                core.require(not unknown_pair_fields, f"{pair_field}: unsupported fields: {', '.join(unknown_pair_fields)}")
                entry_id = core.clean_text(pair.get("entry_id"))
                same_as = core.clean_text(pair.get("same_as"))
                core.require(entry_id in refs and same_as in refs, f"{pair_field} must cite two entries in this order")
                core.require(entry_id != same_as, f"{pair_field} cannot cite itself")
                core.require(entry_id not in duplicate_entries, f"{pair_field} repeats a duplicate entry")
                duplicate_entries.add(entry_id)
                canonical_entries.add(same_as)
            core.require(
                not (duplicate_entries & canonical_entries),
                f"{field}.same_transactions cannot contain chains",
            )

        legs = order.get("legs")
        if legs not in (None, "", []):
            core.require(isinstance(legs, list) and len(legs) >= 2, f"{field}.legs requires at least two legs")
            core.require(
                order.get("pricing") in (None, ""),
                f"{field}.pricing must be empty when pricing is defined per leg",
            )
            referenced_settlements: set[str] = set()
            display_labels_by_direction: dict[str, list[str]] = {}
            for leg_position, leg in enumerate(legs):
                leg_field = f"{field}.legs[{leg_position}]"
                core.require(isinstance(leg, dict), f"{leg_field} must be an object")
                allowed_leg_fields = {
                    "leg_id",
                    "display_label",
                    "direction",
                    "allocation_amount",
                    "pricing",
                    "payout_entry_ids",
                    "recovery_entry_ids",
                }
                unknown_leg_fields = sorted(set(leg) - allowed_leg_fields)
                core.require(not unknown_leg_fields, f"{leg_field}: unsupported fields: {', '.join(unknown_leg_fields)}")
                leg_direction = core.canonical_direction(leg.get("direction"), field=f"{leg_field}.direction")
                _, leg_payout_currency = core.direction_currencies(
                    leg_direction,
                    field=f"{leg_field}.direction",
                )
                if isinstance(leg, dict):
                    leg["direction"] = leg_direction
                display_label = core.clean_text(leg.get("display_label"))
                if "display_label" in leg:
                    core.require(bool(display_label), f"{leg_field}.display_label cannot be blank")
                    core.require(
                        len(display_label) <= 80 and "\n" not in display_label and "\r" not in display_label,
                        f"{leg_field}.display_label must be a single line of at most 80 characters",
                    )
                    leg["display_label"] = display_label
                display_labels_by_direction.setdefault(leg_direction, []).append(display_label)
                allocation = core.parse_decimal(leg.get("allocation_amount"), field=f"{leg_field}.allocation_amount")
                core.require(allocation is not None and allocation > 0, f"{leg_field}.allocation_amount must be positive")
                if isinstance(leg, dict):
                    leg["allocation_amount"] = core.decimal_text(allocation)
                for reference_field in ("payout_entry_ids", "recovery_entry_ids"):
                    raw_references = leg.get(reference_field, [])
                    core.require(isinstance(raw_references, list), f"{leg_field}.{reference_field} must be a list")
                    entry_references = [str(item) for item in raw_references]
                    core.require(
                        len(entry_references) == len(set(entry_references)),
                        f"{leg_field}.{reference_field} repeats an entry",
                    )
                    core.require(
                        set(entry_references) <= set(refs),
                        f"{leg_field}.{reference_field} cites an entry outside this order",
                    )
                    core.require(
                        not (referenced_settlements & set(entry_references)),
                        f"{leg_field}.{reference_field} repeats an entry used by another leg",
                    )
                    referenced_settlements.update(entry_references)
                pricing_scopes += 1
                if _validate_pricing_scope(
                    leg.get("pricing"),
                    field=f"{leg_field}.pricing",
                    source_labels=set(source_labels),
                    payout_currency=leg_payout_currency,
                    require_complete=require_complete,
                ):
                    missing_pricing_scopes.append(f"{leg_field}.pricing")
            for repeated_direction, display_labels in display_labels_by_direction.items():
                if len(display_labels) < 2:
                    continue
                core.require(
                    all(display_labels),
                    f"{field}.legs sharing direction {repeated_direction} require display_label on every leg",
                )
                core.require(
                    len({label.casefold() for label in display_labels}) == len(display_labels),
                    f"{field}.legs sharing direction {repeated_direction} require distinct display_label values",
                )
        else:
            pricing_scopes += 1
            if _validate_pricing_scope(
                order.get("pricing"),
                field=f"{field}.pricing",
                source_labels=set(source_labels),
                payout_currency=order_payout_currency,
                require_complete=require_complete,
            ):
                missing_pricing_scopes.append(f"{field}.pricing")

        if isinstance(legs, list) and legs:
            unknown_pricing = all(
                isinstance(leg.get("pricing"), Mapping)
                and isinstance(leg["pricing"].get("expected"), Mapping)
                and core.clean_text(leg["pricing"]["expected"].get("kind")).casefold()
                == "unknown"
                for leg in legs
            )
        else:
            pricing = order.get("pricing")
            unknown_pricing = (
                isinstance(pricing, Mapping)
                and isinstance(pricing.get("expected"), Mapping)
                and core.clean_text(pricing["expected"].get("kind")).casefold()
                == "unknown"
            )
        if unknown_pricing:
            unknown_pricing_orders += 1
        if is_single_entry_order and unknown_pricing:
            single_unknown_pricing_orders += 1
        meaningful_chat_context = any(
            bool(core.clean_text(message_by_label[label].get("text")))
            for label in source_label_set
        )
        if (
            own_fund_source_only
            and not meaningful_chat_context
            and (blank_direction or unknown_pricing)
        ):
            core.require(
                bool(order_note)
                and any(marker in order_note for marker in MISSING_CONTEXT_NOTE_MARKERS),
                f"{field}: isolated fund evidence requires a specific missing-context note",
            )
        completed_ordinary_sides = {
            core.clean_text(entry_records[entry_id].get("side")).casefold()
            for entry_id in refs
            if entry_records[entry_id].get("result") == "completed"
        }
        order["lifecycle"] = _normalize_order_lifecycle(
            order.get("lifecycle"),
            field=f"{field}.lifecycle",
            source_labels=source_labels,
            message_by_label=message_by_label,
            entry_records=entry_records,
            entry_source_labels=entry_source_labels,
            entry_ids=refs,
            allow_derived_transition=order_id
            in (auto_transition_order_ids or set()),
        )
        if (
            not blank_direction
            and not unknown_pricing
            and {"payment", "payout"} <= completed_ordinary_sides
        ):
            core.require(
                meaningful_chat_context,
                f"{field}: completed ordinary order requires meaningful chat context, not fund images alone",
            )
        source_positions = [
            int(match.group(1))
            for label in source_labels
            if (match := re.fullmatch(r"S(\d+)", label)) is not None
        ]
        pricing_value = order.get("pricing")
        has_confirmed_pricing = (
            isinstance(pricing_value, Mapping)
            and isinstance(pricing_value.get("terms"), Mapping)
            and isinstance(pricing_value.get("expected"), Mapping)
            and core.clean_text(pricing_value["expected"].get("kind")).casefold()
            in {"explicit", "calculated_from_terms"}
        )
        entry_sides = {
            core.clean_text(entry_records[entry_id].get("side")).casefold()
            for entry_id in refs
        }
        if source_positions:
            order_continuity_facts.append(
                {
                    "order_id": order_id,
                    "order_position": position,
                    "customer_nickname": core.clean_text(order.get("customer_nickname")),
                    "customer_key": core.clean_text(order.get("customer_nickname")).casefold(),
                    "direction": core.clean_text(order.get("direction")),
                    "source_start_position": min(source_positions),
                    "source_end_position": max(source_positions),
                    "has_completed_payment": "payment" in completed_ordinary_sides,
                    "has_completed_payout": "payout" in completed_ordinary_sides,
                    "has_confirmed_pricing": has_confirmed_pricing,
                    "has_unknown_pricing": unknown_pricing,
                    "has_complex_flow": bool(legs)
                    or bool(same_transactions)
                    or bool(entry_sides & {"payment_refund", "recovery"}),
                }
            )
        if (
            is_single_entry_order
            and blank_direction
            and unknown_pricing
            and own_fund_source_only
        ):
            mass_degenerate_orders += 1
        order_degradation_signals.append(
            (
                is_single_entry_order,
                blank_direction,
                unknown_pricing,
                own_fund_source_only,
                order_note,
            )
        )

    order_count = len(orders)
    repeated_note_values = {
        note for note, count in order_note_counts.items() if count > 1
    }
    repeated_note_orders = sum(
        bool(note) and note in repeated_note_values
        for *_, note in order_degradation_signals
    )
    multi_signal_degenerate_orders = sum(
        sum((single, blank, unknown_pricing, own_source, note in repeated_note_values))
        >= 3
        for single, blank, unknown_pricing, own_source, note in order_degradation_signals
    )
    mass_degenerate_order_ratio = (
        mass_degenerate_orders / order_count if order_count else 0.0
    )
    multi_signal_degenerate_order_ratio = (
        multi_signal_degenerate_orders / order_count if order_count else 0.0
    )
    suspected_bulk_order_creation = int(
        order_count >= 10
        and (
            mass_degenerate_order_ratio >= 0.8
            or multi_signal_degenerate_order_ratio >= 0.7
        )
    )

    advanced_relation_order_ids: set[str] = set()
    settlement_allocations = decision.get("settlement_allocations", [])
    core.require(
        isinstance(settlement_allocations, list),
        "settlement_allocations must be a list",
    )
    allocated_source_entries: set[str] = set()
    for position, relation in enumerate(settlement_allocations):
        field = f"{group.get('group_key')}.settlement_allocations[{position}]"
        core.require(isinstance(relation, dict), f"{field} must be an object")
        unknown_relation_fields = sorted(
            set(relation) - {"entry_id", "allocations", "source_messages"}
        )
        core.require(
            not unknown_relation_fields,
            f"{field}: unsupported fields: {', '.join(unknown_relation_fields)}",
        )
        source_entry_id = core.clean_text(relation.get("entry_id"))
        core.require(source_entry_id in entry_records, f"{field}.entry_id is unknown")
        core.require(
            source_entry_id not in assigned_entries,
            f"{field}.entry_id must not also appear in an order.entry_ids",
        )
        core.require(
            source_entry_id not in allocated_source_entries,
            f"{field}.entry_id is already used by another settlement allocation",
        )
        source_entry = entry_records[source_entry_id]
        core.require(
            source_entry.get("side") in {"payout", "recovery"},
            f"{field}.entry_id must identify a payout or recovery entry",
        )
        core.require(
            source_entry.get("result") == "completed"
            and source_entry.get("amount_state") == "clear",
            f"{field}.entry_id must be a completed entry with a clear amount",
        )
        source_amount = core.parse_decimal(
            source_entry.get("amount"),
            field=f"{field}.source_amount",
        )
        core.require(
            source_amount is not None and source_amount > 0,
            f"{field}.source_amount must be positive",
        )
        source_currency = core.normalize_currency(
            source_entry.get("currency"),
            field=f"{field}.source_currency",
        )
        allocations = relation.get("allocations")
        core.require(
            isinstance(allocations, list) and len(allocations) >= 2,
            f"{field}.allocations must contain at least two target orders",
        )
        target_order_ids: set[str] = set()
        target_customer_nicknames: set[str] = set()
        allocated_amounts = []
        for allocation_position, allocation in enumerate(allocations):
            allocation_field = f"{field}.allocations[{allocation_position}]"
            core.require(
                isinstance(allocation, dict),
                f"{allocation_field} must be an object",
            )
            unknown_allocation_fields = sorted(set(allocation) - {"order_id", "amount"})
            core.require(
                not unknown_allocation_fields,
                f"{allocation_field}: unsupported fields: {', '.join(unknown_allocation_fields)}",
            )
            target_order_id = core.clean_text(allocation.get("order_id"))
            core.require(
                target_order_id in order_by_id,
                f"{allocation_field}.order_id is unknown",
            )
            core.require(
                target_order_id not in target_order_ids,
                f"{field}.allocations repeats target order {target_order_id}",
            )
            target_order_ids.add(target_order_id)
            target_order = order_by_id[target_order_id]
            core.require(
                target_order.get("legs") in (None, "", []),
                f"{allocation_field}: settlement allocation does not support multi-leg target orders",
            )
            target_customer_nickname = core.clean_text(target_order.get("customer_nickname"))
            core.require(
                bool(target_customer_nickname),
                f"{allocation_field}: target customer_nickname must be known",
            )
            target_customer_nicknames.add(target_customer_nickname.casefold())
            _, target_payout_currency = core.direction_currencies(
                target_order.get("direction"),
                field=f"{allocation_field}.target_direction",
            )
            core.require(
                target_payout_currency == source_currency,
                f"{allocation_field}: target payout currency must be {source_currency}",
            )
            target_entries = [
                entry_records[str(item)]
                for item in target_order.get("entry_ids", [])
                if str(item) in entry_records
            ]
            core.require(
                any(item.get("side") in {"payment", "payment_refund"} for item in target_entries),
                f"{allocation_field}: target order must contain its own payment evidence",
            )
            amount = core.parse_decimal(
                allocation.get("amount"),
                field=f"{allocation_field}.amount",
            )
            core.require(
                amount is not None and amount > 0,
                f"{allocation_field}.amount must be positive",
            )
            allocation["order_id"] = target_order_id
            allocation["amount"] = core.decimal_text(amount)
            allocated_amounts.append(amount)
        core.require(
            len(target_customer_nicknames) == 1,
            f"{field}.allocations must target orders for the same customer",
        )
        advanced_relation_order_ids.update(target_order_ids)
        allocated_total = sum(allocated_amounts, start=source_amount * 0)
        core.require(
            allocated_total == source_amount,
            f"{field}.allocations total {core.decimal_text(allocated_total)} "
            f"must equal source amount {core.decimal_text(source_amount)}",
        )
        source_messages = relation.get("source_messages", [])
        core.require(isinstance(source_messages, list), f"{field}.source_messages must be a list")
        source_labels = [str(item) for item in source_messages]
        core.require(
            len(source_labels) == len(set(source_labels)),
            f"{field}.source_messages repeats a label",
        )
        core.require(
            set(source_labels) <= set(message_by_label),
            f"{field}.source_messages contains an unknown label",
        )
        relation["entry_id"] = source_entry_id
        allocated_source_entries.add(source_entry_id)
        assigned_entries.add(source_entry_id)

    if require_complete:
        unassigned_entries = sorted(entry_ids - assigned_entries)
        core.require(
            not unassigned_entries,
            f"fund entries must be explicitly assigned to an order: {unassigned_entries}",
        )

    links = decision.get("balance_links")
    core.require(isinstance(links, list), "balance_links must be a list")
    for position, link in enumerate(links):
        field = f"{group.get('group_key')}.balance_links[{position}]"
        core.require(isinstance(link, dict), f"{field} must be an object")
        allowed_link_fields = {
            "source_order_id",
            "target_order_id",
            "kind",
            "amount",
            "currency",
            "source_messages",
            "already_in_expected",
        }
        unknown_link_fields = sorted(set(link) - allowed_link_fields)
        core.require(not unknown_link_fields, f"{field}: unsupported fields: {', '.join(unknown_link_fields)}")
        source_order_id = core.clean_text(link.get("source_order_id"))
        target_order_id = core.clean_text(link.get("target_order_id"))
        core.require(source_order_id in order_ids, f"{field}.source_order_id is unknown")
        core.require(target_order_id in order_ids, f"{field}.target_order_id is unknown")
        core.require(source_order_id != target_order_id, f"{field} cannot link an order to itself")
        advanced_relation_order_ids.update({source_order_id, target_order_id})
        kind = core.clean_text(link.get("kind")).casefold()
        core.require(kind in {"shortfall_carryover", "overpayment_carryover"}, f"{field}.kind is unsupported")
        link["kind"] = kind
        amount = core.parse_decimal(link.get("amount"), field=f"{field}.amount")
        core.require(amount is not None and amount > 0, f"{field}.amount must be positive")
        link["amount"] = core.decimal_text(amount)
        link["currency"] = core.normalize_currency(link.get("currency"), field=f"{field}.currency")
        already_in_expected = link.get("already_in_expected", False)
        core.require(isinstance(already_in_expected, bool), f"{field}.already_in_expected must be boolean")
        if "already_in_expected" in link:
            link["already_in_expected"] = already_in_expected
        source_messages = link.get("source_messages", [])
        core.require(isinstance(source_messages, list), f"{field}.source_messages must be a list")
        source_labels = [str(item) for item in source_messages]
        core.require(len(source_labels) == len(set(source_labels)), f"{field}.source_messages repeats a label")
        core.require(set(source_labels) <= set(message_by_label), f"{field}.source_messages contains an unknown label")

    order_continuity_review_candidates = _order_continuity_review_candidates(
        order_continuity_facts,
        advanced_relation_order_ids=advanced_relation_order_ids,
    )

    return {
        "available_media": len(available),
        "missing_media": len(missing),
        "classified_media": len(media_decisions),
        "reference_media": reference_count,
        "fund_media": fund_media_count,
        "fund_entries": len(entry_ids),
        "orders": len(orders),
        "open_orders": open_order_count,
        "single_entry_orders": single_entry_orders,
        "single_entry_order_ratio": single_entry_orders / order_count if order_count else 0.0,
        "blank_direction_orders": blank_direction_orders,
        "blank_direction_order_ratio": blank_direction_orders / order_count if order_count else 0.0,
        "unknown_pricing_orders": unknown_pricing_orders,
        "unknown_pricing_order_ratio": unknown_pricing_orders / order_count if order_count else 0.0,
        "own_fund_source_only_orders": own_fund_source_only_orders,
        "own_fund_source_only_order_ratio": own_fund_source_only_orders / order_count if order_count else 0.0,
        "repeated_order_notes": sum(
            count - 1 for count in order_note_counts.values() if count > 1
        ),
        "repeated_note_orders": repeated_note_orders,
        "repeated_note_order_ratio": repeated_note_orders / order_count if order_count else 0.0,
        "role_side_conflicts": 0,
        "mass_degenerate_orders": mass_degenerate_orders,
        "mass_degenerate_order_ratio": mass_degenerate_order_ratio,
        "single_unknown_pricing_orders": single_unknown_pricing_orders,
        "single_unknown_pricing_order_ratio": (
            single_unknown_pricing_orders / order_count if order_count else 0.0
        ),
        "multi_signal_degenerate_orders": multi_signal_degenerate_orders,
        "multi_signal_degenerate_order_ratio": multi_signal_degenerate_order_ratio,
        "suspected_bulk_order_creation": suspected_bulk_order_creation,
        "order_continuity_review_candidate_count": len(
            order_continuity_review_candidates
        ),
        "order_continuity_review_candidates": order_continuity_review_candidates,
        "order_to_fund_entry_ratio": order_count / len(entry_ids) if entry_ids else 0.0,
        "pricing_scopes": pricing_scopes,
        "missing_pricing_scopes": len(missing_pricing_scopes),
        "unassigned_entries": len(entry_ids - assigned_entries),
        "settlement_allocations": len(settlement_allocations),
        "transfer_payees": transfer_payees,
        "unknown_payees": len(unknown_payee_entry_ids),
        "unknown_payee_warning": int(unknown_payee_warning),
        "unknown_payee_review_required": int(unknown_payee_review_required),
        "unreviewed_unknown_payees": len(unreviewed_unknown_payees),
        "evidence_hashes_computed": evidence_hashes_computed,
        "evidence_hashes_reused": evidence_hashes_reused,
        **_validate_media_observation_cache(
            group,
            decision,
            require_complete=require_complete,
            verify_files=rehash_labels is None,
        ),
    }


def _review_media_queue(
    group: Mapping[str, Any],
    decision: Mapping[str, Any],
    start: int,
    end: int,
    *,
    media_labels: Mapping[str, str] | None = None,
    inventory: Mapping[str, tuple[Mapping[str, Any], Mapping[str, Any]]] | None = None,
    media_hashes: Mapping[str, str] | None = None,
    content_hashes_computed: int = 0,
    content_hashes_reused: int = 0,
) -> dict[str, Any]:
    if inventory is None:
        inventory = _media_inventory(group)
    if media_labels is None:
        media_labels = {
            str(media.get("media_id") or ""): label
            for label, (_, media) in inventory.items()
        }
    finance_mode = decision.get("contract_version") == finance_materials.DECISION_CONTRACT
    store_mode = decision.get("contract_version") == store_ledger.DECISION_CONTRACT
    batch_policy = _adaptive_media_batch_policy(
        decision,
        finance_mode=finance_mode,
    )
    parallel_limit = int(batch_policy["recommended_parallel_limit"])
    classified = set(decision.get("media_decisions", {}))
    cache = decision.get("media_observation_cache", {})
    cache_entries = (
        cache.get("entries", {})
        if isinstance(cache, Mapping) and isinstance(cache.get("entries", {}), Mapping)
        else {}
    )
    batches: list[dict[str, Any]] = []
    non_parallel_items: list[dict[str, Any]] = []
    missing_labels: list[str] = []
    current_items: list[dict[str, Any]] = []
    representative_items: dict[str, dict[str, Any]] = {}
    duplicate_aliases: list[dict[str, Any]] = []
    cache_hits_by_hash: dict[str, dict[str, Any]] = {}

    def flush_current() -> None:
        nonlocal current_items
        if not current_items:
            return
        batches.append(
            {
                "batch_id": f"V{start + 1:05d}-{len(batches) + 1:03d}",
                "view_detail": "original",
                "parallel": len(current_items) > 1,
                "items": current_items,
            }
        )
        current_items = []

    messages = list(group.get("messages", []))
    already_classified = 0
    for message_position in range(start, end):
        message = messages[message_position]
        message_label = f"S{message_position + 1:05d}"
        for media in message.get("media", []):
            if not isinstance(media, Mapping) or not core.media_is_evidence(media):
                continue
            label = media_labels.get(str(media.get("media_id") or ""))
            if not label:
                continue
            if label in classified:
                already_classified += 1
                continue
            if media.get("availability") != "available":
                missing_labels.append(label)
                continue
            source_variant = core.clean_text(media.get("variant")) or core.clean_text(
                media.get("source_field")
            )
            quality_warning = (
                "thumbnail_only"
                if "thumbnail" in source_variant.casefold()
                else None
            )
            item = {
                "label": label,
                "message_label": message_label,
                "path": media.get("path"),
                "kind": media.get("kind"),
                "mime_type": media.get("mime_type"),
                "byte_size": media.get("byte_size"),
                "source_variant": source_variant or None,
                "sender": message.get("sender_name") or message.get("sender_id"),
                "role": message.get("role"),
            }
            if quality_warning:
                item["quality_warning"] = quality_warning
            mime_type = core.clean_text(media.get("mime_type")).casefold()
            kind = core.clean_text(media.get("kind")).casefold()
            path = core.clean_text(media.get("path"))
            parallel_viewable = bool(path) and (
                kind == "image" or mime_type.startswith("image/")
            )
            if not parallel_viewable:
                flush_current()
                non_parallel_items.append(
                    {
                        **item,
                        "reason": "not_a_directly_viewable_image" if path else "missing_path",
                    }
                )
                continue
            digest = core.clean_text((media_hashes or {}).get(label))
            if not digest:
                digest, computed = _media_content_hash_for_label(
                    group,
                    decision,
                    label,
                    inventory=inventory,
                    verify_file=False,
                )
                if computed:
                    content_hashes_computed += 1
                else:
                    content_hashes_reused += 1
            item["content_sha256"] = digest
            cached_entry = cache_entries.get(digest)
            cached_observation = (
                cached_entry.get("observation")
                if isinstance(cached_entry, Mapping)
                and isinstance(cached_entry.get("observation"), Mapping)
                else None
            )
            if (
                cached_observation is not None
                and cached_observation.get("review_status") != "recheck_required"
            ):
                cache_group = cache_hits_by_hash.setdefault(
                    digest,
                    {
                        "content_sha256": digest,
                        "labels": [],
                        "observation": copy.deepcopy(cached_observation),
                    },
                )
                cache_group["labels"].append(
                    {
                        "label": label,
                        "message_label": message_label,
                        "sender": item["sender"],
                        "role": item["role"],
                    }
                )
                continue
            representative = representative_items.get(digest)
            if representative is not None:
                representative.setdefault("same_content_labels", []).append(label)
                duplicate_aliases.append(
                    {
                        "label": label,
                        "message_label": message_label,
                        "representative_label": representative["label"],
                        "content_sha256": digest,
                        "sender": item["sender"],
                        "role": item["role"],
                    }
                )
                continue
            representative_items[digest] = item
            if cached_observation is not None:
                item["quality_warning"] = "cached_recheck_required"
                item["recheck_reasons"] = list(
                    cached_observation.get("recheck_reasons", [])
                )
                flush_current()
                current_items = [item]
                flush_current()
                continue
            if quality_warning:
                flush_current()
                current_items = [item]
                flush_current()
                continue
            current_items.append(item)
            if len(current_items) == parallel_limit:
                flush_current()
    flush_current()
    queued_images = sum(len(batch["items"]) for batch in batches)
    cache_hits = list(cache_hits_by_hash.values())
    cache_hit_labels = sum(len(item["labels"]) for item in cache_hits)
    pending_recheck_batches: list[dict[str, Any]] = []
    pending_recheck_count = 0
    for digest, raw_entry in cache_entries.items():
        if not isinstance(raw_entry, Mapping):
            continue
        observation = raw_entry.get("observation")
        source_labels = raw_entry.get("source_labels")
        if (
            not isinstance(observation, Mapping)
            or observation.get("review_status") != "recheck_required"
            or not isinstance(source_labels, list)
        ):
            continue
        covered = [
            str(label)
            for label in source_labels
            if str(label) in classified and str(label) in inventory
        ]
        if not covered:
            continue
        pending_recheck_count += 1
        if (
            digest in representative_items
            or len(pending_recheck_batches) == MEDIA_RECHECK_QUEUE_LIMIT
        ):
            continue
        representative_label = covered[0]
        _, representative_media = inventory[representative_label]
        pending_recheck_batches.append(
            {
                "batch_id": f"R{len(pending_recheck_batches) + 1:04d}",
                "view_detail": "original",
                "parallel": False,
                "items": [
                    {
                        "label": representative_label,
                        "path": representative_media.get("path"),
                        "content_sha256": digest,
                        "same_content_labels": covered[1:],
                        "quality_warning": "recheck_required",
                        "recheck_reasons": list(observation.get("recheck_reasons", [])),
                    }
                ],
            }
        )
    return {
        "contract_version": MEDIA_QUEUE_CONTRACT,
        "mode": (
            "finance_materials"
            if finance_mode
            else "store_ledger"
            if store_mode
            else "orders"
        ),
        "recommended_parallel_limit": parallel_limit,
        "batch_policy": batch_policy,
        "queued_images": queued_images,
        "covered_unclassified_images": (
            queued_images + len(duplicate_aliases) + cache_hit_labels
        ),
        "content_hashes_computed": content_hashes_computed,
        "content_hashes_reused": content_hashes_reused,
        "cache_hit_groups": len(cache_hits),
        "cache_hit_labels": cache_hit_labels,
        "cache_hits": cache_hits,
        "duplicate_hash_groups": len(
            {item["content_sha256"] for item in duplicate_aliases}
        ),
        "duplicate_alias_labels": len(duplicate_aliases),
        "duplicate_aliases": duplicate_aliases,
        "pending_recheck_count": pending_recheck_count,
        "pending_recheck_batches": pending_recheck_batches,
        "non_parallel_items": non_parallel_items,
        "missing_labels": missing_labels,
        "already_classified_in_page": already_classified,
        "batches": batches,
    }


def _empty_ocr_candidate_cache() -> dict[str, Any]:
    return {
        "contract_version": OCR_CANDIDATE_CACHE_CONTRACT,
        "candidate_contract": OCR_CANDIDATE_CONTRACT,
        "entries": {},
    }


def _normalize_ocr_candidate(
    value: object,
    *,
    digest: str,
) -> dict[str, Any]:
    core.require(isinstance(value, Mapping), "OCR candidate must be an object")
    allowed = {
        "contract_version",
        "content_sha256",
        "status",
        "backend",
        "backend_version",
        "authoritative",
        "text",
        "average_confidence",
        "line_count",
        "elapsed_ms",
        "truncated",
        "error",
    }
    unknown = sorted(set(value) - allowed)
    core.require(not unknown, "OCR candidate has unsupported fields: " + ", ".join(unknown))
    core.require(
        value.get("contract_version") == OCR_CANDIDATE_CONTRACT,
        "OCR candidate has an unsupported contract",
    )
    core.require(
        core.clean_text(value.get("content_sha256")) == digest,
        "OCR candidate content hash does not match its cache key",
    )
    status = core.clean_text(value.get("status")).casefold()
    core.require(status in {"ok", "empty", "error"}, "OCR candidate status is invalid")
    backend = core.clean_text(value.get("backend"))
    core.require(bool(backend), "OCR candidate backend is required")
    backend_version = core.clean_text(value.get("backend_version")) or None
    core.require(
        value.get("authoritative") is False,
        "OCR candidates must be explicitly non-authoritative",
    )
    text_value = str(value.get("text") or "").replace("\x00", "").strip()
    truncated = bool(value.get("truncated"))
    if len(text_value) > OCR_CANDIDATE_TEXT_LIMIT:
        text_value = text_value[:OCR_CANDIDATE_TEXT_LIMIT].rstrip()
        truncated = True
    confidence = value.get("average_confidence")
    core.require(
        confidence is None
        or (
            isinstance(confidence, (int, float))
            and not isinstance(confidence, bool)
            and 0 <= float(confidence) <= 1
        ),
        "OCR candidate average_confidence must be null or 0..1",
    )
    line_count = value.get("line_count", 0)
    elapsed_ms = value.get("elapsed_ms", 0)
    core.require(
        isinstance(line_count, int) and not isinstance(line_count, bool) and line_count >= 0,
        "OCR candidate line_count must be a non-negative integer",
    )
    core.require(
        isinstance(elapsed_ms, int) and not isinstance(elapsed_ms, bool) and elapsed_ms >= 0,
        "OCR candidate elapsed_ms must be a non-negative integer",
    )
    error = core.clean_text(value.get("error"))[:300] or None
    core.require(status == "error" or error is None, "successful OCR candidate cannot contain an error")
    return {
        "contract_version": OCR_CANDIDATE_CONTRACT,
        "content_sha256": digest,
        "status": status,
        "backend": backend,
        "backend_version": backend_version,
        "authoritative": False,
        "text": text_value,
        "average_confidence": float(confidence) if confidence is not None else None,
        "line_count": line_count,
        "elapsed_ms": elapsed_ms,
        "truncated": truncated,
        **({"error": error} if error is not None else {}),
    }


def _load_ocr_candidate_cache(work: Path) -> dict[str, Any]:
    path = _ocr_cache_path(work)
    if not path.is_file():
        return _empty_ocr_candidate_cache()
    try:
        raw = _load_json(path)
        if (
            raw.get("contract_version") != OCR_CANDIDATE_CACHE_CONTRACT
            or raw.get("candidate_contract") != OCR_CANDIDATE_CONTRACT
            or not isinstance(raw.get("entries"), Mapping)
        ):
            return _empty_ocr_candidate_cache()
        entries: dict[str, Any] = {}
        for raw_digest, raw_candidate in raw["entries"].items():
            digest = str(raw_digest)
            if re.fullmatch(r"[0-9a-f]{64}", digest) is None:
                continue
            try:
                entries[digest] = _normalize_ocr_candidate(
                    raw_candidate,
                    digest=digest,
                )
            except (TypeError, ValueError):
                continue
        return {
            "contract_version": OCR_CANDIDATE_CACHE_CONTRACT,
            "candidate_contract": OCR_CANDIDATE_CONTRACT,
            "entries": entries,
        }
    except (OSError, TypeError, ValueError, json.JSONDecodeError):
        # This cache is derived and non-authoritative. Corruption must never
        # block direct visual review.
        return _empty_ocr_candidate_cache()


def _invoke_ocr_worker(requests: list[dict[str, str]]) -> dict[str, Any]:
    worker = Path(__file__).resolve().parent / "ocr_candidates.py"
    if not worker.is_file():
        return {"status": "unavailable", "reason": "ocr_worker_missing", "results": {}}
    payload = {
        "contract_version": OCR_WORKER_CONTRACT,
        "candidate_contract": OCR_CANDIDATE_CONTRACT,
        "text_limit": OCR_CANDIDATE_TEXT_LIMIT,
        "requests": requests,
    }
    try:
        completed = subprocess.run(
            [sys.executable, str(worker)],
            input=json.dumps(payload, ensure_ascii=False),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=OCR_WORKER_TIMEOUT_SECONDS,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        reason = "ocr_worker_timeout" if isinstance(exc, subprocess.TimeoutExpired) else "ocr_worker_failed"
        return {"status": "error", "reason": reason, "results": {}}
    if completed.returncode != 0:
        return {"status": "error", "reason": "ocr_worker_failed", "results": {}}
    try:
        response = json.loads(completed.stdout)
    except json.JSONDecodeError:
        return {"status": "error", "reason": "ocr_worker_invalid_output", "results": {}}
    if not isinstance(response, Mapping) or response.get("contract_version") != OCR_WORKER_CONTRACT:
        return {"status": "error", "reason": "ocr_worker_invalid_contract", "results": {}}
    results = response.get("results", {})
    if not isinstance(results, Mapping):
        return {"status": "error", "reason": "ocr_worker_invalid_results", "results": {}}
    return {
        "status": core.clean_text(response.get("status")) or "error",
        "reason": core.clean_text(response.get("reason")) or None,
        "backend": core.clean_text(response.get("backend")) or None,
        "backend_version": core.clean_text(response.get("backend_version")) or None,
        "results": dict(results),
    }


def _queue_ocr_item_references(media_queue: Mapping[str, Any]) -> list[dict[str, Any]]:
    references: list[dict[str, Any]] = []
    for field in ("batches", "pending_recheck_batches"):
        batches = media_queue.get(field, [])
        if not isinstance(batches, list):
            continue
        for batch in batches:
            if not isinstance(batch, Mapping) or not isinstance(batch.get("items"), list):
                continue
            references.extend(item for item in batch["items"] if isinstance(item, dict))
    return references


def _compact_ocr_candidate(candidate: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "contract_version": candidate.get("contract_version"),
        "status": candidate.get("status"),
        "backend": candidate.get("backend"),
        "backend_version": candidate.get("backend_version"),
        "authoritative": False,
        "text": candidate.get("text"),
        "average_confidence": candidate.get("average_confidence"),
        "line_count": candidate.get("line_count"),
        "elapsed_ms": candidate.get("elapsed_ms"),
        "truncated": bool(candidate.get("truncated")),
        **({"error": candidate.get("error")} if candidate.get("error") else {}),
    }


def _attach_ocr_candidates(
    work: Path,
    run: Mapping[str, Any],
    media_queue: dict[str, Any],
) -> None:
    config = _ocr_candidate_config_for_run(run)
    state: dict[str, Any] = {
        **config,
        "status": "disabled" if not config["enabled"] else "enabled",
        "representative_images": 0,
        "candidate_request_limit": OCR_CANDIDATE_PAGE_LIMIT,
        "deferred_representatives": 0,
        "candidate_count": 0,
        "cache_entry_count": 0,
        "candidate_elapsed_ms": 0,
        "average_candidate_ms": 0.0,
        "error_candidate_count": 0,
        "backends": [],
    }
    media_queue["ocr_candidates"] = state
    if not config["enabled"]:
        return

    item_references = _queue_ocr_item_references(media_queue)
    requests_by_digest: dict[str, dict[str, str]] = {}
    for item in item_references:
        digest = core.clean_text(item.get("content_sha256"))
        path = core.clean_text(item.get("path"))
        if re.fullmatch(r"[0-9a-f]{64}", digest) is None or not path:
            continue
        requests_by_digest.setdefault(digest, {"content_sha256": digest, "path": path})
    state["representative_images"] = len(requests_by_digest)
    if not requests_by_digest:
        state["status"] = "no_representatives"
        return
    active_requests = dict(
        list(requests_by_digest.items())[:OCR_CANDIDATE_PAGE_LIMIT]
    )
    state["deferred_representatives"] = len(requests_by_digest) - len(active_requests)

    cache = _load_ocr_candidate_cache(work)
    entries = cache["entries"]
    missing = [
        request
        for digest, request in active_requests.items()
        if digest not in entries
    ]
    worker_response: dict[str, Any] | None = None
    changed = False
    if missing:
        worker_response = _invoke_ocr_worker(missing)
        raw_results = worker_response.get("results", {})
        if isinstance(raw_results, Mapping):
            for digest, request in active_requests.items():
                if digest in entries or digest not in raw_results:
                    continue
                try:
                    entries[digest] = _normalize_ocr_candidate(
                        raw_results[digest],
                        digest=digest,
                    )
                    changed = True
                except (TypeError, ValueError):
                    continue
    if changed:
        cache_path = _ocr_cache_path(work)
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        core.atomic_json(cache_path, cache)

    resolved = {
        digest: entries[digest]
        for digest in active_requests
        if digest in entries
    }
    for item in item_references:
        digest = core.clean_text(item.get("content_sha256"))
        candidate = resolved.get(digest)
        if candidate is not None:
            item["ocr_candidate"] = _compact_ocr_candidate(candidate)

    backends = sorted(
        {
            core.clean_text(candidate.get("backend"))
            for candidate in resolved.values()
            if core.clean_text(candidate.get("backend"))
        }
    )
    state["candidate_count"] = len(resolved)
    state["cache_entry_count"] = len(entries)
    state["candidate_elapsed_ms"] = sum(
        int(candidate.get("elapsed_ms") or 0) for candidate in resolved.values()
    )
    state["average_candidate_ms"] = (
        state["candidate_elapsed_ms"] / len(resolved) if resolved else 0.0
    )
    state["error_candidate_count"] = sum(
        candidate.get("status") == "error" for candidate in resolved.values()
    )
    state["backends"] = backends
    if len(resolved) == len(active_requests):
        state["status"] = (
            "ready_limited"
            if state["deferred_representatives"]
            else "ready"
        )
    elif resolved:
        state["status"] = "partial"
    elif worker_response is not None:
        worker_status = core.clean_text(worker_response.get("status"))
        state["status"] = worker_status if worker_status in {"unavailable", "error"} else "unavailable"
        reason = core.clean_text(worker_response.get("reason"))
        if reason:
            state["reason"] = reason
    else:
        state["status"] = "unavailable"


def _omit_ocr_candidates_for_output_budget(media_queue: dict[str, Any]) -> None:
    omitted = 0
    for item in _queue_ocr_item_references(media_queue):
        if item.pop("ocr_candidate", None) is not None:
            omitted += 1
    state = media_queue.get("ocr_candidates")
    if isinstance(state, dict) and omitted:
        state["status"] = "omitted_for_output_budget"
        state["candidate_count"] = 0
        state["omitted_candidate_count"] = omitted


def _compact_page(
    group: Mapping[str, Any],
    start: int,
    end: int,
    *,
    label_by_id: Mapping[str, str] | None = None,
    message_by_id: Mapping[str, Mapping[str, Any]] | None = None,
    media_labels: Mapping[str, str] | None = None,
) -> tuple[list[dict[str, Any]], int]:
    messages = list(group.get("messages", []))
    if label_by_id is None:
        label_by_id, _ = _message_labels(group)
    if message_by_id is None:
        message_by_id = {
            str(message.get("message_id") or ""): message
            for message in messages
        }
    if media_labels is None:
        media_labels = {
            str(media.get("media_id") or ""): label
            for label, (_, media) in _media_inventory(group).items()
        }
    result: list[dict[str, Any]] = []
    collapsed = 0
    for index in range(start, end):
        message = messages[index]
        media_items = []
        for media in message.get("media", []):
            media_id = str(media.get("media_id") or "")
            if media_id not in media_labels:
                continue
            media_items.append(
                {
                    "label": media_labels[media_id],
                    "kind": media.get("kind"),
                    "availability": media.get("availability"),
                    "path": media.get("path"),
                }
            )
        if message.get("excluded_from_accounting") and not media_items:
            collapsed += 1
            continue
        reply_id = str(message.get("reply_to_message_id") or "")
        reply = None
        if reply_id in label_by_id:
            target = message_by_id.get(reply_id)
            target_media_labels = [
                media_labels[str(media.get("media_id") or "")]
                for media in (target.get("media", []) if target else [])
                if str(media.get("media_id") or "") in media_labels
            ]
            reply = {
                "label": label_by_id[reply_id],
                "sender": target.get("sender_name") if target else None,
                "text": target.get("text") if target else None,
                "media_labels": target_media_labels,
            }
        result.append(
            {
                "label": f"S{index + 1:05d}",
                "timestamp": message.get("timestamp"),
                "sender": message.get("sender_name") or message.get("sender_id"),
                "role": message.get("role"),
                "text": message.get("text"),
                "accounting_context_only": bool(message.get("accounting_context_only")),
                "reply": reply,
                "media": media_items,
            }
        )
    return result, collapsed


def _review_page_token(
    group: Mapping[str, Any],
    start: int,
    end: int,
    compact: list[dict[str, Any]] | None = None,
    collapsed: int | None = None,
    group_fingerprint: str | None = None,
) -> str:
    if compact is None or collapsed is None:
        compact, collapsed = _compact_page(group, start, end)
    return core.fingerprint_json(
        {
            "contract_version": REVIEW_PAGE_CONTRACT,
            "group_fingerprint": group_fingerprint or core.fingerprint_json(group),
            "page_start": start,
            "page_end": end,
            "messages": compact,
            "collapsed_system_or_empty_messages": collapsed,
        }
    )


def _commit_review_page(
    decision: dict[str, Any],
    batch: Mapping[str, Any],
    group: Mapping[str, Any],
) -> tuple[int, int] | None:
    page_commit = batch.get("page_commit")
    if page_commit is None:
        return None
    assert isinstance(page_commit, Mapping)
    start = page_commit["page_start"]
    end = page_commit["page_end"]
    messages = list(group.get("messages", []))
    committed_through = int(decision.get("reviewed_through") or 0)
    core.require(
        start == committed_through,
        f"page_commit must start at reviewed_through={committed_through}; fetch review next again",
    )
    core.require(start < end <= len(messages), "page_commit range is empty or outside the group")
    core.require(
        end - start <= MAX_PAGE_SIZE,
        f"page_commit cannot exceed {MAX_PAGE_SIZE} messages",
    )
    expected_token = _review_page_token(group, start, end)
    core.require(
        page_commit["page_token"] == expected_token,
        "page_commit token does not match the exact page returned by review next",
    )
    decision["reviewed_through"] = end
    decision["read_complete"] = end == len(messages)
    return start, end


def _open_order_carry_messages(
    group: Mapping[str, Any],
    open_orders: object,
) -> list[dict[str, Any]]:
    if not isinstance(open_orders, list):
        return []
    labels: set[str] = set()
    for item in open_orders:
        if not isinstance(item, Mapping):
            continue
        source_messages = item.get("source_messages", [])
        if isinstance(source_messages, list):
            labels.update(str(value) for value in source_messages)
    positions = sorted(
        int(match.group(1)) - 1
        for label in labels
        if (match := re.fullmatch(r"S(\d{5})", label)) is not None
    )
    result: list[dict[str, Any]] = []
    for position in positions:
        if 0 <= position < len(group.get("messages", [])):
            compact, _ = _compact_page(group, position, position + 1)
            result.extend(compact)
    return result


def _roll_forward_carry_messages(
    group: Mapping[str, Any],
    decision: Mapping[str, Any],
    carry_ids: object,
) -> list[dict[str, Any]]:
    if not isinstance(carry_ids, list) or not carry_ids:
        return []
    wanted = {str(item) for item in carry_ids}
    collection = (
        list(decision.get("records", []))
        + list(decision.get("balance_snapshots", []))
        if decision.get("contract_version") == store_ledger.DECISION_CONTRACT
        else decision.get("orders", [])
    )
    sources: list[str] = []
    for item in collection:
        if not isinstance(item, Mapping) or core.clean_text(item.get("id")) not in wanted:
            continue
        raw_sources = item.get("source_messages", [])
        if isinstance(raw_sources, list):
            sources.extend(str(value) for value in raw_sources)
        lifecycle = item.get("lifecycle")
        if isinstance(lifecycle, Mapping) and isinstance(
            lifecycle.get("source_messages"), list
        ):
            sources.extend(str(value) for value in lifecycle["source_messages"])
    return _open_order_carry_messages(
        group,
        [{"source_messages": list(dict.fromkeys(sources))}],
    )


def _reply_target_carry_messages(
    group: Mapping[str, Any], start: int, end: int
) -> list[dict[str, Any]]:
    messages = list(group.get("messages", []))
    position_by_id = {
        core.clean_text(message.get("message_id")): position
        for position, message in enumerate(messages)
        if isinstance(message, Mapping)
        and core.clean_text(message.get("message_id"))
    }
    positions: list[int] = []
    for message in messages[start:end]:
        if not isinstance(message, Mapping):
            continue
        target_position = position_by_id.get(
            core.clean_text(message.get("reply_to_message_id"))
        )
        if target_position is not None and not start <= target_position < end:
            positions.append(target_position)
    result: list[dict[str, Any]] = []
    for position in dict.fromkeys(positions):
        compact, _ = _compact_page(group, position, position + 1)
        result.extend(compact)
    return result


def _render_result_json(result: Mapping[str, Any], *, compact: bool) -> str:
    if compact:
        return json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n"
    return json.dumps(result, ensure_ascii=False, indent=2) + "\n"


def _review_page_result(
    *,
    group: Mapping[str, Any],
    group_key: str,
    run_group: Mapping[str, Any],
    decision: Mapping[str, Any],
    start: int,
    end: int,
    group_fingerprint: str,
    semantic_fingerprint: str,
    open_field: str,
    open_records: list[Any],
    carry_messages: list[dict[str, Any]],
    label_by_id: Mapping[str, str] | None = None,
    message_by_id: Mapping[str, Mapping[str, Any]] | None = None,
    media_labels: Mapping[str, str] | None = None,
    media_inventory: Mapping[
        str, tuple[Mapping[str, Any], Mapping[str, Any]]
    ] | None = None,
    media_hashes: Mapping[str, str] | None = None,
    content_hashes_computed: int = 0,
    content_hashes_reused: int = 0,
) -> dict[str, Any]:
    messages = list(group.get("messages", []))
    compact, collapsed = _compact_page(
        group,
        start,
        end,
        label_by_id=label_by_id,
        message_by_id=message_by_id,
        media_labels=media_labels,
    )
    page_token = (
        _review_page_token(
            group,
            start,
            end,
            compact,
            collapsed,
            group_fingerprint,
        )
        if start < end
        else None
    )
    combined_carry: list[dict[str, Any]] = []
    seen_carry_labels: set[str] = set()
    for item in carry_messages + _reply_target_carry_messages(group, start, end):
        label = core.clean_text(item.get("label"))
        if label and label in seen_carry_labels:
            continue
        if label:
            seen_carry_labels.add(label)
        combined_carry.append(item)
    result = {
        "contract_version": REVIEW_PAGE_CONTRACT,
        "group_key": group_key,
        "run_label": run_group.get("run_label"),
        "page_start": start,
        "page_end": end,
        "message_count": len(messages),
        "messages": compact,
        "collapsed_system_or_empty_messages": collapsed,
        "returned_messages": len(compact),
        "page_reaches_group_end": end == len(messages),
        "done": bool(decision.get("read_complete")),
        "commit_required": start < end,
        "carry_messages": combined_carry,
        "roll_forward_carry": copy.deepcopy(
            run_group.get("roll_forward_carry", [])
        ),
        "media_queue": _review_media_queue(
            group,
            decision,
            start,
            end,
            media_labels=media_labels,
            inventory=media_inventory,
            media_hashes=media_hashes,
            content_hashes_computed=content_hashes_computed,
            content_hashes_reused=content_hashes_reused,
        ),
        "controlled_editing": _controlled_editing(run_group),
        "page_token": page_token,
        "semantic_fingerprint": semantic_fingerprint,
    }
    result[open_field] = open_records
    return result


RISK_DIAGNOSTIC_FIELDS = (
    "blank_direction_orders",
    "unknown_pricing_orders",
    "single_unknown_pricing_orders",
    "own_fund_source_only_orders",
    "repeated_note_orders",
    "multi_signal_degenerate_orders",
    "suspected_bulk_order_creation",
    "order_continuity_review_candidate_count",
    "unassigned_entries",
    "unknown_payees",
    "unknown_payee_warning",
    "pending_people",
    "unassigned_material_media",
    "pending_store_records",
    "unassigned_store_media",
    "media_recheck_required",
)
RISK_FLAG_FIELDS = (
    "blank_direction_orders",
    "unknown_pricing_orders",
    "repeated_note_orders",
    "multi_signal_degenerate_orders",
    "suspected_bulk_order_creation",
    "order_continuity_review_candidate_count",
    "unassigned_entries",
    "unknown_payee_warning",
    "pending_people",
    "unassigned_material_media",
    "pending_store_records",
    "unassigned_store_media",
    "media_recheck_required",
)


def _review_risk_summary(reports: list[dict[str, Any]]) -> dict[str, Any]:
    totals = {
        field: sum(int(report.get(field) or 0) for report in reports)
        for field in RISK_DIAGNOSTIC_FIELDS
    }
    flagged_groups = sum(
        any(int(report.get(field) or 0) for field in RISK_FLAG_FIELDS)
        for report in reports
    )
    return {
        "groups": len(reports),
        "flagged_groups": flagged_groups,
        "totals": totals,
    }


def _upgrade_legacy_checkpoints(
    work: Path,
    run: dict[str, Any],
    run_group: Mapping[str, Any],
    normalized: Mapping[str, Any],
    group: Mapping[str, Any],
    decision_path: Path,
    decision: Mapping[str, Any],
) -> dict[str, Any]:
    core.require(
        decision.get("contract_version") == LEGACY_DECISION_CONTRACT,
        f"upgrade-checkpoints only accepts {LEGACY_DECISION_CONTRACT}",
    )
    core.require(
        core.clean_text(run_group.get("edit_mode")) == LEGACY_EDIT_CONTROL_MODE,
        "legacy task is not using the controlled 3.1 edit protocol",
    )
    control = decision.get("edit_control")
    core.require(isinstance(control, Mapping), "legacy decision is missing edit_control")
    core.require(
        core.clean_text(control.get("mode")) == LEGACY_EDIT_CONTROL_MODE,
        "legacy decision has an unsupported edit_control mode",
    )
    legacy_fingerprint = _legacy_semantic_fingerprint_3_1(decision)
    core.require(
        core.clean_text(control.get("approved_semantic_fingerprint"))
        == legacy_fingerprint,
        "legacy decision semantic fields changed outside its approved batches",
    )

    previous_reviewed_through = int(decision.get("reviewed_through") or 0)
    candidate = copy.deepcopy(dict(decision))
    candidate["contract_version"] = DECISION_CONTRACT
    candidate["reviewed_through"] = 0
    candidate["read_complete"] = not bool(group.get("messages", []))
    candidate["open_orders"] = []
    candidate["sealed"] = False
    candidate["sealed_decision_fingerprint"] = None
    statistics = _validate_decision(
        normalized,
        group,
        candidate,
        require_complete=False,
        capture_hashes=False,
        rehash_labels=set(),
    )
    current_fingerprint = _semantic_fingerprint(candidate)
    candidate["edit_control"] = _new_edit_control(current_fingerprint)
    core.atomic_json(decision_path, candidate)
    run_group["edit_mode"] = EDIT_CONTROL_MODE
    run_group["sealed"] = False
    run["status"] = "reviewing"
    core.atomic_json(_run_path(work), run)
    return {
        "group_key": run_group.get("group_key"),
        "run_label": run_group.get("run_label"),
        "upgraded_from": LEGACY_DECISION_CONTRACT,
        "contract_version": DECISION_CONTRACT,
        "previous_reviewed_through": previous_reviewed_through,
        "reviewed_through": candidate["reviewed_through"],
        "preserved_media_decisions": len(candidate.get("media_decisions", {})),
        "preserved_orders": len(candidate.get("orders", [])),
        "semantic_fingerprint": current_fingerprint,
        **statistics,
    }


def review_command(args: argparse.Namespace) -> dict[str, Any]:
    work = args.work.resolve()
    run = _load_run(work)
    normalized = _load_snapshot(work, run)
    groups = _group_index(normalized)
    if args.review_action == "status":
        selected = (
            [_select_run_group(run, args.group)]
            if args.group
            else [item for item in run["groups"] if isinstance(item, Mapping)]
        )
        reports = []
        for run_group in selected:
            decision = _load_json(_decision_path(work, run_group))
            semantic_fingerprint = _semantic_fingerprint(decision)
            control = decision.get("edit_control")
            approved_fingerprint = (
                core.clean_text(control.get("approved_semantic_fingerprint"))
                if isinstance(control, Mapping)
                else ""
            )
            open_field = _open_field_for_decision(decision)
            report = {
                    "group_key": run_group.get("group_key"),
                    "run_label": run_group.get("run_label"),
                    "reviewed_through": decision.get("reviewed_through"),
                    "read_complete": bool(decision.get("read_complete")),
                    "message_count": decision.get("message_count"),
                    "classified_media": len(decision.get("media_decisions", {})),
                    "evidence_media_count": decision.get("evidence_media_count"),
                    "sealed": bool(decision.get("sealed")),
                    "controlled_editing": _controlled_editing(run_group),
                    "semantic_fingerprint": semantic_fingerprint,
                    "uncontrolled_semantic_changes": bool(
                        _controlled_editing(run_group)
                        and approved_fingerprint != semantic_fingerprint
                    ),
                }
            report.update(
                _validate_media_observation_cache(
                    groups[str(run_group["group_key"])],
                    decision,
                    require_complete=False,
                    verify_files=False,
                )
            )
            report[open_field] = copy.deepcopy(decision.get(open_field, []))
            reports.append(report)
        return {
            "runtime_contract": run.get("runtime_contract") or "legacy-unversioned",
            "amount_policy": run.get("amount_policy") or "legacy-unspecified",
            "workbook_compatibility": (
                run.get("workbook_compatibility") or "legacy-unspecified"
            ),
            "ocr_candidates": _ocr_candidate_config_for_run(run),
            "groups": reports,
        }

    if args.review_action == "audit":
        selected = (
            [_select_run_group(run, args.group)]
            if args.group
            else [item for item in run["groups"] if isinstance(item, Mapping)]
        )
        audit_reports: list[dict[str, Any]] = []
        for run_group in selected:
            group_key = str(run_group["group_key"])
            decision = _load_json(_decision_path(work, run_group))
            _require_controlled_semantics(run_group, decision)
            statistics = _validate_decision(
                normalized,
                groups[group_key],
                copy.deepcopy(decision),
                require_complete=False,
                capture_hashes=False,
                rehash_labels=set(),
            )
            flags = {
                field: statistics[field]
                for field in RISK_FLAG_FIELDS
                if int(statistics.get(field) or 0)
            }
            audit_reports.append(
                {
                    "group_key": group_key,
                    "run_label": run_group.get("run_label"),
                    "sealed": bool(decision.get("sealed")),
                    "flags": flags,
                    **statistics,
                }
            )
        return {
            "summary": _review_risk_summary(audit_reports),
            "groups": audit_reports,
        }

    run_group = _select_run_group(
        run,
        args.group,
        prefer_incomplete=args.review_action == "next",
    )
    group_key = str(run_group["group_key"])
    group = groups[group_key]
    decision_path = _decision_path(work, run_group)
    decision = _load_json(decision_path)

    if args.review_action == "upgrade-checkpoints":
        return _upgrade_legacy_checkpoints(
            work,
            run,
            run_group,
            normalized,
            group,
            decision_path,
            decision,
        )

    if args.review_action == "apply-batch":
        core.require(decision.get("sealed") is not True, "sealed group cannot accept a review batch")
        controlled = _controlled_editing(run_group)
        if controlled:
            _require_controlled_semantics(run_group, decision)
        current_fingerprint = _semantic_fingerprint(decision)
        batch = _load_review_batch(args.input.resolve())
        expected_open_field = _open_field_for_decision(decision)
        other_open_fields = {
            "open_orders",
            "open_exchanges",
            "open_people",
            "open_records",
        } - {expected_open_field}
        if batch.get("page_commit") is not None:
            core.require(
                expected_open_field in batch,
                f"page commit for this group must include {expected_open_field}",
            )
            supplied_other_open_fields = sorted(other_open_fields & set(batch))
            core.require(
                not supplied_other_open_fields,
                "page commit for this group cannot include "
                + ", ".join(supplied_other_open_fields),
            )
        input_fingerprint = core.fingerprint_json(batch)
        control = decision.get("edit_control")
        last_batch = (
            control.get("last_batch")
            if controlled and isinstance(control, Mapping)
            else None
        )
        if (
            isinstance(last_batch, Mapping)
            and core.clean_text(last_batch.get("id")) == batch["batch_id"]
        ):
            core.require(
                core.clean_text(last_batch.get("input_fingerprint"))
                == input_fingerprint,
                f"review batch id was already used with different content: {batch['batch_id']}",
            )
            core.require(
                core.clean_text(last_batch.get("result_fingerprint"))
                == current_fingerprint,
                "last applied batch no longer matches the current decision",
            )
            statistics = _validate_decision(
                normalized,
                group,
                copy.deepcopy(decision),
                require_complete=False,
                capture_hashes=False,
                rehash_labels=set(),
            )
            return {
                "group_key": group_key,
                "run_label": run_group.get("run_label"),
                "sealed": False,
                "controlled_editing": True,
                "batch_id": batch["batch_id"],
                "idempotent_replay": True,
                "semantic_fingerprint": current_fingerprint,
                "reviewed_through": decision.get("reviewed_through"),
                "read_complete": bool(decision.get("read_complete")),
                **statistics,
            }
        core.require(
            batch["base_fingerprint"] == current_fingerprint,
            "review batch is stale: base_fingerprint does not match the current decision; "
            "refresh review status and rebuild the batch",
        )
        auto_transition_order_ids: set[str] = set()
        if decision.get("contract_version") in {
            DECISION_CONTRACT,
            large_daily.DECISION_CONTRACT,
        }:
            existing_orders = {
                core.clean_text(item.get("id")): item
                for item in decision.get("orders", [])
                if isinstance(item, Mapping)
            }
            for update in batch.get("orders", []):
                if not isinstance(update, Mapping) or "lifecycle" in update:
                    continue
                order_id = core.clean_text(update.get("id"))
                existing = existing_orders.get(order_id)
                lifecycle = (
                    existing.get("lifecycle")
                    if isinstance(existing, Mapping)
                    else None
                )
                if (
                    isinstance(lifecycle, Mapping)
                    and lifecycle.get("status") == "pending_next_day"
                ):
                    auto_transition_order_ids.add(order_id)
        candidate = _merge_review_batch(decision, batch)
        committed_page = _commit_review_page(candidate, batch, group)
        statistics = _validate_decision(
            normalized,
            group,
            candidate,
            require_complete=False,
            capture_hashes=True,
            rehash_labels={
                str(label)
                for label in batch.get("media_decisions", {})
            },
            auto_transition_order_ids=auto_transition_order_ids,
        )
        observation_apply_statistics = _apply_media_observations(
            candidate,
            batch,
            group,
        )
        statistics.update(
            _validate_media_observation_cache(
                group,
                candidate,
                require_complete=False,
                verify_files=False,
            )
        )
        statistics.update(observation_apply_statistics)
        if "roll_forward_carry" in run_group:
            run_group["roll_forward_carry"] = _carry_ids(candidate)
        result_fingerprint = _semantic_fingerprint(candidate)
        prior_batch_count = (
            control.get("batch_count")
            if controlled
            and isinstance(control, Mapping)
            and isinstance(control.get("batch_count"), int)
            else 0
        )
        candidate["edit_control"] = {
            "mode": EDIT_CONTROL_MODE,
            "approved_semantic_fingerprint": result_fingerprint,
            "batch_count": prior_batch_count + 1,
            "last_batch": {
                "id": batch["batch_id"],
                "input_fingerprint": input_fingerprint,
                "base_fingerprint": current_fingerprint,
                "result_fingerprint": result_fingerprint,
                "applied_at": datetime.now(timezone.utc).isoformat(),
            },
        }
        core.atomic_json(decision_path, candidate)
        run_group["edit_mode"] = EDIT_CONTROL_MODE
        run_group["sealed"] = False
        core.atomic_json(_run_path(work), run)
        return {
            "group_key": group_key,
            "run_label": run_group.get("run_label"),
            "sealed": False,
            "controlled_editing": True,
            "batch_id": batch["batch_id"],
            "idempotent_replay": False,
            "semantic_fingerprint": result_fingerprint,
            "committed_page": (
                {"page_start": committed_page[0], "page_end": committed_page[1]}
                if committed_page is not None
                else None
            ),
            "reviewed_through": candidate.get("reviewed_through"),
            "read_complete": bool(candidate.get("read_complete")),
            **statistics,
        }

    if args.review_action == "next":
        expected_contracts = _expected_contracts_for_run_group(run_group)
        core.require(
            decision.get("contract_version") in expected_contracts,
            "review next decision contract does not match this run; "
            "start a fresh work directory for this paging protocol",
        )
        core.require(
            _require_controlled_semantics(run_group, decision),
            "review next requires the controlled apply-batch protocol",
        )
        exact_limit = args.limit
        if exact_limit is not None:
            core.require(
                1 <= exact_limit <= MAX_PAGE_SIZE,
                f"--limit must be 1..{MAX_PAGE_SIZE}",
            )
        start = int(decision.get("reviewed_through") or 0)
        messages = list(group.get("messages", []))
        maximum_count = exact_limit if exact_limit is not None else DEFAULT_PAGE_SIZE
        maximum_end = min(start + maximum_count, len(messages))
        group_fingerprint = core.fingerprint_json(group)
        semantic_fingerprint = _semantic_fingerprint(decision)
        open_field = _open_field_for_decision(decision)
        open_records = copy.deepcopy(decision.get(open_field, []))
        carry_messages = _open_order_carry_messages(group, open_records)
        carry_messages.extend(
            _roll_forward_carry_messages(
                group,
                decision,
                run_group.get("roll_forward_carry", []),
            )
        )
        carry_messages = list(
            {
                core.clean_text(item.get("label")) or core.fingerprint_json(item): item
                for item in carry_messages
            }.values()
        )
        label_by_id, _ = _message_labels(group)
        message_by_id = {
            str(message.get("message_id") or ""): message
            for message in messages
        }
        media_inventory = _media_inventory(group)
        media_labels = {
            str(media.get("media_id") or ""): label
            for label, (_, media) in media_inventory.items()
        }
        media_hashes: dict[str, str] = {}
        media_hash_was_computed: set[str] = set()
        media_position_by_label: dict[str, int] = {}
        classified_media_labels = set(decision.get("media_decisions", {}))
        for message_position in range(start, maximum_end):
            message = messages[message_position]
            for media in message.get("media", []):
                if not isinstance(media, Mapping) or not core.media_is_evidence(media):
                    continue
                label = media_labels.get(str(media.get("media_id") or ""))
                if (
                    not label
                    or label in classified_media_labels
                    or media.get("availability") != "available"
                ):
                    continue
                mime_type = core.clean_text(media.get("mime_type")).casefold()
                kind = core.clean_text(media.get("kind")).casefold()
                if kind != "image" and not mime_type.startswith("image/"):
                    continue
                digest, computed = _media_content_hash_for_label(
                    group,
                    decision,
                    label,
                    inventory=media_inventory,
                    verify_file=False,
                )
                media_hashes[label] = digest
                media_position_by_label[label] = message_position
                if computed:
                    media_hash_was_computed.add(label)

        def build_page(end: int) -> dict[str, Any]:
            page_hash_labels = {
                label
                for label, position in media_position_by_label.items()
                if position < end
            }
            return _review_page_result(
                group=group,
                group_key=group_key,
                run_group=run_group,
                decision=decision,
                start=start,
                end=end,
                group_fingerprint=group_fingerprint,
                semantic_fingerprint=semantic_fingerprint,
                open_field=open_field,
                open_records=open_records,
                carry_messages=carry_messages,
                label_by_id=label_by_id,
                message_by_id=message_by_id,
                media_labels=media_labels,
                media_inventory=media_inventory,
                media_hashes=media_hashes,
                content_hashes_computed=len(
                    page_hash_labels & media_hash_was_computed
                ),
                content_hashes_reused=len(
                    page_hash_labels - media_hash_was_computed
                ),
            )

        def finalize_page(end: int) -> dict[str, Any]:
            page = build_page(end)
            _attach_ocr_candidates(work, run, page["media_queue"])
            return page

        if exact_limit is not None or start == maximum_end:
            return finalize_page(maximum_end)

        selected_page: dict[str, Any] | None = None
        first_page_chars: int | None = None
        for end in range(start + 1, maximum_end + 1):
            candidate = build_page(end)
            output_chars = len(_render_result_json(candidate, compact=True))
            if first_page_chars is None:
                first_page_chars = output_chars
            if output_chars <= DEFAULT_REVIEW_PAGE_OUTPUT_CHAR_BUDGET:
                selected_page = candidate

        core.require(
            selected_page is not None,
            "review next cannot return even one complete message within the "
            f"{DEFAULT_REVIEW_PAGE_OUTPUT_CHAR_BUDGET}-character output budget "
            f"(first complete page needs {first_page_chars} characters); "
            "do not commit a truncated page",
        )
        selected_end = int(selected_page["page_end"])
        selected_page = finalize_page(selected_end)
        while (
            len(_render_result_json(selected_page, compact=True))
            > DEFAULT_REVIEW_PAGE_OUTPUT_CHAR_BUDGET
            and selected_end > start + 1
        ):
            selected_end -= 1
            selected_page = finalize_page(selected_end)
        if (
            len(_render_result_json(selected_page, compact=True))
            > DEFAULT_REVIEW_PAGE_OUTPUT_CHAR_BUDGET
        ):
            _omit_ocr_candidates_for_output_budget(selected_page["media_queue"])
        core.require(
            len(_render_result_json(selected_page, compact=True))
            <= DEFAULT_REVIEW_PAGE_OUTPUT_CHAR_BUDGET,
            "review next cannot include one complete message and its media queue "
            f"within the {DEFAULT_REVIEW_PAGE_OUTPUT_CHAR_BUDGET}-character output budget; "
            "do not commit a truncated page",
        )
        return selected_page

    controlled = _require_controlled_semantics(run_group, decision)
    prior_fingerprint = _semantic_fingerprint(decision)
    statistics = _validate_decision(
        normalized,
        group,
        decision,
        require_complete=args.review_action == "seal",
        capture_hashes=True,
    )
    current_fingerprint = _semantic_fingerprint(decision)
    if controlled:
        control = decision["edit_control"]
        control["approved_semantic_fingerprint"] = current_fingerprint
        last_batch = control.get("last_batch")
        if (
            isinstance(last_batch, dict)
            and last_batch.get("result_fingerprint") == prior_fingerprint
        ):
            last_batch["result_fingerprint"] = current_fingerprint
    if args.review_action == "seal":
        decision["sealed"] = True
        decision["sealed_decision_fingerprint"] = current_fingerprint
    elif prior_fingerprint != current_fingerprint or decision.get("sealed_decision_fingerprint") != current_fingerprint:
        decision["sealed"] = False
        decision["sealed_decision_fingerprint"] = None
    core.atomic_json(decision_path, decision)
    run_group["sealed"] = bool(decision.get("sealed"))
    if "roll_forward_carry" in run_group:
        run_group["roll_forward_carry"] = _carry_ids(decision)
    core.atomic_json(_run_path(work), run)
    return {
        "group_key": group_key,
        "run_label": run_group.get("run_label"),
        "sealed": bool(decision.get("sealed")),
        "controlled_editing": controlled,
        "semantic_fingerprint": current_fingerprint,
        **statistics,
    }


def _entry_event_type(entry: Mapping[str, Any]) -> str:
    side = core.clean_text(entry.get("side")).casefold()
    if core.clean_text(entry.get("kind")).casefold() == "cash":
        return "cash_payment" if side in {"payment", "payment_refund"} else "cash_payout"
    if side == "recovery":
        return "payout_recovery"
    if side == "payout":
        return "payout_screenshot"
    return "payment_screenshot"


def _translated_pricing(
    value: object,
    *,
    message_by_label: Mapping[str, Mapping[str, Any]],
) -> dict[str, Any]:
    core.require(isinstance(value, Mapping), "pricing must be an object")
    pricing = {
        key: copy.deepcopy(item)
        for key, item in value.items()
        if key != "source_messages"
    }
    pricing["source_message_ids"] = [
        str(message_by_label[str(label)]["message_id"])
        for label in value.get("source_messages", [])
    ]
    return pricing


def _translated_lifecycle(
    value: object,
    *,
    message_by_label: Mapping[str, Mapping[str, Any]],
) -> dict[str, Any]:
    core.require(isinstance(value, Mapping), "lifecycle must be an object")
    result = {
        key: copy.deepcopy(item)
        for key, item in value.items()
        if key not in {"source_messages", "history"}
    }
    result["source_message_ids"] = [
        str(message_by_label[str(label)]["message_id"])
        for label in value.get("source_messages", [])
    ]
    history: list[dict[str, Any]] = []
    for raw_event in value.get("history", []):
        core.require(isinstance(raw_event, Mapping), "lifecycle history must contain objects")
        event = {
            key: copy.deepcopy(item)
            for key, item in raw_event.items()
            if key != "source_messages"
        }
        event["source_message_ids"] = [
            str(message_by_label[str(label)]["message_id"])
            for label in raw_event.get("source_messages", [])
        ]
        history.append(event)
    result["history"] = history
    return result


def _translated_advanced_order(
    raw: Mapping[str, Any],
    *,
    event_by_entry: Mapping[str, str],
    message_by_label: Mapping[str, Mapping[str, Any]],
) -> dict[str, Any]:
    result: dict[str, Any] = {}
    same_transactions = raw.get("same_transactions")
    if same_transactions not in (None, "", []):
        core.require(isinstance(same_transactions, list), "same_transactions must be a list")
        result["same_transactions"] = [
            {
                "event_id": event_by_entry[str(item.get("entry_id") or "")],
                "same_as": event_by_entry[str(item.get("same_as") or "")],
            }
            for item in same_transactions
        ]
    legs = raw.get("legs")
    if legs not in (None, "", []):
        core.require(isinstance(legs, list), "legs must be a list")
        compiled_legs = []
        for item in legs:
            core.require(isinstance(item, Mapping), "each leg must be an object")
            leg = {
                key: copy.deepcopy(value)
                for key, value in item.items()
                if key
                not in {
                    "payout_entry_ids",
                    "recovery_entry_ids",
                    "pricing",
                }
            }
            leg["pricing"] = _translated_pricing(
                item.get("pricing"),
                message_by_label=message_by_label,
            )
            leg["payout_event_ids"] = [
                event_by_entry[str(value)] for value in item.get("payout_entry_ids", [])
            ]
            if item.get("recovery_entry_ids") not in (None, []):
                leg["recovery_event_ids"] = [
                    event_by_entry[str(value)] for value in item.get("recovery_entry_ids", [])
                ]
            compiled_legs.append(leg)
        result["legs"] = compiled_legs
    return result


def _compile_decisions_v3(
    normalized: Mapping[str, Any],
    decisions: Mapping[str, Mapping[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    events: list[dict[str, Any]] = []
    plan_groups: list[dict[str, Any]] = []
    for group in normalized.get("groups", []):
        group_key = str(group["group_key"])
        decision = decisions[group_key]
        inventory = _media_inventory(group)
        media_decisions = decision["media_decisions"]
        message_label_by_id, message_by_label = _message_labels(group)
        del message_label_by_id
        event_by_entry: dict[str, str] = {}
        side_by_entry: dict[str, str] = {}
        for label, (message, media) in inventory.items():
            media_id = str(media.get("media_id") or "")
            message_id = str(message.get("message_id") or "")
            if media.get("availability") == "missing":
                continue
            media_decision = media_decisions[label]
            if media_decision["classification"] == "reference":
                events.append(
                    {
                        "event_id": f"{media_id}#reference",
                        "group_key": group_key,
                        "message_id": message_id,
                        "media_id": media_id,
                        "type": "irrelevant_media",
                        "note": media_decision.get("note"),
                    }
                )
                continue
            for position, entry in enumerate(media_decision["entries"], start=1):
                entry_id = f"{label}.{position}"
                event_id = f"{media_id}#entry:{position}"
                amount_state = str(entry["amount_state"])
                event = {
                    "event_id": event_id,
                    "group_key": group_key,
                    "message_id": message_id,
                    "media_id": media_id,
                    "type": _entry_event_type(entry),
                    "evidence_sha256": media_decision["evidence_sha256"],
                    "ocr": {
                        "amount": entry.get("amount"),
                        "amount_text": entry.get("amount_text"),
                        "amount_basis": entry.get("amount_basis"),
                        "currency": entry.get("currency"),
                        "payee": entry["payee"],
                        "payee_state": entry.get("payee_state"),
                        "status_text": entry.get("status_text"),
                        "status_class": STATUS_CLASS[str(entry["result"])],
                        "status_class_confidence": "high",
                        "amount_completeness": "complete" if amount_state == "clear" else amount_state,
                        "confidence": "high" if amount_state == "clear" else "low",
                    },
                }
                event["flow_side"] = entry["side"]
                if isinstance(entry.get("side_exception"), Mapping):
                    side_exception = entry["side_exception"]
                    event["side_exception"] = {
                        "kind": side_exception["kind"],
                        "source_message_ids": [
                            str(message_by_label[str(source_label)]["message_id"])
                            for source_label in side_exception["source_messages"]
                        ],
                        "detail": side_exception["detail"],
                    }
                side_by_entry[entry_id] = str(entry["side"])
                if entry.get("note"):
                    event["note"] = entry["note"]
                events.append(event)
                event_by_entry[entry_id] = event_id

        plan_orders: list[dict[str, Any]] = []
        case_by_order_id: dict[str, str] = {}
        for raw in decision["orders"]:
            short_order_id = str(raw["id"])
            case_id = f"{group_key}:{short_order_id}"
            case_by_order_id[short_order_id] = case_id
            entry_ids = [str(item) for item in raw["entry_ids"]]
            order: dict[str, Any] = {
                "case_id": case_id,
                "event_ids": [event_by_entry[item] for item in entry_ids],
                "source_message_ids": [
                    str(message_by_label[str(item)]["message_id"])
                    for item in raw["source_messages"]
                ],
            }
            for field in (
                "customer_nickname",
                "direction",
            ):
                order[field] = copy.deepcopy(raw.get(field))
            if raw.get("fund_type") not in (None, ""):
                order["fund_type"] = copy.deepcopy(raw.get("fund_type"))
            for field in ("note",):
                if raw.get(field) not in (None, ""):
                    order[field] = copy.deepcopy(raw[field])
            if raw.get("pricing") not in (None, ""):
                order["pricing"] = _translated_pricing(
                    raw.get("pricing"),
                    message_by_label=message_by_label,
                )
            if raw.get("lifecycle") not in (None, ""):
                order["lifecycle"] = _translated_lifecycle(
                    raw.get("lifecycle"),
                    message_by_label=message_by_label,
                )
            event_sides = {
                event_by_entry[entry_id]: side_by_entry[entry_id]
                for entry_id in entry_ids
                if entry_id in side_by_entry
            }
            order["event_sides"] = event_sides
            order.update(
                _translated_advanced_order(
                    raw,
                    event_by_entry=event_by_entry,
                    message_by_label=message_by_label,
                )
            )
            plan_orders.append(order)
        settlement_allocations = []
        for raw_relation in decision.get("settlement_allocations", []):
            relation = {
                "source_event_id": event_by_entry[str(raw_relation["entry_id"])],
                "allocations": [
                    {
                        "target_case_id": case_by_order_id[str(item["order_id"])],
                        "amount": copy.deepcopy(item["amount"]),
                    }
                    for item in raw_relation["allocations"]
                ],
            }
            if raw_relation.get("source_messages"):
                relation["source_message_ids"] = [
                    str(message_by_label[str(item)]["message_id"])
                    for item in raw_relation["source_messages"]
                ]
            settlement_allocations.append(relation)
        balance_links = []
        for raw_link in decision.get("balance_links", []):
            link = {
                key: copy.deepcopy(value)
                for key, value in raw_link.items()
                if key not in {"source_order_id", "target_order_id", "source_messages"}
            }
            link["source_case_id"] = case_by_order_id[core.clean_text(raw_link.get("source_order_id"))]
            link["target_case_id"] = case_by_order_id[core.clean_text(raw_link.get("target_order_id"))]
            if raw_link.get("source_messages"):
                link["source_message_ids"] = [
                    str(message_by_label[str(item)]["message_id"])
                    for item in raw_link["source_messages"]
                ]
            balance_links.append(link)
        plan_groups.append(
            {
                "group_key": group_key,
                "orders": plan_orders,
                "balance_links": balance_links,
                "settlement_allocations": settlement_allocations,
            }
        )
    events_fingerprint = core.fingerprint_json(events)
    plan = {
        "contract_version": simple_ledger.SIMPLE_PLAN_CONTRACT,
        "normalized_source_fingerprint": normalized.get("source_fingerprint"),
        "events_fingerprint": events_fingerprint,
        "groups": plan_groups,
    }
    return events, plan


AMENDMENTS_CONTRACT = "group-chat-confirmed-amendments/1.0"
ROLL_FORWARD_CONTRACT = "group-chat-roll-forward/1.0"


def _successor_period(
    run: Mapping[str, Any],
    normalized: Mapping[str, Any],
    requested_date: str,
) -> dict[str, str | None]:
    accounting_date = _normalize_accounting_date(requested_date)
    assert accounting_date is not None
    prior_from = core.clean_text(run.get("accounting_from"))
    prior_to = core.clean_text(run.get("accounting_to"))
    if prior_from or prior_to:
        core.require(
            bool(prior_from) and bool(prior_to),
            "predecessor accounting window is incomplete",
        )
        start = datetime.strptime(prior_from, "%Y-%m-%d %H:%M")
        end = datetime.strptime(prior_to, "%Y-%m-%d %H:%M")
        duration = end - start
        core.require(duration > timedelta(0), "predecessor accounting window is invalid")
        successor_start = end
        successor_end = successor_start + duration
        core.require(
            accounting_date == successor_start.date().isoformat(),
            "roll-forward date must be the calendar date on which the immediately following window starts",
        )
        return {
            "date": accounting_date,
            "accounting_from": successor_start.strftime("%Y-%m-%d %H:%M"),
            "accounting_to": successor_end.strftime("%Y-%m-%d %H:%M"),
        }

    prior_date = core.clean_text(run.get("accounting_date"))
    if prior_date:
        expected = (
            datetime.strptime(prior_date, "%Y-%m-%d").date() + timedelta(days=1)
        ).isoformat()
    else:
        dates = [
            str(message.get("timestamp") or "")[:10]
            for group in normalized.get("groups", [])
            for message in group.get("messages", [])
            if isinstance(message, Mapping)
            and not message.get("accounting_context_only")
        ]
        core.require(bool(dates), "cannot infer the predecessor accounting date")
        expected = (
            datetime.strptime(max(dates), "%Y-%m-%d").date() + timedelta(days=1)
        ).isoformat()
    core.require(
        accounting_date == expected,
        f"roll-forward date must immediately follow the predecessor period: expected {expected}",
    )
    return {"date": accounting_date, "accounting_from": None, "accounting_to": None}


def _load_finished_decisions(
    work: Path,
    run: Mapping[str, Any],
    normalized: Mapping[str, Any],
) -> dict[str, dict[str, Any]]:
    core.require(run.get("status") == "finished", "predecessor run must be finished")
    groups = _group_index(normalized)
    decisions: dict[str, dict[str, Any]] = {}
    for run_group in run.get("groups", []):
        core.require(isinstance(run_group, Mapping), "predecessor run group is invalid")
        group_key = core.clean_text(run_group.get("group_key"))
        core.require(group_key in groups, f"predecessor group is missing: {group_key}")
        original = _load_json(_decision_path(work, run_group))
        _require_controlled_semantics(run_group, original)
        core.require(original.get("sealed") is True, f"predecessor group is not sealed: {group_key}")
        core.require(
            original.get("sealed_decision_fingerprint")
            == _semantic_fingerprint(original),
            f"predecessor sealed decision changed: {group_key}",
        )
        decision = copy.deepcopy(original)
        _validate_decision(
            normalized,
            groups[group_key],
            decision,
            require_complete=True,
            capture_hashes=False,
            rehash_labels=set(),
        )
        decisions[group_key] = decision
    return decisions


def _render_predecessor_for_comparison(
    previous_work: Path,
    *,
    template: Path,
    decisions: Mapping[str, Mapping[str, Any]] | None = None,
) -> Path:
    temporary_root = Path(
        tempfile.mkdtemp(prefix="roll-forward-compare-", dir=previous_work.parent)
    )
    clone = temporary_root / "work"
    clone.mkdir()
    shutil.copy2(_run_path(previous_work), _run_path(clone))
    shutil.copytree(previous_work / "decisions", clone / "decisions")
    (clone / "snapshot").mkdir()
    shutil.copy2(_snapshot_path(previous_work), _snapshot_path(clone))
    if decisions is not None:
        clone_run = _load_run(clone)
        run_groups = {
            core.clean_text(item.get("group_key")): item
            for item in clone_run.get("groups", [])
            if isinstance(item, Mapping)
        }
        core.require(
            set(decisions) == set(run_groups),
            "amended predecessor decisions do not match the published groups",
        )
        for group_key, source_decision in decisions.items():
            candidate = copy.deepcopy(dict(source_decision))
            candidate["sealed"] = True
            candidate["sealed_decision_fingerprint"] = None
            candidate["edit_control"] = _new_edit_control(
                _semantic_fingerprint(candidate)
            )
            candidate["sealed_decision_fingerprint"] = _semantic_fingerprint(
                candidate
            )
            core.atomic_json(
                _decision_path(clone, run_groups[group_key]), candidate
            )
    output = temporary_root / "expected.xlsx"
    try:
        finish_run(
            SimpleNamespace(
                work=clone,
                output=output,
                template=template,
                replace_predecessor=False,
            )
        )
    except Exception:
        shutil.rmtree(temporary_root, ignore_errors=True)
        raise
    return output


def _predecessor_workbook_drift(
    previous_work: Path,
    previous_run: Mapping[str, Any],
    previous_output: Path,
) -> tuple[bool, str, str]:
    core.require(previous_output.is_file(), f"predecessor workbook is missing: {previous_output}")
    registered = Path(str(previous_run.get("output") or "")).resolve()
    core.require(
        registered == previous_output,
        "--previous-output must be the workbook published by the predecessor run",
    )
    actual_hash = core.sha256_file(previous_output)
    saved_hash = core.clean_text(previous_run.get("output_sha256"))
    saved_business = core.clean_text(
        previous_run.get("output_business_fingerprint")
    )
    actual_business = _workbook_business_fingerprint(previous_output)
    if saved_hash and actual_hash == saved_hash:
        return False, actual_hash, actual_business
    if saved_business:
        return actual_business != saved_business, actual_hash, actual_business

    expected = _render_predecessor_for_comparison(
        previous_work,
        template=Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
    )
    temporary_root = expected.parent
    try:
        legacy = _workbook_uses_legacy_order_layout(previous_output)
        expected_business = _workbook_business_fingerprint(
            expected, legacy=legacy
        )
        actual_legacy_business = _workbook_business_fingerprint(
            previous_output, legacy=legacy
        )
        return (
            actual_legacy_business != expected_business,
            actual_hash,
            actual_business,
        )
    finally:
        shutil.rmtree(temporary_root, ignore_errors=True)


def _require_amendments_rebuild_workbook(
    previous_work: Path,
    previous_run: Mapping[str, Any],
    previous_output: Path,
    decisions: Mapping[str, Mapping[str, Any]],
) -> None:
    expected = _render_predecessor_for_comparison(
        previous_work,
        template=Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
        decisions=decisions,
    )
    temporary_root = expected.parent
    try:
        legacy = (
            not bool(
                core.clean_text(previous_run.get("output_business_fingerprint"))
            )
            and _workbook_uses_legacy_order_layout(previous_output)
        )
        expected_business = _workbook_business_fingerprint(
            expected, legacy=legacy
        )
        actual_business = _workbook_business_fingerprint(
            previous_output, legacy=legacy
        )
        core.require(
            expected_business == actual_business,
            "confirmed amendments do not fully explain the predecessor workbook business-cell changes",
        )
    finally:
        shutil.rmtree(temporary_root, ignore_errors=True)


def _set_amendment_field(target: dict[str, Any], field_path: str, value: Any) -> None:
    parts = [part for part in field_path.split(".") if part]
    core.require(bool(parts), "confirmed amendment field is required")
    current: Any = target
    for part in parts[:-1]:
        if isinstance(current, list):
            core.require(part.isdigit(), "list amendment path segment must be an index")
            position = int(part)
            core.require(0 <= position < len(current), "amendment list index is outside the object")
            current = current[position]
        else:
            core.require(isinstance(current, dict) and part in current, f"unknown amendment path: {field_path}")
            current = current[part]
    final = parts[-1]
    if isinstance(current, list):
        core.require(final.isdigit(), "list amendment path segment must be an index")
        position = int(final)
        core.require(0 <= position < len(current), "amendment list index is outside the object")
        current[position] = copy.deepcopy(value)
    else:
        core.require(isinstance(current, dict), f"invalid amendment path: {field_path}")
        current[final] = copy.deepcopy(value)


def _apply_confirmed_amendments(
    path: Path | None,
    *,
    previous_output_sha256: str,
    normalized: Mapping[str, Any],
    decisions: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    if path is None:
        return {"count": 0, "sha256": None, "path": None}
    resolved = path.resolve()
    payload = _load_json(resolved)
    core.require(
        payload.get("contract_version") == AMENDMENTS_CONTRACT,
        f"unsupported confirmed amendments contract; expected {AMENDMENTS_CONTRACT}",
    )
    core.require(
        core.clean_text(payload.get("previous_output_sha256"))
        == previous_output_sha256,
        "confirmed amendments are not bound to the current predecessor workbook hash",
    )
    amendments = payload.get("amendments")
    core.require(isinstance(amendments, list) and amendments, "confirmed amendments must be nonempty")
    groups = _group_index(normalized)
    allowed_roots = {
        "orders": {
            "customer_nickname",
            "direction",
            "pricing",
            "note",
            "fund_type",
            "lifecycle",
            "legs",
        },
        "records": {
            "record_type",
            "posting_status",
            "posting_at",
            "voucher_number",
            "voucher_date",
            "source_messages",
            "media_roles",
            "representative_media_label",
            "movements",
            "transfer",
            "status_reason",
            "note",
        },
    }
    applied: list[dict[str, str]] = []
    for position, amendment in enumerate(amendments):
        field = f"confirmed amendments[{position}]"
        core.require(isinstance(amendment, Mapping), f"{field} must be an object")
        unknown = sorted(
            set(amendment)
            - {"group_key", "collection", "id", "field", "value", "reason", "basis", "source_messages"}
        )
        core.require(not unknown, f"{field} has unsupported fields: {', '.join(unknown)}")
        group_key = core.clean_text(amendment.get("group_key"))
        collection = core.clean_text(amendment.get("collection")).casefold()
        item_id = core.clean_text(amendment.get("id"))
        field_path = core.clean_text(amendment.get("field"))
        reason = core.clean_text(amendment.get("reason"))
        basis = core.clean_text(amendment.get("basis"))
        core.require(group_key in decisions and group_key in groups, f"{field}.group_key is unknown")
        core.require(collection in allowed_roots, f"{field}.collection must be orders or records")
        core.require(bool(item_id) and bool(reason) and bool(basis), f"{field} requires id, reason, and basis")
        root = field_path.split(".", 1)[0]
        core.require(root in allowed_roots[collection], f"{field}.field is not amendable")
        source_value = amendment.get("source_messages")
        core.require(isinstance(source_value, list) and source_value, f"{field}.source_messages is required")
        source_labels = [str(item) for item in source_value]
        _, message_by_label = _message_labels(groups[group_key])
        core.require(
            len(source_labels) == len(set(source_labels))
            and set(source_labels) <= set(message_by_label),
            f"{field}.source_messages contains an unknown predecessor message",
        )
        collection_value = decisions[group_key].get(collection)
        core.require(isinstance(collection_value, list), f"{field}.collection is unavailable for this mode")
        matches = [
            item
            for item in collection_value
            if isinstance(item, dict) and core.clean_text(item.get("id")) == item_id
        ]
        core.require(len(matches) == 1, f"{field}.id must identify exactly one object")
        _set_amendment_field(matches[0], field_path, amendment.get("value"))
        applied.append(
            {
                "group_key": group_key,
                "collection": collection,
                "id": item_id,
                "field": field_path,
                "reason": reason,
                "basis": basis,
            }
        )
    return {
        "count": len(applied),
        "sha256": core.sha256_file(resolved),
        "path": str(resolved),
        "items": applied,
    }


def _relocate_snapshot_paths(
    normalized: dict[str, Any], old_root: Path, new_root: Path
) -> None:
    old_resolved = old_root.resolve()
    for group in normalized.get("groups", []):
        for message in group.get("messages", []):
            for media in message.get("media", []):
                if not isinstance(media, dict) or not media.get("path"):
                    continue
                path = Path(str(media["path"])).resolve()
                try:
                    relative = path.relative_to(old_resolved)
                except ValueError:
                    continue
                media["path"] = str((new_root / relative).resolve())


def _label_maps(
    previous_group: Mapping[str, Any], cumulative_group: Mapping[str, Any]
) -> tuple[dict[str, str], dict[str, str], set[str]]:
    _, previous_messages = _message_labels(previous_group)
    _, cumulative_messages = _message_labels(cumulative_group)
    cumulative_message_labels = {
        core.clean_text(message.get("stable_message_id")): label
        for label, message in cumulative_messages.items()
    }
    message_map: dict[str, str] = {}
    for label, message in previous_messages.items():
        stable_id = core.clean_text(message.get("stable_message_id"))
        core.require(stable_id in cumulative_message_labels, "predecessor message disappeared during merge")
        message_map[label] = cumulative_message_labels[stable_id]

    previous_media = _media_inventory(previous_group)
    cumulative_media = _media_inventory(cumulative_group)
    cumulative_media_by_stable = {
        core.clean_text(media.get("stable_media_id")): (label, media)
        for label, (_, media) in cumulative_media.items()
    }
    media_map: dict[str, str] = {}
    reusable_media_labels: set[str] = set()
    for label, (_, media) in previous_media.items():
        stable_id = core.clean_text(media.get("stable_media_id"))
        core.require(stable_id in cumulative_media_by_stable, "predecessor media disappeared during merge")
        cumulative_label, cumulative_item = cumulative_media_by_stable[stable_id]
        media_map[label] = cumulative_label
        previous_hash = core.clean_text(media.get("blob_sha256"))
        cumulative_hash = core.clean_text(cumulative_item.get("blob_sha256"))
        if previous_hash and previous_hash == cumulative_hash:
            reusable_media_labels.add(cumulative_label)
    return message_map, media_map, reusable_media_labels


def _remap_decision_labels(
    value: Any,
    *,
    message_map: Mapping[str, str],
    media_map: Mapping[str, str],
) -> Any:
    if isinstance(value, list):
        return [
            _remap_decision_labels(
                item, message_map=message_map, media_map=media_map
            )
            for item in value
        ]
    if isinstance(value, dict):
        result: dict[str, Any] = {}
        for key, item in value.items():
            mapped_key = media_map.get(key, message_map.get(key, key))
            result[mapped_key] = _remap_decision_labels(
                item, message_map=message_map, media_map=media_map
            )
        return result
    if isinstance(value, str):
        if value in message_map:
            return message_map[value]
        if value in media_map:
            return media_map[value]
        match = re.fullmatch(r"(M\d{4})(\.\d+)", value)
        if match and match.group(1) in media_map:
            return media_map[match.group(1)] + match.group(2)
    return copy.deepcopy(value)


def _carry_ids(decision: Mapping[str, Any]) -> list[str]:
    if decision.get("contract_version") == store_ledger.DECISION_CONTRACT:
        result = [
            core.clean_text(item.get("id"))
            for item in decision.get("records", [])
            if isinstance(item, Mapping)
            and (
                item.get("posting_status") != "posted"
                or bool(core.clean_text(item.get("status_reason")))
                or bool(core.clean_text(item.get("note")))
                or item.get("record_type") == "internal_transfer"
            )
        ]
        closing_snapshots = [
            item
            for item in decision.get("balance_snapshots", [])
            if isinstance(item, Mapping) and item.get("kind") == "closing"
        ]
        if closing_snapshots:
            latest_date = max(
                core.clean_text(item.get("date")) for item in closing_snapshots
            )
            result.extend(
                core.clean_text(item.get("id"))
                for item in closing_snapshots
                if core.clean_text(item.get("date")) == latest_date
            )
        return list(dict.fromkeys(item for item in result if item))

    def pricing_unknown(value: object) -> bool:
        if not isinstance(value, Mapping):
            return True
        expected = value.get("expected")
        return not isinstance(expected, Mapping) or core.clean_text(
            expected.get("kind")
        ).casefold() == "unknown"

    result: list[str] = []
    for item in decision.get("orders", []):
        if not isinstance(item, Mapping):
            continue
        lifecycle = item.get("lifecycle")
        legs = item.get("legs")
        exceptional = (
            not isinstance(lifecycle, Mapping)
            or lifecycle.get("status") != "completed"
            or bool(core.clean_text(item.get("note")))
            or not bool(core.clean_text(item.get("direction")))
            or (
                any(
                    pricing_unknown(leg.get("pricing"))
                    for leg in legs
                    if isinstance(leg, Mapping)
                )
                if isinstance(legs, list) and legs
                else pricing_unknown(item.get("pricing"))
            )
        )
        if exceptional:
            result.append(core.clean_text(item.get("id")))
    return result


def _migrate_legacy_large_decision(
    normalized: Mapping[str, Any],
    group: Mapping[str, Any],
    legacy: Mapping[str, Any],
) -> dict[str, Any]:
    """Convert a sealed exchange-only decision into reviewable full orders."""

    core.require(
        legacy.get("contract_version") == large_daily.LEGACY_DECISION_CONTRACT,
        "legacy large migration received the wrong contract",
    )
    core.require(
        legacy.get("open_exchanges") in (None, []),
        "legacy large run still has open exchanges and cannot be rolled forward",
    )
    _, message_by_label = _message_labels(group)
    orders: list[dict[str, Any]] = []
    for position, raw_exchange in enumerate(legacy.get("exchanges", [])):
        field = f"legacy large exchanges[{position}]"
        core.require(isinstance(raw_exchange, Mapping), f"{field} must be an object")
        entry_ids = [str(item) for item in raw_exchange.get("entry_ids", [])]
        core.require(
            bool(entry_ids),
            f"{field} has no source fund entries; regenerate this legacy group with the full-order contract before roll-forward",
        )
        sources = [str(item) for item in raw_exchange.get("source_messages", [])]
        customer_names = list(
            dict.fromkeys(
                core.clean_text(
                    message_by_label[label].get("sender_name")
                    or message_by_label[label].get("sender_id")
                )
                for label in sources
                if label in message_by_label
                and message_by_label[label].get("role") == "客户候选"
                and core.clean_text(
                    message_by_label[label].get("sender_name")
                    or message_by_label[label].get("sender_id")
                )
            )
        )
        migration_note = "旧大额合同迁移：需结合前序聊天复核客户和完整订单边界。"
        original_note = core.clean_text(raw_exchange.get("note"))
        orders.append(
            {
                "id": core.clean_text(raw_exchange.get("id")),
                "entry_ids": entry_ids,
                "source_messages": sources,
                "customer_nickname": (
                    customer_names[0] if len(customer_names) == 1 else ""
                ),
                "direction": core.clean_text(raw_exchange.get("direction")),
                "pricing": {
                    "source_messages": sources,
                    "terms": {
                        "rate": raw_exchange.get("rate"),
                        "operator": raw_exchange.get("operator"),
                    },
                    "expected": {
                        "kind": "explicit",
                        "amount": raw_exchange.get("target_amount"),
                    },
                },
                "fund_type": raw_exchange.get("fund_type"),
                "note": "\n".join(
                    item for item in (original_note, migration_note) if item
                ),
            }
        )

    candidate = copy.deepcopy(dict(legacy))
    candidate["contract_version"] = large_daily.DECISION_CONTRACT
    candidate["group_mode"] = "large"
    candidate["orders"] = orders
    candidate["open_orders"] = []
    candidate["balance_links"] = []
    candidate["settlement_allocations"] = []
    candidate.setdefault("unknown_payee_reviewed_entry_ids", [])
    candidate.pop("exchanges", None)
    candidate.pop("open_exchanges", None)
    candidate["sealed"] = False
    candidate["sealed_decision_fingerprint"] = None
    candidate["edit_control"] = _new_edit_control(_semantic_fingerprint(candidate))
    _validate_decision(
        normalized,
        group,
        candidate,
        require_complete=True,
        capture_hashes=False,
        rehash_labels=set(),
    )
    candidate["edit_control"] = _new_edit_control(_semantic_fingerprint(candidate))
    candidate["sealed"] = True
    candidate["sealed_decision_fingerprint"] = _semantic_fingerprint(candidate)
    return candidate


def _migrate_decision(
    normalized: Mapping[str, Any],
    previous_group: Mapping[str, Any],
    cumulative_group: Mapping[str, Any],
    previous_decision: Mapping[str, Any],
    *,
    group_mode: str,
) -> tuple[dict[str, Any], list[str], int]:
    message_map, media_map, reusable_media_labels = _label_maps(
        previous_group, cumulative_group
    )
    candidate = _remap_decision_labels(
        previous_decision, message_map=message_map, media_map=media_map
    )
    core.require(isinstance(candidate, dict), "migrated decision must be an object")
    media_decisions = candidate.get("media_decisions")
    if isinstance(media_decisions, dict):
        for label in list(media_decisions):
            if label not in reusable_media_labels:
                media_decisions.pop(label)
    expected = _decision_template(normalized, cumulative_group, group_mode=group_mode)
    for field in (
        "normalized_source_fingerprint",
        "group_fingerprint",
        "group_key",
        "group_name",
        "platform",
        "message_count",
        "evidence_media_count",
        "group_mode",
    ):
        if field in expected:
            candidate[field] = copy.deepcopy(expected[field])
        else:
            candidate.pop(field, None)
    previous_stable_ids = [
        core.clean_text(message.get("stable_message_id"))
        for message in previous_group.get("messages", [])
    ]
    cumulative_stable_ids = [
        core.clean_text(message.get("stable_message_id"))
        for message in cumulative_group.get("messages", [])
    ]
    core.require(
        cumulative_stable_ids[: len(previous_stable_ids)] == previous_stable_ids,
        "predecessor chronology is not a stable prefix after merging snapshots",
    )
    reviewed_through = len(previous_stable_ids)
    candidate["reviewed_through"] = reviewed_through
    candidate["read_complete"] = reviewed_through == len(cumulative_stable_ids)
    candidate["sealed"] = False
    candidate["sealed_decision_fingerprint"] = None
    candidate["edit_control"] = _new_edit_control(_semantic_fingerprint(candidate))
    _validate_decision(
        normalized,
        cumulative_group,
        candidate,
        require_complete=False,
        capture_hashes=False,
        rehash_labels=set(),
    )
    candidate["edit_control"] = _new_edit_control(_semantic_fingerprint(candidate))
    carry = _carry_ids(candidate)
    new_count = len(cumulative_stable_ids) - reviewed_through
    if not carry and new_count == 0:
        candidate["sealed"] = True
        candidate["sealed_decision_fingerprint"] = _semantic_fingerprint(candidate)
    return candidate, carry, new_count


def _cumulative_accounting_window(
    predecessor: Mapping[str, Any], incoming: Mapping[str, Any]
) -> tuple[dict[str, str | bool], list[dict[str, Any]]]:
    periods: list[dict[str, Any]] = []
    previous_periods = predecessor.get("accounting_periods")
    if isinstance(previous_periods, list):
        periods.extend(copy.deepcopy(previous_periods))
    elif isinstance(predecessor.get("accounting_window"), Mapping):
        periods.append(copy.deepcopy(dict(predecessor["accounting_window"])))
    else:
        predecessor_timestamps = [
            core.clean_text(message.get("timestamp"))
            for group in predecessor.get("groups", [])
            for message in group.get("messages", [])
            if isinstance(message, Mapping)
            and not message.get("accounting_context_only")
        ]
        if predecessor_timestamps:
            zone = ZoneInfo("Asia/Bangkok")
            predecessor_start_date = min(
                datetime.fromisoformat(value).astimezone(zone).date()
                for value in predecessor_timestamps
            )
            predecessor_end_date = max(
                datetime.fromisoformat(value).astimezone(zone).date()
                for value in predecessor_timestamps
            ) + timedelta(days=1)
            periods.append(
                {
                    "start": datetime.combine(
                        predecessor_start_date, datetime.min.time(), tzinfo=zone
                    ).isoformat(),
                    "end": datetime.combine(
                        predecessor_end_date, datetime.min.time(), tzinfo=zone
                    ).isoformat(),
                    "reply_context_preserved": bool(
                        predecessor.get("accounting_window", {}).get(
                            "reply_context_preserved", False
                        )
                    )
                    if isinstance(predecessor.get("accounting_window"), Mapping)
                    else False,
                    "inferred": True,
                }
            )
    if isinstance(incoming.get("accounting_window"), Mapping):
        periods.append(copy.deepcopy(dict(incoming["accounting_window"])))
    starts = [core.clean_text(item.get("start")) for item in periods if item.get("start")]
    ends = [core.clean_text(item.get("end")) for item in periods if item.get("end")]
    if not starts or not ends:
        timestamps = [
            core.clean_text(message.get("timestamp"))
            for document in (predecessor, incoming)
            for group in document.get("groups", [])
            for message in group.get("messages", [])
            if isinstance(message, Mapping)
        ]
        core.require(bool(timestamps), "cumulative accounting window has no messages")
        zone = ZoneInfo("Asia/Bangkok")
        start_date = min(datetime.fromisoformat(value).astimezone(zone).date() for value in timestamps)
        end_date = max(datetime.fromisoformat(value).astimezone(zone).date() for value in timestamps) + timedelta(days=1)
        starts = [datetime.combine(start_date, datetime.min.time(), tzinfo=zone).isoformat()]
        ends = [datetime.combine(end_date, datetime.min.time(), tzinfo=zone).isoformat()]
    return (
        {
            "start": min(starts),
            "end": max(ends),
            "reply_context_preserved": True,
        },
        periods,
    )


def roll_forward_run(args: argparse.Namespace) -> dict[str, Any]:
    previous_work = args.previous_work.resolve()
    work = args.work.resolve()
    previous_output = args.previous_output.resolve()
    core.require(previous_work.is_dir(), f"predecessor work directory is missing: {previous_work}")
    core.require(not work.exists(), f"work directory already exists; roll-forward requires a new path: {work}")
    inputs = [path.resolve() for path in args.inputs]
    for path in inputs:
        core.require(path.exists(), f"input does not exist: {path}")

    previous_run = _load_run(previous_work)
    group_mode = core.clean_text(previous_run.get("group_mode")).casefold()
    core.require(group_mode != finance_materials.GROUP_MODE, "finance mode does not support roll-forward")
    core.require(
        group_mode in {"small", "large", store_ledger.GROUP_MODE},
        "predecessor mode does not support roll-forward",
    )
    previous_normalized = _load_snapshot(previous_work, previous_run)
    predecessor_roster = Path(
        str(
            previous_run.get("roster_path")
            or Path(__file__).resolve().parent.parent / "config" / "roster.yaml"
        )
    ).resolve()
    core.require(
        predecessor_roster.is_file(),
        f"predecessor roster is missing: {predecessor_roster}",
    )
    predecessor_roster_sha256 = core.clean_text(
        previous_run.get("roster_sha256")
    )
    if predecessor_roster_sha256:
        core.require(
            core.sha256_file(predecessor_roster)
            == predecessor_roster_sha256,
            "predecessor roster changed; restore the recorded parsing settings before roll-forward",
        )
    previous_decisions = _load_finished_decisions(
        previous_work, previous_run, previous_normalized
    )
    period = _successor_period(previous_run, previous_normalized, args.date)
    drifted, previous_output_sha256, previous_business = _predecessor_workbook_drift(
        previous_work, previous_run, previous_output
    )
    amendments = _apply_confirmed_amendments(
        args.confirmed_amendments,
        previous_output_sha256=previous_output_sha256,
        normalized=previous_normalized,
        decisions=previous_decisions,
    )
    core.require(
        not drifted or amendments["count"] > 0,
        "predecessor workbook business cells changed; provide --confirmed-amendments bound to its current hash",
    )
    core.require(
        drifted or amendments["count"] == 0,
        "confirmed amendments were provided but the predecessor workbook has no business-cell changes",
    )
    previous_groups = _group_index(previous_normalized)
    for group_key, decision in previous_decisions.items():
        _validate_decision(
            previous_normalized,
            previous_groups[group_key],
            decision,
            require_complete=True,
            capture_hashes=False,
            rehash_labels=set(),
        )
    if drifted:
        _require_amendments_rebuild_workbook(
            previous_work,
            previous_run,
            previous_output,
            previous_decisions,
        )
    legacy_migrations = 0
    if group_mode == "large":
        for group_key, decision in list(previous_decisions.items()):
            if decision.get("contract_version") != large_daily.LEGACY_DECISION_CONTRACT:
                continue
            previous_decisions[group_key] = _migrate_legacy_large_decision(
                previous_normalized,
                previous_groups[group_key],
                decision,
            )
            legacy_migrations += 1
    _assign_stable_identities(previous_normalized)
    previous_groups = _group_index(previous_normalized)

    work.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="roll-forward-", dir=work.parent) as temporary_name:
        temporary_root = Path(temporary_name)
        staging = temporary_root / "work"
        prior_ocr = _ocr_candidate_config_for_run(previous_run)
        start_run(
            SimpleNamespace(
                work=staging,
                inputs=inputs,
                mode=group_mode,
                contains=previous_run.get("group_name_contains", "小额"),
                date=(
                    period["date"]
                    if group_mode == "large" or not period["accounting_from"]
                    else None
                ),
                accounting_from=period["accounting_from"],
                accounting_to=period["accounting_to"],
                timezone=previous_run.get("accounting_timezone", "Asia/Bangkok"),
                roster=predecessor_roster,
                line_backup=[],
                line_android_backup=[],
                line_self_name=previous_run.get("line_self_name", "LINE_SELF"),
                ocr_candidates=bool(prior_ocr["enabled"]),
            )
        )
        incoming_run = _load_run(staging)
        incoming = _load_snapshot(staging, incoming_run)
        incoming_run_sha256 = core.sha256_file(_run_path(staging))
        incoming_snapshot_sha256 = core.sha256_file(_snapshot_path(staging))
        incoming_source_fingerprint = core.clean_text(
            incoming.get("source_fingerprint")
        )
        _relocate_snapshot_paths(incoming, staging, work)
        cumulative, merge_statistics = _merge_normalized_snapshots(
            previous_normalized, incoming
        )
        cumulative_window, accounting_periods = _cumulative_accounting_window(
            previous_normalized, incoming
        )
        cumulative["accounting_window"] = cumulative_window
        cumulative["accounting_periods"] = accounting_periods

        cumulative_groups = _group_index(cumulative)
        group_reports: list[dict[str, Any]] = []
        used_labels: set[str] = set()
        for group in cumulative["groups"]:
            group_key = core.clean_text(group.get("group_key"))
            if group_key in previous_decisions:
                decision, carry, new_count = _migrate_decision(
                    cumulative,
                    previous_groups[group_key],
                    group,
                    previous_decisions[group_key],
                    group_mode=group_mode,
                )
            else:
                decision = _decision_template(cumulative, group, group_mode=group_mode)
                carry = []
                new_count = len(group.get("messages", []))
            filename = _decision_filename(group_key)
            core.atomic_json(staging / "decisions" / filename, decision)
            base_label = f"[{_platform_prefix(group.get('platform'))}] {group.get('group_name')}"
            run_label = base_label
            if run_label in used_labels:
                run_label = f"{base_label} ({core.stable_token(group_key, length=6)})"
            used_labels.add(run_label)
            group_reports.append(
                {
                    "group_key": group_key,
                    "group_name": group.get("group_name"),
                    "platform": group.get("platform"),
                    "run_label": run_label,
                    "messages": len(group.get("messages", [])),
                    "evidence_media": len(_media_inventory(group)),
                    "decision_file": f"decisions/{filename}",
                    "edit_mode": EDIT_CONTROL_MODE,
                    "group_mode": group_mode,
                    "sealed": bool(decision.get("sealed")),
                    "roll_forward_carry": carry,
                    "new_messages": new_count,
                }
            )

        snapshot_path = _snapshot_path(staging)
        core.atomic_json(snapshot_path, cumulative)
        core.load_normalized(snapshot_path)
        run = copy.deepcopy(incoming_run)
        run.update(
            {
                "status": "reviewing",
                "groups": group_reports,
                "normalized_source_fingerprint": cumulative["source_fingerprint"],
                "snapshot_sha256": core.sha256_file(snapshot_path),
                "accounting_date": period["date"],
                "input_roots": [str(path) for path in inputs],
                "roll_forward": {
                    "contract_version": ROLL_FORWARD_CONTRACT,
                    "predecessor_work": str(previous_work),
                    "predecessor_run_sha256": core.sha256_file(_run_path(previous_work)),
                    "predecessor_snapshot_sha256": core.sha256_file(
                        _snapshot_path(previous_work)
                    ),
                    "predecessor_output": str(previous_output),
                    "predecessor_output_sha256": previous_output_sha256,
                    "predecessor_output_business_fingerprint": previous_business,
                    "incoming_run_sha256": incoming_run_sha256,
                    "incoming_snapshot_sha256": incoming_snapshot_sha256,
                    "incoming_source_fingerprint": incoming_source_fingerprint,
                    "incoming_input_roots": [str(path) for path in inputs],
                    "inherited_parsing_settings": {
                        "group_mode": group_mode,
                        "timezone": previous_run.get(
                            "accounting_timezone", "Asia/Bangkok"
                        ),
                        "roster_path": str(predecessor_roster),
                        "roster_sha256": core.sha256_file(predecessor_roster),
                        "line_self_name": previous_run.get(
                            "line_self_name", "LINE_SELF"
                        ),
                        "ocr_candidates": bool(prior_ocr["enabled"]),
                    },
                    "current_date": period["date"],
                    "current_accounting_from": period["accounting_from"],
                    "current_accounting_to": period["accounting_to"],
                    "merge": merge_statistics,
                    "confirmed_amendments": amendments,
                    "legacy_contract_migrations": legacy_migrations,
                },
            }
        )
        run.pop("finished_at", None)
        run.pop("output", None)
        run.pop("output_sha256", None)
        run.pop("output_business_fingerprint", None)
        if period["accounting_from"]:
            run["accounting_from"] = period["accounting_from"]
            run["accounting_to"] = period["accounting_to"]
        else:
            run.pop("accounting_from", None)
            run.pop("accounting_to", None)
        core.atomic_json(_run_path(staging), run)
        staging.replace(work)

    return {
        "work": str(work),
        "mode": group_mode,
        "date": period["date"],
        "groups": len(group_reports),
        "messages": cumulative["statistics"]["messages"],
        "carry_objects": sum(len(item["roll_forward_carry"]) for item in group_reports),
        "workbook_business_drift_confirmed": bool(drifted),
        "confirmed_amendments": amendments["count"],
        "legacy_contract_migrations": legacy_migrations,
        **merge_statistics,
    }


def _business_cell_value(value: object) -> object:
    if isinstance(value, datetime):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _workbook_uses_legacy_order_layout(path: Path) -> bool:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for worksheet in workbook.worksheets:
            headers = {
                core.clean_text(cell.value)
                for cell in next(worksheet.iter_rows(min_row=1, max_row=1), ())
                if core.clean_text(cell.value)
            }
            if {"记录类型", "订单编号"} <= headers:
                return not {"订单状态", "完成时间"} <= headers
        return False
    finally:
        workbook.close()


def _workbook_business_material(path: Path, *, legacy: bool = False) -> dict[str, Any]:
    """Read semantic workbook content while deliberately ignoring formatting."""

    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        sheets: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            first_headers = [
                core.clean_text(worksheet.cell(1, column).value)
                for column in range(1, worksheet.max_column + 1)
            ]
            is_order_sheet = "记录类型" in first_headers and "订单编号" in first_headers
            if is_order_sheet:
                title_row = next(
                    (
                        row
                        for row in range(2, worksheet.max_row + 1)
                        if worksheet.cell(row, 1).value == "资金汇总"
                    ),
                    None,
                )
                detail_end = (title_row - 1) if title_row else worksheet.max_row
                allowed_headers = [header for header in first_headers if header]
                if legacy:
                    allowed_headers = [
                        header
                        for header in allowed_headers
                        if header not in {"订单状态", "完成时间"}
                    ]
                header_columns = {
                    header: first_headers.index(header) + 1 for header in allowed_headers
                }
                detail_rows = []
                for row in range(2, detail_end + 1):
                    material = {
                        header: _business_cell_value(
                            worksheet.cell(row, column).value
                        )
                        for header, column in header_columns.items()
                    }
                    if any(value not in (None, "") for value in material.values()):
                        detail_rows.append(material)
                summary_rows: list[dict[str, Any]] = []
                if title_row:
                    summary_header_row = title_row + 1
                    summary_headers = [
                        core.clean_text(worksheet.cell(summary_header_row, column).value)
                        for column in range(1, worksheet.max_column + 1)
                    ]
                    allowed_summary_headers = [
                        header for header in summary_headers if header
                    ]
                    if legacy:
                        allowed_summary_headers = [
                            header
                            for header in allowed_summary_headers
                            if header != "统计日期"
                        ]
                    summary_columns = {
                        header: summary_headers.index(header) + 1
                        for header in allowed_summary_headers
                    }
                    for row in range(summary_header_row + 1, worksheet.max_row + 1):
                        material = {}
                        for header, column in summary_columns.items():
                            cell_value = _business_cell_value(
                                worksheet.cell(row, column).value
                            )
                            if legacy and header == "状态":
                                cell_value = {
                                    "已完成": "已确认",
                                    "待次日继续": "待确认",
                                    "已取消": "待确认",
                                }.get(cell_value, cell_value)
                            if legacy and header in {"换出合计", "换入合计"} and cell_value == "—":
                                cell_value = "未确认"
                            material[header] = cell_value
                        if any(value not in (None, "") for value in material.values()):
                            summary_rows.append(material)
                sheet_material: dict[str, Any] = {
                    "title": worksheet.title,
                    "detail_headers": allowed_headers,
                    "detail_rows": detail_rows,
                    "summary_rows": summary_rows,
                }
            else:
                last_row = max(
                    (
                        row
                        for row in range(1, worksheet.max_row + 1)
                        if any(
                            worksheet.cell(row, column).value not in (None, "")
                            for column in range(1, worksheet.max_column + 1)
                        )
                    ),
                    default=0,
                )
                last_column = max(
                    (
                        column
                        for column in range(1, worksheet.max_column + 1)
                        if any(
                            worksheet.cell(row, column).value not in (None, "")
                            for row in range(1, last_row + 1)
                        )
                    ),
                    default=0,
                )
                sheet_material = {
                    "title": worksheet.title,
                    "cells": [
                        [
                            _business_cell_value(worksheet.cell(row, column).value)
                            for column in range(1, last_column + 1)
                        ]
                        for row in range(1, last_row + 1)
                    ],
                }
            sheets.append(sheet_material)
        return {"sheets": sheets}
    finally:
        workbook.close()


def _workbook_business_fingerprint(path: Path, *, legacy: bool = False) -> str:
    return core.fingerprint_json(_workbook_business_material(path, legacy=legacy))


def _protected_replace_enabled(args: argparse.Namespace) -> bool:
    return bool(getattr(args, "replace_predecessor", False))


def _validate_finish_destination(
    run: Mapping[str, Any], output: Path, *, replace_predecessor: bool
) -> None:
    if not replace_predecessor:
        core.require(not output.exists(), f"output already exists; choose a new file: {output}")
        return
    lineage = run.get("roll_forward")
    core.require(
        isinstance(lineage, Mapping),
        "--replace-predecessor is only available for a roll-forward work directory",
    )
    predecessor_output = Path(
        str(lineage.get("predecessor_output") or "")
    ).resolve()
    core.require(
        output == predecessor_output,
        "--replace-predecessor may only target the predecessor workbook registered in run.json",
    )
    core.require(output.is_file(), f"registered predecessor workbook is missing: {output}")
    expected_hash = core.clean_text(lineage.get("predecessor_output_sha256"))
    core.require(
        core.sha256_file(output) == expected_hash,
        "predecessor workbook changed after roll-forward started; rebuild the successor run",
    )


def _publish_finished_workbook(
    temporary_workbook: Path,
    output: Path,
    *,
    run: Mapping[str, Any],
    replace_predecessor: bool,
) -> None:
    _validate_finish_destination(
        run, output, replace_predecessor=replace_predecessor
    )
    temporary_workbook.replace(output)


def _record_finished_output(run: dict[str, Any], work: Path, output: Path) -> None:
    run["status"] = "finished"
    run["finished_at"] = datetime.now(timezone.utc).isoformat()
    run["output"] = str(output)
    run["output_sha256"] = core.sha256_file(output)
    run["output_business_fingerprint"] = _workbook_business_fingerprint(output)
    core.atomic_json(_run_path(work), run)


def finish_run(args: argparse.Namespace) -> dict[str, Any]:
    work = args.work.resolve()
    output = args.output.resolve()
    run = _load_run(work)
    replace_predecessor = _protected_replace_enabled(args)
    _validate_finish_destination(
        run, output, replace_predecessor=replace_predecessor
    )
    normalized = _load_snapshot(work, run)
    groups = _group_index(normalized)
    decisions: dict[str, Mapping[str, Any]] = {}
    review_statistics: dict[str, dict[str, Any]] = {}
    finish_risk_reports: list[dict[str, Any]] = []
    for run_group in run["groups"]:
        group_key = str(run_group["group_key"])
        decision = _load_json(_decision_path(work, run_group))
        _require_controlled_semantics(run_group, decision)
        core.require(decision.get("sealed") is True, f"group is not sealed: {run_group.get('run_label')}")
        current_fingerprint = _semantic_fingerprint(decision)
        core.require(
            decision.get("sealed_decision_fingerprint") == current_fingerprint,
            f"sealed decision changed; review and seal again: {run_group.get('run_label')}",
        )
        statistics = _validate_decision(
            normalized,
            groups[group_key],
            decision,
            require_complete=True,
            capture_hashes=False,
        )
        decisions[group_key] = decision
        review_statistics[group_key] = statistics
        finish_risk_reports.append(
            {
                "group_key": group_key,
                "run_label": run_group.get("run_label"),
                **statistics,
            }
        )

    group_mode = core.clean_text(run.get("group_mode")).casefold()
    finance_run = group_mode == finance_materials.GROUP_MODE
    store_run = group_mode == store_ledger.GROUP_MODE
    large_run = group_mode == "large"
    decision_contracts = {
        core.clean_text(decision.get("contract_version"))
        for decision in decisions.values()
    }
    if store_run:
        core.require(
            decision_contracts == {store_ledger.DECISION_CONTRACT},
            "store-ledger decisions must all use the supported store contract",
        )
        ledger, ledger_statistics = store_ledger.compile_ledger(normalized, decisions)
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="finish-", dir=work) as temporary_name:
            temporary = Path(temporary_name)
            ledger_path = temporary / "store_ledger.json"
            workbook_path = temporary / "workbook.xlsx"
            core.atomic_json(ledger_path, ledger)
            workbook = store_ledger.build_workbook(ledger)
            workbook.save(workbook_path)
            workbook.close()
            errors = store_ledger.check_workbook(workbook_path, ledger)
            core.require(
                not errors,
                "store-ledger workbook verification failed: " + "; ".join(errors[:10]),
            )
            _publish_finished_workbook(
                workbook_path,
                output,
                run=run,
                replace_predecessor=replace_predecessor,
            )
        _record_finished_output(run, work, output)
        return {
            "output": str(output),
            **ledger_statistics,
            "missing_media": sum(
                int(item.get("missing_media") or 0)
                for item in review_statistics.values()
            ),
            "risk_report": _review_risk_summary(finish_risk_reports),
        }
    if finance_run:
        core.require(
            decision_contracts == {finance_materials.DECISION_CONTRACT},
            "finance-material decisions must all use the supported finance contract",
        )
        ledger, ledger_statistics = finance_materials.compile_ledger(normalized, decisions)
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="finish-", dir=work) as temporary_name:
            temporary = Path(temporary_name)
            ledger_path = temporary / "finance_materials.json"
            workbook_path = temporary / "workbook.xlsx"
            core.atomic_json(ledger_path, ledger)
            workbook = finance_materials.build_workbook(ledger)
            workbook.save(workbook_path)
            workbook.close()
            errors = finance_materials.check_workbook(workbook_path, ledger)
            core.require(
                not errors,
                "finance-material workbook verification failed: " + "; ".join(errors[:10]),
            )
            _publish_finished_workbook(
                workbook_path,
                output,
                run=run,
                replace_predecessor=replace_predecessor,
            )
        _record_finished_output(run, work, output)
        return {
            "output": str(output),
            **ledger_statistics,
            "material_images": sum(
                int(item.get("document_media") or 0)
                + int(item.get("profile_media") or 0)
                for item in review_statistics.values()
            ),
            "missing_media": sum(
                int(item.get("missing_media") or 0)
                for item in review_statistics.values()
            ),
            "risk_report": _review_risk_summary(finish_risk_reports),
        }
    if large_run:
        core.require(
            len(decision_contracts) == 1
            and decision_contracts
            <= {
                large_daily.DECISION_CONTRACT,
                large_daily.LEGACY_DECISION_CONTRACT,
            },
            "large-group decisions must all use the same supported contract",
        )

    if large_run and decision_contracts == {large_daily.LEGACY_DECISION_CONTRACT}:
        accounting_date = core.clean_text(run.get("accounting_date"))
        core.require(
            bool(accounting_date),
            "large-group runs require one accounting date",
        )
        ledger, ledger_statistics = large_daily.compile_daily_ledger(
            normalized,
            decisions,
            accounting_date=accounting_date,
        )
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="finish-", dir=work) as temporary_name:
            temporary = Path(temporary_name)
            ledger_path = temporary / "large_daily.json"
            workbook_path = temporary / "workbook.xlsx"
            core.atomic_json(ledger_path, ledger)
            workbook = build_workbook.build_large_workbook(
                build_workbook.validate_large_daily(ledger),
                args.template.resolve(),
            )
            workbook.save(workbook_path)
            workbook.close()
            errors = check_workbook.check(workbook_path, ledger_path)
            core.require(
                not errors,
                "workbook verification failed: " + "; ".join(errors[:10]),
            )
            _publish_finished_workbook(
                workbook_path,
                output,
                run=run,
                replace_predecessor=replace_predecessor,
            )
        _record_finished_output(run, work, output)
        return {
            "output": str(output),
            "groups": ledger_statistics["groups"],
            "exchanges": ledger_statistics["exchanges"],
            "summary_rows": ledger_statistics["summary_rows"],
            "fund_entries": sum(
                int(item.get("fund_entries") or 0)
                for item in review_statistics.values()
            ),
            "risk_report": _review_risk_summary(finish_risk_reports),
        }

    events, plan = _compile_decisions_v3(normalized, decisions)
    events_fingerprint = core.fingerprint_json(events)
    core.require(plan["events_fingerprint"] == events_fingerprint, "internal events fingerprint mismatch")
    orders, ledger_statistics = simple_ledger.compile_simple_ledger(
        normalized,
        events,
        events_fingerprint,
        plan,
    )
    large_daily_statistics: dict[str, int] | None = None
    if large_run:
        accounting_date = core.clean_text(run.get("accounting_date"))
        core.require(bool(accounting_date), "large-group runs require one accounting date")
        orders, large_daily_statistics = large_daily.compile_order_daily_ledger(
            orders,
            accounting_date=accounting_date,
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="finish-", dir=work) as temporary_name:
        temporary = Path(temporary_name)
        orders_path = temporary / "orders.json"
        workbook_path = temporary / "workbook.xlsx"
        core.atomic_json(orders_path, orders)
        validated_groups = build_workbook.validate_orders(orders)
        workbook = (
            build_workbook.build_large_order_workbook(
                validated_groups,
                args.template.resolve(),
            )
            if large_run
            else build_workbook.build_workbook(
                validated_groups,
                args.template.resolve(),
            )
        )
        workbook.save(workbook_path)
        workbook.close()
        errors = check_workbook.check(workbook_path, orders_path)
        core.require(not errors, "workbook verification failed: " + "; ".join(errors[:10]))
        _publish_finished_workbook(
            workbook_path,
            output,
            run=run,
            replace_predecessor=replace_predecessor,
        )

    _record_finished_output(run, work, output)
    result = {
        "output": str(output),
        "groups": ledger_statistics["groups"],
        "orders": ledger_statistics["orders"],
        "fund_entries": sum(item["fund_entries"] for item in review_statistics.values()),
        "pending_orders": ledger_statistics["pending_orders"],
        "warnings": ledger_statistics["warnings"],
        "risk_report": _review_risk_summary(finish_risk_reports),
    }
    if large_run:
        assert large_daily_statistics is not None
        result["groups"] = large_daily_statistics["groups"]
        result["orders"] = large_daily_statistics["orders"]
        result["summary_rows"] = large_daily_statistics["summary_rows"]
    return result


def main(argv: list[str] | None = None) -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")
    sys.stderr.reconfigure(encoding="utf-8", errors="backslashreplace")
    args = parse_args(argv)
    try:
        if args.command == "start":
            result = start_run(args)
        elif args.command == "roll-forward":
            result = roll_forward_run(args)
        elif args.command == "review":
            result = review_command(args)
        else:
            result = finish_run(args)
    except (
        OSError,
        ValueError,
        KeyError,
        TypeError,
        json.JSONDecodeError,
        ZoneInfoNotFoundError,
    ) as exc:
        print(f"Reconciliation failed: {exc}", file=sys.stderr)
        return 2
    compact = args.command == "review" and args.review_action == "next"
    sys.stdout.write(_render_result_json(result, compact=compact))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
