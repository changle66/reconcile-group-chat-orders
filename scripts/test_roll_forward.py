from __future__ import annotations

import copy
import json
import tempfile
import unittest
from argparse import Namespace
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import build_workbook
import core
import large_daily
import reconcile
import store_ledger


TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx"


def _epoch(value: str) -> str:
    return str(int(datetime.fromisoformat(value).timestamp()))


class RollForwardTests(unittest.TestCase):
    def _finished_small(
        self, root: Path, *, pending: bool
    ) -> tuple[Path, Path]:
        source = root / "day1"
        photos = source / "photos"
        photos.mkdir(parents=True)
        (photos / "payment.jpg").write_bytes(b"day-one-payment")
        messages = [
            {
                "id": 1,
                "type": "message",
                "date_unixtime": _epoch("2026-05-29T10:00:00+07:00"),
                "from": "Alice",
                "from_id": "user:alice",
                "text": "付款100 CNY，按5换500 THB",
                "photo": "photos/payment.jpg",
            }
        ]
        if not pending:
            (photos / "payout.jpg").write_bytes(b"day-one-payout")
            messages.append(
                {
                    "id": 2,
                    "type": "message",
                    "date_unixtime": _epoch("2026-05-29T10:05:00+07:00"),
                    "from": "QQ财务3",
                    "from_id": "user:staff",
                    "text": "已回500 THB",
                    "photo": "photos/payout.jpg",
                }
            )
        (source / "result.json").write_text(
            json.dumps(
                {"id": "roll-small", "name": "跨日小额群", "messages": messages},
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        work = root / "run1"
        reconcile.start_run(
            reconcile.parse_args(
                [
                    "start",
                    str(source),
                    "--work",
                    str(work),
                    "--mode",
                    "small",
                    "--contains",
                    "小额",
                    "--date",
                    "2026-05-29",
                    "--no-ocr-candidates",
                ]
            )
        )
        page = reconcile.review_command(
            Namespace(work=work, review_action="next", group=None, limit=500)
        )
        media_decisions = {
            "M0001": {
                "classification": "fund",
                "viewed_original": True,
                "entries": [
                    {
                        "amount": "100",
                        "currency": "CNY",
                        "payee": "张三",
                        "payee_state": "visible",
                        "side": "payment",
                        "result": "completed",
                        "amount_state": "clear",
                    }
                ],
            }
        }
        entry_ids = ["M0001.1"]
        source_messages = ["S00001"]
        if not pending:
            media_decisions["M0002"] = {
                "classification": "fund",
                "viewed_original": True,
                "entries": [
                    {
                        "amount": "500",
                        "currency": "THB",
                        "payee": "206-4-xxx781",
                        "payee_state": "visible",
                        "side": "payout",
                        "result": "completed",
                        "amount_state": "clear",
                    }
                ],
            }
            entry_ids.append("M0002.1")
            source_messages.append("S00002")
        batch = root / "day1-batch.json"
        core.atomic_json(
            batch,
            {
                "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
                "batch_id": "day1",
                "base_fingerprint": page["semantic_fingerprint"],
                "page_commit": {
                    "page_start": page["page_start"],
                    "page_end": page["page_end"],
                    "page_token": page["page_token"],
                },
                "open_orders": [],
                "media_decisions": media_decisions,
                "orders": [
                    {
                        "id": "O001",
                        "entry_ids": entry_ids,
                        "source_messages": source_messages,
                        "customer_nickname": "Alice",
                        "direction": "CNY->THB",
                        "pricing": {
                            "source_messages": ["S00001"],
                            "terms": {"rate": "5", "operator": "multiply"},
                            "expected": {"kind": "explicit", "amount": "500"},
                        },
                    }
                ],
            },
        )
        reconcile.review_command(
            Namespace(
                work=work,
                review_action="apply-batch",
                group="跨日小额群",
                input=batch,
            )
        )
        reconcile.review_command(
            Namespace(work=work, review_action="seal", group="跨日小额群")
        )
        output = root / "累计小额.xlsx"
        reconcile.finish_run(Namespace(work=work, output=output, template=TEMPLATE))
        return work, output

    def _finished_legacy_large(self, root: Path) -> tuple[Path, Path]:
        source = root / "legacy-large-day1"
        photos = source / "photos"
        photos.mkdir(parents=True)
        (photos / "payment.jpg").write_bytes(b"legacy-large-payment")
        (photos / "payout.jpg").write_bytes(b"legacy-large-payout")
        (source / "result.json").write_text(
            json.dumps(
                {
                    "id": "roll-large",
                    "name": "跨日大额固定群",
                    "messages": [
                        {
                            "id": 1,
                            "type": "message",
                            "date_unixtime": _epoch("2026-05-29T10:00:00+07:00"),
                            "from": "Alice",
                            "from_id": "user:alice",
                            "text": "微信付款100 CNY，按5换500 THB",
                            "photo": "photos/payment.jpg",
                        },
                        {
                            "id": 2,
                            "type": "message",
                            "date_unixtime": _epoch("2026-05-29T10:05:00+07:00"),
                            "from": "QQ财务3",
                            "from_id": "user:staff",
                            "text": "已回500 THB",
                            "photo": "photos/payout.jpg",
                        },
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        work = root / "legacy-large-run1"
        reconcile.start_run(
            reconcile.parse_args(
                [
                    "start",
                    str(source),
                    "--work",
                    str(work),
                    "--mode",
                    "large",
                    "--date",
                    "2026-05-29",
                    "--no-ocr-candidates",
                ]
            )
        )
        run = reconcile._load_run(work)
        run_group = run["groups"][0]
        decision_path = reconcile._decision_path(work, run_group)
        decision = reconcile._load_json(decision_path)
        decision["contract_version"] = large_daily.LEGACY_DECISION_CONTRACT
        for field in ("orders", "open_orders", "balance_links", "settlement_allocations"):
            decision.pop(field, None)
        decision["exchanges"] = []
        decision["open_exchanges"] = []
        decision["edit_control"] = reconcile._new_edit_control(
            reconcile._semantic_fingerprint(decision)
        )
        core.atomic_json(decision_path, decision)
        page = reconcile.review_command(
            Namespace(work=work, review_action="next", group=None, limit=500)
        )
        batch = root / "legacy-large-batch.json"
        core.atomic_json(
            batch,
            {
                "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
                "batch_id": "legacy-large-day1",
                "base_fingerprint": page["semantic_fingerprint"],
                "page_commit": {
                    "page_start": page["page_start"],
                    "page_end": page["page_end"],
                    "page_token": page["page_token"],
                },
                "open_exchanges": [],
                "media_decisions": {
                    "M0001": {
                        "classification": "fund",
                        "viewed_original": True,
                        "entries": [
                            {
                                "amount": "100",
                                "currency": "CNY",
                                "payee": "张三",
                                "payee_state": "visible",
                                "side": "payment",
                                "result": "completed",
                                "amount_state": "clear",
                            }
                        ],
                    },
                    "M0002": {
                        "classification": "fund",
                        "viewed_original": True,
                        "entries": [
                            {
                                "amount": "500",
                                "currency": "THB",
                                "payee": "206-4-xxx781",
                                "payee_state": "visible",
                                "side": "payout",
                                "result": "completed",
                                "amount_state": "clear",
                            }
                        ],
                    },
                },
                "exchanges": [
                    {
                        "id": "L001",
                        "source_messages": ["S00001", "S00002"],
                        "entry_ids": ["M0001.1", "M0002.1"],
                        "fund_type": "wechat",
                        "direction": "CNY->THB",
                        "source_amount": "100",
                        "rate": "5",
                        "operator": "multiply",
                        "target_amount": "500",
                    }
                ],
            },
        )
        reconcile.review_command(
            Namespace(
                work=work,
                review_action="apply-batch",
                group="跨日大额固定群",
                input=batch,
            )
        )
        reconcile.review_command(
            Namespace(work=work, review_action="seal", group="跨日大额固定群")
        )
        output = root / "累计大额.xlsx"
        reconcile.finish_run(Namespace(work=work, output=output, template=TEMPLATE))
        return work, output

    def _day_two_source(
        self, root: Path, *, payout: bool, reply_to_day_one: bool = False
    ) -> Path:
        source = root / ("day2-payout" if payout else "day2-text")
        source.mkdir()
        message: dict[str, object] = {
            "id": 3,
            "type": "message",
            "date_unixtime": _epoch("2026-05-30T11:00:00+07:00"),
            "from": "QQ财务3" if payout else "Alice",
            "from_id": "user:staff" if payout else "user:alice",
            "text": "已回500 THB" if payout else "第二天已复核",
        }
        if payout:
            photos = source / "photos"
            photos.mkdir()
            (photos / "payout.jpg").write_bytes(b"day-two-payout")
            message["photo"] = "photos/payout.jpg"
        if reply_to_day_one:
            message["reply_to_message_id"] = 1
        (source / "result.json").write_text(
            json.dumps(
                {
                    "id": "roll-small",
                    "name": "跨日小额群",
                    "messages": [message],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        return source

    def _roll(
        self,
        root: Path,
        work: Path,
        output: Path,
        source: Path,
        *,
        amendments: Path | None = None,
    ) -> tuple[Path, dict]:
        successor = root / ("run2-amended" if amendments else "run2")
        result = reconcile.roll_forward_run(
            Namespace(
                previous_work=work,
                inputs=[source],
                previous_output=output,
                work=successor,
                date="2026-05-30",
                confirmed_amendments=amendments,
            )
        )
        return successor, result

    def _commit_successor_page(
        self,
        root: Path,
        work: Path,
        *,
        close_pending: bool,
    ) -> None:
        page = reconcile.review_command(
            Namespace(work=work, review_action="next", group=None, limit=500)
        )
        payload: dict[str, object] = {
            "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
            "batch_id": "day2",
            "base_fingerprint": page["semantic_fingerprint"],
            "page_commit": {
                "page_start": page["page_start"],
                "page_end": page["page_end"],
                "page_token": page["page_token"],
            },
            "open_orders": [],
        }
        if close_pending:
            payload["media_decisions"] = {
                "M0002": {
                    "classification": "fund",
                    "viewed_original": True,
                    "entries": [
                        {
                            "amount": "500",
                            "currency": "THB",
                            "payee": "206-4-xxx781",
                            "payee_state": "visible",
                            "side": "payout",
                            "result": "completed",
                            "amount_state": "clear",
                        }
                    ],
                }
            }
            payload["orders"] = [
                {
                    "id": "O001",
                    "entry_ids": ["M0001.1", "M0002.1"],
                    "source_messages": ["S00001", "S00002"],
                    "customer_nickname": "Alice",
                    "direction": "CNY->THB",
                    "pricing": {
                        "source_messages": ["S00001"],
                        "terms": {"rate": "5", "operator": "multiply"},
                        "expected": {"kind": "explicit", "amount": "500"},
                    },
                }
            ]
        batch = root / "day2-batch.json"
        core.atomic_json(batch, payload)
        reconcile.review_command(
            Namespace(
                work=work,
                review_action="apply-batch",
                group="跨日小额群",
                input=batch,
            )
        )
        reconcile.review_command(
            Namespace(work=work, review_action="seal", group="跨日小额群")
        )

    def test_small_pending_order_completes_in_place_and_preserves_predecessor(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, output = self._finished_small(root, pending=True)
            predecessor_run_bytes = reconcile._run_path(work).read_bytes()
            predecessor_snapshot_bytes = reconcile._snapshot_path(work).read_bytes()
            successor, result = self._roll(
                root, work, output, self._day_two_source(root, payout=True)
            )
            self.assertEqual(result["carry_objects"], 1)
            run = reconcile._load_run(successor)
            decision = reconcile._load_json(
                reconcile._decision_path(successor, run["groups"][0])
            )
            self.assertEqual(decision["orders"][0]["id"], "O001")
            self.assertEqual(
                decision["orders"][0]["lifecycle"]["status"],
                "pending_next_day",
            )
            self.assertEqual(decision["open_orders"], [])
            self._commit_successor_page(root, successor, close_pending=True)
            decision = reconcile._load_json(
                reconcile._decision_path(successor, reconcile._load_run(successor)["groups"][0])
            )
            lifecycle = decision["orders"][0]["lifecycle"]
            self.assertEqual(lifecycle["status"], "completed")
            self.assertEqual(lifecycle["completion_at"][:10], "2026-05-30")
            self.assertEqual(
                [item["status"] for item in lifecycle["history"]],
                ["pending_next_day", "completed"],
            )
            reconcile.finish_run(
                Namespace(
                    work=successor,
                    output=output,
                    template=TEMPLATE,
                    replace_predecessor=True,
                )
            )
            self.assertEqual(reconcile._run_path(work).read_bytes(), predecessor_run_bytes)
            self.assertEqual(
                reconcile._snapshot_path(work).read_bytes(), predecessor_snapshot_bytes
            )
            workbook = load_workbook(output, read_only=True, data_only=False)
            try:
                sheet = workbook.worksheets[0]
                self.assertEqual(
                    sheet.cell(2, core.HEADERS.index("订单状态") + 1).value,
                    "已完成",
                )
                self.assertEqual(
                    sheet.cell(2, core.HEADERS.index("完成时间") + 1).value.date().isoformat(),
                    "2026-05-30",
                )
            finally:
                workbook.close()

    def test_explicit_pending_lifecycle_is_not_closed_by_two_partial_sides(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, _ = self._finished_small(root, pending=False)
            run = reconcile._load_run(work)
            normalized = reconcile._load_snapshot(work, run)
            group = normalized["groups"][0]
            decision = reconcile._load_json(
                reconcile._decision_path(work, run["groups"][0])
            )
            pending_at = group["messages"][-1]["timestamp"]
            decision["orders"][0]["lifecycle"] = {
                "status": "pending_next_day",
                "completion_at": "",
                "reason": "聊天明确说明尚未全部换完，次日继续",
                "source_messages": ["S00002"],
                "history": [
                    {
                        "status": "pending_next_day",
                        "at": pending_at,
                        "reason": "聊天明确说明尚未全部换完，次日继续",
                        "source_messages": ["S00002"],
                    }
                ],
            }
            reconcile._validate_decision(
                normalized,
                group,
                decision,
                require_complete=True,
                capture_hashes=False,
                rehash_labels=set(),
            )
            self.assertEqual(
                decision["orders"][0]["lifecycle"]["status"],
                "pending_next_day",
            )

    def test_legacy_large_contract_migrates_to_reviewable_full_order(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, output = self._finished_legacy_large(root)
            source = root / "legacy-large-day2"
            source.mkdir()
            (source / "result.json").write_text(
                json.dumps(
                    {
                        "id": "roll-large",
                        "name": "跨日大额固定群",
                        "messages": [
                            {
                                "id": 3,
                                "type": "message",
                                "date_unixtime": _epoch(
                                    "2026-05-30T09:00:00+07:00"
                                ),
                                "from": "Alice",
                                "from_id": "user:alice",
                                "text": "第二日复核旧单",
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            successor = root / "legacy-large-run2"
            result = reconcile.roll_forward_run(
                Namespace(
                    previous_work=work,
                    inputs=[source],
                    previous_output=output,
                    work=successor,
                    date="2026-05-30",
                    confirmed_amendments=None,
                )
            )
            self.assertEqual(result["legacy_contract_migrations"], 1)
            run = reconcile._load_run(successor)
            decision = reconcile._load_json(
                reconcile._decision_path(successor, run["groups"][0])
            )
            self.assertEqual(
                decision["contract_version"], large_daily.DECISION_CONTRACT
            )
            self.assertEqual(decision["orders"][0]["id"], "L001")
            self.assertEqual(
                decision["orders"][0]["lifecycle"]["status"], "completed"
            )
            self.assertIn("旧大额合同迁移", decision["orders"][0]["note"])
            self.assertEqual(run["groups"][0]["roll_forward_carry"], ["L001"])

    def test_large_summary_uses_occurrence_date_and_three_amount_states(self) -> None:
        orders = [
            {
                "fund_type": "bank_card",
                "direction": "CNY->THB",
                "actual_rate": "5",
                "rate_operator": "multiply",
                "reconciliation": {"status": "matched"},
                "lifecycle": {
                    "status": "completed",
                    "completion_at": "2026-09-05T10:00:00+07:00",
                    "history": [
                        {"status": "pending_next_day", "at": "2026-09-04T10:00:00+07:00"},
                        {"status": "completed", "at": "2026-09-05T10:00:00+07:00"},
                    ],
                },
                "flows": [
                    {"side": "payment", "amount": "100", "included": True, "status": "completed", "message_time": "2026-09-04T10:00:00+07:00"},
                    {"side": "payout", "amount": "500", "included": True, "status": "completed", "message_time": "2026-09-05T10:00:00+07:00"},
                ],
            },
            {
                "fund_type": "cash",
                "direction": "CNY->THB",
                "actual_rate": "6",
                "rate_operator": "multiply",
                "reconciliation": {"status": "pending"},
                "flows": [
                    {"side": "payment", "amount": "100", "included": True, "status": "completed", "message_time": "2026-09-06T10:00:00+07:00"},
                    {"side": "payment_refund", "amount": "100", "included": True, "status": "completed", "message_time": "2026-09-06T11:00:00+07:00"},
                ],
            },
            {
                "fund_type": "usdt",
                "direction": "USDT->THB",
                "actual_rate": "32",
                "rate_operator": "multiply",
                "reconciliation": {"status": "pending"},
                "flows": [
                    {"side": "payment", "amount": None, "included": False, "status": "unknown", "message_time": "2026-09-07T10:00:00+07:00"},
                ],
            },
            {
                "fund_type": "wechat",
                "direction": "CNY->THB",
                "actual_rate": "5",
                "rate_operator": "multiply",
                "reconciliation": {"status": "pending"},
                "lifecycle": {
                    "status": "cancelled",
                    "completion_at": "",
                    "history": [
                        {"status": "cancelled", "at": "2026-09-08T10:00:00+07:00"}
                    ],
                },
                "flows": [
                    {"side": "payment", "amount": "100", "included": False, "status": "failed", "message_time": "2026-09-08T10:00:00+07:00"},
                ],
            },
        ]
        summaries = large_daily.compile_order_summaries(
            orders, group_key="line:large", accounting_date="2026-09-07"
        )
        by_date = {item["statistics_date"]: item for item in summaries}
        self.assertEqual(by_date["2026-09-04"]["source_total"], "100")
        self.assertEqual(by_date["2026-09-04"]["target_total_state"], "absent")
        self.assertEqual(by_date["2026-09-05"]["source_total_state"], "absent")
        self.assertEqual(by_date["2026-09-05"]["target_total"], "500")
        self.assertEqual(by_date["2026-09-06"]["source_total"], "0")
        self.assertEqual(by_date["2026-09-07"]["source_total_state"], "unknown")
        self.assertNotIn("2026-09-08", by_date)
        rows = build_workbook.large_summary_rows({"daily_summaries": summaries})
        rendered = {row[0]: row for row in rows}
        self.assertEqual(rendered["2026-09-04"][7], "—")
        self.assertEqual(rendered["2026-09-06"][6], 0)
        self.assertEqual(rendered["2026-09-07"][6], "未确认")

    def test_stable_merge_deduplicates_whatsapp_and_enriches_media(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            image = root / "evidence.jpg"
            image.write_bytes(b"same-evidence")

            def document(message_id: str, media: dict) -> dict:
                return {
                    "contract_version": core.NORMALIZED_CONTRACT,
                    "timezone": "Asia/Bangkok",
                    "source_fingerprint": "sha256:" + ("1" if "old" in message_id else "2") * 64,
                    "groups": [
                        {
                            "group_key": "whatsapp:group",
                            "group_name": "小额群",
                            "platform": "WhatsApp",
                            "messages": [
                                {
                                    "message_id": message_id,
                                    "timestamp": "2026-09-05T10:00:00+07:00",
                                    "sender_name": "Alice",
                                    "sender_id": "name:alice",
                                    "role": "客户候选",
                                    "text": "付款",
                                    "reply_to_message_id": None,
                                    "media": [media],
                                    "source_sequence": 1,
                                    "excluded_from_accounting": False,
                                }
                            ],
                        }
                    ],
                }

            missing = {
                "media_id": "old-media",
                "kind": "image",
                "mime_type": "image/jpeg",
                "availability": "missing",
                "path": None,
                "byte_size": None,
                "missing_kind": "not_exported",
            }
            available = {
                "media_id": "new-media",
                "kind": "image",
                "mime_type": "image/jpeg",
                "availability": "available",
                "path": str(image),
                "byte_size": image.stat().st_size,
                "blob_sha256": core.sha256_file(image),
            }
            previous = document("snapshot-old", missing)
            reconcile._assign_stable_identities(previous)
            merged, stats = reconcile._merge_normalized_snapshots(
                previous,
                document("snapshot-new", available),
            )
            self.assertEqual(stats["messages_deduplicated"], 1)
            self.assertEqual(stats["media_enriched"], 1)
            self.assertEqual(len(merged["groups"][0]["messages"]), 1)
            self.assertEqual(
                merged["groups"][0]["messages"][0]["media"][0]["availability"],
                "available",
            )
            _, _, reusable = reconcile._label_maps(
                previous["groups"][0], merged["groups"][0]
            )
            self.assertEqual(reusable, set())

            previous_available = document("snapshot-old", copy.deepcopy(available))
            reconcile._assign_stable_identities(previous_available)
            same_media, _ = reconcile._merge_normalized_snapshots(
                previous_available,
                document("snapshot-new", copy.deepcopy(available)),
            )
            _, _, reusable = reconcile._label_maps(
                previous_available["groups"][0], same_media["groups"][0]
            )
            self.assertEqual(reusable, {"M0001"})

            changed_text = document("snapshot-new", copy.deepcopy(available))
            changed_text["groups"][0]["messages"][0]["text"] = "付款文字被修改"
            with self.assertRaisesRegex(ValueError, "message identity conflict"):
                reconcile._merge_normalized_snapshots(previous, changed_text)

    def test_stable_merge_accepts_line_adapter_switch_and_rejects_conflict(self) -> None:
        def document(message_id: str, text: str) -> dict:
            return {
                "contract_version": core.NORMALIZED_CONTRACT,
                "timezone": "Asia/Bangkok",
                "source_fingerprint": "sha256:" + "3" * 64,
                "groups": [
                    {
                        "group_key": "line:c123",
                        "group_name": "门店开票群",
                        "platform": "LINE",
                        "messages": [
                            {
                                "message_id": message_id,
                                "timestamp": "2026-09-05T10:00:00+07:00",
                                "sender_name": "员工甲",
                                "sender_id": "u123",
                                "role": "内部人员",
                                "text": text,
                                "reply_to_message_id": None,
                                "media": [],
                                "source_sequence": 1,
                                "excluded_from_accounting": False,
                            }
                        ],
                    }
                ],
            }

        merged, stats = reconcile._merge_normalized_snapshots(
            document("line:c123:ios-zid", "+100 USD"),
            document("line:c123:android-local", "+100 USD"),
        )
        self.assertEqual(stats["messages_deduplicated"], 1)
        self.assertEqual(len(merged["groups"][0]["messages"]), 1)
        with self.assertRaisesRegex(ValueError, "message identity conflict"):
            reconcile._merge_normalized_snapshots(
                document("line:c123:same-id", "+100 USD"),
                document("line:c123:same-id", "-100 USD"),
            )

        def reply_document(target: str) -> dict:
            value = document("reply", "更正")
            value["groups"][0]["messages"] = [
                {
                    "message_id": "a",
                    "timestamp": "2026-09-05T09:58:00+07:00",
                    "sender_name": "员工甲",
                    "sender_id": "u123",
                    "role": "内部人员",
                    "text": "票据A",
                    "reply_to_message_id": None,
                    "media": [],
                    "source_sequence": 1,
                    "excluded_from_accounting": False,
                },
                {
                    "message_id": "b",
                    "timestamp": "2026-09-05T09:59:00+07:00",
                    "sender_name": "员工甲",
                    "sender_id": "u123",
                    "role": "内部人员",
                    "text": "票据B",
                    "reply_to_message_id": None,
                    "media": [],
                    "source_sequence": 2,
                    "excluded_from_accounting": False,
                },
                {
                    "message_id": "reply",
                    "timestamp": "2026-09-05T10:00:00+07:00",
                    "sender_name": "员工甲",
                    "sender_id": "u123",
                    "role": "内部人员",
                    "text": "更正",
                    "reply_to_message_id": target,
                    "media": [],
                    "source_sequence": 3,
                    "excluded_from_accounting": False,
                },
            ]
            return value

        with self.assertRaisesRegex(ValueError, "reply relationships"):
            reconcile._merge_normalized_snapshots(
                reply_document("a"), reply_document("b")
            )

        multi_group = reply_document("a")
        multi_group["groups"].append(
            {
                "group_key": "line:other",
                "group_name": "另一个群",
                "platform": "LINE",
                "messages": [
                    {
                        "message_id": "a",
                        "timestamp": "2026-09-05T11:00:00+07:00",
                        "sender_name": "其他人",
                        "sender_id": "u999",
                        "role": "客户候选",
                        "text": "另一个群复用了相同原始ID",
                        "reply_to_message_id": None,
                        "media": [],
                        "source_sequence": 1,
                        "excluded_from_accounting": False,
                    }
                ],
            }
        )
        empty = {
            "contract_version": core.NORMALIZED_CONTRACT,
            "timezone": "Asia/Bangkok",
            "source_fingerprint": "sha256:" + "5" * 64,
            "groups": [],
        }
        canonical, _ = reconcile._merge_normalized_snapshots(multi_group, empty)
        first_group = next(
            item for item in canonical["groups"] if item["group_key"] == "line:c123"
        )
        self.assertEqual(
            first_group["messages"][2]["reply_to_message_id"],
            first_group["messages"][0]["stable_message_id"],
        )

    def test_date_gap_is_rejected_before_successor_work_is_created(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, output = self._finished_small(root, pending=False)
            source = self._day_two_source(root, payout=False)
            successor = root / "gap-run"
            with self.assertRaisesRegex(ValueError, "immediately follow"):
                reconcile.roll_forward_run(
                    Namespace(
                        previous_work=work,
                        inputs=[source],
                        previous_output=output,
                        work=successor,
                        date="2026-06-01",
                        confirmed_amendments=None,
                    )
                )
            self.assertFalse(successor.exists())

    def test_business_drift_requires_hash_bound_semantic_amendment(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, output = self._finished_small(root, pending=False)
            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                worksheet = workbook.worksheets[0]
                worksheet.cell(2, core.HEADERS.index("客户昵称") + 1).value = "Alice-复核"
                worksheet.cell(2, core.HEADERS.index("备注") + 1).value = "人工确认备注"
                workbook.save(output)
            finally:
                workbook.close()
            source = self._day_two_source(root, payout=False)
            blocked = root / "blocked-run"
            with self.assertRaisesRegex(ValueError, "business cells changed"):
                reconcile.roll_forward_run(
                    Namespace(
                        previous_work=work,
                        inputs=[source],
                        previous_output=output,
                        work=blocked,
                        date="2026-05-30",
                        confirmed_amendments=None,
                    )
                )
            self.assertFalse(blocked.exists())
            amendments = root / "amendments.json"
            core.atomic_json(
                amendments,
                {
                    "contract_version": reconcile.AMENDMENTS_CONTRACT,
                    "previous_output_sha256": core.sha256_file(output),
                    "amendments": [
                        {
                            "group_key": "telegram:roll-small",
                            "collection": "orders",
                            "id": "O001",
                            "field": "customer_nickname",
                            "value": "Alice-复核",
                            "reason": "旧表人工确认客户显示名",
                            "basis": "S00001 的客户消息",
                            "source_messages": ["S00001"],
                        }
                    ],
                },
            )
            with self.assertRaisesRegex(ValueError, "do not fully explain"):
                self._roll(root, work, output, source, amendments=amendments)
            payload = reconcile._load_json(amendments)
            payload["amendments"].append(
                {
                    "group_key": "telegram:roll-small",
                    "collection": "orders",
                    "id": "O001",
                    "field": "note",
                    "value": "人工确认备注",
                    "reason": "旧表人工补充已确认备注",
                    "basis": "S00001 的换汇说明",
                    "source_messages": ["S00001"],
                }
            )
            core.atomic_json(amendments, payload)
            successor, result = self._roll(
                root, work, output, source, amendments=amendments
            )
            self.assertTrue(result["workbook_business_drift_confirmed"])
            decision = reconcile._load_json(
                reconcile._decision_path(successor, reconcile._load_run(successor)["groups"][0])
            )
            self.assertEqual(decision["orders"][0]["customer_nickname"], "Alice-复核")

    def test_formatting_only_drift_is_accepted(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, output = self._finished_small(root, pending=False)
            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                workbook.worksheets[0].column_dimensions["A"].width = 42
                workbook.save(output)
            finally:
                workbook.close()
            successor, result = self._roll(
                root,
                work,
                output,
                self._day_two_source(
                    root, payout=False, reply_to_day_one=True
                ),
            )
            self.assertTrue(successor.is_dir())
            self.assertFalse(result["workbook_business_drift_confirmed"])
            page = reconcile.review_command(
                Namespace(
                    work=successor,
                    review_action="next",
                    group=None,
                    limit=500,
                )
            )
            self.assertEqual(page["roll_forward_carry"], [])
            prior_context = next(
                item for item in page["carry_messages"] if item["label"] == "S00001"
            )
            self.assertTrue(prior_context["media"][0]["path"])

    def test_atomic_publish_failure_leaves_predecessor_unchanged(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work, output = self._finished_small(root, pending=False)
            successor, _ = self._roll(
                root, work, output, self._day_two_source(root, payout=False)
            )
            self._commit_successor_page(root, successor, close_pending=False)
            before = output.read_bytes()
            with patch.object(Path, "replace", side_effect=OSError("locked by WPS")):
                with self.assertRaisesRegex(OSError, "locked by WPS"):
                    reconcile.finish_run(
                        Namespace(
                            work=successor,
                            output=output,
                            template=TEMPLATE,
                            replace_predecessor=True,
                        )
                    )
            self.assertEqual(output.read_bytes(), before)
            self.assertEqual(reconcile._load_run(successor)["status"], "reviewing")

    def test_store_pending_record_posts_next_day_with_original_voucher_date(self) -> None:
        group = {
            "group_key": "line:store",
            "group_name": "曼谷门店开票群",
            "platform": "LINE",
            "messages": [
                {
                    "message_id": "m1",
                    "timestamp": "2026-09-04T10:00:00+07:00",
                    "source_sequence": 1,
                    "sender_id": "u1",
                    "sender_name": "员工甲",
                    "role": "内部人员",
                    "text": "票号0041268，明天继续",
                    "reply_to_message_id": None,
                    "excluded_from_accounting": False,
                    "media": [],
                },
                {
                    "message_id": "m2",
                    "timestamp": "2026-09-05T09:00:00+07:00",
                    "source_sequence": 2,
                    "sender_id": "u1",
                    "sender_name": "员工甲",
                    "role": "内部人员",
                    "text": "昨天票据已入账 +100 USD",
                    "reply_to_message_id": "m1",
                    "excluded_from_accounting": False,
                    "media": [],
                },
            ],
        }
        normalized = {
            "contract_version": core.NORMALIZED_CONTRACT,
            "timezone": "Asia/Bangkok",
            "source_fingerprint": "sha256:" + "4" * 64,
            "accounting_window": {
                "start": "2026-09-04T00:00:00+07:00",
                "end": "2026-09-06T00:00:00+07:00",
            },
            "groups": [group],
        }
        expected = reconcile._decision_template(
            normalized, group, group_mode=store_ledger.GROUP_MODE
        )
        pending = copy.deepcopy(expected)
        pending["reviewed_through"] = 2
        pending["read_complete"] = True
        pending["records"] = [
            {
                "id": "R001",
                "record_type": "cash_movement",
                "posting_status": "pending",
                "posting_at": "",
                "voucher_number": "0041268",
                "voucher_date": "2026-09-04",
                "source_messages": ["S00001"],
                "media_roles": [],
                "representative_media_label": "",
                "movements": [
                    {
                        "currency": "USD",
                        "sign": "+",
                        "amount": "100",
                        "basis": "chat_explicit",
                        "source_messages": ["S00001"],
                    }
                ],
                "status_reason": "次日继续",
            }
        ]
        store_ledger.validate_decision(
            normalized,
            group,
            pending,
            expected,
            require_complete=True,
            capture_hashes=False,
            rehash_labels=set(),
        )
        posted_record = copy.deepcopy(pending["records"][0])
        posted_record.pop("status_history")
        posted_record.update(
            {
                "posting_status": "posted",
                "posting_at": "2026-09-05T09:00:00+07:00",
                "source_messages": ["S00001", "S00002"],
                "status_reason": "",
            }
        )
        candidate = store_ledger.merge_review_batch(
            pending, {"records": [posted_record]}
        )
        store_ledger.validate_decision(
            normalized,
            group,
            candidate,
            expected,
            require_complete=True,
            capture_hashes=False,
            rehash_labels=set(),
        )
        ledger, _ = store_ledger.compile_ledger(
            normalized, {"line:store": candidate}
        )
        record = ledger["groups"][0]["records"][0]
        self.assertEqual(record["voucher_date"], "2026-09-04")
        self.assertEqual(record["posting_at"][:10], "2026-09-05")
        self.assertEqual(
            [item["status"] for item in record["status_history"]],
            ["pending", "posted"],
        )
        self.assertEqual(
            {row["date"] for row in ledger["groups"][0]["daily_reconciliation"]},
            {"2026-09-05"},
        )


if __name__ == "__main__":
    unittest.main()
