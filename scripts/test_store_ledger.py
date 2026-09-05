from __future__ import annotations

import base64
import json
import tempfile
import unittest
from argparse import Namespace
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import core
import reconcile
import store_ledger


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl6qWQAAAAASUVORK5CYII="
)


def _message(message_id: str, timestamp: str, text: str, sequence: int) -> dict:
    return {
        "message_id": message_id,
        "timestamp": timestamp,
        "source_sequence": sequence,
        "sender_id": "staff",
        "sender_name": "内部人员",
        "role": "内部人员",
        "text": text,
        "reply_to_message_id": None,
        "excluded_from_accounting": False,
        "media": [],
    }


class StoreLedgerTests(unittest.TestCase):
    def _start_fixture(self, root: Path) -> Path:
        source = root / "source"
        photos = source / "photos"
        photos.mkdir(parents=True)
        (photos / "ticket.png").write_bytes(PNG_1X1)
        (source / "result.json").write_text(
            json.dumps(
                {
                    "id": "bangkok-store",
                    "name": "曼谷门店开票群",
                    "messages": [
                        {
                            "id": 1,
                            "type": "message",
                            "date_unixtime": "1780000000",
                            "from": "门店员工甲",
                            "from_id": "user:staff-a",
                            "text": "票号 0040099",
                            "photo": "photos/ticket.png",
                        },
                        {
                            "id": 2,
                            "type": "message",
                            "date_unixtime": "1780000060",
                            "from": "门店员工乙",
                            "from_id": "user:staff-b",
                            "text": "USD ➕100✖️32.65🟰3265",
                            "reply_to_message_id": 1,
                        },
                        {
                            "id": 3,
                            "type": "message",
                            "date_unixtime": "1780000120",
                            "from": "门店员工乙",
                            "from_id": "user:staff-b",
                            "text": "日结 USD 100，THB 6735",
                        },
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        work = root / "run"
        report = reconcile.start_run(
            reconcile.parse_args(
                [
                    "start",
                    str(source),
                    "--work",
                    str(work),
                    "--mode",
                    store_ledger.GROUP_MODE,
                    "--no-ocr-candidates",
                ]
            )
        )
        self.assertEqual(report["groups"], 1)
        self.assertEqual(report["amount_policy"], store_ledger.AMOUNT_POLICY)
        run = reconcile._load_run(work)
        self.assertEqual(run["group_mode"], store_ledger.GROUP_MODE)
        self.assertEqual(run["group_name_contains"], store_ledger.GROUP_NAME_MARKER)
        self.assertTrue(reconcile._large_group_name("曼谷门店开票群"))
        return work

    def test_public_workflow_preserves_reply_media_and_publishes_store_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            work = self._start_fixture(root)
            page = reconcile.review_command(
                Namespace(work=work, review_action="next", group="曼谷门店开票群", limit=500)
            )
            self.assertIn("open_records", page)
            self.assertNotIn("open_orders", page)
            self.assertEqual(page["media_queue"]["mode"], "store_ledger")
            self.assertEqual(page["messages"][1]["reply"]["media_labels"], ["M0001"])
            posting_at = page["messages"][1]["timestamp"]
            closing_at = page["messages"][2]["timestamp"]
            date = posting_at[:10]
            batch = root / "store-batch.json"
            core.atomic_json(
                batch,
                {
                    "contract_version": reconcile.REVIEW_BATCH_CONTRACT,
                    "batch_id": "store-page-001",
                    "base_fingerprint": page["semantic_fingerprint"],
                    "page_commit": {
                        "page_start": page["page_start"],
                        "page_end": page["page_end"],
                        "page_token": page["page_token"],
                    },
                    "open_records": [],
                    "media_decisions": {
                        "M0001": {
                            "classification": "voucher",
                            "viewed_original": True,
                            "facts": {
                                "voucher_number": "0040099",
                                "voucher_date": date,
                                "movements": [
                                    {"currency": "USD", "sign": "+", "amount": "100", "rate": "32.65"},
                                    {"currency": "THB", "sign": "-", "amount": "32.65"},
                                ],
                            },
                        }
                    },
                    "media_observations": {
                        "M0001": {
                            "contract_version": reconcile.MEDIA_OBSERVATION_CONTRACT,
                            "classification": "voucher",
                            "review_status": "clear",
                            "viewed_original": True,
                            "recheck_reasons": [],
                        }
                    },
                    "records": [
                        {
                            "id": "R001",
                            "record_type": "exchange",
                            "posting_status": "posted",
                            "posting_at": posting_at,
                            "voucher_number": "0040099",
                            "voucher_date": date,
                            "source_messages": ["S00001", "S00002"],
                            "media_roles": [{"label": "M0001", "role": "final"}],
                            "representative_media_label": "M0001",
                            "movements": [
                                {
                                    "currency": "USD",
                                    "sign": "+",
                                    "amount": "100",
                                    "rate": "32.65",
                                    "basis": "ticket",
                                    "source_messages": ["S00001"],
                                },
                                {
                                    "currency": "THB",
                                    "sign": "-",
                                    "amount": "3265",
                                    "basis": "direct_reply_correction",
                                    "source_messages": ["S00002"],
                                    "note": "直接引用票据的完整算式更正票面笔误",
                                },
                            ],
                        }
                    ],
                    "balance_snapshots": [
                        {
                            "id": "B001",
                            "date": closing_at[:10],
                            "kind": "closing",
                            "source_messages": ["S00003"],
                            "media_labels": [],
                            "balances": [
                                {"currency": "USD", "amount": "100"},
                                {"currency": "THB", "amount": "6735"},
                            ],
                        }
                    ],
                },
            )
            applied = reconcile.review_command(
                Namespace(
                    work=work,
                    review_action="apply-batch",
                    group="曼谷门店开票群",
                    input=batch,
                )
            )
            self.assertTrue(applied["read_complete"])
            self.assertEqual(applied["records"], 1)
            reconcile.review_command(
                Namespace(work=work, review_action="seal", group="曼谷门店开票群")
            )
            output = root / "门店流水.xlsx"
            result = reconcile.finish_run(
                Namespace(work=work, output=output, template=root / "unused.xlsx")
            )
            self.assertEqual(result["posted_records"], 1)
            self.assertEqual(result["balance_mismatches"], 0)
            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["曼谷门店开票群"])
                worksheet = workbook["曼谷门店开票群"]
                self.assertEqual(worksheet.cell(2, 10).value, 100)
                self.assertEqual(worksheet.cell(3, 10).value, -3265)
                self.assertEqual(len(worksheet._images), 1)
                self.assertEqual(worksheet.cell(7, 10).value, "期初待确认")
            finally:
                workbook.close()

    def test_transfer_pairing_pending_exclusion_and_daily_mismatch(self) -> None:
        group_a = {
            "group_key": "line:bangkok",
            "group_name": "曼谷门店开票群",
            "platform": "LINE",
            "messages": [
                _message("a1", "2026-09-03T20:00:00+07:00", "+100 USD，日结100", 1),
                _message("a2", "2026-09-04T20:00:00+07:00", "-20 USD，调拨CNY，日结75", 2),
            ],
        }
        group_b = {
            "group_key": "line:pattaya",
            "group_name": "芭提雅门店开票群",
            "platform": "LINE",
            "messages": [
                _message("b1", "2026-09-04T20:01:00+07:00", "+1500 CNY 曼谷调拨", 1)
            ],
        }
        normalized = {
            "contract_version": core.NORMALIZED_CONTRACT,
            "timezone": "Asia/Bangkok",
            "source_fingerprint": "sha256:" + "1" * 64,
            "groups": [group_a, group_b],
        }

        def decision_for(group: dict, records: list[dict], snapshots: list[dict]) -> dict:
            decision = reconcile._decision_template(
                normalized, group, group_mode=store_ledger.GROUP_MODE
            )
            decision["reviewed_through"] = len(group["messages"])
            decision["read_complete"] = True
            decision["records"] = records
            decision["balance_snapshots"] = snapshots
            expected = reconcile._decision_template(
                normalized, group, group_mode=store_ledger.GROUP_MODE
            )
            store_ledger.validate_decision(
                normalized,
                group,
                decision,
                expected,
                require_complete=True,
                capture_hashes=False,
                rehash_labels=set(),
            )
            return decision

        common = {
            "posting_status": "posted",
            "voucher_number": "",
            "voucher_date": "",
            "media_roles": [],
            "representative_media_label": "",
            "status_reason": "",
            "note": "",
        }
        decision_a = decision_for(
            group_a,
            [
                {
                    **common,
                    "id": "A001",
                    "record_type": "cash_movement",
                    "posting_at": "2026-09-03T20:00:00+07:00",
                    "source_messages": ["S00001"],
                    "movements": [
                        {
                            "currency": "USD",
                            "sign": "+",
                            "amount": "100",
                            "basis": "chat_explicit",
                            "source_messages": ["S00001"],
                        }
                    ],
                },
                {
                    **common,
                    "id": "A002",
                    "record_type": "cash_movement",
                    "posting_at": "2026-09-04T20:00:00+07:00",
                    "source_messages": ["S00002"],
                    "movements": [
                        {
                            "currency": "USD",
                            "sign": "-",
                            "amount": "20",
                            "basis": "chat_explicit",
                            "source_messages": ["S00002"],
                        }
                    ],
                },
                {
                    **common,
                    "id": "A003",
                    "record_type": "internal_transfer",
                    "posting_at": "2026-09-04T20:00:00+07:00",
                    "source_messages": ["S00002"],
                    "movements": [
                        {
                            "currency": "CNY",
                            "sign": "-",
                            "amount": "1500",
                            "basis": "chat_explicit",
                            "source_messages": ["S00002"],
                        }
                    ],
                    "transfer": {
                        "transfer_id": "T001",
                        "counterparty_group_name": "芭提雅门店开票群",
                        "direction": "send",
                    },
                },
                {
                    **common,
                    "id": "A004",
                    "record_type": "cash_movement",
                    "posting_status": "pending",
                    "posting_at": "",
                    "source_messages": ["S00002"],
                    "status_reason": "12号取现金，尚未完成",
                    "movements": [
                        {
                            "currency": "THB",
                            "sign": "-",
                            "amount": "294000",
                            "basis": "chat_explicit",
                            "source_messages": ["S00002"],
                        }
                    ],
                },
            ],
            [
                {
                    "id": "BA1",
                    "date": "2026-09-03",
                    "kind": "closing",
                    "source_messages": ["S00001"],
                    "media_labels": [],
                    "balances": [{"currency": "USD", "amount": "100"}],
                },
                {
                    "id": "BA2",
                    "date": "2026-09-04",
                    "kind": "closing",
                    "source_messages": ["S00002"],
                    "media_labels": [],
                    "balances": [{"currency": "USD", "amount": "75"}],
                },
            ],
        )
        decision_b = decision_for(
            group_b,
            [
                {
                    **common,
                    "id": "B001",
                    "record_type": "internal_transfer",
                    "posting_at": "2026-09-04T20:01:00+07:00",
                    "source_messages": ["S00001"],
                    "movements": [
                        {
                            "currency": "CNY",
                            "sign": "+",
                            "amount": "1500",
                            "basis": "chat_explicit",
                            "source_messages": ["S00001"],
                        }
                    ],
                    "transfer": {
                        "transfer_id": "T001",
                        "counterparty_group_name": "曼谷门店开票群",
                        "direction": "receive",
                    },
                }
            ],
            [],
        )
        ledger, stats = store_ledger.compile_ledger(
            normalized,
            {"line:bangkok": decision_a, "line:pattaya": decision_b},
        )
        self.assertEqual(stats["matched_transfers"], 1)
        self.assertEqual(stats["pending_records"], 1)
        self.assertEqual(stats["balance_mismatches"], 1)
        group = ledger["groups"][0]
        pending = next(record for record in group["records"] if record["id"] == "A004")
        self.assertIsNone(pending["movements"][0]["actual_amount"])
        usd_day_two = next(
            row
            for row in group["daily_reconciliation"]
            if row["date"] == "2026-09-04" and row["currency"] == "USD"
        )
        self.assertEqual(usd_day_two["opening_balance"], "100")
        self.assertEqual(usd_day_two["calculated_closing"], "80")
        self.assertEqual(usd_day_two["observed_closing"], "75")
        self.assertEqual(usd_day_two["difference"], "-5")
        self.assertEqual(usd_day_two["status"], "存在差额")
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "store.xlsx"
            workbook = store_ledger.build_workbook(ledger)
            workbook.save(output)
            workbook.close()
            self.assertEqual(store_ledger.check_workbook(output, ledger), [])

    def test_direct_reply_correction_rejects_incomplete_formula(self) -> None:
        group = {
            "group_key": "line:test",
            "group_name": "测试门店开票群",
            "platform": "LINE",
            "messages": [
                _message("m1", "2026-09-04T10:00:00+07:00", "票据", 1),
                _message("m2", "2026-09-04T10:01:00+07:00", "应该是3265", 2),
            ],
        }
        group["messages"][0]["media"] = [
            {
                "media_id": "media-1",
                "kind": "image",
                "mime_type": "image/png",
                "availability": "available",
                "path": __file__,
                "byte_size": Path(__file__).stat().st_size,
                "blob_sha256": core.sha256_file(Path(__file__)),
            }
        ]
        group["messages"][1]["reply_to_message_id"] = "m1"
        normalized = {
            "timezone": "Asia/Bangkok",
            "source_fingerprint": "sha256:" + "2" * 64,
            "groups": [group],
        }
        decision = reconcile._decision_template(
            normalized, group, group_mode=store_ledger.GROUP_MODE
        )
        decision["reviewed_through"] = 2
        decision["read_complete"] = True
        decision["media_decisions"] = {
            "M0001": {
                "classification": "voucher",
                "viewed_original": True,
                "facts": {"movements": [{"currency": "THB", "sign": "-", "amount": "32.65"}]},
                "evidence_sha256": core.sha256_file(Path(__file__)),
            }
        }
        decision["records"] = [
            {
                "id": "R1",
                "record_type": "exchange",
                "posting_status": "posted",
                "posting_at": "2026-09-04T10:01:00+07:00",
                "source_messages": ["S00001", "S00002"],
                "media_roles": [{"label": "M0001", "role": "final"}],
                "representative_media_label": "M0001",
                "movements": [
                    {
                        "currency": "THB",
                        "sign": "-",
                        "amount": "3265",
                        "basis": "direct_reply_correction",
                        "source_messages": ["S00002"],
                        "note": "更正",
                    }
                ],
            }
        ]
        expected = reconcile._decision_template(
            normalized, group, group_mode=store_ledger.GROUP_MODE
        )
        with self.assertRaisesRegex(ValueError, "complete formula"):
            store_ledger.validate_decision(
                normalized,
                group,
                decision,
                expected,
                require_complete=True,
                capture_hashes=False,
                rehash_labels=set(),
            )

    def test_line_emoji_formula_is_normalized(self) -> None:
        group = {
            "group_key": "line:test",
            "group_name": "测试门店开票群",
            "platform": "LINE",
            "messages": [
                _message("m1", "2026-09-04T10:00:00+07:00", "票据", 1),
                _message("m2", "2026-09-04T10:01:00+07:00", "USD ➕100✖️32.65🟰3265", 2),
            ],
        }
        group["messages"][0]["media"] = [
            {
                "media_id": "media-1",
                "kind": "image",
                "mime_type": "image/png",
                "availability": "available",
                "path": __file__,
            }
        ]
        group["messages"][1]["reply_to_message_id"] = "m1"
        self.assertTrue(
            store_ledger._direct_reply_formula_exists(
                group,
                ["S00002"],
                {"M0001"},
                "3265",
            )
        )

    def test_ios_route_uses_store_group_pattern_automatically(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            backup = root / "ios-backup"
            backup.mkdir()
            work = root / "run"

            def fake_extract(arguments: list[str]) -> int:
                self.assertIn("--group-pattern", arguments)
                pattern_index = arguments.index("--group-pattern") + 1
                self.assertEqual(arguments[pattern_index], store_ledger.GROUP_NAME_MARKER)
                self.assertNotIn("--all-group-chats", arguments)
                output = Path(arguments[arguments.index("-o") + 1])
                output.parent.mkdir(parents=True, exist_ok=True)
                output.write_text(
                    json.dumps(
                        {
                            "contract_version": core.NORMALIZED_CONTRACT,
                            "timezone": "Asia/Bangkok",
                            "source_fingerprint": "sha256:" + "0" * 64,
                            "groups": [
                                {
                                    "group_key": "line:store-group",
                                    "group_name": "曼谷门店开票群",
                                    "platform": "LINE",
                                    "messages": [],
                                }
                            ],
                        },
                        ensure_ascii=False,
                    ),
                    encoding="utf-8",
                )
                return 0

            with patch("reconcile._line_has_matching_group", return_value=True), patch(
                "reconcile.extract_line_ios.main", side_effect=fake_extract
            ):
                documents = reconcile._line_documents(
                    work,
                    [backup],
                    contains=store_ledger.GROUP_NAME_MARKER,
                    group_mode=store_ledger.GROUP_MODE,
                    timezone_name="Asia/Bangkok",
                    roster_path=root / "roster.yaml",
                    self_name="LINE_SELF",
                )
            self.assertEqual(len(documents), 1)
            self.assertEqual(documents[0]["groups"][0]["group_name"], "曼谷门店开票群")

    def test_date_filter_keeps_reply_target_as_context_without_posting_it(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source"
            photos = source / "photos"
            photos.mkdir(parents=True)
            (photos / "previous.png").write_bytes(PNG_1X1)
            previous = datetime.fromisoformat("2026-09-03T23:59:00+07:00")
            current = datetime.fromisoformat("2026-09-04T00:01:00+07:00")
            (source / "result.json").write_text(
                json.dumps(
                    {
                        "id": "store-window",
                        "name": "曼谷门店开票群",
                        "messages": [
                            {
                                "id": 1,
                                "type": "message",
                                "date_unixtime": str(int(previous.timestamp())),
                                "from": "员工甲",
                                "from_id": "staff:a",
                                "text": "前一日票据",
                                "photo": "photos/previous.png",
                            },
                            {
                                "id": 2,
                                "type": "message",
                                "date_unixtime": str(int(current.timestamp())),
                                "from": "员工乙",
                                "from_id": "staff:b",
                                "text": "USD +100 × 32.65 = 3265",
                                "reply_to_message_id": 1,
                            },
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            work = root / "run"
            reconcile.start_run(
                reconcile.parse_args(
                    [
                        "start",
                        str(source),
                        "--work",
                        str(work),
                        "--mode",
                        store_ledger.GROUP_MODE,
                        "--date",
                        "2026-09-04",
                    ]
                )
            )
            run = reconcile._load_run(work)
            normalized = reconcile._load_snapshot(work, run)
            group = normalized["groups"][0]
            self.assertEqual(len(group["messages"]), 2)
            self.assertTrue(group["messages"][0]["accounting_context_only"])
            self.assertFalse(group["messages"][1].get("accounting_context_only", False))
            page = reconcile.review_command(
                Namespace(work=work, review_action="next", group="曼谷门店开票群", limit=500)
            )
            self.assertTrue(page["messages"][0]["accounting_context_only"])
            self.assertEqual(page["messages"][1]["reply"]["media_labels"], ["M0001"])

            decision = reconcile._load_json(
                reconcile._decision_path(work, run["groups"][0])
            )
            decision["reviewed_through"] = 2
            decision["read_complete"] = True
            decision["records"] = [
                {
                    "id": "context-only",
                    "record_type": "cash_movement",
                    "posting_status": "pending",
                    "posting_at": "",
                    "source_messages": ["S00001"],
                    "media_roles": [],
                    "representative_media_label": "",
                    "movements": [
                        {
                            "currency": "USD",
                            "sign": "+",
                            "amount": "100",
                            "basis": "chat_explicit",
                            "source_messages": ["S00001"],
                        }
                    ],
                    "status_reason": "仅用于验证区间外上下文",
                }
            ]
            expected = reconcile._decision_template(
                normalized, group, group_mode=store_ledger.GROUP_MODE
            )
            with self.assertRaisesRegex(ValueError, "context-only"):
                store_ledger.validate_decision(
                    normalized,
                    group,
                    decision,
                    expected,
                    require_complete=False,
                    capture_hashes=False,
                    rehash_labels=set(),
                )


if __name__ == "__main__":
    unittest.main()
