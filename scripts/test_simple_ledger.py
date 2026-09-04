#!/usr/bin/env python3
"""Focused regression tests for the default simple bookkeeping mode."""

from __future__ import annotations

import json
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

import build_workbook
import check_workbook
import core
import simple_ledger


GROUP_KEY = "telegram:4909790695"


def message(
    message_id: str,
    *,
    sender_id: str,
    sender_name: str,
    role: str,
    timestamp: str,
    media_id: str | None = None,
    blob_sha256: str | None = None,
) -> dict[str, object]:
    media: list[dict[str, object]] = []
    if media_id:
        media.append(
            {
                "media_id": media_id,
                "kind": "image",
                "availability": "available",
                "path": f"C:/evidence/{media_id.replace(':', '_')}.jpg",
                "blob_sha256": blob_sha256 or f"sha256:{media_id}",
            }
        )
    return {
        "message_id": message_id,
        "sender_id": sender_id,
        "sender_name": sender_name,
        "role": role,
        "timestamp": timestamp,
        "text": "",
        "media": media,
    }


def fund_event(
    event_id: str,
    *,
    message_id: str,
    media_id: str,
    event_type: str,
    amount: str,
    currency: str,
    status_text: str | None,
    status_class: str,
    confidence: str = "high",
    payee: str | None = "测试收款方",
    flow_side: str | None = None,
) -> dict[str, object]:
    effective_payee = (
        "现金" if event_type in {"cash_payment", "cash_payout"} else payee
    )
    effective_payee_state = (
        "cash"
        if event_type in {"cash_payment", "cash_payout"}
        else "not_shown"
        if effective_payee in (None, "", "未显示")
        else "unreadable"
        if effective_payee == "无法辨认"
        else "visible"
    )
    effective_side = flow_side or {
        "payment_screenshot": "payment",
        "cash_payment": "payment",
        "payout_screenshot": "payout",
        "cash_payout": "payout",
        "payment_refund": "payment_refund",
        "payout_recovery": "recovery",
    }.get(event_type)
    event = {
        "event_id": event_id,
        "group_key": GROUP_KEY,
        "message_id": message_id,
        "media_id": media_id,
        "type": event_type,
        "ocr": {
            "amount": amount,
            "currency": currency,
            "payee": effective_payee,
            "payee_state": effective_payee_state,
            "payee_type": "crypto" if currency == "USDT" else "bank",
            "status_text": status_text,
            "status_class": status_class,
            "status_class_confidence": "high",
            "amount_completeness": "complete",
            "confidence": confidence,
        },
    }
    if effective_side is not None:
        event["flow_side"] = effective_side
    return event


def base_fixture(
    *,
    payment_confidence: str = "high",
    payment_status: str = "completed",
) -> tuple[dict, list[dict]]:
    normalized = {
        "timezone": "Asia/Bangkok",
        "source_fingerprint": "sha256:normalized-test",
        "groups": [
            {
                "group_key": GROUP_KEY,
                "platform": "Telegram",
                "group_name": "QQ钱庄～小额出🐱🐱🐱群1",
                "messages": [
                    message(
                        f"{GROUP_KEY}:65727",
                        sender_id="user5570170493",
                        sender_name="梅鮪花黔",
                        role="客户候选",
                        timestamp="2026-08-08T19:08:27+07:00",
                    ),
                    message(
                        f"{GROUP_KEY}:65735",
                        sender_id="user5570170493",
                        sender_name="梅鮪花黔",
                        role="客户候选",
                        timestamp="2026-08-08T19:14:52+07:00",
                        media_id=f"{GROUP_KEY}:65735#media:0",
                    ),
                    message(
                        f"{GROUP_KEY}:65740",
                        sender_id="user6372534512",
                        sender_name="QQ～财务4",
                        role="内部人员",
                        timestamp="2026-08-08T19:31:40+07:00",
                        media_id=f"{GROUP_KEY}:65740#media:0",
                    ),
                ],
            }
        ],
    }
    events = [
        fund_event(
            f"{GROUP_KEY}:65735#payment_screenshot#1",
            message_id=f"{GROUP_KEY}:65735",
            media_id=f"{GROUP_KEY}:65735#media:0",
            event_type="payment_screenshot",
            amount="923",
            currency="USDT",
            status_text="失败" if payment_status == "failed" else None,
            status_class=payment_status,
            confidence=payment_confidence,
        ),
        fund_event(
            f"{GROUP_KEY}:65740#payout_screenshot#1",
            message_id=f"{GROUP_KEY}:65740",
            media_id=f"{GROUP_KEY}:65740#media:0",
            event_type="payout_screenshot",
            amount="30000",
            currency="THB",
            status_text="成功",
            status_class="completed",
        ),
    ]
    return normalized, events


def legacy_simple_fixture(
    *,
    payment_confidence: str = "high",
    payment_status: str = "completed",
) -> tuple[dict, list[dict], dict]:
    """Return the one plan 1.0 fixture kept for compatibility coverage."""
    normalized, events = base_fixture(
        payment_confidence=payment_confidence,
        payment_status=payment_status,
    )
    return normalized, events, {
        "contract_version": simple_ledger.LEGACY_SIMPLE_PLAN_CONTRACT,
        "normalized_source_fingerprint": normalized["source_fingerprint"],
        "events_fingerprint": "sha256:events-test",
        "groups": [
            {
                "group_key": GROUP_KEY,
                "orders": [
                    {
                        "event_ids": [event["event_id"] for event in events],
                        "event_sides": {
                            events[0]["event_id"]: "payment",
                            events[1]["event_id"]: "payout",
                        },
                        "customer_nickname": "梅鮪花黔",
                        "direction": "USDT->THB",
                        "rate": "32.5",
                        "expected_payout": "30000",
                    }
                ],
            }
        ],
    }


def current_contract_fixture(
    *,
    payment_confidence: str = "high",
    payment_status: str = "completed",
) -> tuple[dict, list[dict], dict]:
    normalized, events = base_fixture(
        payment_confidence=payment_confidence,
        payment_status=payment_status,
    )
    payment_message_id = str(events[0]["message_id"])
    payout_message_id = str(events[1]["message_id"])
    context_message_id = str(normalized["groups"][0]["messages"][0]["message_id"])
    events[1]["ocr"]["payee"] = "206-4-xxx781"
    plan = {
        "contract_version": simple_ledger.SIMPLE_PLAN_CONTRACT,
        "normalized_source_fingerprint": normalized["source_fingerprint"],
        "events_fingerprint": "sha256:events-test",
        "groups": [
            {
                "group_key": GROUP_KEY,
                "orders": [
                    {
                        "event_ids": [event["event_id"] for event in events],
                        "event_sides": {
                            events[0]["event_id"]: "payment",
                            events[1]["event_id"]: "payout",
                        },
                        "source_message_ids": [
                            context_message_id,
                            payment_message_id,
                            payout_message_id,
                        ],
                        "customer_nickname": "梅鮪花黔",
                        "direction": "USDT->THB",
                        "pricing": {
                            "source_message_ids": [context_message_id],
                            "terms": {"rate": "32.5", "operator": "multiply"},
                            "expected": {"kind": "explicit", "amount": "30000"},
                        },
                    }
                ],
            }
        ],
    }
    return normalized, events, plan


def fixture(
    *,
    payment_confidence: str = "high",
    payment_status: str = "completed",
) -> tuple[dict, list[dict], dict]:
    """Return the plan 2.0 fixture used by all current business tests."""
    return current_contract_fixture(
        payment_confidence=payment_confidence,
        payment_status=payment_status,
    )


def append_fund(
    normalized: dict,
    events: list[dict],
    *,
    sequence: int,
    sender_id: str,
    sender_name: str,
    role: str,
    timestamp: str,
    event_type: str,
    amount: str,
    currency: str,
    blob_sha256: str | None = None,
) -> dict:
    message_id = f"{GROUP_KEY}:{sequence}"
    media_id = f"{message_id}#media:0"
    normalized["groups"][0]["messages"].append(
        message(
            message_id,
            sender_id=sender_id,
            sender_name=sender_name,
            role=role,
            timestamp=timestamp,
            media_id=media_id,
            blob_sha256=blob_sha256,
        )
    )
    event = fund_event(
        f"{message_id}#{event_type}#1",
        message_id=message_id,
        media_id=media_id,
        event_type=event_type,
        amount=amount,
        currency=currency,
        status_text="成功",
        status_class="completed",
        payee="206-4-xxx781" if currency == "THB" else "测试收款方",
    )
    events.append(event)
    return event


class SimpleLedgerTests(unittest.TestCase):
    def test_v3_formula_applies_only_explicit_rounding(self) -> None:
        pricing = {
            "source_message_ids": ["message-1"],
            "terms": {"rate": "1.005", "operator": "multiply"},
            "expected": {"kind": "calculated_from_terms"},
        }
        exact = simple_ledger._pricing_authority_result(
            pricing,
            payment_total=core.parse_decimal("1", field="payment"),
            payment_currency="CNY",
            payout_currency="THB",
            field="order",
        )
        self.assertEqual(str(exact["expected"]), "1.005")

        rounded_pricing = deepcopy(pricing)
        rounded_pricing["terms"]["rounding"] = {
            "unit": "0.01",
            "mode": "half_up",
            "currency": "THB",
        }
        rounded = simple_ledger._pricing_authority_result(
            rounded_pricing,
            payment_total=core.parse_decimal("1", field="payment"),
            payment_currency="CNY",
            payout_currency="THB",
            field="order",
        )
        self.assertEqual(str(rounded["expected"]), "1.01")

    def test_event_loader_requires_and_normalizes_payee(self) -> None:
        normalized, events, _ = fixture()
        with tempfile.TemporaryDirectory(prefix="simple-events-payee-") as temporary:
            events_dir = Path(temporary)
            event_path = events_dir / "events.jsonl"
            events[0]["ocr"]["payee"] = ""
            event_path.write_text(
                "\n".join(json.dumps(event, ensure_ascii=False) for event in events) + "\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "ocr.payee is required"):
                core.load_events(events_dir, normalized)

            events[0]["ocr"]["payee"] = "  某收款人  "
            event_path.write_text(
                "\n".join(json.dumps(event, ensure_ascii=False) for event in events) + "\n",
                encoding="utf-8",
            )
            loaded, _ = core.load_events(events_dir, normalized)
            self.assertEqual(loaded[0]["ocr"]["payee"], "某收款人")

            events[0]["type"] = "cash_payment"
            events[0]["ocr"]["payee"] = None
            events[0]["ocr"]["payee_state"] = "cash"
            event_path.write_text(
                "\n".join(json.dumps(event, ensure_ascii=False) for event in events) + "\n",
                encoding="utf-8",
            )
            loaded, _ = core.load_events(events_dir, normalized)
            self.assertEqual(loaded[0]["ocr"]["payee"], "现金")

    def test_legacy_simple_plan_contract_remains_supported(self) -> None:
        normalized, events, plan = legacy_simple_fixture()

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        self.assertEqual(orders["groups"][0]["orders"][0]["order_status"], "completed")

    def test_explicit_model_judgments_compile_without_business_inference(self) -> None:
        normalized, events, plan = fixture()

        orders, statistics = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["order_id"], "20260808-001")
        self.assertNotIn("customer_id", order)
        self.assertEqual(order["direction"], "USDT->THB")
        self.assertEqual(order["payment_total"], "923")
        self.assertEqual(order["actual_payout_total"], "30000")
        self.assertEqual(order["expected_payout"], "30000")
        self.assertEqual(order["actual_rate"], "32.5")
        self.assertEqual(order["review_result"], "")
        self.assertEqual(order["order_status"], "completed")
        self.assertTrue(order["flows"][0]["included"])
        self.assertNotIn("inclusion_basis", order["flows"][0])
        self.assertEqual(statistics["groups"], 1)
        self.assertEqual(statistics["pending_orders"], 0)

    def test_uncertain_side_is_unknown_not_zero_and_never_becomes_overpayment(self) -> None:
        normalized, events, plan = fixture(payment_confidence="low")
        plan["groups"][0]["orders"][0]["pricing"]["expected"] = {
            "kind": "calculated_from_terms"
        }

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertIsNone(order["payment_total"])
        self.assertIsNone(order["expected_payout"])
        self.assertEqual(order["actual_payout_total"], "30000")
        self.assertEqual(order["review_result"], "待确认")
        self.assertEqual(order["order_status"], "pending_evidence")
        self.assertNotIn("多转", order["review_result"])

    def test_explicitly_failed_payment_is_not_counted(self) -> None:
        normalized, events, plan = fixture(payment_status="failed")

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        payment = order["flows"][0]
        self.assertFalse(payment["included"])
        self.assertEqual(payment["status"], "failed")
        self.assertFalse(payment["display_in_workbook"])
        self.assertIsNone(build_workbook.flow_row_note(payment))
        self.assertNotIn("客户付款", [row[0] for row in build_workbook.order_rows(order)])
        self.assertIsNone(order["payment_total"])
        self.assertEqual(order["review_result"], "待确认")
        self.assertNotIn("失败", order["anomaly_note"])

    def test_explicit_failure_text_cannot_be_overridden_as_completed(self) -> None:
        normalized, events, plan = fixture()
        events[0]["ocr"]["status_text"] = "失败"
        events[0]["ocr"]["status_class"] = "completed"
        events[0]["ocr"]["status_class_confidence"] = "high"

        with self.assertRaisesRegex(
            ValueError,
            "explicit failure status must use status_class=failed",
        ):
            simple_ledger.compile_simple_ledger(
                normalized,
                events,
                "sha256:events-test",
                plan,
            )

    def test_model_unknown_status_is_not_counted_even_when_amount_is_clear(self) -> None:
        normalized, events, plan = fixture(payment_status="unknown")

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertFalse(order["flows"][0]["included"])
        self.assertIsNone(order["payment_total"])
        self.assertEqual(order["review_result"], "待确认")

    def test_explicit_side_controls_flow_when_screenshot_types_disagree(self) -> None:
        normalized, events, plan = fixture()
        events[0]["type"] = "payout_screenshot"
        events[1]["type"] = "payment_screenshot"

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["flows"][0]["flow_type"], "客户付款")
        self.assertEqual(order["flows"][1]["flow_type"], "内部回款")
        self.assertEqual(order["payment_total"], "923")
        self.assertEqual(order["actual_payout_total"], "30000")

    def test_identical_image_hash_is_not_a_business_duplicate_without_model_declaration(self) -> None:
        normalized, events, plan = fixture()
        original_message = normalized["groups"][0]["messages"][1]
        original_message["media"][0]["blob_sha256"] = "sha256:identical-image"
        duplicate_message_id = f"{GROUP_KEY}:65736"
        duplicate_media_id = f"{duplicate_message_id}#media:0"
        duplicate_message = message(
            duplicate_message_id,
            sender_id="user5570170493",
            sender_name="梅鮪花黔",
            role="客户候选",
            timestamp="2026-08-08T19:15:00+07:00",
            media_id=duplicate_media_id,
            blob_sha256="sha256:identical-image",
        )
        normalized["groups"][0]["messages"].append(duplicate_message)
        duplicate_event = fund_event(
            f"{duplicate_message_id}#payment_screenshot#1",
            message_id=duplicate_message_id,
            media_id=duplicate_media_id,
            event_type="payment_screenshot",
            amount="923",
            currency="USDT",
            status_text=None,
            status_class="completed",
        )
        events.append(duplicate_event)
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["event_ids"].insert(1, duplicate_event["event_id"])
        raw_order["event_sides"][duplicate_event["event_id"]] = "payment"
        raw_order["source_message_ids"].append(duplicate_event["message_id"])

        orders, statistics = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["payment_total"], "1846")
        self.assertTrue(order["flows"][1]["included"])
        self.assertIsNone(order["flows"][1]["duplicate_of"])
        self.assertFalse(any("duplicate" in warning for warning in statistics["warnings"]))

    def test_unassigned_fund_event_is_rejected_instead_of_becoming_an_auto_order(self) -> None:
        normalized, events, plan = fixture()
        plan["groups"][0]["orders"][0]["event_ids"] = [events[0]["event_id"]]
        plan["groups"][0]["orders"][0]["event_sides"] = {
            events[0]["event_id"]: "payment"
        }

        with self.assertRaisesRegex(ValueError, "must be explicitly assigned by the model"):
            simple_ledger.compile_simple_ledger(
                normalized,
                events,
                "sha256:events-test",
                plan,
            )

    def test_current_plan_requires_explicit_event_side(self) -> None:
        normalized, events, plan = fixture()
        plan["groups"][0]["orders"][0].pop("event_sides")

        with self.assertRaisesRegex(ValueError, "requires an explicit side for every event"):
            simple_ledger.compile_simple_ledger(
                normalized,
                events,
                "sha256:events-test",
                plan,
            )

    def test_internal_sender_can_relay_customer_payment_with_explicit_side(self) -> None:
        normalized, events, plan = fixture()
        payment_message = normalized["groups"][0]["messages"][1]
        payment_message["sender_id"] = "user6372534512"
        payment_message["sender_name"] = "QQ～财务4"
        payment_message["role"] = "内部人员"
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["event_sides"] = {
            events[0]["event_id"]: "payment",
            events[1]["event_id"]: "payout",
        }
        raw_order["customer_nickname"] = "梅鮪花黔"
        context_message_id = raw_order["source_message_ids"][0]
        events[0]["side_exception"] = {
            "kind": "relayed_customer_payment",
            "source_message_ids": [context_message_id],
            "detail": "内部人员代客户转发付款凭证",
        }

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["flows"][0]["side"], "payment")
        self.assertNotIn("sender_role", order["flows"][0])
        self.assertEqual(order["payment_total"], "923")
        self.assertEqual(order["order_status"], "completed")

    def test_current_contract_rejects_internal_sender_payment_without_side_exception(self) -> None:
        normalized, events, plan = current_contract_fixture()
        payment_message = normalized["groups"][0]["messages"][1]
        payment_message["sender_id"] = "user6372534512"
        payment_message["sender_name"] = "QQ～财务4"
        payment_message["role"] = "内部人员"

        with self.assertRaisesRegex(
            ValueError,
            "internal sender.*ordinary fund event.*side=payout",
        ):
            simple_ledger.compile_simple_ledger(
                normalized,
                events,
                "sha256:events-test",
                plan,
            )

    def test_current_contract_order_side_must_match_validated_event_side(self) -> None:
        normalized, events, plan = current_contract_fixture()
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["event_sides"][events[0]["event_id"]] = "payout"

        with self.assertRaisesRegex(
            ValueError,
            "event_sides.*must match event.flow_side",
        ):
            simple_ledger.compile_simple_ledger(
                normalized,
                events,
                "sha256:events-test",
                plan,
            )

    def test_current_contract_side_exception_basis_must_belong_to_order(self) -> None:
        normalized, events, plan = current_contract_fixture()
        payment_message = normalized["groups"][0]["messages"][1]
        payment_message["sender_id"] = "user6372534512"
        payment_message["sender_name"] = "QQ～财务4"
        payment_message["role"] = "内部人员"
        context_message_id = str(normalized["groups"][0]["messages"][0]["message_id"])
        events[0]["side_exception"] = {
            "kind": "relayed_customer_payment",
            "source_message_ids": [context_message_id],
            "detail": "内部人员代客户转发付款凭证",
        }
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["source_message_ids"].remove(context_message_id)

        with self.assertRaisesRegex(
            ValueError,
            "side_exception.source_message_ids must belong to simple order",
        ):
            simple_ledger.compile_simple_ledger(
                normalized,
                events,
                "sha256:events-test",
                plan,
            )

    def test_current_contract_rejects_pending_for_normal_processing_status(self) -> None:
        for status_text in ("待区块确认", "Pending", "Processing", "Confirming", "Unconfirmed"):
            with self.subTest(status_text=status_text):
                normalized, events, plan = current_contract_fixture()
                events[0]["ocr"]["status_text"] = status_text
                events[0]["ocr"]["status_class"] = "pending"

                with self.assertRaisesRegex(
                    ValueError,
                    "normal processing status.*status_class=completed",
                ):
                    simple_ledger.compile_simple_ledger(
                        normalized,
                        events,
                        "sha256:events-test",
                        plan,
                    )

    def test_current_contract_counts_normal_processing_status_as_completed(self) -> None:
        normalized, events, plan = current_contract_fixture()
        events[0]["ocr"]["status_text"] = "确认中"
        events[0]["ocr"]["status_class"] = "completed"

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )
        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["payment_total"], "923")
        self.assertTrue(order["flows"][0]["included"])

    def test_current_contract_does_not_count_explicit_failed_fund_attempt(self) -> None:
        normalized, events, plan = current_contract_fixture()
        events[0]["ocr"]["status_text"] = "交易失败"
        events[0]["ocr"]["status_class"] = "failed"

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )
        order = orders["groups"][0]["orders"][0]
        self.assertIsNone(order["payment_total"])
        self.assertFalse(order["flows"][0]["included"])

    def test_different_screenshots_of_same_transaction_are_counted_once(self) -> None:
        normalized, events, plan = fixture()
        duplicate = append_fund(
            normalized,
            events,
            sequence=65736,
            sender_id="user5570170493",
            sender_name="梅鮪花黔",
            role="客户候选",
            timestamp="2026-08-08T19:15:00+07:00",
            event_type="payment_screenshot",
            amount="923",
            currency="USDT",
            blob_sha256="sha256:different-crop",
        )
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["event_ids"].insert(1, duplicate["event_id"])
        raw_order["event_sides"][duplicate["event_id"]] = "payment"
        raw_order["source_message_ids"].append(duplicate["message_id"])
        raw_order["same_transactions"] = [
            {"event_id": duplicate["event_id"], "same_as": events[0]["event_id"]}
        ]

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        duplicate_flow = order["flows"][1]
        self.assertEqual(order["payment_total"], "923")
        self.assertFalse(duplicate_flow["included"])
        self.assertEqual(duplicate_flow["duplicate_basis"], "declared_same_transaction")
        self.assertFalse(duplicate_flow["display_in_workbook"])
        self.assertIsNone(build_workbook.flow_row_note(duplicate_flow))
        self.assertEqual(
            [row[0] for row in build_workbook.order_rows(order)].count("客户付款"),
            1,
        )
        self.assertNotIn("同一笔交易", order["anomaly_note"])

    def test_declared_same_transaction_summary_without_payee_does_not_downgrade_order(self) -> None:
        normalized, events, plan = fixture()
        summary = events[0]
        summary["ocr"]["payee"] = "未显示"
        summary["ocr"]["payee_state"] = "not_shown"
        detail = append_fund(
            normalized,
            events,
            sequence=65736,
            sender_id="user5570170493",
            sender_name="梅鮪花黔",
            role="客户候选",
            timestamp="2026-08-08T19:15:00+07:00",
            event_type="payment_screenshot",
            amount="923",
            currency="USDT",
            blob_sha256="sha256:detail-page",
        )
        detail["ocr"]["payee"] = "TZ7nsfaXoJCsNytYFcT5yEZLXMAH87QKUg"
        detail["ocr"]["payee_state"] = "visible"
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["event_ids"].insert(1, detail["event_id"])
        raw_order["event_sides"][detail["event_id"]] = "payment"
        raw_order["source_message_ids"].append(detail["message_id"])
        raw_order["same_transactions"] = [
            {"event_id": summary["event_id"], "same_as": detail["event_id"]}
        ]

        orders, statistics = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        summary_flow = order["flows"][0]
        self.assertFalse(summary_flow["included"])
        self.assertEqual(summary_flow["duplicate_basis"], "declared_same_transaction")
        self.assertEqual(summary_flow["payee"], "未显示")
        self.assertNotIn("evidence_pending", summary_flow)
        self.assertEqual(order["payment_total"], "923")
        self.assertEqual(order["order_status"], "completed")
        self.assertEqual(order["review_result"], "")
        self.assertNotIn("收款方未显示", order["anomaly_note"] or "")
        self.assertEqual(statistics["pending_orders"], 0)

    def test_missing_payee_is_recorded_without_downgrading_order(self) -> None:
        for payee in ("未显示", "无法辨认"):
            with self.subTest(payee=payee):
                normalized, events, plan = fixture()
                events[0]["ocr"]["payee"] = payee
                events[0]["ocr"]["payee_state"] = {
                    "未显示": "not_shown",
                    "无法辨认": "unreadable",
                }[payee]

                orders, statistics = simple_ledger.compile_simple_ledger(
                    normalized,
                    events,
                    "sha256:events-test",
                    plan,
                )

                order = orders["groups"][0]["orders"][0]
                payment = order["flows"][0]
                self.assertEqual(payment["payee"], payee)
                self.assertNotIn("evidence_pending", payment)
                self.assertEqual(order["order_status"], "completed")
                self.assertEqual(order["review_result"], "")
                self.assertNotIn("收款方", order["anomaly_note"] or "")
                self.assertEqual(statistics["pending_orders"], 0)

    def test_simple_mode_requires_and_records_payee(self) -> None:
        normalized, events, plan = fixture()
        events[0]["ocr"]["payee"] = "某收款人"
        events[0]["ocr"]["payee_type"] = "bank"

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        payment = orders["groups"][0]["orders"][0]["flows"][0]
        self.assertEqual(payment["payee"], "某收款人")
        self.assertNotIn("payee_type", payment)
        self.assertNotIn("payee_source", payment)
        payment_row = next(
            row
            for row in build_workbook.order_rows(orders["groups"][0]["orders"][0])
            if row[0] == "客户付款"
        )
        self.assertEqual(core.HEADERS[-2], "收款方")
        self.assertEqual(core.HEADERS[-1], "聊天消息时间")
        self.assertNotIn("客户标识", core.HEADERS)
        self.assertEqual(len(payment_row), len(core.HEADERS))
        self.assertEqual(payment_row[core.HEADERS.index("收款方")], "某收款人")

        events[0]["ocr"]["payee"] = " "
        with self.assertRaisesRegex(ValueError, "payee is required"):
            simple_ledger.compile_simple_ledger(
                normalized,
                events,
                "sha256:events-test",
                plan,
            )

        orders["groups"][0]["orders"][0]["flows"][0]["payee"] = ""
        with self.assertRaisesRegex(ValueError, "payee is required"):
            build_workbook.validate_orders(orders)

        orders["groups"][0]["orders"][0]["flows"][0]["payee"] = "群内收款方"
        with self.assertRaisesRegex(ValueError, "generic placeholder"):
            build_workbook.validate_orders(orders)

    def test_detailed_payee_text_round_trips_to_workbook_unchanged(self) -> None:
        for payee in ("王少秋（**秋）", "206-4-xxx781"):
            with self.subTest(payee=payee):
                normalized, events, plan = fixture()
                events[0]["ocr"]["payee"] = payee

                orders, _ = simple_ledger.compile_simple_ledger(
                    normalized,
                    events,
                    "sha256:events-test",
                    plan,
                )

                payment = orders["groups"][0]["orders"][0]["flows"][0]
                self.assertEqual(payment["payee"], payee)
                payment_row = next(
                    row
                    for row in build_workbook.order_rows(orders["groups"][0]["orders"][0])
                    if row[0] == "客户付款"
                )
                self.assertEqual(payment_row[core.HEADERS.index("收款方")], payee)

    def test_cash_fund_flow_records_cash_as_payee(self) -> None:
        normalized, events, plan = fixture()
        events[0]["type"] = "cash_payment"
        events[0]["ocr"]["payee"] = None
        events[0]["ocr"]["payee_state"] = "cash"

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        payment = orders["groups"][0]["orders"][0]["flows"][0]
        self.assertEqual(payment["payee"], "现金")

    def test_cash_payout_with_formatted_amount_text_uses_normalized_amount(self) -> None:
        normalized, events, plan = fixture()
        payout = events[1]
        payout["type"] = "cash_payout"
        payout["ocr"]["amount"] = "100000"
        payout["ocr"]["amount_text"] = "100,000"
        payout["ocr"]["payee"] = None
        payout["ocr"]["payee_state"] = "cash"
        plan["groups"][0]["orders"][0]["pricing"]["expected"] = {
            "kind": "explicit",
            "amount": "100000",
        }

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        payout_flow = next(flow for flow in order["flows"] if flow["event_id"] == payout["event_id"])
        self.assertEqual(payout_flow["amount"], "100000")
        self.assertTrue(payout_flow["included"])
        self.assertEqual(order["actual_payout_total"], "100000")
        self.assertEqual(order["order_status"], "completed")

    def test_division_rate_is_calculated_and_displayed_explicitly(self) -> None:
        normalized, events, plan = fixture()
        events[0]["ocr"]["amount"] = "1000"
        events[0]["ocr"]["currency"] = "CNY"
        events[1]["ocr"]["amount"] = "200"
        events[1]["ocr"]["currency"] = "USDT"
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["direction"] = "CNY->USDT"
        raw_order["pricing"] = {
            "source_message_ids": raw_order["pricing"]["source_message_ids"],
            "terms": {"rate": "5", "operator": "divide"},
            "expected": {"kind": "calculated_from_terms"},
        }

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["expected_payout"], "200")
        self.assertEqual(order["actual_rate_display"], "÷5")
        self.assertEqual(order["review_result"], "")

    def test_customer_requested_network_fee_and_other_fees_are_recorded_and_applied(self) -> None:
        normalized, events, plan = fixture()
        events[0]["ocr"]["amount"] = "102"
        events[1]["ocr"]["amount"] = "3245"
        raw_order = plan["groups"][0]["orders"][0]
        fees = [
            {
                "kind": "delivery_fee",
                "amount": "2",
                "currency": "USDT",
                "treatment": "added_to_payment",
            },
            {
                "kind": "service_fee",
                "amount": "4",
                "currency": "THB",
                "treatment": "deducted_from_payout",
            },
            {
                "kind": "network_fee",
                "amount": "1",
                "currency": "THB",
                "treatment": "deducted_from_payout",
                "customer_requested": True,
            },
        ]
        raw_order["pricing"] = {
            "source_message_ids": raw_order["pricing"]["source_message_ids"],
            "terms": {
                "rate": "32.5",
                "operator": "multiply",
                "fees": fees,
                "rounding": {"unit": "5", "mode": "down", "currency": "THB"},
            },
            "expected": {"kind": "calculated_from_terms"},
        }

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["expected_payout"], "3245")
        self.assertEqual(len(order["fee_adjustments"]), 3)
        self.assertEqual(
            sum(fee["kind"] == "network_fee" for fee in order["fee_adjustments"]),
            1,
        )
        self.assertEqual(order["rounding"], {"unit": "5", "mode": "down", "currency": "THB"})
        self.assertEqual(order["anomaly_note"], "")
        pricing_rows = build_workbook.pricing_detail_rows(order)
        self.assertEqual(
            [row[0] for row in pricing_rows],
            ["配送费", "手续费", "网络费用", "舍入"],
        )
        self.assertEqual(
            [row[core.HEADERS.index("备注")] for row in pricing_rows[:3]],
            [None, None, None],
        )
        self.assertIsNone(pricing_rows[3][core.HEADERS.index("备注")])
        self.assertEqual(order["review_result"], "")

    def test_deducted_delivery_fee_row_has_no_normal_process_note(self) -> None:
        normalized, events, plan = fixture()
        events[0]["ocr"]["amount"] = "3102"
        events[1]["ocr"]["amount"] = "100000"
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["pricing"] = {
            "source_message_ids": raw_order["pricing"]["source_message_ids"],
            "terms": {
                "rate": "32.4",
                "operator": "multiply",
                "fees": [
                    {
                        "kind": "delivery_fee",
                        "amount": "500",
                        "currency": "THB",
                        "treatment": "deducted_from_payout",
                    }
                ],
            },
            "expected": {"kind": "explicit", "amount": "100000"},
        }

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["expected_payout"], "100000")
        self.assertEqual(order["actual_payout_total"], "100000")
        self.assertEqual(order["review_result"], "")
        fee_row = build_workbook.pricing_detail_rows(order)[0]
        self.assertEqual(fee_row[0], "配送费")
        self.assertEqual(fee_row[core.HEADERS.index("收款方实际到账金额")], 500)
        self.assertEqual(fee_row[core.HEADERS.index("流水币种")], "THB")
        self.assertIsNone(fee_row[core.HEADERS.index("备注")])

    def test_platform_displayed_network_fee_is_not_recorded_or_applied(self) -> None:
        normalized, events, plan = fixture()
        events[0]["ocr"]["amount"] = "100"
        events[1]["ocr"]["amount"] = "3250"
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["pricing"] = {
            "source_message_ids": raw_order["pricing"]["source_message_ids"],
            "terms": {
                "rate": "32.5",
                "operator": "multiply",
                "fees": [
                    {
                        "kind": "network_fee",
                        "amount": "1",
                        "currency": "THB",
                        "treatment": "deducted_from_payout",
                    }
                ],
            },
            "expected": {"kind": "calculated_from_terms"},
        }

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["expected_payout"], "3250")
        self.assertEqual(order["fee_adjustments"], [])
        self.assertEqual(build_workbook.pricing_detail_rows(order), [])
        self.assertEqual(order["review_result"], "")

    def test_customer_requested_network_fee_must_be_deducted(self) -> None:
        normalized, events, plan = fixture()
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["pricing"]["terms"]["fees"] = [
            {
                "kind": "network_fee",
                "amount": "1",
                "currency": "THB",
                "treatment": "included_in_quote",
                "customer_requested": True,
            }
        ]

        with self.assertRaisesRegex(ValueError, "must deduct a customer-requested network fee"):
            simple_ledger.compile_simple_ledger(
                normalized,
                events,
                "sha256:events-test",
                plan,
            )

    def test_explicit_unknown_customer_is_not_replaced_from_sender_metadata(self) -> None:
        normalized, events, plan = fixture()
        second_payment = append_fund(
            normalized,
            events,
            sequence=65736,
            sender_id="another-customer",
            sender_name="另一位客户",
            role="客户候选",
            timestamp="2026-08-08T19:15:00+07:00",
            event_type="payment_screenshot",
            amount="77",
            currency="USDT",
        )
        events[1]["ocr"]["amount"] = "32500"
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["event_ids"].insert(1, second_payment["event_id"])
        raw_order["event_sides"][second_payment["event_id"]] = "payment"
        raw_order["source_message_ids"].append(second_payment["message_id"])
        raw_order["customer_nickname"] = ""
        raw_order["pricing"]["expected"] = {
            "kind": "explicit",
            "amount": "32500",
        }

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["payment_total"], "1000")
        self.assertNotIn("customer_id", order)
        self.assertEqual(order["review_result"], "待确认")
        self.assertEqual(order["order_status"], "pending_identity")
        self.assertIn("客户无法唯一确认", order["anomaly_note"])

    def test_explicit_unknown_direction_is_not_derived_from_flow_currencies(self) -> None:
        normalized, events, plan = fixture()
        plan["groups"][0]["orders"][0]["direction"] = ""

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["direction"], "")
        self.assertEqual(order["review_result"], "待确认")
        self.assertNotEqual(order["order_status"], "completed")

    def test_one_payment_can_close_two_exchange_legs_without_anomaly(self) -> None:
        normalized, events, plan = fixture()
        events[0]["ocr"]["amount"] = "1000"
        events[1]["ocr"]["amount"] = "19500"
        second_payout = append_fund(
            normalized,
            events,
            sequence=65741,
            sender_id="user6372534512",
            sender_name="QQ～财务4",
            role="内部人员",
            timestamp="2026-08-08T19:32:40+07:00",
            event_type="payout_screenshot",
            amount="2000",
            currency="CNY",
        )
        raw_order = plan["groups"][0]["orders"][0]
        raw_order["event_ids"].append(second_payout["event_id"])
        raw_order["event_sides"][second_payout["event_id"]] = "payout"
        raw_order["source_message_ids"].append(second_payout["message_id"])
        pricing_source_ids = raw_order["pricing"]["source_message_ids"]
        raw_order.pop("pricing")
        raw_order["legs"] = [
            {
                "leg_id": "thb",
                "direction": "USDT->THB",
                "allocation_amount": "600",
                "pricing": {
                    "source_message_ids": pricing_source_ids,
                    "terms": {"rate": "32.5", "operator": "multiply"},
                    "expected": {"kind": "calculated_from_terms"},
                },
                "payout_event_ids": [events[1]["event_id"]],
                "recovery_event_ids": [],
            },
            {
                "leg_id": "cny",
                "direction": "USDT->CNY",
                "allocation_amount": "400",
                "pricing": {
                    "source_message_ids": pricing_source_ids,
                    "terms": {"rate": "5", "operator": "multiply"},
                    "expected": {"kind": "calculated_from_terms"},
                },
                "payout_event_ids": [second_payout["event_id"]],
                "recovery_event_ids": [],
            },
        ]

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["direction"], "USDT->THB\nUSDT->CNY")
        self.assertEqual(order["actual_rate_display"], "USDT->THB：32.5\nUSDT->CNY：5")
        self.assertEqual(order["review_result"], "")
        self.assertEqual(order["anomaly_note"], "")
        self.assertEqual(order["note"], "")
        self.assertEqual(order["order_status"], "completed")
        rows = build_workbook.order_rows(order)
        leg_rows = [row for row in rows if row[0] == "换汇明细"]
        self.assertEqual(leg_rows, [])
        summary_row = next(row for row in rows if row[0] == "订单汇总")
        self.assertIsNone(summary_row[core.HEADERS.index("备注")])
        self.assertEqual(
            [row[0] for row in rows if row[0] in {"客户付款", "内部回款"}],
            ["客户付款", "内部回款", "内部回款"],
        )

    def test_bern_style_split_labels_only_conflicted_crypto_legs_pending(self) -> None:
        normalized, events, plan = current_contract_fixture()
        events[0]["ocr"]["amount"] = "6000"
        events[1]["ocr"]["amount"] = "32000"
        cash_payout = append_fund(
            normalized,
            events,
            sequence=65741,
            sender_id="user6372534512",
            sender_name="QQ～财务4",
            role="内部人员",
            timestamp="2026-08-24T15:48:31+07:00",
            event_type="cash_payout",
            amount="158860",
            currency="THB",
        )
        trx_initial = append_fund(
            normalized,
            events,
            sequence=65742,
            sender_id="user6372534512",
            sender_name="QQ～财务4",
            role="内部人员",
            timestamp="2026-08-24T15:37:15+07:00",
            event_type="payout_screenshot",
            amount="7.75",
            currency="TRX",
        )
        trx_followup = append_fund(
            normalized,
            events,
            sequence=65743,
            sender_id="user6372534512",
            sender_name="QQ～财务4",
            role="内部人员",
            timestamp="2026-08-24T17:41:43+07:00",
            event_type="payout_screenshot",
            amount="44",
            currency="TRX",
        )
        for event in (cash_payout, trx_initial, trx_followup):
            event["flow_side"] = "payout"
            event["ocr"]["payee_state"] = (
                "cash" if event["type"] == "cash_payout" else "visible"
            )

        raw_order = plan["groups"][0]["orders"][0]
        new_events = [cash_payout, trx_initial, trx_followup]
        raw_order["event_ids"].extend(event["event_id"] for event in new_events)
        raw_order["event_sides"].update(
            {event["event_id"]: "payout" for event in new_events}
        )
        raw_order["source_message_ids"].extend(
            event["message_id"] for event in new_events
        )
        context_message_id = raw_order["source_message_ids"][0]
        raw_order["direction"] = ""
        raw_order.pop("pricing")
        raw_order["note"] = (
            "20 USDT 的 TRX 部分同时出现预期49、声称到账30和凭证44，"
            "没有最终数值确认。"
        )
        raw_order["legs"] = [
            {
                "leg_id": "thb_transfer",
                "display_label": "USDT->THB（转账）",
                "direction": "USDT->THB",
                "allocation_amount": "1000",
                "pricing": {
                    "source_message_ids": [context_message_id],
                    "terms": {"rate": "32", "operator": "multiply"},
                    "expected": {"kind": "explicit", "amount": "32000"},
                },
                "payout_event_ids": [events[1]["event_id"]],
                "recovery_event_ids": [],
            },
            {
                "leg_id": "thb_cash",
                "display_label": "USDT->THB（现金）",
                "direction": "USDT->THB",
                "allocation_amount": "4980",
                "pricing": {
                    "source_message_ids": [context_message_id],
                    "terms": {"rate": "31.9", "operator": "multiply"},
                    "expected": {"kind": "explicit", "amount": "158860"},
                },
                "payout_event_ids": [cash_payout["event_id"]],
                "recovery_event_ids": [],
            },
            {
                "leg_id": "trx_initial",
                "display_label": "USDT->TRX（首段3U）",
                "direction": "USDT->TRX",
                "allocation_amount": "3",
                "pricing": {
                    "source_message_ids": [context_message_id],
                    "expected": {"kind": "unknown", "reason": "not_stated"},
                },
                "payout_event_ids": [trx_initial["event_id"]],
                "recovery_event_ids": [],
            },
            {
                "leg_id": "trx_followup",
                "display_label": "USDT->TRX（后段17U）",
                "direction": "USDT->TRX",
                "allocation_amount": "17",
                "pricing": {
                    "source_message_ids": [context_message_id],
                    "expected": {
                        "kind": "unknown",
                        "reason": "conflicting_authority",
                    },
                },
                "payout_event_ids": [trx_followup["event_id"]],
                "recovery_event_ids": [],
            },
        ]

        orders, statistics = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(
            order["direction"],
            "USDT->THB（转账）\nUSDT->THB（现金）\nUSDT->TRX（首段3U）\nUSDT->TRX（后段17U）",
        )
        self.assertEqual(
            order["actual_rate_display"],
            "USDT->THB（转账）：32\nUSDT->THB（现金）：31.9\n"
            "USDT->TRX（首段3U）：待确认\nUSDT->TRX（后段17U）：待确认",
        )
        self.assertEqual(
            order["review_result"],
            "USDT->TRX（首段3U）：待确认\nUSDT->TRX（后段17U）：待确认",
        )
        self.assertEqual(order["order_status"], "pending_pricing")
        self.assertEqual(statistics["pending_orders"], 1)
        rows = build_workbook.order_rows(order)
        summary = rows[0]
        self.assertEqual(summary[core.HEADERS.index("备注")], raw_order["note"])
        self.assertNotIn("群聊中的最终金额", summary[core.HEADERS.index("备注")])
        payout_directions = [
            row[core.HEADERS.index("换汇方向")]
            for row in rows
            if row[0] == "内部回款"
        ]
        self.assertEqual(
            payout_directions,
            [
                "USDT->THB（转账）",
                "USDT->TRX（首段3U）",
                "USDT->THB（现金）",
                "USDT->TRX（后段17U）",
            ],
        )

    def test_order_row_note_uses_only_model_note(self) -> None:
        self.assertEqual(
            build_workbook.order_row_note(
                {"note": "具体问题", "anomaly_note": "自动生成的重复说明"}
            ),
            "具体问题",
        )
        self.assertEqual(
            build_workbook.order_row_note({"note": "", "anomaly_note": "实际异常"}),
            None,
        )
        self.assertEqual(
            build_workbook.order_row_note(
                {
                    "note": "具体冲突只写一次",
                    "reconciliation": {
                        "status": "pending",
                        "detail": "脚本通用待确认说明",
                    },
                }
            ),
            "具体冲突只写一次",
        )

    def test_payment_refund_and_payout_recovery_use_net_amounts(self) -> None:
        normalized, events, plan = fixture()
        events[0]["ocr"]["amount"] = "1000"
        events[1]["ocr"]["amount"] = "30000"
        refund = append_fund(
            normalized,
            events,
            sequence=65736,
            sender_id="user6372534512",
            sender_name="QQ～财务4",
            role="内部人员",
            timestamp="2026-08-08T19:20:00+07:00",
            event_type="payout_screenshot",
            amount="100",
            currency="USDT",
        )
        recovery = append_fund(
            normalized,
            events,
            sequence=65741,
            sender_id="user5570170493",
            sender_name="梅鮪花黔",
            role="客户候选",
            timestamp="2026-08-08T19:35:00+07:00",
            event_type="payment_screenshot",
            amount="750",
            currency="THB",
        )
        raw_order = plan["groups"][0]["orders"][0]
        context_message_id = raw_order["source_message_ids"][0]
        refund["flow_side"] = "payment_refund"
        refund["side_exception"] = {
            "kind": "explicit_payment_refund",
            "source_message_ids": [context_message_id],
            "detail": "聊天明确该笔为付款退款",
        }
        recovery["flow_side"] = "recovery"
        recovery["side_exception"] = {
            "kind": "explicit_recovery",
            "source_message_ids": [context_message_id],
            "detail": "聊天明确该笔为回款追回",
        }
        raw_order["event_ids"] = [event["event_id"] for event in events]
        raw_order["event_sides"] = {
            events[0]["event_id"]: "payment",
            events[1]["event_id"]: "payout",
            refund["event_id"]: "payment_refund",
            recovery["event_id"]: "recovery",
        }
        raw_order["source_message_ids"].extend(
            [refund["message_id"], recovery["message_id"]]
        )
        raw_order["pricing"] = {
            "source_message_ids": [context_message_id],
            "terms": {"rate": "32.5", "operator": "multiply"},
            "expected": {"kind": "calculated_from_terms"},
        }

        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        order = orders["groups"][0]["orders"][0]
        self.assertEqual(order["payment_total"], "900")
        self.assertEqual(order["actual_payout_total"], "29250")
        self.assertEqual(order["expected_payout"], "29250")
        self.assertEqual([flow["flow_type"] for flow in order["flows"]], [
            "客户付款", "付款退款", "内部回款", "回款追回"
        ])
        self.assertEqual(order["review_result"], "")

    def test_explicit_cross_order_shortfall_and_overpayment_are_applied(self) -> None:
        for kind, source_payout, target_payout, expected_adjustment in [
            ("shortfall_carryover", "3200", "3300", "50"),
            ("overpayment_carryover", "3300", "3200", "-50"),
        ]:
            with self.subTest(kind=kind):
                normalized, _, plan = fixture()
                normalized["groups"][0]["messages"] = []
                events: list[dict] = []
                first_payment = append_fund(
                    normalized, events, sequence=70001, sender_id="customer-1", sender_name="客户甲",
                    role="客户候选", timestamp="2026-08-08T10:00:00+07:00",
                    event_type="payment_screenshot", amount="100", currency="USDT",
                )
                first_payout = append_fund(
                    normalized, events, sequence=70002, sender_id="staff-1", sender_name="财务",
                    role="内部人员", timestamp="2026-08-08T10:05:00+07:00",
                    event_type="payout_screenshot", amount=source_payout, currency="THB",
                )
                second_payment = append_fund(
                    normalized, events, sequence=70003, sender_id="customer-1", sender_name="客户甲",
                    role="客户候选", timestamp="2026-08-08T11:00:00+07:00",
                    event_type="payment_screenshot", amount="100", currency="USDT",
                )
                second_payout = append_fund(
                    normalized, events, sequence=70004, sender_id="staff-1", sender_name="财务",
                    role="内部人员", timestamp="2026-08-08T11:05:00+07:00",
                    event_type="payout_screenshot", amount=target_payout, currency="THB",
                )
                plan["groups"][0]["orders"] = [
                    {
                        "case_id": "first",
                        "event_ids": [first_payment["event_id"], first_payout["event_id"]],
                        "event_sides": {
                            first_payment["event_id"]: "payment",
                            first_payout["event_id"]: "payout",
                        },
                        "source_message_ids": [
                            first_payment["message_id"],
                            first_payout["message_id"],
                        ],
                        "customer_nickname": "客户甲",
                        "direction": "USDT->THB",
                        "pricing": {
                            "source_message_ids": [first_payment["message_id"]],
                            "terms": {"rate": "32.5", "operator": "multiply"},
                            "expected": {"kind": "calculated_from_terms"},
                        },
                    },
                    {
                        "case_id": "second",
                        "event_ids": [second_payment["event_id"], second_payout["event_id"]],
                        "event_sides": {
                            second_payment["event_id"]: "payment",
                            second_payout["event_id"]: "payout",
                        },
                        "source_message_ids": [
                            second_payment["message_id"],
                            second_payout["message_id"],
                        ],
                        "customer_nickname": "客户甲",
                        "direction": "USDT->THB",
                        "pricing": {
                            "source_message_ids": [second_payment["message_id"]],
                            "terms": {"rate": "32.5", "operator": "multiply"},
                            "expected": {"kind": "calculated_from_terms"},
                        },
                    },
                ]
                plan["groups"][0]["balance_links"] = [
                    {
                        "source_case_id": "first",
                        "target_case_id": "second",
                        "kind": kind,
                        "amount": "50",
                        "currency": "THB",
                        "source_message_ids": [
                            first_payment["message_id"],
                            second_payment["message_id"],
                        ],
                    }
                ]

                orders, _ = simple_ledger.compile_simple_ledger(
                    normalized,
                    events,
                    "sha256:events-test",
                    plan,
                )

                first, second = orders["groups"][0]["orders"]
                self.assertEqual(second["balance_adjustment"], expected_adjustment)
                self.assertEqual(second["expected_payout"], target_payout)
                self.assertEqual(second["review_result"], "")
                self.assertIn("20260808-002", first["anomaly_note"])
                self.assertIn("20260808-001", second["anomaly_note"])
                self.assertEqual(orders["groups"][0]["balance_links"][0]["status"], "applied")

                plan["groups"][0]["balance_links"][0]["amount"] = "40"
                invalid_orders, _ = simple_ledger.compile_simple_ledger(
                    normalized,
                    events,
                    "sha256:events-test",
                    plan,
                )
                invalid_first, invalid_second = invalid_orders["groups"][0]["orders"]
                self.assertEqual(invalid_first["review_result"], "待确认")
                self.assertEqual(invalid_second["review_result"], "待确认")
                self.assertEqual(
                    invalid_orders["groups"][0]["balance_links"][0]["status"],
                    "pending",
                )

    def test_workbook_freezes_centers_and_colors_review_results(self) -> None:
        normalized, events, plan = fixture()
        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )
        short_order = deepcopy(orders["groups"][0]["orders"][0])
        short_order["reconciliation"] = {
            "status": "short",
            "difference": "-10",
            "currency": "THB",
            "reason": None,
            "detail": None,
        }
        short_order["review_result"] = "少转 10 THB"
        over_order = deepcopy(short_order)
        over_order["order_id"] = "20260808-002"
        over_order["reconciliation"] = {
            "status": "over",
            "difference": "20",
            "currency": "THB",
            "reason": None,
            "detail": None,
        }
        over_order["review_result"] = "多转 20 THB"
        orders["groups"][0]["orders"] = [short_order, over_order]

        with tempfile.TemporaryDirectory(prefix="simple-ledger-format-") as temporary:
            root = Path(temporary)
            orders_path = root / "orders.json"
            workbook_path = root / "ledger.xlsx"
            core.atomic_json(orders_path, orders)
            groups = build_workbook.validate_orders(orders)
            workbook = build_workbook.build_workbook(
                groups,
                Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
            )
            try:
                workbook.save(workbook_path)
            finally:
                workbook.close()

            self.assertEqual(check_workbook.check(workbook_path, orders_path), [])
            workbook = load_workbook(workbook_path, read_only=False, data_only=False)
            try:
                worksheet = workbook.worksheets[0]
                self.assertEqual(worksheet.freeze_panes, "A2")
                for row in worksheet.iter_rows(
                    min_row=1,
                    max_row=worksheet.max_row,
                    min_col=1,
                    max_col=len(core.HEADERS),
                ):
                    for cell in row:
                        self.assertEqual(cell.alignment.horizontal, "center")
                        self.assertEqual(cell.alignment.vertical, "center")

                payee_column = core.HEADERS.index("收款方") + 1
                payee_cells = [
                    worksheet.cell(row, payee_column)
                    for row in range(2, worksheet.max_row + 1)
                    if worksheet.cell(row, payee_column).value is not None
                ]
                self.assertTrue(payee_cells)
                for cell in payee_cells:
                    self.assertEqual(cell.data_type, "s")
                    self.assertEqual(cell.number_format, "@")

                review_column = core.HEADERS.index("核对结果") + 1
                review_cells = {
                    str(cell.value).split()[0]: cell
                    for cell in (
                        worksheet.cell(row, review_column)
                        for row in range(2, worksheet.max_row + 1)
                    )
                    if cell.value
                }
                self.assertTrue(
                    str(review_cells["少转"].font.color.rgb)
                    .upper()
                    .endswith(build_workbook.REVIEW_SHORT_FONT_RGB)
                )
                self.assertTrue(
                    str(review_cells["多转"].font.color.rgb)
                    .upper()
                    .endswith(build_workbook.REVIEW_OVER_FONT_RGB)
                )
                formulas = [
                    str(formula)
                    for conditional in worksheet.conditional_formatting
                    for rule in conditional.rules
                    for formula in (rule.formula or [])
                ]
                self.assertTrue(any("少转" in formula for formula in formulas))
                self.assertTrue(any("多转" in formula for formula in formulas))
                self.assertTrue(
                    all(
                        check_workbook.WPS_CONDITIONAL_FORMULA_RE.fullmatch(formula)
                        for formula in formulas
                    )
                )
                self.assertFalse(
                    any(
                        cell.data_type == "f"
                        for row in worksheet.iter_rows()
                        for cell in row
                    )
                )
                actual_received_column = (
                    core.HEADERS.index("收款方实际到账金额") + 1
                )
                actual_received_letter = worksheet.cell(
                    1,
                    actual_received_column,
                ).column_letter
                self.assertGreaterEqual(
                    worksheet.column_dimensions[actual_received_letter].width or 0,
                    18,
                )
            finally:
                workbook.close()

    def test_simple_output_builds_and_checks_workbook(self) -> None:
        normalized, events, plan = fixture()
        pricing_terms = plan["groups"][0]["orders"][0]["pricing"]["terms"]
        pricing_terms["fees"] = [
            {
                "kind": "service_fee",
                "amount": "10",
                "currency": "THB",
                "treatment": "included_in_quote",
            }
        ]
        pricing_terms["rounding"] = {
            "unit": "1",
            "mode": "half_up",
            "currency": "THB",
        }
        orders, _ = simple_ledger.compile_simple_ledger(
            normalized,
            events,
            "sha256:events-test",
            plan,
        )

        with tempfile.TemporaryDirectory(prefix="simple-ledger-workbook-") as temporary:
            root = Path(temporary)
            orders_path = root / "orders.json"
            workbook_path = root / "ledger.xlsx"
            core.atomic_json(orders_path, orders)
            groups = build_workbook.validate_orders(orders)
            workbook = build_workbook.build_workbook(
                groups,
                Path(__file__).resolve().parent.parent / "assets" / "模版.xlsx",
            )
            try:
                workbook.save(workbook_path)
            finally:
                workbook.close()

            self.assertEqual(
                check_workbook.check(workbook_path, orders_path),
                [],
            )
            workbook = load_workbook(workbook_path, read_only=False, data_only=False)
            try:
                worksheet = workbook.worksheets[0]
                headers = [worksheet.cell(1, column).value for column in range(1, len(core.HEADERS) + 1)]
                self.assertEqual(headers, core.HEADERS)
                self.assertEqual(len(headers), len(core.HEADERS))
                self.assertNotIn("客户付款币种", headers)
                self.assertNotIn("内部回款币种", headers)
                self.assertIn("付款合计", headers)
                self.assertIn("汇率", headers)
                self.assertIn("收款方实际到账金额", headers)
                self.assertNotIn("流水金额", headers)
                self.assertNotIn("异常备注", headers)
                fee_row = next(row for row in worksheet.iter_rows(values_only=True) if row[0] == "手续费")
                self.assertEqual(
                    fee_row[core.HEADERS.index("备注")],
                    None,
                )
                self.assertIsNone(fee_row[core.HEADERS.index("收款方")])
                fund_row_types = {"客户付款", "付款退款", "内部回款", "回款追回", "未归单资金图片"}
                for row in worksheet.iter_rows(min_row=2, values_only=True):
                    if row[0] in fund_row_types:
                        self.assertTrue(row[core.HEADERS.index("收款方")])
            finally:
                workbook.close()

            workbook = load_workbook(workbook_path, read_only=False, data_only=False)
            try:
                worksheet = workbook.worksheets[0]
                fund_row = next(
                    row_index
                    for row_index in range(2, worksheet.max_row + 1)
                    if worksheet.cell(row_index, 1).value in fund_row_types
                )
                worksheet.cell(fund_row, core.HEADERS.index("收款方") + 1).value = None
                workbook.save(workbook_path)
            finally:
                workbook.close()
            self.assertTrue(
                any(
                    "fund-flow row requires payee" in error
                    for error in check_workbook.check(workbook_path, orders_path)
                )
            )


if __name__ == "__main__":
    unittest.main()
